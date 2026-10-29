#!/usr/bin/env python3
# SOURCE: ba-kit/requirement-to-flow/.claude/skills/business-analyst/scripts/selftest-basic-design.py @ 447770e
# DA TACH KHOI BAN GOC requirement-to-flow (D2, 27/10/2026): tu sinh fixture bang build-basic-design.py khi
# khong truyen workbook + case check 16-24 (cot S "API gọi", ten / mo ta khong ma, Giai đoạn 1/2). Case check 1-15
# van dung chung.
"""Self-test cho Quality Gate O4 — chung minh gate KHONG pass rong.

Cach dung:
    python3 selftest-basic-design.py                     # tu sinh fixture (inventory + recon + spec) -> --asis
    python3 selftest-basic-design.py <workbook-da-PASS.xlsx> [--before <backup.xlsx>] [--asis]

--asis: gate chay voi --asis + them case check 16-24 (AS-IS, kit system-to-doc).
Case nao thieu dieu kien tien quyet (vd workbook AS-IS khong co error) -> SKIP, khong tinh la loi.

Lay 1 workbook da PASS gate, tiem tung loi da biet vao ban copy, roi xac nhan
gate bat dung check ky vong. Gate PASS ma khong co self-test nay = chua chung
minh duoc gi: mot gate hong van "PASS" moi thu.

Bo test nay da tung tim ra 2 loi that cua gate:
  - check 14 CRASH khi border side = None (AttributeError, khong phai FAIL)
  - check 13 do im.width thay vi anchor.ext -> bo sot anh bi phong to trong Excel
"""
import argparse
import importlib.util
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.styles import Border, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
GATE = os.path.join(HERE, "verify-basic-design.py")
sys.path.insert(0, HERE)
import inv_schema as S  # noqa: E402

# Dung lai chinh logic cua gate de xac dinh working sheet.
# Tu viet lai se sai: sheet `Screen Index` co cot F = "Sheet Type", nen
# ws["F6"] cua no cung bang "WORKING_SCREEN" — phai loai sheet he thong.
_spec = importlib.util.spec_from_file_location("bd_gate", GATE)
_gate = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_gate)


def unmerge_at(ws, r):
    for rng in [str(x) for x in ws.merged_cells.ranges]:
        if ws[rng.split(":")[0]].row == r:
            ws.unmerge_cells(rng)


def screen_sheets(wb):
    return _gate.working_sheets(wb)


def by_sid(wb, sid):
    """Sheet man tim theo Screen ID o F1 (ten sheet = ten man, khong phai SC-xxx)."""
    return next(ws for ws in screen_sheets(wb) if str(ws["F1"].value) == sid)


def first_screen_sheet(wb):
    ws = screen_sheets(wb)
    return ws[0] if ws else None


# moi case: (ten, check ky vong, ham tiem loi)
# ham tiem loi tra ve chuoi ly do khi thieu dieu kien -> SKIP (khong tinh la loi)
def _no_border(wb):
    if "Screen Error message" not in wb.sheetnames:
        return "Screen Error message"
    es = wb["Screen Error message"]
    for r in range(_gate.CATALOG_HEADER_ROW + 1, es.max_row + 1):
        v = es.cell(r, 1).value
        if v and not str(v).startswith(("🔴", "🟡", "🔵", "🟢")):
            for c in range(1, 11):
                es.cell(r, c).border = Border()
            return None
    return "row trong Screen Error message"


def _header_fill(wb):
    if "Screen Index" not in wb.sheetnames or wb["Screen Index"].cell(6, 2).value is None:
        return "row trong Screen Index"
    for c in range(1, 13):
        wb["Screen Index"].cell(6, c).fill = PatternFill("solid", fgColor="FF1F4E78")


def _drop_err_table(wb):
    ws = first_screen_sheet(wb)
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 8).value
        if v and "ERROR SCENARIOS" in str(v):
            for rr in range(r, r + 12):
                unmerge_at(ws, rr)
                for c in range(8, 19):
                    ws.cell(rr, c).value = None
            return None
    return f"bang ERROR SCENARIOS o {ws.title}"


def _touch_sample(wb):
    if "Sample" not in wb.sheetnames:
        return "sheet Sample"
    wb["Sample"]["B1"] = "BI GHI BAY"


def _blow_up_image(wb):
    ws = first_screen_sheet(wb)
    if not (ws and ws._images):
        return "anh UI o screen sheet dau"
    im = ws._images[0]
    im.anchor.ext.cx = int(1600 * 9525)
    im.anchor.ext.cy = int(1200 * 9525)


def _skip_chg_id(wb):
    if "Change History" not in wb.sheetnames or wb["Change History"].cell(6, 1).value is None:
        return "row trong Change History"
    h = wb["Change History"]
    r = 6
    while h.cell(r, 1).value is not None:
        r += 1
    h.cell(r - 1, 1).value = "CHG-0099"


def _drop_index_row(wb):
    if "Screen Index" not in wb.sheetnames or wb["Screen Index"].cell(6, 2).value is None:
        return "row trong Screen Index"
    wb["Screen Index"].delete_rows(6)


def _ghost_code(wb):
    first_screen_sheet(wb)["Q9"] = "loi -> E_ZZ_999"


def _strip_ev(wb):
    # xoa EV o cot R cua item data row dau tien (sau header H="#")
    ws = first_screen_sheet(wb)
    hdr = next((r for r in range(1, 41)
                if str(ws.cell(r, 8).value or "").strip() == "#"), None)
    if hdr is None:
        return f"header bang item o {ws.title}"
    for r in range(hdr + 1, ws.max_row + 1):
        if "ERROR SCENARIOS" in str(ws.cell(r, 8).value or ""):
            break
        name = str(ws.cell(r, 9).value or "") + str(ws.cell(r, 10).value or "")
        body = str(ws.cell(r, 16).value or "") + str(ws.cell(r, 17).value or "")
        if name.strip() and body.strip():
            ws.cell(r, 18).value = "xem man hinh"
            return None
    return f"item data row co mo ta o {ws.title}"


def _item_row(wb, pred):
    """-> (ws, row) cua item dau tien thoa pred(ws, r) trong bang item."""
    for ws in screen_sheets(wb):
        for r in range(9, ws.max_row + 1):
            if "ERROR SCENARIOS" in str(ws.cell(r, 8).value or ""):
                break
            if ws.cell(r, 9).value and pred(ws, r):
                return ws, r
    return None, None


def _blank_api(wb):
    ws, r = _item_row(wb, lambda w, r: str(w.cell(r, 11).value).lower() == "button")
    if not ws:
        return "item Button"
    ws.cell(r, 19).value = None


def _legacy_word(wb):
    ws, r = _item_row(wb, lambda w, r: True)
    if not ws:
        return "item"
    ws.cell(r, 17).value = "Hanh vi khi nhan: " + S.LEGACY_WORD


def _blank_cell(wb):
    ws, r = _item_row(wb, lambda w, r: True)
    if not ws:
        return "item"
    ws.cell(r, 13).value = None   # M required


def _blank_err_catalog(wb):
    es = wb["Screen Error message"]
    for r in range(_gate.CATALOG_HEADER_ROW + 1, es.max_row + 1):
        if es.cell(r, 2).value:          # dong loi (banner khong co cot B)
            es.cell(r, 10).value = None  # Nguon
            return None
    return "dong loi trong Screen Error message"


def _blank_err_scenario(wb):
    for ws in screen_sheets(wb):
        for r in range(1, ws.max_row + 1):
            if "ERROR SCENARIOS" in str(ws.cell(r, 8).value or ""):
                ws.cell(r + 2, 16).value = None   # Message hien thi (VN) cua dong loi dau
                return None
    return "bang ERROR SCENARIOS"


def _draft_status(wb):
    ws = first_screen_sheet(wb)
    ws["E4"], ws["F4"] = "Status", "DRAFT"


def _status_col(wb):
    wb["Screen Index"].cell(5, 12).value = "Status"


def _maxlen_dash(wb):
    ws, r = _item_row(wb, lambda w, r: str(w.cell(r, 11).value).lower() == "textbox")
    if not ws:
        return "item Textbox"
    ws.cell(r, 14).value = "ー"


def _api_in_desc(wb):
    ws, r = _item_row(wb, lambda w, r: True)
    if not ws:
        return "item"
    ws.cell(r, 17).value = "Hành vi khi nhấn: gọi POST /api/orders"


def _handler_name(wb):
    ws, r = _item_row(wb, lambda w, r: True)
    if not ws:
        return "item"
    ws.cell(r, 9).value = "fetchData"


def _sid_sheet_name(wb):
    ws = first_screen_sheet(wb)
    idx = wb["Screen Index"]
    r = next(r for r in range(6, idx.max_row + 1) if idx.cell(r, 5).value == ws.title)
    ws.title = str(ws["F1"].value)
    idx.cell(r, 5).value = ws.title


CASES = [
    ("xoa duong ke data row", 14, _no_border),
    ("data row to fill header", 14, _header_fill),
    ("xoa bang ERROR SCENARIOS", 15, _drop_err_table),
    ("ghi bay vao Sample", 1, _touch_sample),
    ("phong to anh vuot khung", 13, _blow_up_image),
    ("Change ID nhay coc", 12, _skip_chg_id),
    ("xoa row Screen Index", 5, _drop_index_row),
    ("tro Message Code ma", 6, _ghost_code),
]
CASE_ASIS = [("xoa EV o cot R item dau", 16, _strip_ev),
             ("xoa o 'API gọi' cua nut", 17, _blank_api),
             ("con chu cu trong mo ta (D1)", 18, _legacy_word),
             ("bo trong o required cua item", 19, _blank_cell),
             ("bo trong Nguon o Screen Error message", 20, _blank_err_catalog),
             ("bo trong message o ERROR SCENARIOS", 20, _blank_err_scenario),
             ("ghi lai Status / DRAFT o E4/F4", 21, _draft_status),
             ("Screen Index co cot Status", 21, _status_col),
             ("maxlength textbox = ー", 22, _maxlen_dash),
             ("mo ta item co POST /api/...", 23, _api_in_desc),
             ("ten item la ten ham fetchData", 23, _handler_name),
             ("ten sheet = SC-xxx", 24, _sid_sheet_name)]


def make_fixture(tmp):
    """Inventory + fe-api-calls + crawl + spec gia -> build-basic-design.py -> (workbook, before)."""
    import csv
    import json
    v = os.path.join(tmp, "ver1_011026_baseline")
    i = os.path.join(v, "_internal")
    os.makedirs(os.path.join(i, "crawl"))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cols in S.SHEETS.items():
        wb.create_sheet(name).append(cols)
    wb["10_Site"].append(["WEB-01", "Admin", "https://stg.example.jp", "STG", "test_user", "READ_ONLY", "REPO-01", 2, ""])
    # Screen Name lan ma / route -> ten sheet phai la "ログイン"
    wb["02_Screen"].append(["SC-001", "WEB-01", "/login", "SC-001 ログイン (/login)", "Form", "User", "—", "F-001",
                            "EV-0001", 2, "Confirmed", ""])
    wb["02_Screen"].append(["SC-002", "WEB-01", "/orders/:id", "注文詳細", "Detail", "User", "SC-001", "F-002", "EV-0002", 3,
                            "Confirmed", ""])
    wb["08_API_Fields"].append(["API-009", "Request", "body", "password", "string", "Yes", "maxLength 64", "",
                                "REPO-02:src/dto/login.ts#L8"])
    wb["04_DB_Columns"].append(["users", "nick_name", "varchar(30)", "", "", "YES", "", "30", "", "", "", "", "DB-01", ""])
    wb.save(os.path.join(i, "inventory.xlsx"))
    cols = ["Screen Route", "Screen File", "Trigger", "UI Label", "Method", "Path", "API ID", "Via", "Trigger Ref",
            "Code Ref", "Confidence"]
    with open(os.path.join(i, "fe-api-calls.csv"), "w", encoding="utf8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(cols)
        w.writerow(["/orders/[id]", "REPO-01:src/pages/orders/[id].tsx", "mount", "—", "GET", "/api/orders/:id",
                    "API-001", "axios", "REPO-01:src/pages/orders/[id].tsx#L5", "REPO-01:src/services/o.ts#L2", "static"])
        w.writerow(["/orders/[id]", "REPO-01:src/pages/orders/[id].tsx", "handleSave → onClick", "保存", "POST",
                    "/api/orders", "API-002", "axios", "REPO-01:src/pages/orders/[id].tsx#L9", "REPO-01:src/services/o.ts#L3",
                    "static"])
        w.writerow(["/orders/[id]", "REPO-01:src/pages/orders/[id].tsx", "handleDelete → onClick", "削除", "DELETE",
                    "/api/orders/:id", "API-003", "axios", "REPO-01:src/pages/orders/[id].tsx#L12",
                    "REPO-01:src/services/o.ts#L4", "static"])
        # man co spec items day du + trigger khong khop (nhan "—") -> KHONG duoc them dong
        w.writerow(["/login", "REPO-01:src/pages/login.tsx", "submitLogin → onClick", "—", "POST", "/api/session",
                    "API-010", "axios", "REPO-01:src/pages/login.tsx#L31", "REPO-01:src/services/a.ts#L2", "static"])
        # nut chi co trong code, khong co nhan -> ten "Cần xác nhận", ten ham chi o cot R
        w.writerow(["/orders/[id]", "REPO-01:src/pages/orders/[id].tsx", "fetchStats → onClick", "—", "GET",
                    "/api/stats", "API-004", "axios", "REPO-01:src/pages/orders/[id].tsx#L15",
                    "REPO-01:src/services/o.ts#L5", "static"])
    page2 = "https://stg.example.jp/orders/1"
    json.dump([{"url": "https://stg.example.jp/login", "evId": "EV-0001", "items": []},
               {"url": page2, "evId": "EV-0002", "items": [
                   {"tag": "input", "type": "text", "label": "メモ", "maxlength": 50, "required": False},
                   {"tag": "button", "type": "submit", "label": "保存"},
                   {"tag": "button", "type": "button", "label": "印刷"},
                   {"tag": "a", "type": None, "label": "一覧へ戻る"}]}],
              open(os.path.join(i, "crawl", "pages.json"), "w", encoding="utf8"), ensure_ascii=False)
    json.dump([{"page": page2, "url": "https://stg.example.jp/api/stats", "status": 500},
               {"page": page2, "url": "https://stg.example.jp/api/log", "status": "REQUEST_BLOCKED", "noise": "read-only-block"}],
              open(os.path.join(i, "crawl", "network.json"), "w"))
    json.dump([{"url": page2, "text": "Failed to fetch", "noise": "read-only-block"}],
              open(os.path.join(i, "crawl", "console.json"), "w"))
    os.makedirs(os.path.join(i, "crawl", "evidence"))
    try:   # Giai đoạn 2: anh chup cua crawl + "mark" o item -> builder tu danh so theo cot #
        from PIL import Image
        Image.new("RGB", (1400, 900), "white").save(os.path.join(i, "crawl", "evidence", "EV-0001.png"))
    except ImportError:
        pass
    spec = {"screens": {"SC-002": {"title": "注文の詳細"}, "SC-001": dict(purpose="ログインする", items=[
        {"name_jp": "メールアドレス", "name_vn": "Email", "type": "Textbox", "status": "Enable", "required": "Yes",
         "maxlength": "256", "default": "ー", "p": "・必須\n1. 未入力 → E_SC001_001", "q": "・Bắt buộc\n1. Chưa nhập → E_SC001_001",
         "r": "REPO-01:src/pages/login.tsx#L20", "mark": [0.5, 0.2]},
        {"name_jp": "パスワード", "name_vn": "Mật khẩu", "type": "Password", "status": "Enable", "field": "password",
         "p": "・パスワード", "q": "・Mật khẩu", "r": "REPO-01:src/pages/login.tsx#L24", "mark": [0.5, 0.3]},
        {"name_jp": "ニックネーム", "name_vn": "Biệt danh", "type": "Textbox", "status": "Enable", "required": "No",
         "field": "nickName", "p": "・ニックネーム", "q": "・Biệt danh", "r": "REPO-01:src/pages/login.tsx#L26",
         "mark": [0.5, 0.4]},
        {"name_jp": "ログイン", "name_vn": "Đăng nhập", "type": "Button", "status": "Enable", "next": "SC-002",
         "api": "POST /api/login (API-009)", "p": "・ログイン", "q": "・Đăng nhập", "r": "REPO-01:src/pages/login.tsx#L30",
         "mark": [0.5, 0.5]}],
        errors=[{"item": "メールアドレス", "category": "🔴 Validation", "display": "Inline", "code": "E_SC001_001",
                 "msg_jp": "メールアドレスを入力してください。", "msg_vn": "Vui lòng nhập email.",
                 "action_jp": "未入力で「ログイン」を押す", "action_vn": "Để trống rồi nhấn 「ログイン」",
                 "next": "Không gửi form, focus ô email", "source": "REPO-01:src/i18n/ja.json#L3"}])}}
    json.dump(spec, open(os.path.join(i, "spec.json"), "w", encoding="utf8"), ensure_ascii=False)
    out = os.path.join(v, "05_Basic_Design", "BasicDesign_WEB-01_ver1.xlsx")
    before = os.path.join(i, "bd_before_WEB-01.xlsx")
    p = subprocess.run([sys.executable, os.path.join(HERE, "build-basic-design.py"), "--inventory",
                        os.path.join(i, "inventory.xlsx"), "--site", "WEB-01", "--version", "ver1",
                        "--fe-api-calls", os.path.join(i, "fe-api-calls.csv"), "--crawl", os.path.join(i, "crawl"),
                        "--spec", os.path.join(i, "spec.json"), "--q-start", "10", "--out", out, "--before", before,
                        "--date", "2026-10-27"], capture_output=True, text=True)
    if p.returncode:
        sys.exit("FAIL: build-basic-design.py loi:\n" + p.stdout + p.stderr)
    return out, before, json.loads(p.stdout)


def fixture_checks(out, stat):
    """build-basic-design: cot S dung nguon fe-api-calls, nut chi co trong code co dong, loi web vao ERROR SCENARIOS."""
    wb = openpyxl.load_workbook(out)
    ws = by_sid(wb, "SC-002")
    rows = {str(ws.cell(r, 10).value): r for r in range(9, ws.max_row + 1) if ws.cell(r, 10).value}
    api = {k: str(ws.cell(r, 19).value) for k, r in rows.items()}
    txt = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    ok = [("Mở màn -> GET (API-001)", api.get("Mở màn") == "Mở màn → GET /api/orders/:id (API-001)"),
          ("nut 保存 -> POST (API-002)", api.get("保存") == "POST /api/orders (API-002)"),
          ("nut 削除 chi co trong code -> dong item + code-ref", api.get("削除") == "DELETE /api/orders/:id (API-003)"
           and "#L12" in str(ws.cell(rows.get("削除", 1), 18).value)),
          ("nut 印刷 khong co bang chung -> Cần xác nhận — Q-010", api.get("印刷") == S.NEED_CONFIRM + " — Q-010"),
          ("link -> Không gọi API", api.get("一覧へ戻る", "").startswith("Không gọi API")),
          ("loi web HTTP 500 vao ERROR SCENARIOS, noise bi bo", "HTTP 500" in txt and "Failed to fetch" not in txt),
          ("khong freeze panel (R13)", all(w.freeze_panes is None for w in wb.worksheets)),
          ("stdout dem o Cần xác nhận + next_q", stat.get("confirm_cells", 0) >= 1 and stat.get("next_q", 0) > 10)]
    ws1 = by_sid(wb, "SC-001")
    r1 = {str(ws1.cell(r, 9).value): r for r in range(9, ws1.max_row + 1) if ws1.cell(r, 9).value}
    pw, nk, lg = r1.get("パスワード", 1), r1.get("ニックネーム", 1), r1.get("ログイン", 1)
    err1 = next((r for r in range(9, ws1.max_row + 1) if "ERROR SCENARIOS" in str(ws1.cell(r, 8).value or "")),
                ws1.max_row + 1)
    idx = wb["Screen Index"]
    names = {str(w["F1"].value): w.title for w in screen_sheets(wb)}
    fx = next((r for r, v in rows.items() if v and str(ws.cell(v, 9).value).startswith(S.NEED_CONFIRM_JP)), None)
    item_txt = " ".join(str(ws.cell(r, c).value) for r in rows.values() for c in (9, 10, 16, 17))
    ok += [("spec co items + trigger khong khop -> khong them dong (4 item, khong ten ham)",
            sum(1 for r in range(9, ws1.max_row + 1) if isinstance(ws1.cell(r, 8).value, int)
                and "ERROR" not in str(ws1.cell(r - 1, 8).value or "") and r < err1) == 4 and "submitLogin" not in " ".join(str(c.value) for row in ws1.iter_rows() for c in row if c.value)),
           ("ten sheet = ten man (bo SC-xxx / route), title spec",
            names == {"SC-001": "ログイン", "SC-002": "注文の詳細"}),
           ("Screen Index: khong cot Status, B co hyperlink toi sheet",
            "Status" not in [c.value for c in idx[5]] and idx["B6"].hyperlink is not None
            and "ログイン" in str(idx["B6"].hyperlink.location)),
           ("sheet khong con o Status E4/F4", ws1["E4"].value is None and ws1["F4"].value is None),
           ("field password -> 08_API_Fields: Yes / 64 + nguon o R",
            ws1.cell(pw, 13).value == "Yes" and ws1.cell(pw, 14).value == "64"
            and "08_API_Fields API-009.password" in str(ws1.cell(pw, 18).value)),
           ("field nickName ~ nick_name -> 04_DB_Columns maxlength 30",
            ws1.cell(nk, 13).value == "No" and ws1.cell(nk, 14).value == "30"
            and "04_DB_Columns users.nick_name" in str(ws1.cell(nk, 18).value)),
           ("next SC-002 -> 「注文の詳細」 o P / Q",
            "「注文の詳細」へ遷移" in str(ws1.cell(lg, 16).value)
            and "Bấm → chuyển sang màn 「注文の詳細」" in str(ws1.cell(lg, 17).value)),
           ("nut code khong nhan -> ten Cần xác nhận, ten ham chi o R",
            fx is not None and "fetchStats" not in item_txt and "fetchStats" in str(ws.cell(rows[fx], 18).value)),
           ("Textbox web khong co required -> Cần xác nhận (khong ghi ー)",
            str(ws.cell(rows.get("メモ", 1), 13).value).startswith(S.NEED_CONFIRM))]
    try:
        import PIL  # noqa: F401
        ok.append(("Giai đoạn 2: mark + evidence/EV-0001.png -> anh danh so chen vao sheet",
                   len(ws1._images) == 1 and stat.get("images") == 1 and os.path.isfile(
                       os.path.join(os.path.dirname(os.path.dirname(out)), "_internal", "bd", "annotated", "SC-001.png"))))
    except ImportError:
        pass
    bad = 0
    for name, good in ok:
        bad += not good
        print(f"{'build: ' + name:60} {'OK' if good else 'THAT BAI'}")
    return bad


CODE_PAGE = """import { useRouter } from 'next/router';
import { removeOrder, exportCsv } from '../services/o';
export default function Orders() {
  const router = useRouter();
  const handleRemove = async () => { await removeOrder(1); };
  const goEdit = () => { router.push('/orders/1/edit'); };
  return (<div>
    <DesktopDatePicker label="from" />
    <IconButton aria-label="delete" onClick={handleRemove}><DeleteIcon /></IconButton>
    <Button onClick={goEdit}>編集</Button>
    <Button onClick={() => exportCsv()}>CSV出力</Button>
  </div>);
}
"""


def code_checks(tmp):
    """Nut khong khop nhan -> tra code (fe-screen-files.json), nhieu EV / o, nhan co gia tri dong."""
    import csv
    import json
    v = os.path.join(tmp, "ver1_021026_code")
    i, repo = os.path.join(v, "_internal"), os.path.join(tmp, "repo")
    code = os.path.join(i, "recon", "code", "REPO-01")
    for d in (code, os.path.join(i, "crawl"), os.path.join(repo, "src", "pages")):
        os.makedirs(d, exist_ok=True)
    open(os.path.join(repo, "src", "pages", "orders.tsx"), "w", encoding="utf8").write(CODE_PAGE)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cols in S.SHEETS.items():
        wb.create_sheet(name).append(cols)
    wb["10_Site"].append(["WEB-01", "Admin", "https://stg.example.jp", "STG", "test_user", "READ_ONLY", "REPO-01", 1, ""])
    wb["02_Screen"].append(["SC-001", "WEB-01", "/orders", "注文一覧", "List", "User", "—", "F-001",
                            "EV-0003; EV-0004", 1, "Confirmed", ""])
    wb["11_Repo"].append(["REPO-01", repo, "FE", "Next.js", "WEB-01", ""])
    wb.save(os.path.join(i, "inventory.xlsx"))
    ref = "REPO-01:src/pages/orders.tsx#L%d"
    with open(os.path.join(code, "fe-api-calls.csv"), "w", encoding="utf8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["Screen Route", "Screen File", "Trigger", "UI Label", "Method", "Path", "API ID", "Via",
                    "Trigger Ref", "Code Ref", "Confidence"])
        w.writerow(["/orders", "REPO-01:src/pages/orders.tsx", "handleRemove → onClick", "—", "DELETE",
                    "/api/orders/:id", "API-003", "x", ref % 5, "REPO-01:src/services/o.ts#L2", "static"])
        w.writerow(["/orders", "REPO-01:src/pages/orders.tsx", "onClick", "—", "GET", "/api/orders/csv", "API-004",
                    "x", ref % 11, "REPO-01:src/services/o.ts#L3", "static"])
        w.writerow(["/orders", "REPO-01:src/pages/orders.tsx", "handleSave → onClick", "保存", "POST", "/api/orders",
                    "API-002", "x", ref % 20, "REPO-01:src/services/o.ts#L4", "static"])
    json.dump({"/orders": ["src/pages/orders.tsx"]}, open(os.path.join(code, "fe-screen-files.json"), "w"))
    json.dump([{"url": "https://stg.example.jp/orders", "evId": "EV-0004", "items": [
        {"tag": "button", "type": "button", "label": "Choose date, selected date is Oct 28, 2026"},
        {"tag": "button", "type": "button", "label": "delete"},
        {"tag": "button", "type": "button", "label": "編集"},
        {"tag": "button", "type": "button", "label": "CSV出力"},
        {"tag": "button", "type": "button", "label": "保存する"},
        {"tag": "button", "type": "button", "label": "印刷"},
        {"tag": "button", "type": "button", "label": "10/28"},
        {"tag": "button", "type": "button", "label": "10/29"}]}],
        open(os.path.join(i, "crawl", "pages.json"), "w", encoding="utf8"), ensure_ascii=False)
    out = os.path.join(v, "05_Basic_Design", "BasicDesign_WEB-01_ver1.xlsx")
    p = subprocess.run([sys.executable, os.path.join(HERE, "build-basic-design.py"), "--inventory",
                        os.path.join(i, "inventory.xlsx"), "--site", "WEB-01", "--version", "ver1",
                        "--fe-api-calls", os.path.join(code, "fe-api-calls.csv"), "--crawl", os.path.join(i, "crawl"),
                        "--q-start", "1", "--out", out, "--date", "2026-10-28"], capture_output=True, text=True)
    if p.returncode:
        print("build (code): loi\n" + p.stdout + p.stderr)
        return 1
    stat = json.loads(p.stdout)
    ws = by_sid(openpyxl.load_workbook(out), "SC-001")
    rows = {str(ws.cell(r, 10).value): r for r in range(9, ws.max_row + 1) if ws.cell(r, 10).value}
    api = {k: str(ws.cell(r, 19).value) for k, r in rows.items()}
    ref_of = {k: str(ws.cell(r, 18).value) for k, r in rows.items()}
    txt = " ".join(str(c.value) for row in ws.iter_rows(min_row=9, min_col=9, max_col=17) for c in row if c.value)
    ok = [("Screenshot EV 'EV-0003; EV-0004' -> ghep trang EV-0004", "delete" in rows),
          ("nhan datepicker bo ngay da chon -> 'Choose date' + code-ref", api.get("Choose date", "").startswith(
              "Không gọi API — mở bộ chọn") and "orders.tsx#L8" in ref_of.get("Choose date", "")),
          ("nhan khong chua gia tri ngay (2026 / Oct 28)", "2026" not in txt and "Oct 28" not in txt),
          ("o ngay 10/28, 10/29 -> 1 dong （日付）", "（日付）" in rows and "10/28" not in txt),
          ("icon aria-label=delete -> handler -> DELETE (API-003)", api.get("delete") == "DELETE /api/orders/:id (API-003)"
           and "#L9" in ref_of.get("delete", "")),
          ("router.push -> Không gọi API — chuyển trang", api.get("編集", "").startswith("Không gọi API — chuyển trang")),
          ("onClick viet thang trong JSX -> Trigger Ref cung dong (API-004)",
           api.get("CSV出力") == "GET /api/orders/csv (API-004)"),
          ("nhan gan dung 保存する ~ 保存 -> POST (API-002)", api.get("保存する") == "POST /api/orders (API-002)"),
          ("khong co trong code -> Cần xác nhận — Q-001", api.get("印刷") == S.NEED_CONFIRM + " — Q-001"),
          ("stdout code_matched / confirm_cells", stat.get("code_matched") == 4 and stat.get("confirm_cells") == 2)]
    bad = 0
    for name, good in ok:
        bad += not good
        print(f"{'code: ' + name:60} {'OK' if good else 'THAT BAI'}")
    if bad:
        print(json.dumps({"api": api, "stat": stat}, ensure_ascii=False))
    return bad


def gate_levels(path, extra=()):
    """-> (exit, {check: PASS|FAIL|WARN})."""
    p = subprocess.run([sys.executable, GATE, path, "--asis"] + list(extra), capture_output=True, text=True)
    lv = {}
    for line in p.stdout.splitlines():
        c = [x.strip() for x in line.split("|")]
        if len(c) > 4 and c[1].isdigit():
            lv[int(c[1])] = c[3].strip("*")
    return p.returncode, lv


def phase_checks(tmp, out2):
    """Giai đoạn 1 (khong --crawl): build chay, gate PASS, check 10/13 PASS; Giai đoạn 2 thieu anh -> check 10 WARN."""
    import json
    v = os.path.join(tmp, "ver1_011026_baseline")
    i = os.path.join(v, "_internal")
    out1 = os.path.join(v, "05_Basic_Design", "BasicDesign_WEB-01_ver1_phase1.xlsx")
    p = subprocess.run([sys.executable, os.path.join(HERE, "build-basic-design.py"), "--inventory",
                        os.path.join(i, "inventory.xlsx"), "--site", "WEB-01", "--version", "ver1",
                        "--fe-api-calls", os.path.join(i, "fe-api-calls.csv"), "--spec", os.path.join(i, "spec.json"),
                        "--q-start", "50", "--out", out1, "--questions-out", os.path.join(i, "q1.csv"),
                        "--date", "2026-10-29"], capture_output=True, text=True)
    if p.returncode:
        print("phase 1: build loi\n" + p.stdout + p.stderr)
        return 1
    stat = json.loads(p.stdout)
    wb = openpyxl.load_workbook(out1)
    e1, l1 = gate_levels(out1)
    e1b, l1b = gate_levels(out1, ["--phase", "2"])
    _, l2 = gate_levels(out2)
    ok = [("Giai đoạn 1: stdout phase=1, khong anh", stat.get("phase") == 1 and stat.get("images") == 0),
          ("Giai đoạn 1: A9 = 'Chưa có ảnh — bổ sung ở Giai đoạn 2 (Playwright)'",
           all(str(w["A9"].value) == "Chưa có ảnh — bổ sung ở Giai đoạn 2 (Playwright)" for w in screen_sheets(wb))),
          ("Giai đoạn 1: gate --asis exit 0, check 10/13 PASS (khong WARN)",
           e1 == 0 and l1.get(10) == "PASS" and l1.get(13) == "PASS"),
          ("Giai đoạn 1 ep --phase 2 -> check 10 WARN", l1b.get(10) == "WARN"),
          ("Giai đoạn 2: man thieu anh -> check 10 WARN", l2.get(10) == "WARN"),
          ("check 20-24 PASS tren workbook dung", all(l1.get(k) == "PASS" for k in range(20, 25))
           and all(l2.get(k) == "PASS" for k in range(20, 25)))]
    bad = 0
    for name, good in ok:
        bad += not good
        print(f"{'phase: ' + name:60} {'OK' if good else 'THAT BAI'}")
    if bad:
        print(json.dumps({"l1": l1, "l1b": l1b, "l2": l2, "stat": stat}, ensure_ascii=False))
    return bad


def run_gate(path, before, screens, asis=False):
    cmd = [sys.executable, GATE, path, "--expect-screens", screens]
    if before:
        cmd += ["--before", before]
    if asis:
        cmd.append("--asis")
    p = subprocess.run(cmd, capture_output=True, text=True)
    out = p.stdout + p.stderr
    if "Traceback" in out:
        return "CRASH", []
    fails = []
    for line in out.splitlines():
        if line.startswith("| ") and "**FAIL**" in line:
            try:
                fails.append(int(line.split("|")[1].strip()))
            except ValueError:
                pass
    return p.returncode, fails


def next_checks():
    """"next" = ー / trong -> khong co man dich (khong hoi); nhieu SC -> noi ten man."""
    sp = importlib.util.spec_from_file_location("bdb", os.path.join(HERE, "build-basic-design.py"))
    m = importlib.util.module_from_spec(sp)
    sp.loader.exec_module(m)
    ctx = {"titles": {"SC-004": "予約管理", "SC-012": "店舗一覧"}, "routes": {}, "fields": None}
    ok = ("ー" in m.NO_NEXT and m.screen_ref("SC-004 / SC-012", ctx) == "予約管理」・「店舗一覧"
          and m.screen_ref("ー（マスタータイプは SC-012）", ctx) == "店舗一覧" and m.screen_ref("SC-999", ctx) is None)
    print("%-34s %s" % ("next: ー / nhieu SC / SC la", "OK" if ok else "THAT BAI"))
    return 0 if ok else 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook", nargs="?", help="bo trong -> tu sinh fixture bang build-basic-design.py (--asis)")
    ap.add_argument("--before")
    ap.add_argument("--asis", action="store_true",
                    help="gate chay --asis + them case check 16-24")
    args = ap.parse_args()
    pre_bad = 0
    fixture_dir = None
    if not args.workbook:
        fixture_dir = tempfile.mkdtemp(prefix="selftest-bd-")
        args.workbook, args.before, stat = make_fixture(fixture_dir)
        args.asis = True
        pre_bad = fixture_checks(args.workbook, stat) + code_checks(fixture_dir) + \
            phase_checks(fixture_dir, args.workbook) + next_checks()
    cases = CASES + (CASE_ASIS if args.asis else [])

    wb = openpyxl.load_workbook(args.workbook)
    screens = ",".join(str(ws["F1"].value) for ws in screen_sheets(wb))
    if not screens:
        sys.exit("FAIL: workbook khong co working sheet nao")

    code, fails = run_gate(args.workbook, args.before, screens, args.asis)
    print(f"{'BASELINE (khong tiem loi)':34} exit={code} fails={fails or '[]'}")
    if code != 0:
        sys.exit("FAIL: workbook dau vao phai PASS gate truoc khi self-test")

    tmp = tempfile.mkdtemp()
    bad = skip = 0
    print()
    print(f"{'LOI DA TIEM':34} {'KY VONG':8} {'BAT DUOC':12} KQ")
    print("-" * 74)
    for name, want, inject in cases:
        p = os.path.join(tmp, f"t_{want}_{abs(hash(name)) % 9999}.xlsx")
        shutil.copy(args.workbook, p)
        w = openpyxl.load_workbook(p)
        why = inject(w)
        if why:
            skip += 1
            print(f"{name:34} #{want:<7} {'-':12} SKIP — workbook khong co {why}")
            continue
        w.save(p)
        code, fails = run_gate(p, args.before, screens, args.asis)
        ok = code == 1 and want in fails
        if not ok:
            bad += 1
        print(f"{name:34} #{want:<7} {str(fails) if fails != 'CRASH' else 'CRASH':12} "
              f"{'OK' if ok else 'THAT BAI'}")
    shutil.rmtree(tmp, ignore_errors=True)
    if fixture_dir:
        shutil.rmtree(fixture_dir, ignore_errors=True)
    print()
    n = len(cases) - skip
    print(f"{n - bad}/{n} case bat dung check ky vong" + (f" ({skip} SKIP)" if skip else "")
          + (f" · build-basic-design {pre_bad} THAT BAI" if pre_bad else "")
          + f" — {'PASS' if not (bad or pre_bad) else 'FAIL'}")
    sys.exit(1 if bad or pre_bad else 0)


if __name__ == "__main__":
    main()
