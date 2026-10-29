#!/usr/bin/env python3
# SOURCE: ba-kit/requirement-to-flow/.claude/skills/business-analyst/scripts/verify-basic-design.py @ 447770e
# DA TACH KHOI BAN GOC requirement-to-flow (D2, 27/10/2026): kit nay them cot S "API gọi" + check 16-24.
# Check 1-15 van la logic chung — sua check 1-15 thi sua CA HAI noi; check 16-24 chi o day.
# Rieng kit nay (29/10/2026): bo cot / o "Status" (check 4 khong doi F4), check 10/13 theo Phase.
"""Quality Gate O4 — Basic Design master workbook (= gate V-BD cua kit system-to-doc).

Chay 15 check tren master workbook (+ check 16-24 khi co --asis):
  16 item co mo ta -> cot R co EV-xxxx / code-ref REPO-xx:path#Lnn / "Cần xác nhận — Q-xxx"
  17 item dang button / link / form -> cot S "API gọi" co API-xxx, "Không gọi API" hoac "Cần xác nhận — Q-xxx"
  18 khong con chu cu inv_schema.LEGACY_WORD (D1)
  19 khong o trong o cot I..Q cua item (khong ap dung ghi ー)
  20 khong o trong o dong loi: Screen Error message, Common mesage, bang ERROR SCENARIOS
  21 khong con "Status" / DRAFT: Screen Index khong co cot Status, screen sheet E4/F4 trong
  22 item nhap lieu: required (M) / maxlength (N) co gia tri hoac "Cần xác nhận — Q-xxx"
     ("ー" chi khi khong ap dung: maxlength cua select / radio / checkbox / date; required cua item Disable)
  23 ten item (I/J) va mo ta (P/Q) khong chua "GET /api/..", chu "API", ten ham camelCase
  24 ten sheet / Screen Name doc duoc: khong SC-xxx, route /..., camelCase, placeholder "chưa có tiêu đề"
  10/13 theo Phase: Giai đoạn 1 (chua chay Playwright) chua co anh -> PASS; Giai đoạn 2 -> WARN nhu cu
Doc rule day du: .claude/sys-agent/outputs/o5-basic-design.md

Usage:
  python3 verify-basic-design.py <master.xlsx> \
      [--before <backup.xlsx>] \
      [--expect-screens "ID1,ID2"] [--asis] [--phase 1|2|auto] \
      [--out <report.md>]

Exit: 0 = PASS | 1 = co FAIL | 2 = thieu openpyxl / khong doc duoc file
"""
import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import inv_schema as S  # noqa: E402
from gate_report import legacy_word_hits  # noqa: E402

try:
    import openpyxl
    import openpyxl.utils
except ImportError:
    sys.stderr.write(
        "FAIL: thieu openpyxl. Chay: pip install openpyxl\n")
    sys.exit(2)

SYSTEM_SHEETS = {
    "Common mesage",
    "Screen Error message",
    "Screen Index",
    "Change History",
}
TEMPLATE_SHEET = "Sample"

# Cell map — xem workbook-structure.md section 2.1
META = {
    "screen_id": "F1",
    "flow_id": "F2",
    "version": "F3",
    "source_flow_version": "F5",
    "sheet_type": "F6",
}
PLACEHOLDERS = {"", "TEMPLATE", "TBD", "N/A", "[PLACEHOLDER]", "NONE", "-"}

SCREEN_INDEX_HEADER_ROW = 5
CHANGE_HISTORY_HEADER_ROW = 5
CATALOG_HEADER_ROW = 3

# Vung anh UI = cot A-F. Tong column width 122 ky tu ~ 854 px.
HEADER_FILLS = {"FFDAEEF3", "FF1F4E78"}   # fill cua header/title — data row cam dung
ERR_TITLE_MARK = "ERROR SCENARIOS"

IMG_MAX_W = 840   # px, chua le 14px
IMG_MAX_H = 900   # px

SAMPLE_LIKE_RE = re.compile(r"(?i)(^|[_ -])(sample|example|demo|template|mau|copy)([_ -]|$)")

EV_RE = re.compile(r"\bEV-\d{4,}\b|\bREPO-\d{2,}:\S+#L\d+")   # EV hoac code-ref (code-first, R3)
CONFIRM_RE = re.compile(r"(%s|%s) — Q-\d{3,}" % (S.NEED_CONFIRM, S.NEED_CONFIRM_JP))
API_CALL_COL = 19          # S — "API gọi" (D2: ngay sau cot R)
API_CALL_HDR = "API gọi"
CLICK_TYPES = ("button", "hyperlink", "link", "form", "submit")
API_CALL_RE = re.compile(r"\bAPI-\d{3,}\b|Không gọi API|(%s|%s) — Q-\d{3,}" % (S.NEED_CONFIRM, S.NEED_CONFIRM_JP))

CODE_RE = re.compile(r"\b(?:E_[A-Z0-9]+(?:_[A-Z])?_\d{3}|COMMON_[A-Z]+_\d{3})\b")

PHASE1_MARK = "bổ sung ở Giai đoạn 2"   # A9 cua sheet chua co anh khi build Giai đoạn 1 (build-basic-design.py)
# dong bo voi build-basic-design.py INPUT_TYPES / NO_MAXLEN / CODE_RX
INPUT_TYPES = {"textbox", "text", "input", "textarea", "select", "dropdown", "combobox", "radio", "checkbox",
               "date", "datepicker", "time", "number", "password", "email"}
NO_MAXLEN = {"select", "dropdown", "combobox", "radio", "checkbox", "date", "datepicker", "time"}
TEXT_MAXLEN = {"textbox", "text", "input", "textarea", "password", "email"}
# shortcut: camelCase bat ca ten rieng kieu iPhone / YouTube — gap thi doi ten hien thi, khong ha check
CODE_TXT_RE = re.compile(r"\b(?:GET|POST|PUT|PATCH|DELETE)\s+/|\bAPI\b|\b[a-z]+[A-Z][A-Za-z0-9]*\b|=>|\w+\(\)")
BAD_NAME_RE = re.compile(r"\bSC-\d+\b|(^|\s)/[A-Za-z0-9_:\[\]{}.\-]|\b[a-z]+[A-Z][A-Za-z0-9]*\b|"
                         r"(?i:chưa có tiêu đề|chưa đặt tên|no title|untitled|\[SCREEN NAME\]|^TEMPLATE$)|"
                         r"%s|%s" % (S.NEED_CONFIRM, S.NEED_CONFIRM_JP))


class Report:
    def __init__(self):
        self.rows = []

    def add(self, num, name, level, detail):
        self.rows.append((num, name, level, detail))

    def counts(self):
        fail = sum(1 for r in self.rows if r[2] == "FAIL")
        warn = sum(1 for r in self.rows if r[2] == "WARN")
        ok = sum(1 for r in self.rows if r[2] == "PASS")
        return ok, fail, warn


def has_border(c):
    """True khi cell co du 4 canh. Side co the la None -> khong duoc crash."""
    b = c.border
    if b is None:
        return False
    for e in ("left", "right", "top", "bottom"):
        side = getattr(b, e, None)
        if side is None or not side.style:
            return False
    return True


def _fold(s):
    """Bo dau + lower de so khop caption khong phu thuoc dau tieng Viet."""
    import unicodedata
    s = unicodedata.normalize("NFD", s.replace("đ", "d").replace("Đ", "D"))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").lower()


def img_size_px(im):
    """Kich thuoc Excel THUC SU render (EMU trong anchor.ext), fallback im.width/height.

    Gan im.width sau khi load KHONG doi anchor.ext -> doc im.width la doc kich
    thuoc goc cua file anh, khong phai kich thuoc hien thi.
    """
    ext = getattr(getattr(im, "anchor", None), "ext", None)
    if ext is not None and getattr(ext, "cx", None) and getattr(ext, "cy", None):
        return ext.cx / 9525.0, ext.cy / 9525.0, "anchor.ext"
    return getattr(im, "width", None), getattr(im, "height", None), "im.width"


def cell(ws, ref):
    v = ws[ref].value
    return "" if v is None else str(v).strip()


def sheet_values(ws):
    """Snapshot gia tri cell de so sanh truoc/sau."""
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                out[c.coordinate] = str(c.value)
    return out


def is_screen_sheet(ws):
    """Sheet co o Sheet Type -> la screen sheet (template hoac working)."""
    return cell(ws, META["sheet_type"]) != ""


def working_sheets(wb):
    return [ws for ws in wb.worksheets
            if ws.title not in SYSTEM_SHEETS
            and ws.title != TEMPLATE_SHEET
            and is_screen_sheet(ws)]


def read_screen_index(wb):
    """-> {sheet_name: {col_letter: value}} + list row loi."""
    if "Screen Index" not in wb.sheetnames:
        return None, []
    ws = wb["Screen Index"]
    # tim cot theo header (kit nay da bo cot "Status" -> UI Source o H thay vi I)
    col = {str(c.value or "").strip(): c.column_letter for c in ws[SCREEN_INDEX_HEADER_ROW] if c.value}
    rows = []
    for r in range(SCREEN_INDEX_HEADER_ROW + 1, ws.max_row + 1):
        screen_id = cell(ws, f"B{r}")
        sheet_name = cell(ws, f"E{r}")
        if not screen_id and not sheet_name:
            continue
        rows.append({
            "row": r,
            "screen_id": screen_id,
            "sheet_name": sheet_name,
            "screen_name": cell(ws, f"D{r}"),
            "ui_source": cell(ws, f"{col.get('UI Source', 'I')}{r}"),
        })
    return ws, rows


def catalog_codes(wb):
    codes = set()
    for name in ("Screen Error message", "Common mesage"):
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        for r in range(CATALOG_HEADER_ROW + 1, ws.max_row + 1):
            v = cell(ws, f"A{r}")
            if v:
                codes.add(v)
    return codes


def referenced_codes(ws):
    """Moi Message Code duoc tro trong cot P / Q cua bang item."""
    found = {}
    for col in ("P", "Q"):
        for r in range(1, ws.max_row + 1):
            v = ws[f"{col}{r}"].value
            if not v:
                continue
            for m in CODE_RE.findall(str(v)):
                found.setdefault(m, f"{col}{r}")
    return found


def change_history_rows(wb):
    if "Change History" not in wb.sheetnames:
        return None
    ws = wb["Change History"]
    n = 0
    for r in range(CHANGE_HISTORY_HEADER_ROW + 1, ws.max_row + 1):
        if cell(ws, f"A{r}"):
            n += 1
    return n


def run(master_path, before_path, expect_screens, approved, rep, asis_mode=False, phase="auto"):
    try:
        wb = openpyxl.load_workbook(master_path)
    except Exception as exc:  # noqa: BLE001
        sys.stderr.write(f"FAIL: khong mo duoc {master_path}: {exc}\n")
        sys.exit(2)

    wb_before = None
    if before_path:
        try:
            wb_before = openpyxl.load_workbook(before_path)
        except Exception as exc:  # noqa: BLE001
            sys.stderr.write(f"FAIL: khong mo duoc backup {before_path}: {exc}\n")
            sys.exit(2)

    ws_list = working_sheets(wb)
    idx_ws, idx_rows = read_screen_index(wb)
    if phase == "auto":   # build Giai đoạn 1 ghi PHASE1_MARK o A9 cua sheet chua co anh
        phase = "1" if any(PHASE1_MARK in cell(ws, "A9") for ws in ws_list) else "2"

    # --- Check 1: Sample immutable -------------------------------------
    if TEMPLATE_SHEET not in wb.sheetnames:
        rep.add(1, "Sample immutable", "FAIL",
                f"Sheet `{TEMPLATE_SHEET}` khong ton tai — da bi xoa hoac doi ten")
    elif wb_before is None:
        rep.add(1, "Sample immutable", "WARN",
                "Khong co --before nen chi kiem tra su ton tai. "
                "Chay lai kem --before de so byte.")
    elif TEMPLATE_SHEET not in wb_before.sheetnames:
        rep.add(1, "Sample immutable", "WARN",
                "Backup khong co sheet Sample — khong so sanh duoc")
    else:
        a = sheet_values(wb_before[TEMPLATE_SHEET])
        b = sheet_values(wb[TEMPLATE_SHEET])
        diff = sorted(set(a) ^ set(b)) + \
            sorted(k for k in set(a) & set(b) if a[k] != b[k])
        if diff:
            rep.add(1, "Sample immutable", "FAIL",
                    f"Sheet Sample bi sua {len(diff)} cell: {diff[:10]}")
        else:
            rep.add(1, "Sample immutable", "PASS", "Sample khong bi cham")

    # --- Check 2: Sheet Type da neutralize -----------------------------
    bad = [ws.title for ws in ws_list
           if cell(ws, META["sheet_type"]) != "WORKING_SCREEN"]
    if bad:
        rep.add(2, "Sheet Type = WORKING_SCREEN", "FAIL",
                "Sheet chua neutralize (F6 != WORKING_SCREEN): " + ", ".join(bad))
    else:
        rep.add(2, "Sheet Type = WORKING_SCREEN", "PASS",
                f"{len(ws_list)} working sheet dung Sheet Type")

    # --- Check 3: Screen ID khong trung -------------------------------
    seen = {}
    dup = []
    for ws in ws_list:
        sid = cell(ws, META["screen_id"])
        if sid and sid in seen:
            dup.append(f"{sid} ({seen[sid]} + {ws.title})")
        seen[sid] = ws.title
    if dup:
        rep.add(3, "Screen ID duy nhat", "FAIL", "Trung Screen ID: " + "; ".join(dup))
    else:
        rep.add(3, "Screen ID duy nhat", "PASS", f"{len(seen)} Screen ID phan biet")

    # --- Check 4: metadata day du -------------------------------------
    missing = []
    for ws in ws_list:
        for key in ("screen_id", "flow_id", "version", "source_flow_version"):
            v = cell(ws, META[key])
            if key == "screen_id" and v.upper() in PLACEHOLDERS:
                missing.append(f"{ws.title}.{META[key]}({key})")
            elif not v:
                missing.append(f"{ws.title}.{META[key]}({key})")
    if missing:
        rep.add(4, "Metadata F1,F2,F3,F5 day du", "FAIL",
                "Thieu/placeholder: " + ", ".join(missing[:12]))
    else:
        rep.add(4, "Metadata F1,F2,F3,F5 day du", "PASS", "Moi working sheet du metadata")

    # --- Check 5: Screen Index dong bo --------------------------------
    if idx_ws is None:
        rep.add(5, "Screen Index dong bo", "FAIL", "Khong co sheet `Screen Index`")
    else:
        by_sheet = {}
        for r in idx_rows:
            by_sheet.setdefault(r["sheet_name"], []).append(r)
        problems = []
        for ws in ws_list:
            hits = by_sheet.get(ws.title, [])
            if len(hits) == 0:
                problems.append(f"{ws.title}: thieu row trong Screen Index")
            elif len(hits) > 1:
                problems.append(f"{ws.title}: co {len(hits)} row trung")
            else:
                sid = cell(ws, META["screen_id"])
                if hits[0]["screen_id"] != sid:
                    problems.append(
                        f"{ws.title}: Screen ID lech "
                        f"(sheet={sid} vs index={hits[0]['screen_id']})")
        orphan = [r["sheet_name"] for r in idx_rows
                  if r["sheet_name"] and r["sheet_name"] not in wb.sheetnames]
        for o in orphan:
            problems.append(f"Screen Index tro toi sheet khong ton tai: {o}")
        if problems:
            rep.add(5, "Screen Index dong bo", "FAIL", "; ".join(problems[:10]))
        else:
            rep.add(5, "Screen Index dong bo", "PASS",
                    f"{len(ws_list)} sheet khop Screen Index")

    # --- Check 6: Message Code ton tai --------------------------------
    known = catalog_codes(wb)
    dangling = []
    total_refs = 0
    for ws in ws_list:
        refs = referenced_codes(ws)
        total_refs += len(refs)
        for code, where in refs.items():
            if code not in known:
                dangling.append(f"{ws.title}!{where} -> {code}")
    if dangling:
        rep.add(6, "Message Code ton tai", "FAIL",
                f"{len(dangling)} code chua co trong catalog: "
                + ", ".join(dangling[:10]))
    else:
        rep.add(6, "Message Code ton tai", "PASS",
                f"{total_refs} reference, tat ca deu co trong catalog")

    # --- Check 7: Change History co row moi ---------------------------
    now = change_history_rows(wb)
    if now is None:
        rep.add(7, "Change History co row moi", "FAIL",
                "Khong co sheet `Change History`")
    elif wb_before is not None:
        was = change_history_rows(wb_before) or 0
        if now <= was:
            rep.add(7, "Change History co row moi", "FAIL",
                    f"Khong co row moi (truoc={was}, sau={now})")
        else:
            rep.add(7, "Change History co row moi", "PASS",
                    f"+{now - was} row")
    elif now == 0:
        rep.add(7, "Change History co row moi", "FAIL", "Change History rong")
    else:
        rep.add(7, "Change History co row moi", "WARN",
                f"{now} row — khong co --before nen khong biet row nao moi")

    # --- Check 8: sheet ngoai scope khong doi -------------------------
    if wb_before is None:
        rep.add(8, "Sheet ngoai scope khong doi", "WARN",
                "Bo qua — can --before de so sanh")
    else:
        in_scope = {ws.title for ws in ws_list} | {
            "Screen Index", "Change History",
            "Common mesage", "Screen Error message"}
        touched, declared = [], []
        for name in wb_before.sheetnames:
            if name in in_scope or name == TEMPLATE_SHEET:
                continue
            changed = None
            if name not in wb.sheetnames:
                changed = f"{name} (bi xoa)"
            elif sheet_values(wb_before[name]) != sheet_values(wb[name]):
                changed = name
            if changed is None:
                continue
            (declared if name in approved else touched).append(changed)
        if touched:
            rep.add(8, "Sheet ngoai scope khong doi", "FAIL",
                    "Sheet ngoai scope bi sua ma KHONG duoc khai bao: "
                    + ", ".join(touched)
                    + ". Neu user da approve mo rong scope, chay lai kem "
                      "--approved-scope \"<ten sheet>\".")
        elif declared:
            rep.add(8, "Sheet ngoai scope khong doi", "WARN",
                    "Thay doi ngoai scope da duoc user approve: "
                    + ", ".join(declared))
        else:
            rep.add(8, "Sheet ngoai scope khong doi", "PASS",
                    "Khong sheet ngoai scope nao bi cham")

    # --- Check 9: khop --expect-screens -------------------------------
    if not expect_screens:
        rep.add(9, "Khop Screen Creation Plan", "WARN",
                "Bo qua — khong truyen --expect-screens")
    else:
        actual = {cell(ws, META["screen_id"]) for ws in ws_list}
        expected = set(expect_screens)
        miss = sorted(expected - actual)
        extra = sorted(actual - expected)
        msgs = []
        if miss:
            msgs.append("THIEU: " + ", ".join(miss))
        if extra:
            msgs.append("THUA: " + ", ".join(extra))
        if msgs:
            rep.add(9, "Khop Screen Creation Plan", "FAIL", " | ".join(msgs))
        else:
            rep.add(9, "Khop Screen Creation Plan", "PASS",
                    f"Dung {len(expected)} screen, THIEU: [] THUA: []")

    # --- Check 10: sheet khong anh phai khai bao ----------------------
    if idx_ws is None:
        rep.add(10, "Khai bao UI Source", "FAIL", "Khong co sheet `Screen Index`")
    else:
        by_sheet = {r["sheet_name"]: r for r in idx_rows}
        undeclared, declared = [], []
        for ws in ws_list:
            has_img = len(getattr(ws, "_images", [])) > 0
            row = by_sheet.get(ws.title)
            ui = (row or {}).get("ui_source", "")
            if not has_img:
                if "NO IMAGE" not in ui.upper():
                    undeclared.append(f"{ws.title} (UI Source={ui or 'rong'})")
                else:
                    declared.append(ws.title)
        if undeclared:
            rep.add(10, "Khai bao UI Source", "FAIL",
                    "Sheet khong co anh nhung khong khai bao `NO IMAGE`: "
                    + ", ".join(undeclared))
        elif declared and phase == "1":
            rep.add(10, "Khai bao UI Source", "PASS",
                    f"Giai đoạn 1 (chua chay Playwright): {len(declared)} sheet chua co anh, da khai bao — "
                    "bo sung o Giai đoạn 2")
        elif declared:
            rep.add(10, "Khai bao UI Source", "WARN",
                    f"{len(declared)} sheet chua co anh, da khai bao dung: "
                    + ", ".join(declared))
        else:
            rep.add(10, "Khai bao UI Source", "PASS",
                    "Moi working sheet deu co anh UI")

    # --- Check 11: khong con sheet vi du tu file mau -------------------
    leftovers = []
    for ws in ws_list:
        sid = cell(ws, META["screen_id"])
        hit = SAMPLE_LIKE_RE.search(ws.title) or SAMPLE_LIKE_RE.search(sid)
        in_plan = sid in set(expect_screens) if expect_screens else False
        if hit and not in_plan:
            leftovers.append(f"{ws.title} (Screen ID={sid or 'rong'})")
    if leftovers:
        rep.add(11, "Khong con sheet vi du tu file mau", "FAIL",
                "Sheet vi du/demo cua template con sot lai trong master du an: "
                + ", ".join(leftovers)
                + ". Khi khoi tao master tu file mau, PHAI xoa moi sheet khong "
                  "thuoc du an (xem output-5-basic-design.md muc 1.2). "
                  "Neu day THUC SU la screen cua du an, them Screen ID vao "
                  "--expect-screens.")
    else:
        rep.add(11, "Khong con sheet vi du tu file mau", "PASS",
                "Khong co sheet vi du nao ngoai Screen Creation Plan")

    # --- Check 12: Change ID lien tuc, khong trung ---------------------
    if "Change History" in wb.sheetnames:
        hs = wb["Change History"]
        ids = []
        for r in range(CHANGE_HISTORY_HEADER_ROW + 1, hs.max_row + 1):
            v = cell(hs, f"A{r}")
            if v:
                ids.append(v)
        bad = []
        seen_ids = set()
        for n, v in enumerate(ids, start=1):
            if v in seen_ids:
                bad.append(f"trung: {v}")
            seen_ids.add(v)
            m = re.fullmatch(r"CHG-(\d{4})", v)
            if not m:
                bad.append(f"sai dinh dang: {v} (can CHG-NNNN)")
            elif int(m.group(1)) != n:
                bad.append(f"nhay coc: {v} (cho doi CHG-{n:04d})")
        if bad:
            rep.add(12, "Change ID lien tuc", "FAIL",
                    f"{len(ids)} row Change History: " + "; ".join(bad[:8]))
        else:
            rep.add(12, "Change ID lien tuc", "PASS",
                    f"{len(ids)} row, CHG-0001..CHG-{len(ids):04d} lien tuc")
    else:
        rep.add(12, "Change ID lien tuc", "FAIL", "Khong co sheet `Change History`")

    # --- Check 13: anh UI khong tran khoi vung A-F --------------------
    bad_img = []
    n_img = 0
    for ws in ws_list:
        for k, im in enumerate(getattr(ws, "_images", []), start=1):
            n_img += 1
            w, h, src = img_size_px(im)
            col = getattr(getattr(im.anchor, "_from", None), "col", None)
            row = getattr(getattr(im.anchor, "_from", None), "row", None)
            pos = f"{ws.title} anh#{k}"
            if col is not None and col != 0:
                bad_img.append(f"{pos}: anchor cot {col + 1} (phai la cot A)")
            if w and w > IMG_MAX_W:
                bad_img.append(
                    f"{pos}: rong {int(w)}px > {IMG_MAX_W}px ({src}) — tran sang cot H:R, "
                    f"de len bang item" + (f" (anchor A{row + 1})" if row is not None else ""))
            if h and h > IMG_MAX_H:
                bad_img.append(f"{pos}: cao {int(h)}px > {IMG_MAX_H}px")
    if bad_img:
        rep.add(13, "Anh UI vua vung A-F", "FAIL",
                f"{len(bad_img)} van de: " + "; ".join(bad_img[:6])
                + ". Resize giu ty le theo output-5-basic-design.md muc 1.1d.")
    elif n_img:
        # anh phai duoc danh so (badge do = cot # cua bang item): caption o A9 phai noi ro
        no_cap = [ws.title for ws in ws_list
                  if getattr(ws, "_images", [])
                  and "so do tren anh" not in _fold(str(ws["A9"].value or ""))]
        if no_cap:   # yeu cau nguoi dung 29/10/2026: anh chen vao Basic Design BAT BUOC danh so -> FAIL
            rep.add(13, "Anh UI vua vung A-F", "FAIL",
                    f"{n_img} anh vua khung, NHUNG {len(no_cap)} sheet thieu caption A9 "
                    f"'so do tren anh = # trong bang item' (anh chua duoc danh so?): "
                    + "; ".join(no_cap[:6]) + ". Chay giai doan 2 (crawl ghi vi tri -> tu danh so) hoac khai 'mark'.")
        else:
            rep.add(13, "Anh UI vua vung A-F", "PASS",
                    f"{n_img} anh, deu nam trong {IMG_MAX_W}x{IMG_MAX_H}px, co caption danh so "
                    f"(do bang anchor.ext — kich thuoc Excel thuc su render)")
    elif phase == "1":
        rep.add(13, "Anh UI vua vung A-F", "PASS",
                "Giai đoạn 1: chua co anh (bo sung o Giai đoạn 2) — khong co gi de kiem")
    else:
        rep.add(13, "Anh UI vua vung A-F", "WARN",
                "Khong co anh nao de kiem (moi sheet deu NO IMAGE)")

    # --- Check 14: style data row -------------------------------------
    def scan_rows(ws, hdr, cols, label, limit=400):
        bad = []
        n = 0
        for r in range(hdr + 1, min(ws.max_row, hdr + limit) + 1):
            if all(ws.cell(row=r, column=c).value is None for c in cols):
                continue
            # bo qua banner (row merged toan bo dai cot)
            merged = any(rng.min_row == r and rng.min_col == cols[0]
                         and rng.max_col >= cols[-1] for rng in ws.merged_cells.ranges)
            if merged:
                continue
            n += 1
            for c in cols:
                cell = ws.cell(row=r, column=c)
                if not has_border(cell):
                    bad.append(f"{label} r{r} cot {c}: thieu duong ke")
                    break
                f = cell.fill
                rgb = f.fgColor.rgb if f and f.fill_type else None
                if rgb in HEADER_FILLS:
                    bad.append(f"{label} r{r}: data row dung fill cua header ({rgb})")
                    break
        return bad, n

    style_bad, style_n = [], 0
    idx_cols = max([c.column for c in wb["Screen Index"][SCREEN_INDEX_HEADER_ROW] if c.value] or [12]) \
        if "Screen Index" in wb.sheetnames else 12
    for nm, hdr, cols in (("Common mesage", CATALOG_HEADER_ROW, range(1, 10)),
                          ("Screen Error message", CATALOG_HEADER_ROW, range(1, 11)),
                          ("Screen Index", SCREEN_INDEX_HEADER_ROW, range(1, idx_cols + 1)),
                          ("Change History", CHANGE_HISTORY_HEADER_ROW, range(1, 17))):
        if nm in wb.sheetnames:
            b, n = scan_rows(wb[nm], hdr, list(cols), nm)
            style_bad += b
            style_n += n
    if style_bad:
        rep.add(14, "Style data row", "FAIL",
                f"{len(style_bad)} van de: " + "; ".join(style_bad[:6])
                + ". Dung bd_styles.py, dung copy style tu row lien truoc "
                  "(workbook-structure.md muc 2.8).")
    else:
        rep.add(14, "Style data row", "PASS",
                f"{style_n} data row deu co duong ke va khong dung fill header")

    # --- Check 15: bang ERROR SCENARIOS trong screen sheet -------------
    err_by_screen = {}
    if "Screen Error message" in wb.sheetnames:
        es = wb["Screen Error message"]
        for r in range(CATALOG_HEADER_ROW + 1, es.max_row + 1):
            code, sid = cell(es, f"A{r}"), cell(es, f"B{r}")
            if code and sid and not code.startswith(("🔴", "🟡", "🔵", "🟢")):
                err_by_screen[sid] = err_by_screen.get(sid, 0) + 1
    miss = []
    for ws in ws_list:
        sid = cell(ws, META["screen_id"])
        want = err_by_screen.get(sid, 0)
        got = 0
        title_row = None
        for r in range(1, ws.max_row + 1):
            v = ws.cell(row=r, column=8).value
            if v and ERR_TITLE_MARK in str(v):
                title_row = r
                break
        if title_row:
            for r in range(title_row + 2, ws.max_row + 1):
                if ws.cell(row=r, column=8).value is None:
                    break
                got += 1
        if want and not title_row:
            miss.append(f"{ws.title}: thieu bang ERROR SCENARIOS ({want} error trong catalog)")
        elif want and got != want:
            miss.append(f"{ws.title}: bang co {got} row, catalog co {want} error")
        elif not want and title_row:
            miss.append(f"{ws.title}: co bang ERROR SCENARIOS nhung man khong co error nao")
    if miss:
        rep.add(15, "Bang ERROR SCENARIOS", "FAIL",
                "; ".join(miss[:6]) + ". Xem workbook-structure.md muc 2.9.")
    else:
        rep.add(15, "Bang ERROR SCENARIOS", "PASS",
                f"{len(ws_list)} sheet, so row bang khop so error trong catalog")

    # ---- Check 16-19 (RIENG cho kit system-to-doc, --asis) ----
    # AS-IS: moi mo ta hanh vi phai tro ve bang chung o cot R (EV / code-ref) hoac khai "Cần xác nhận — Q-xxx".
    # Kit forward khong co vi noi dung la TO-BE (do nguoi quyet).
    if asis_mode:
        miss16, miss17, miss19, miss22, miss23 = [], [], [], [], []
        for ws in ws_list:
            # chi quet data row bang item: sau header (H="#", P co 記述/Description),
            # dung truoc title ERROR SCENARIOS
            hdr = None
            for r in range(1, min(ws.max_row, 40) + 1):
                h = str(ws.cell(row=r, column=8).value or "").strip()
                p = str(ws.cell(row=r, column=16).value or "")
                if h == "#" and ("記述" in p or "Description" in p):
                    hdr = r
                    break
            if hdr is None:
                miss16.append(f"{ws.title}: khong tim thay header bang item (H='#', P='記述')")
                continue
            if str(ws.cell(row=hdr, column=API_CALL_COL).value or "").strip() != API_CALL_HDR:
                miss17.append(f"{ws.title}!S{hdr}: thieu header '{API_CALL_HDR}'")
            for r in range(hdr + 1, ws.max_row + 1):
                if ERR_TITLE_MARK in str(ws.cell(row=r, column=8).value or ""):
                    break
                name = (str(ws.cell(row=r, column=9).value or "")
                        + str(ws.cell(row=r, column=10).value or "")).strip()
                if not name:
                    continue
                miss19 += [f"{ws.title}!{chr(64 + c)}{r}" for c in range(9, 18)
                           if not str(ws.cell(row=r, column=c).value or "").strip()]
                kind = str(ws.cell(row=r, column=11).value or "").strip().lower()
                if kind in INPUT_TYPES:
                    req, mx = cell(ws, f"M{r}"), cell(ws, f"N{r}")
                    st = cell(ws, f"L{r}").lower()
                    if not req or (req == "ー" and "disable" not in st and "read" not in st):
                        miss22.append(f"{ws.title}!M{r} ({kind} required='{req or 'rong'}')")
                    if not mx or (mx == "ー" and kind in TEXT_MAXLEN):
                        miss22.append(f"{ws.title}!N{r} ({kind} maxlength='{mx or 'rong'}')")
                for c in (9, 10, 16, 17):
                    m = CODE_TXT_RE.search(str(ws.cell(row=r, column=c).value or ""))
                    if m:
                        miss23.append(f"{ws.title}!{chr(64 + c)}{r} ('{m.group(0)}')")
                api = str(ws.cell(row=r, column=API_CALL_COL).value or "")
                if kind in CLICK_TYPES and not API_CALL_RE.search(api):
                    miss17.append(f"{ws.title}!S{r} ({kind}='{api or 'rong'}')")
                body = (str(ws.cell(row=r, column=16).value or "")
                        + str(ws.cell(row=r, column=17).value or "")).strip()
                cm = str(ws.cell(row=r, column=18).value or "")  # R
                if body and not (EV_RE.search(cm) or CONFIRM_RE.search(cm)):
                    miss16.append(f"{ws.title}!R{r}")
        if miss16:
            rep.add(16, "AS-IS: item co mo ta phai tro bang chung o cot R", "FAIL",
                    "; ".join(miss16[:8]) +
                    ". Them EV-xxxx / code-ref REPO-xx:path#Lnn, hoac ghi 'Cần xác nhận — Q-xxx' o cot R.")
        else:
            rep.add(16, "AS-IS: item co mo ta phai tro bang chung o cot R", "PASS",
                    f"{len(ws_list)} sheet, moi item co mo ta deu co EV / code-ref / Cần xác nhận — Q-xxx")
        if miss17:
            rep.add(17, "Cot S 'API gọi' cho button / link / form", "FAIL",
                    "; ".join(miss17[:8]) + ". Ghi '<METHOD> <path> (API-xxx)' tu fe-api-calls.csv, "
                    "'Không gọi API' hoac 'Cần xác nhận — Q-xxx'.")
        else:
            rep.add(17, "Cot S 'API gọi' cho button / link / form", "PASS",
                    "Moi item bam duoc deu ghi API goi / khong goi API")
        hits = legacy_word_hits(master_path)
        rep.add(18, "Khong con chu cu (D1)", "FAIL" if hits else "PASS",
                (f"{len(hits)} o con '{S.LEGACY_WORD}': " + ", ".join(hits[:8]) + " — ghi 'Cần xác nhận — Q-xxx'")
                if hits else "Khong con chu cu")
        rep.add(19, "Khong o trong o cot I..Q cua item", "FAIL" if miss19 else "PASS",
                ("; ".join(miss19[:8]) + ". Khong ap dung -> 'ー'; chua co bang chung -> 'Cần xác nhận — Q-xxx'.")
                if miss19 else "Moi item du cot I..Q")

        # 20: o dong loi khong trong (banner = dong merge ca hang -> bo qua)
        def blank_rows(ws, first, cols, label, stop=False):
            out = []
            for r in range(first, ws.max_row + 1):
                vals = [str(ws.cell(row=r, column=c).value or "").strip() for c in cols]
                if not any(vals):
                    if stop:      # bang ERROR SCENARIOS ket thuc o dong trong dau tien
                        break
                    continue
                if any(m.min_row == r and m.min_col == cols[0] and m.max_col >= cols[-1]
                       for m in ws.merged_cells.ranges):
                    continue
                out += [f"{label}!{openpyxl.utils.get_column_letter(c)}{r}" for c, v in zip(cols, vals) if not v]
            return out
        miss20 = []
        for nm, cols in (("Screen Error message", range(1, 11)), ("Common mesage", range(1, 10))):
            if nm in wb.sheetnames:
                miss20 += blank_rows(wb[nm], CATALOG_HEADER_ROW + 1, list(cols), nm)
        for ws in ws_list:
            t = next((r for r in range(1, ws.max_row + 1)
                      if ERR_TITLE_MARK in str(ws.cell(row=r, column=8).value or "")), None)
            if t:
                miss20 += blank_rows(ws, t + 2, [8, 9, 10, 11, 13, 16, 17, 19], ws.title, stop=True)
        rep.add(20, "Khong o trong o dong loi (error message)", "FAIL" if miss20 else "PASS",
                ("; ".join(miss20[:8]) + ". Khong ap dung -> 'ー'; chua biet -> 'Cần xác nhận — Q-xxx' + Open Question.")
                if miss20 else "Screen Error message · Common mesage · ERROR SCENARIOS du o")

        # 21: khong con trang thai "Status" / DRAFT
        miss21 = []
        if idx_ws is not None and any(str(c.value or "").strip().lower() == "status"
                                      for c in idx_ws[SCREEN_INDEX_HEADER_ROW]):
            miss21.append("Screen Index con cot 'Status'")
        miss21 += [f"{ws.title}!E4/F4 ('{cell(ws, 'E4')}'/'{cell(ws, 'F4')}')" for ws in ws_list
                   if cell(ws, "E4").lower() == "status" or "DRAFT" in cell(ws, "F4").upper()]
        rep.add(21, "Khong con trang thai Status / DRAFT", "FAIL" if miss21 else "PASS",
                "; ".join(miss21[:8]) if miss21 else "Screen Index khong co cot Status, sheet khong co o Status")

        rep.add(22, "Item nhap lieu co required / maxlength", "FAIL" if miss22 else "PASS",
                ("; ".join(miss22[:8]) + ". Lay tu validation FE / DTO / 08_API_Fields / 04_DB_Columns (spec 'field'); "
                 "khong co -> 'Cần xác nhận — Q-xxx'.") if miss22 else "Moi item nhap lieu du required / maxlength")
        rep.add(23, "Ten / mo ta item khong chua ma (API, ham)", "FAIL" if miss23 else "PASS",
                ("; ".join(miss23[:8]) + ". Viet theo dieu nguoi dung thay / lam; API o cot S, ten ham / code-ref o cot R.")
                if miss23 else "Cot I/J/P/Q khong co METHOD /api, chu API, ten ham")

        # 24: ten man doc duoc
        miss24 = [f"sheet '{ws.title}'" for ws in ws_list if BAD_NAME_RE.search(ws.title)]
        miss24 += [f"{ws.title}!B1 '{cell(ws, 'B1')}'" for ws in ws_list if BAD_NAME_RE.search(cell(ws, "B1"))]
        miss24 += [f"Screen Index D{r['row']} '{r['screen_name']}'" for r in idx_rows
                   if BAD_NAME_RE.search(r["screen_name"])]
        rep.add(24, "Ten man doc duoc (khong ma / route)", "FAIL" if miss24 else "PASS",
                ("; ".join(miss24[:8]) + ". Them 'title' cho man trong spec (ten ngan, khach doc duoc).")
                if miss24 else f"{len(ws_list)} ten sheet / Screen Name doc duoc")

    return wb, ws_list


def render(rep, master_path, ws_list):
    ok, fail, warn = rep.counts()
    lines = [
        "# Quality Gate O4 — Basic Design",
        "",
        f"- Workbook: `{master_path}`",
        f"- Working sheets: {len(ws_list)}"
        + (" (" + ", ".join(ws.title for ws in ws_list) + ")" if ws_list else ""),
        f"- **{len(rep.rows)} checks · {ok} PASS · {fail} FAIL · {warn} WARN**",
        "",
        "| # | Check | Ket qua | Chi tiet |",
        "|---|---|---|---|",
    ]
    icon = {"PASS": "PASS", "FAIL": "**FAIL**", "WARN": "WARN"}
    for num, name, level, detail in rep.rows:
        safe = detail.replace("|", "\\|")
        lines.append(f"| {num} | {name} | {icon[level]} | {safe} |")
    lines.append("")
    lines.append("**FAIL = 0 moi duoc bao Output 5 Done.** "
                 "Khong duoc ha nguong gate.")
    return "\n".join(lines) + "\n"


def main():
    ap = argparse.ArgumentParser(description="Quality Gate O4 — Basic Design")
    ap.add_argument("master", help="duong dan master workbook (.xlsx)")
    ap.add_argument("--before", help="ban backup truoc khi ghi (de so diff)")
    ap.add_argument("--expect-screens", default="",
                    help="danh sach Screen ID, phan cach dau phay")
    ap.add_argument("--approved-scope", default="",
                    help="ten cac sheet NGOAI scope ma user da approve sua/xoa "
                         "(phan cach dau phay). Chi dung sau khi co approval "
                         "that — xem output-5-basic-design.md muc 5 Buoc 5")
    ap.add_argument("--asis", action="store_true",
                    help="Bat check 16-24 cho kit system-to-doc: mo ta co bang chung o cot R, "
                         "cot S 'API gọi', khong chu cu, khong o trong, khong Status, ten / mo ta khong ma")
    ap.add_argument("--phase", choices=["1", "2", "auto"], default="auto",
                    help="1 = chi doc code (chua co anh -> khong WARN check 10/13); 2 = da chay Playwright; "
                         "auto = doc tu workbook (A9 'bổ sung ở Giai đoạn 2')")
    ap.add_argument("--out", help="ghi report markdown ra file")
    args = ap.parse_args()

    expect = [s.strip() for s in args.expect_screens.split(",") if s.strip()]
    approved = {s.strip() for s in args.approved_scope.split(",") if s.strip()}
    rep = Report()
    _, ws_list = run(args.master, args.before, expect, approved, rep, args.asis, args.phase)

    out = render(rep, args.master, ws_list)
    print(out)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as fh:
            fh.write(out)
        print(f"Report: {args.out}")

    ok, fail, warn = rep.counts()
    print(f"{len(rep.rows)} checks · {ok} PASS · {fail} FAIL · {warn} WARN")
    sys.exit(1 if fail else 0)


if __name__ == "__main__":
    main()
