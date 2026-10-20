#!/usr/bin/env python3
"""Boc 1 file 見積書 (mau cong ty, VD ban FINAL gui KH) thanh templates/template_mitsumori.xlsx cho 見積書 Luong 2.

  python3 extract-mitsumori-template.py --src <見積書.xlsx> [--sheet <ten sheet chi tiet>] \
      [--notes-sheet 前提条件・留意事項] [--out templates/template_mitsumori.xlsx]

Khac extract-estimation-template.py (layout cu co dinh cot K..P): script nay DO COT THEO HEADER (No. · 項目（大/中/小）·
改修内容 · 流用可否 · 必須/オプション · STEP · 要件定義 · UI・UX · 実装 · テスト · 管理 · 小計) nen dung duoc cho moi ban
見積書 doi layout. Ghi vao sheet an `_meta`: vi tri dong mau (section / item / tong / 備考), ban do cot, he so cong doan doc tu
CONG THUC (要件定義 = 実装 x a · UI = 実装 x b (nhieu muc) · テスト = 実装 x c · 管理 = SUM(...) x d), 人月 = 人日 / n.

Bo du lieu cua du an cu: moi o tu dong section tro xuong bi xoa gia tri (giu style), 件名 / 見積日 / 有効期限 / ten KH thanh
placeholder, sheet tien de chi giu tieu de ◆, cac sheet khac bi xoa. Exit 0 OK · 1 khong dung dang · 2 thieu openpyxl.
"""
import argparse
import os
import re
import sys

HDR = {"no": ["No."], "big": ["項目（大）"], "mid": ["項目（中）"], "small": ["項目（小）"], "content": ["改修内容"],
       "reuse": ["流用可否"], "must": ["必須"], "step": ["STEP"], "req": ["要件定義"], "ui": ["UI"], "impl": ["実装"],
       "test": ["テスト"], "mgmt": ["管理"], "sub": ["小計"]}


def find_layout(ws):
    """-> (row header, {key: col}) — header co the tran 2 dong (dong 'No.' + dong ngay duoi)."""
    for r in range(1, min(ws.max_row, 80) + 1):
        if any(str(ws.cell(row=r, column=c).value or "").strip() == "No." for c in range(1, 6)):
            cols = {}
            for rr in (r, r + 1):
                for c in range(1, ws.max_column + 1):
                    v = str(ws.cell(row=rr, column=c).value or "")
                    for k, keys in HDR.items():
                        if k not in cols and any(x in v for x in keys):
                            cols[k] = c
            return r, cols
    return None, {}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--sheet", default=None, help="mac dinh: sheet dau tien co dong 'No.' + cot 実装")
    ap.add_argument("--notes-sheet", default="前提条件・留意事項")
    ap.add_argument("--out", default="templates/template_mitsumori.xlsx")
    a = ap.parse_args()
    try:
        import openpyxl
    except ImportError:
        print("Thieu openpyxl", file=sys.stderr)
        return 2
    wb = openpyxl.load_workbook(a.src)
    ws, hdr, cols = None, None, {}
    for s in ([wb[a.sheet]] if a.sheet else wb.worksheets):
        hdr, cols = find_layout(s)
        if hdr and "impl" in cols:
            ws = s
            break
    miss = [k for k in HDR if k not in cols]
    if not ws or miss:
        print("Khong thay header 見積書 du cot (thieu: %s)" % ", ".join(miss or ["header"]), file=sys.stderr)
        return 1
    L = lambda k: openpyxl.utils.get_column_letter(cols[k])  # noqa: E731
    sec = item = tot = notes = None
    ratios = {"req": set(), "ui": set(), "test": set(), "mgmt": set()}
    month = None
    for r in range(hdr + 2, ws.max_row + 1):
        lead = " ".join(str(ws.cell(row=r, column=c).value or "") for c in range(cols["no"], cols["mid"] + 1)).strip()
        impl = ws.cell(row=r, column=cols["impl"]).value
        if sec is None and lead.startswith("◆"):
            sec = r
        elif item is None and isinstance(impl, (int, float)) and impl > 0:
            item = r
        if lead.startswith("合計") and tot is None:
            tot = r
        if lead.startswith("■") and notes is None and tot:
            notes = r
        for k in ("req", "ui", "test"):
            m = re.match(r"^=%s%d\*([\d.]+)$" % (L("impl"), r), str(ws.cell(row=r, column=cols[k]).value or "").replace(" ", ""))
            if m:
                ratios[k].add(float(m.group(1)))
        m = re.match(r"^=SUM\(%s%d:%s%d\)\*([\d.]+)$" % (L("req"), r, L("test"), r),
                     str(ws.cell(row=r, column=cols["mgmt"]).value or "").replace(" ", ""))
        if m:
            ratios["mgmt"].add(float(m.group(1)))
        m = re.match(r"^=%s\d+/(\d+)$" % L("sub"), str(ws.cell(row=r, column=cols["sub"]).value or "").replace(" ", ""))
        if m:
            month = int(m.group(1))
    if not (sec and item and tot):
        print("Khong thay dong mau section / item / 合計", file=sys.stderr)
        return 1
    one = lambda k: max(ratios[k], key=lambda x: x) if ratios[k] else None  # noqa: E731
    meta = {"source_file": os.path.basename(a.src), "source_sheet": ws.title, "row_header": hdr, "row_section": sec,
            "row_item": item, "row_total": tot, "row_notes": notes or tot + 8,
            "ratio_req": one("req"), "ratio_ui_large": max(ratios["ui"]) if ratios["ui"] else None,
            "ratio_ui": min(ratios["ui"]) if ratios["ui"] else None, "ratio_test": one("test"),
            "ratio_mgmt": one("mgmt"), "md_per_month": month or 20, "notes_sheet": a.notes_sheet}
    meta.update({"col_" + k: L(k) for k in cols})
    if None in (meta["ratio_req"], meta["ratio_ui"], meta["ratio_test"], meta["ratio_mgmt"]):
        print("Khong doc duoc he so tu cong thuc: %s" % ratios, file=sys.stderr)
        return 1
    # xoa du lieu du an cu: giu style, bo gia tri tu dong section tro xuong; dau trang -> placeholder
    for row in ws.iter_rows(min_row=sec):
        for c in row:
            if c.value is not None:
                c.value = None
                c.hyperlink = None
    for row in ws.iter_rows(min_row=1, max_row=hdr - 1):
        for c in row:
            v = str(c.value or "")
            if c.column > 1 and ws.cell(row=c.row, column=2).value == "■件名" and c.column == 3:
                c.value = "{{TITLE}}"
            elif ws.cell(row=c.row, column=2).value == "■見積日" and c.column == 3:
                c.value = "{{DATE}}"
            elif ws.cell(row=c.row, column=2).value == "■有効期限" and c.column == 3:
                c.value = "{{VALID_UNTIL}}"
            elif v.endswith("様") or (c.column == 2 and ws.cell(row=c.row, column=5).value == "御中"):
                c.value = "{{CUSTOMER}}"
    for s in list(wb.worksheets):
        if s is ws:
            continue
        if s.title == a.notes_sheet:
            for row in s.iter_rows():
                for c in row:
                    if c.value is not None and not str(c.value).lstrip().startswith("◆"):
                        c.value = None
                        c.hyperlink = None
            continue
        wb.remove(s)
    ws.title = "見積書"
    mt = wb.create_sheet("_meta")
    for k, v in meta.items():
        mt.append([k, v])
    mt.sheet_state = "hidden"
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    wb.save(a.out)
    print("Da ghi %s — header dong %d · section %d · item %d · 合計 %d · cot 実装 %s · he so 要件 %s / UI %s,%s / テスト %s / 管理 %s · 人月 /%s"
          % (a.out, hdr, sec, item, tot, meta["col_impl"], meta["ratio_req"], meta["ratio_ui_large"], meta["ratio_ui"],
             meta["ratio_test"], meta["ratio_mgmt"], meta["md_per_month"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
