#!/usr/bin/env python3
"""GATE V8 — Bug List (O7).

  python3 verify-bug-list.py <07_BugList/BugList_<sys>_ver<N>.xlsx> --inventory <_internal/inventory.xlsx> [--out report.md]

File nay di ra NGOAI cong ty (bao khach hang) nen nguong bang chung cao.
1 sheet duy nhat "Bug List" (feedback 28/10/2026): bug muc High / Urgent anh huong thao tac nguoi dung, tim thay
khi doc source code (Code review) va / hoac khi crawl website (Playwright). Khong co bug -> sheet chi co header.
Pham vi quet + nguoi nhan khai o 00_Meta cua inventory (bug_scan_scope, bug_recipient).
Chan: sai so sheet · muc ngoai High / Urgent · thieu buoc tai hien · thieu bang chung (screenshot / console-log /
har / code-ref co that) · console error la bao dong gia cua cong cu (noise read-only-block / csp-block) · URL la
API · trung lap · Screen / Module khong tro ve SC- co that · con chu cu.
"""
import argparse
import importlib.util
import os
import re
import sys

import inv_schema as S
from gate_report import Gate, legacy_word_hits

_spec = importlib.util.spec_from_file_location(
    "ve", os.path.join(os.path.dirname(os.path.abspath(__file__)), "verify-evidence.py"))
VE = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(VE)

REQUIRED = ["Title", "Screen / Module", "Category", "Severity", "Repro Steps",
            "Expected", "Actual", "Evidence", "Reproduced", "Detected By", "Env"]


def load_bugs(path):
    """-> (ten sheet, [dong])"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("Thieu openpyxl. Chay: pip install openpyxl")
    wb = load_workbook(path, data_only=True)
    if S.BUG_SHEET not in wb.sheetnames:
        return wb.sheetnames, None
    ws = wb[S.BUG_SHEET]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    rows = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        d = {h: ("" if r[j] is None else str(r[j]).strip())
             for j, h in enumerate(header) if h and j < len(r)}
        d["__row__"] = i
        rows.append(d)
    return wb.sheetnames, rows


def norm(s):
    return re.sub(r"[^a-z0-9]+", " ", (s or "").lower()).strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("buglist")
    ap.add_argument("--inventory", default=None)
    ap.add_argument("--out", default=None)
    a = ap.parse_args()

    names, bugs = load_bugs(a.buglist)
    g = Gate("GATE V8 — Bug List")
    g.check(1, "Dung 1 sheet '%s'" % S.BUG_SHEET, [] if names == [S.BUG_SHEET] else ["sheet hien co: %s" % names])
    if bugs is None:
        return g.emit(a.out)

    ids = [r.get("Bug ID", "") for r in bugs]
    rx = re.compile(r"^BUG-\d{3,}$")
    g.check(2, "Bug ID dung dinh dang BUG-NNN va khong trung",
            [i for i in ids if not rx.match(i)] + sorted({i for i in ids if ids.count(i) > 1}))

    g.check(3, "Moi bug dien du cot bat buoc",
            ["%s.%s" % (r.get("Bug ID", r["__row__"]), c) for r in bugs for c in REQUIRED if not r.get(c, "").strip()])

    bad = []
    for r in bugs:
        tag = r.get("Bug ID", r["__row__"])
        if r.get("Category") not in S.BUG_CATEGORY:
            bad.append("%s Category=%s" % (tag, r.get("Category")))
        if r.get("Severity") not in S.BUG_SEVERITY:
            bad.append("%s Severity=%s (chi %s; muc thap hon khong ghi)" % (tag, r.get("Severity"), "/".join(S.BUG_SEVERITY)))
        if r.get("Detected By") not in S.BUG_DETECTED_BY:
            bad.append("%s Detected By=%s" % (tag, r.get("Detected By")))
        if not re.match(r"^(Yes — \d+/\d+|Intermittent|%s)$" % S.BUG_CODE_ONLY, r.get("Reproduced", "")):
            bad.append("%s Reproduced=%s" % (tag, r.get("Reproduced")))
        if r.get("Pre-existing") not in S.BUG_PRE_EXISTING:
            bad.append("%s Pre-existing=%s" % (tag, r.get("Pre-existing")))
    g.check(4, "Moi cot enum dung gia tri cho phep", bad)

    g.check(5, "Repro Steps co >= 2 buoc danh so",
            [r.get("Bug ID") for r in bugs if len(re.findall(r"(?m)^\s*\d+[.)]", r.get("Repro Steps", ""))) < 2])

    inv = S.load(a.inventory) if a.inventory else None
    if inv:
        evids = {x.get("EV ID", "") for x in inv["05_Evidence"]}
        g.check(6, "Evidence tro toi EV co that trong inventory",
                ["%s -> %s" % (r.get("Bug ID"), r.get("Evidence")) for r in bugs
                 if not [e for e in S.split_ids(r.get("Evidence", "")) if e in evids]])
    else:
        g.warn(6, "Evidence phan giai", "chua truyen --inventory")

    # 7 — evidence >= 1 screenshot / console-log / har / code-ref co that, khong phai noise cong cu;
    #     Detected By khop loai evidence; URL la man, khong phai API
    evs = {x.get("EV ID", ""): x for x in inv["05_Evidence"]} if inv else {}
    repo_paths = {x.get("Repo ID", ""): x.get("Path", "") for x in inv["11_Repo"]} if inv else {}
    root = os.path.dirname(os.path.abspath(a.inventory)) if inv else ""
    bad7 = []
    for r in bugs:
        tag = r.get("Bug ID")
        url = r.get("URL / Route", "")
        if re.search(r"(^|/)api(/|$)", re.sub(r"^https?://[^/]+", "", url), re.I):
            bad7.append("%s URL '%s' la API — bug phai gan voi man + thao tac" % (tag, url))
        if not inv:
            continue
        kinds = set()
        for e in S.split_ids(r.get("Evidence", "")):
            x = evs.get(e, {})
            t = x.get("Type", "")
            if t == "console-log" and any(n in x.get("Note", "") + x.get("Locator", "") for n in S.BUG_NOISE):
                bad7.append("%s %s la bao dong gia cua cong cu quet (%s) — khong phai bug" % (tag, e, "/".join(S.BUG_NOISE)))
                continue
            if t == "code-ref":
                ok, err = VE.resolve_code_ref(x.get("Locator", ""), repo_paths, root, None)
                if not ok:
                    bad7.append("%s %s code-ref khong co that: %s" % (tag, e, err))
                    continue
            if t in S.BUG_EVIDENCE:
                kinds.add("code" if t == "code-ref" else "web")
        by = r.get("Detected By", "")
        if not kinds:
            bad7.append("%s khong co evidence %s" % (tag, "/".join(S.BUG_EVIDENCE)))
        elif "Code review" in by and "code" not in kinds:
            bad7.append("%s Detected By=%s nhung khong co EV code-ref" % (tag, by))
        elif "Playwright" in by and "web" not in kinds:
            bad7.append("%s Detected By=%s nhung khong co EV screenshot/console-log/har" % (tag, by))
        if r.get("Reproduced") == S.BUG_CODE_ONLY and "web" in kinds:
            bad7.append("%s Reproduced='%s' nhung co EV website — ghi 'Yes — n/n'" % (tag, S.BUG_CODE_ONLY))
    g.check(7, "Evidence >= 1 screenshot/console-log/har/code-ref co that, khong phai noise cong cu", bad7)

    g.check(8, "Moi bug ghi Business Impact (tac dong toi thao tac nguoi dung)",
            [r.get("Bug ID") for r in bugs if not r.get("Business Impact", "").strip()])

    seen, dups = {}, []
    for r in bugs:
        key = (norm(r.get("Title")), norm(r.get("URL / Route")))
        if key in seen:
            dups.append("%s ~ %s" % (seen[key], r.get("Bug ID")))
        else:
            seen[key] = r.get("Bug ID")
    g.check(9, "Khong co bug trung lap (cung tieu de + cung route)", dups)

    if inv:
        meta = S.meta(inv)
        g.check(10, "00_Meta inventory khai bug_scan_scope + bug_recipient",
                [k for k in ("bug_scan_scope", "bug_recipient") if meta.get(k, "") in ("", S.NEED_CONFIRM, "—")])
        known = {x.get("Screen ID", "") for x in inv["02_Screen"]}
        bad_ref = []
        for r in bugs:
            refs = re.findall(r"\bSC-\d{3,}\b", r.get("Screen / Module", ""))
            if not refs:
                bad_ref.append("%s '%s' khong co SC- (bug phai gan voi man)" % (r.get("Bug ID"), r.get("Screen / Module", "")))
            bad_ref += ["%s -> %s khong co trong inventory" % (r.get("Bug ID"), x) for x in refs if x not in known]
        g.check(11, "Screen / Module tro ve SC- co that", bad_ref)
    else:
        g.warn(10, "00_Meta pham vi quet", "chua truyen --inventory")
        g.warn(11, "Screen / Module phan giai", "chua truyen --inventory")

    g.check(12, "Khong con chu cu '%s' (khong phan biet hoa thuong)" % S.LEGACY_WORD, legacy_word_hits(a.buglist))

    rc = g.emit(a.out)
    sev = {}
    for r in bugs:
        sev[r.get("Severity", "?")] = sev.get(r.get("Severity", "?"), 0) + 1
    print("BUGS: %d · %s" % (len(bugs), " · ".join("%s=%d" % kv for kv in sorted(sev.items())) or "-"), file=sys.stderr)
    return rc


if __name__ == "__main__":
    sys.exit(main())
