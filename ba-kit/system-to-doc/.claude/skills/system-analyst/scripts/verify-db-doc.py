#!/usr/bin/env python3
"""GATE — O3 Database Documentation (doi chieu voi inventory).

  python3 verify-db-doc.py <DB_Doc.xlsx> --inventory <_internal>/inventory.xlsx \\
      [--erd-png <ERD.png>] [--out <_internal>/gates/v-db.md]

Chan: bang thieu/trung sheet, so cot lech 04_DB_Columns, quan he tro toi bang/cot khong ton tai
(bia), link noi bo gay, o trong trong luoi du lieu, Confidence High khong co code-ref, thieu ERD.
"""
import argparse
import os
import sys

import inv_schema as S
from gate_report import Gate

try:
    from openpyxl import load_workbook
except ImportError:
    print("Thieu openpyxl. Chay: pip install openpyxl", file=sys.stderr)
    sys.exit(2)

import importlib.util

_spec = importlib.util.spec_from_file_location(
    "vapi", os.path.join(os.path.dirname(os.path.abspath(__file__)), "verify-api-doc.py"))
_vapi = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_vapi)
link_target = _vapi.link_target


def grid(ws, first, second, max_row=80):
    """Tim header (cot A=first, B=second) -> (hdr_row, header list, [(row, values)])."""
    hdr = None
    for r in range(1, min(ws.max_row, max_row) + 1):
        if ws.cell(r, 1).value == first and ws.cell(r, 2).value == second:
            hdr = r
            break
    if not hdr:
        return None, [], []
    header = []
    while ws.cell(hdr, len(header) + 1).value not in (None, ""):
        header.append(str(ws.cell(hdr, len(header) + 1).value))
    rows = []
    for r in range(hdr + 1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, len(header) + 1)]
        if all(x in (None, "") for x in vals):
            break
        rows.append((r, vals))
    return hdr, header, rows


def blanks(name, rows):
    return ["%s!row%d col%d" % (name, r, c + 1) for r, vals in rows for c, v in enumerate(vals)
            if v is None or str(v).strip() == ""]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("doc")
    ap.add_argument("--inventory", required=True)
    ap.add_argument("--erd-png")
    ap.add_argument("--out")
    a = ap.parse_args()

    data = S.load(a.inventory)
    tables = [r["Table"] for r in data["03_DB_Tables"] if r.get("Table")]
    colmap = {}
    for c in data["04_DB_Columns"]:
        if c.get("Table"):
            colmap.setdefault(c["Table"], set()).add(c.get("Column", ""))
    evtype = {e.get("EV ID"): e.get("Type", "") for e in data["05_Evidence"]}
    wb = load_workbook(a.doc)
    g = Gate("GATE O3 — Database Documentation")

    req = ["00_Overview", "01_Relationships", "02_ERD"]
    g.check(1, "Co sheet 00_Overview, 01_Relationships, 02_ERD", [s for s in req if s not in wb.sheetnames])
    if any(s not in wb.sheetnames for s in req):
        return g.emit(a.out)

    # sheet bang: B2 = ten bang (A2 = TABLE)
    tsheets = {}
    dup = []
    for ws in wb.worksheets:
        if ws.title in req:
            continue
        if ws.cell(2, 1).value == "TABLE":
            t = str(ws.cell(2, 2).value)
            if t in tsheets:
                dup.append(t)
            tsheets.setdefault(t, []).append(ws)
    bad = ["%s thieu sheet" % t for t in tables if t not in tsheets]
    bad += ["%s co %d sheet" % (t, len(tsheets[t])) for t in sorted(set(dup))]
    bad += ["sheet %s cho bang khong co trong 03" % tsheets[t][0].title for t in tsheets if t not in tables]
    g.check(2, "Moi bang trong 03_DB_Tables co dung 1 sheet chi tiet", bad)

    # 3 overview index
    hdr, header, orows = grid(wb["00_Overview"], "No", "Table")
    names = [str(v[1]) for _, v in orows]
    bad = [] if hdr else ["khong thay bang (No · Table)"]
    bad += ["%s thieu" % t for t in tables if names.count(t) == 0]
    bad += ["%s trung" % t for t in sorted(set(names)) if names.count(t) > 1]
    bad += ["%s khong co trong 03" % t for t in sorted(set(names) - set(tables))]
    g.check(3, "00_Overview co dung 1 dong / bang", bad)

    # 4 so cot
    bad, grids = [], {}
    for t in tables:
        if t not in tsheets:
            continue
        ws = tsheets[t][0]
        h, hd, rows = grid(ws, "No", "Column")
        grids[t] = (ws, hd, rows)
        doc_cols = [str(v[1]) for _, v in rows if str(v[1]) != S.UNKNOWN]
        inv_cols = colmap.get(t, set())
        if len(doc_cols) != len(inv_cols) or set(doc_cols) != inv_cols:
            miss = sorted(inv_cols - set(doc_cols))
            extra = sorted(set(doc_cols) - inv_cols)
            bad.append("%s: doc=%d, 04=%d%s%s" % (t, len(doc_cols), len(inv_cols),
                                                 (" thieu " + ",".join(miss[:5])) if miss else "",
                                                 (" thua " + ",".join(extra[:5])) if extra else ""))
    g.check(4, "So cot moi sheet bang = so dong 04_DB_Columns", bad)

    # 5 quan he: dau mut ton tai
    h, rh, rrows = grid(wb["01_Relationships"], "No", "From Table")
    bad = [] if h else ["khong thay bang (No · From Table)"]
    for r, vals in rrows:
        rec = dict(zip(rh, vals))
        if rec.get("From Table") == "—":
            continue
        for tk, ck in (("From Table", "From Column"), ("To Table", "To Column")):
            tt, cc = str(rec.get(tk)), str(rec.get(ck))
            if tt not in tables:
                bad.append("row%d %s=%s khong ton tai" % (r, tk, tt))
            elif cc not in colmap.get(tt, set()):
                bad.append("row%d %s.%s khong ton tai" % (r, tt, cc))
    g.check(5, "Moi quan he tro toi bang + cot co that (khong bia)", bad)

    # 6 hyperlink
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                tg = link_target(c)
                if not tg:
                    continue
                where = "%s!%s" % (ws.title, c.coordinate)
                if tg[0] in ("BAD", "EXTERNAL"):
                    bad.append("%s -> link khong hop le %s" % (where, tg[1]))
                elif tg[0] not in wb.sheetnames:
                    bad.append("%s -> sheet '%s' khong ton tai" % (where, tg[0]))
                else:
                    tv = wb[tg[0]][tg[1]].value
                    if tv in (None, ""):
                        bad.append("%s -> %s!%s rong" % (where, tg[0], tg[1]))
                    elif tg[1] == "B2" and str(c.value) in tables and str(tv) != str(c.value):
                        bad.append("%s (%s) -> sheet cua bang %s" % (where, c.value, tv))
    g.check(6, "Moi hyperlink noi bo tro toi sheet+o ton tai (dung bang)", bad)

    # 7 o trong
    bad = blanks("00_Overview", orows) + blanks("01_Relationships", rrows)
    for t, (ws, hd, rows) in grids.items():
        bad += blanks(ws.title, rows)
    g.check(7, "Khong co o trong trong luoi du lieu (UNKNOWN / — thay cho trong)", bad)

    # 8 Confidence High -> co code-ref
    bad = []
    for t, (ws, hd, rows) in grids.items():
        if "Confidence" not in hd or "Evidence" not in hd:
            continue
        ci, ei = hd.index("Confidence"), hd.index("Evidence")
        for r, vals in rows:
            if str(vals[ci]) == "High" and not any(evtype.get(e) == "code-ref" for e in S.split_ids(vals[ei])):
                bad.append("%s!row%d %s" % (ws.title, r, vals[1]))
    for r in data["03_DB_Tables"]:
        if r.get("Confidence") == "High" and not any(evtype.get(e) == "code-ref"
                                                      for e in S.split_ids(r.get("Evidence"))):
            bad.append("03 %s" % r.get("Table"))
    g.check(8, "Dong Confidence=High co Evidence loai code-ref", bad)

    # 9 ERD
    imgs = getattr(wb["02_ERD"], "_images", [])
    if a.erd_png:
        bad = []
        if not os.path.isfile(a.erd_png):
            bad.append("khong thay file %s" % a.erd_png)
        if not imgs:
            bad.append("02_ERD khong nhung anh")
        g.check(9, "ERD PNG ton tai va duoc nhung vao 02_ERD", bad)
    else:
        g.check(9, "ERD duoc nhung vao 02_ERD", [] if imgs else ["chua co ERD (khong truyen --erd-png)"],
                level="WARN")
    return g.emit(a.out)


if __name__ == "__main__":
    sys.exit(main())
