#!/usr/bin/env python3
"""Ve code map Website -> FE repo -> nhom API -> BE repo -> bang DB (PNG + Markdown/Mermaid).

  python3 render-codemap.py --inventory <_internal>/inventory.xlsx --code <_internal>/recon/code
          --png <ver>/02_API/CodeMap_<sys>_ver<N>.png --md <ver>/02_API/CodeMap_<sys>_ver<N>.md
          [--json <_internal>/codemap.json] [--max-tables 25] [--max-groups 60] [--dpi 150]

Net LIEN = ghi nhan trong inventory (10_Site.FE Repo, 02_Screen.Site, 07_API.Repo /
Called By Screens / Related Tables). Net DUT = chi suy ra tu recon (fe-api-calls khop path API,
table-refs cua file handler) -> phai kiem chung truoc khi ghi Confirmed.
Node la NHOM (khong phai tung endpoint); nhan canh = so luong. Chi can matplotlib.
"""
import argparse
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import inv_schema as S  # noqa: E402

EMPTY = {"", "UNKNOWN", "—", "-", "N/A", "NONE", "TBD"}
LANES = ["Websites", "FE repos", "API groups", "BE repos", "DB tables"]
LANE_FILL = ["#EEF3FB", "#EEF7F0", "#FAFAF7", "#F6F1FA", "#F5F5F5"]
BATCH_FILL = "#FFF6E5"
INK = "#222222"
DOC_COLOR = "#333333"
INF_COLOR = "#2F6DB5"


def ids(v):
    return [x for x in S.split_ids(v) if x.upper() not in EMPTY]


def val(v):
    v = (v or "").strip()
    return "" if v.upper() in EMPTY else v


def load_inventory(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("Thieu openpyxl. Chay: pip install openpyxl", file=sys.stderr)
        sys.exit(2)
    names = load_workbook(path, read_only=True).sheetnames
    want = {k: S.SHEETS[k] for k in ("00_Meta", "02_Screen", "03_DB_Tables", "07_API",
                                      "10_Site", "11_Repo") if k in names}
    data = S.load(path, want)
    for k in ("00_Meta", "02_Screen", "03_DB_Tables", "07_API", "10_Site", "11_Repo"):
        data.setdefault(k, [])
    return data


def load_json(p, default):
    try:
        with open(p, encoding="utf8") as fh:
            return json.load(fh)
    except (OSError, ValueError):
        return default


def load_recon(code_dir):
    out = {}
    if not code_dir or not os.path.isdir(code_dir):
        return out
    for rid in sorted(os.listdir(code_dir)):
        d = os.path.join(code_dir, rid)
        if not (os.path.isdir(d) and re.match(r"^REPO-\d+$", rid)):
            continue
        out[rid] = {
            "calls": load_json(os.path.join(d, "fe-api-calls.json"), []),
            "routes": load_json(os.path.join(d, "fe-routes.json"), []),
            "refs": load_json(os.path.join(d, "table-refs.json"), {}).get("files", {}),
            "stack": load_json(os.path.join(d, "stack.json"), {}),
        }
    return out


def npath(p):
    p = (p or "").strip()
    m = re.match(r"^https?://[^/]+(/.*)?$", p)
    if m:
        p = m.group(1) or "/"
    p = p.split("?")[0].split("#")[0]
    p = re.sub(r"\{[^}]+\}|:[A-Za-z_]\w*\*?|\[[^\]]+\]|<[^>]+>", ":p", p)
    p = re.sub(r"/{2,}", "/", "/" + p.lstrip("/")).rstrip("/").lower()
    return p or "/"


def path_match(fe, api):
    a, b = npath(fe), npath(api)
    if a == b:
        return True
    lo, sh = (a, b) if len(a) > len(b) else (b, a)
    return len(sh) > 1 and sh != "/:p" and lo.endswith(sh)


def route_rx(route):
    r = npath(route)
    return re.compile("^" + re.escape(r).replace(":p", "[^/]+") + "$")


def parse_loc(loc):
    m = re.match(r"^(REPO-\d+):(.+?)(?:#L\d+(?:-L?\d+)?)?$", (loc or "").strip())
    return (m.group(1), m.group(2)) if m else (None, None)


def file_tables(refs, rel, depth=2, seen=None):
    seen = seen if seen is not None else set()
    if rel in seen or rel not in refs:
        return set()
    seen.add(rel)
    out = {t["table"] for t in refs[rel].get("tables", []) if not t.get("how", "").startswith("model def")}
    if depth > 0:
        for imp in refs[rel].get("imports", []):
            out |= file_tables(refs, imp, depth - 1, seen)
    return out


def file_closure(refs, rel, depth=2):
    out, frontier = {rel}, [rel]
    for _ in range(depth):
        nxt = []
        for f in frontier:
            for imp in refs.get(f, {}).get("imports", []):
                if imp not in out:
                    out.add(imp)
                    nxt.append(imp)
        frontier = nxt
    return out


# ---------------------------------------------------------------- graph
class Graph:
    def __init__(self):
        self.nodes = {}      # id -> {lane, label, kind}
        self.edges = {}      # (a,b) -> {"doc": n, "inf": n}

    def node(self, nid, lane, label, kind="normal"):
        if nid not in self.nodes:
            self.nodes[nid] = {"id": nid, "lane": lane, "label": label, "kind": kind}
        return nid

    def edge(self, a, b, documented, n=1):
        if a not in self.nodes or b not in self.nodes or a == b:
            return
        e = self.edges.setdefault((a, b), {"doc": 0, "inf": 0})
        e["doc" if documented else "inf"] += n


def build(data, recon, max_tables, max_groups):
    g = Graph()
    meta = S.meta(data)
    sites = {r.get("Site ID"): r for r in data["10_Site"] if r.get("Site ID")}
    repos = {r.get("Repo ID"): r for r in data["11_Repo"] if r.get("Repo ID")}
    screens = {r.get("Screen ID"): r for r in data["02_Screen"] if r.get("Screen ID")}
    apis = [r for r in data["07_API"] if r.get("API ID")]
    inv_tables = [val(r.get("Table")) for r in data["03_DB_Tables"] if val(r.get("Table"))]

    for rid in recon:
        repos.setdefault(rid, {"Repo ID": rid, "Kind": (recon[rid]["stack"].get("kind") or ""),
                               "Path": "", "Stack": "", "For Site": "", "__recon_only__": True})
    for sid, s in sites.items():
        g.node("site:" + sid, 0, "%s\n%s" % (sid, val(s.get("Site Name")) or val(s.get("URL")) or ""))
    fe_ids, be_ids = [], []
    for rid, r in repos.items():
        kind = (val(r.get("Kind")) or "").upper()
        stack = val(r.get("Stack"))
        lab = "%s (%s)%s" % (rid, kind or "?", ("\n" + stack[:26]) if stack else "")
        if kind in ("FE", "FULLSTACK"):
            g.node("fe:" + rid, 1, lab)
            fe_ids.append(rid)
        if kind in ("BE", "FULLSTACK", "BATCH", "OTHER", ""):
            g.node("be:" + rid, 3, lab)
            be_ids.append(rid)

    # site -> FE repo (documented: 10_Site.FE Repo, 11_Repo.For Site, 00_Meta.source_repos)
    site_fe = {sid: set() for sid in sites}
    for sid, s in sites.items():
        for rid in ids(s.get("FE Repo")):
            site_fe[sid].add(rid)
    for rid, r in repos.items():
        for sid in ids(r.get("For Site")):
            site_fe.setdefault(sid, set()).add(rid)
    for m in re.finditer(r"(REPO-\d+)=[^;]*?\(FE\s*(?:->|→)\s*(WEB-\d+)", meta.get("source_repos", "")):
        site_fe.setdefault(m.group(2), set()).add(m.group(1))
    for sid, rset in site_fe.items():
        for rid in rset:
            if "fe:" + rid not in g.nodes and rid in repos:
                g.node("fe:" + rid, 1, "%s (FE)" % rid)
                fe_ids.append(rid)
            g.edge("site:" + sid, "fe:" + rid, True)

    # API groups
    grp_of = {}
    for a in apis:
        kind = (val(a.get("Kind")) or "API").upper()
        repo = val(a.get("Repo")) or "?"
        grp = val(a.get("Group")) or ("Batch" if kind == "BATCH" else "Queue" if kind == "QUEUE" else "(no group)")
        is_batch = kind in ("BATCH", "QUEUE")
        nid = "grp:%s:%s:%s" % ("B" if is_batch else "A", repo, grp)
        grp_of[a["API ID"]] = nid
        if nid not in g.nodes:
            g.node(nid, 2, grp, "batch" if is_batch else "normal")
            g.nodes[nid]["apis"] = []
            g.nodes[nid]["repo"] = repo
        g.nodes[nid]["apis"].append(a["API ID"])
    for nid, n in g.nodes.items():
        if nid.startswith("grp:"):
            n["label"] = "%s\n%d %s · %s" % (n["label"][:28], len(n["apis"]),
                                             "job" if n["kind"] == "batch" else "API", n["repo"])

    # tables
    def tnode(t):
        nid = "tbl:" + t
        g.node(nid, 4, t if t in inv_tables else t + " *")
        return nid
    for t in inv_tables:
        tnode(t)

    # documented API edges
    detail_doc_api = {}
    for a in apis:
        gid = grp_of[a["API ID"]]
        repo = val(a.get("Repo"))
        if repo and "be:" + repo in g.nodes:
            g.edge(gid, "be:" + repo, True)
        for sc in ids(a.get("Called By Screens")):
            detail_doc_api.setdefault(sc, set()).add(a["API ID"])
            sid = val(screens.get(sc, {}).get("Site"))
            fes = site_fe.get(sid) or set()
            if fes:
                for rid in fes:
                    g.edge("fe:" + rid, gid, True)
            elif sid and "site:" + sid in g.nodes:
                g.edge("site:" + sid, gid, True)
        for t in ids(a.get("Related Tables")):
            tn = tnode(t)
            if repo and "be:" + repo in g.nodes:
                g.edge("be:" + repo, tn, True)
            else:
                g.edge(gid, tn, True)

    # inferred: FE calls -> API groups
    api_by_id = {a["API ID"]: a for a in apis}
    call_hits = {}   # (fe repo, file) -> set(api ids)
    for rid, rc in recon.items():
        for c in rc["calls"]:
            for a in apis:
                if (val(a.get("Kind")) or "API").upper() not in ("API", "WEBHOOK"):
                    continue
                am = (val(a.get("Method")) or "ANY").upper()
                cm = (c.get("method") or "ANY").upper()
                if am not in ("ANY", cm) and cm != "ANY":
                    continue
                if path_match(c.get("path") or c.get("url"), val(a.get("Path / Schedule"))):
                    call_hits.setdefault((rid, c.get("file")), set()).add(a["API ID"])
    fe_grp = {}
    for (rid, _f), aset in call_hits.items():
        for aid in aset:
            fe_grp.setdefault((rid, grp_of[aid]), set()).add(aid)
    for (rid, gid), aset in fe_grp.items():
        if "fe:" + rid not in g.nodes:
            g.node("fe:" + rid, 1, "%s (FE?)" % rid)
        g.edge("fe:" + rid, gid, False, len(aset))

    # inferred: handler file -> tables
    api_inf_tables = {}
    for a in apis:
        rid, rel = parse_loc(a.get("Handler"))
        if not rid or rid not in recon:
            continue
        ts = file_tables(recon[rid]["refs"], rel)
        if ts:
            api_inf_tables[a["API ID"]] = ts
            repo = val(a.get("Repo")) or rid
            for t in ts:
                tn = tnode(t)
                if "be:" + repo in g.nodes:
                    g.edge("be:" + repo, tn, False)
                else:
                    g.edge(grp_of[a["API ID"]], tn, False)

    # cap tables / groups
    hidden = cap_lane(g, 4, max_tables, "tbl:+more", "tables")
    cap_lane(g, 2, max_groups, "grp:+more", "groups")

    # detail: screen -> API -> tables
    detail = []
    for sc, s in sorted(screens.items()):
        sid = val(s.get("Site"))
        doc = sorted(detail_doc_api.get(sc, set()))
        inf = set()
        route = val(s.get("URL / Route"))
        if route:
            for rid in site_fe.get(sid, set()):
                rc = recon.get(rid)
                if not rc:
                    continue
                files = set()
                for fr in rc["routes"]:
                    if route_rx(fr.get("route", "")).match(npath(route)):
                        files |= component_files(rc, fr)
                for f in list(files):
                    files |= file_closure(rc["refs"], f)
                for f in files:
                    inf |= call_hits.get((rid, f), set())
        inf = sorted(inf - set(doc))
        tdoc = sorted({t for aid in doc for t in ids(api_by_id[aid].get("Related Tables"))})
        tinf = sorted({t for aid in doc + inf for t in api_inf_tables.get(aid, set())} - set(tdoc))
        detail.append({"screen": sc, "site": sid, "route": route, "api_doc": doc, "api_inf": inf,
                       "tables_doc": tdoc, "tables_inf": tinf})
    batch = []
    for a in apis:
        if (val(a.get("Kind")) or "").upper() in ("BATCH", "QUEUE"):
            tdoc = ids(a.get("Related Tables"))
            batch.append({"api": a["API ID"], "kind": val(a.get("Kind")),
                          "schedule": val(a.get("Path / Schedule")) or "UNKNOWN",
                          "tables_doc": tdoc,
                          "tables_inf": sorted(api_inf_tables.get(a["API ID"], set()) - set(tdoc))})
    return g, detail, batch, hidden


def component_files(rc, fr):
    comp = fr.get("component") or ""
    known = set(rc["refs"]) | {c.get("file") for c in rc["calls"]}
    if "/" in comp or re.search(r"\.(tsx|jsx|ts|js|vue|svelte)$", comp):
        base = os.path.normpath(os.path.join(os.path.dirname(fr.get("file") or ""), comp)).replace(os.sep, "/")
        hits = {f for f in known if f == comp or f.startswith(base) or f.startswith(comp)}
        return hits or {comp}
    stem = comp.split(".")[-1]
    return {f for f in known if f and os.path.splitext(os.path.basename(f))[0] == stem} or set()


def cap_lane(g, lane, cap, more_id, word):
    lane_nodes = [n for n in g.nodes.values() if n["lane"] == lane]
    if len(lane_nodes) <= cap:
        return []
    deg = {n["id"]: 0 for n in lane_nodes}
    for (a, b), e in g.edges.items():
        for x in (a, b):
            if x in deg:
                deg[x] += e["doc"] + e["inf"]
    keep = set(sorted(deg, key=lambda k: (-deg[k], k))[:cap - 1])
    drop = [k for k in deg if k not in keep]
    g.node(more_id, lane, "+%d more %s\n(xem bang chi tiet)" % (len(drop), word), "more")
    for (a, b), e in list(g.edges.items()):
        if a in drop or b in drop:
            na, nb = (more_id if a in drop else a), (more_id if b in drop else b)
            del g.edges[(a, b)]
            if na != nb:
                t = g.edges.setdefault((na, nb), {"doc": 0, "inf": 0})
                t["doc"] += e["doc"]
                t["inf"] += e["inf"]
    for k in drop:
        del g.nodes[k]
    return drop


# ---------------------------------------------------------------- layout + draw
def order_lanes(g):
    lanes = {i: sorted([n for n in g.nodes.values() if n["lane"] == i],
                       key=lambda n: (n["kind"] in ("batch", "more"), n["id"])) for i in range(5)}
    nbr = {}
    for (a, b) in g.edges:
        nbr.setdefault(a, []).append(b)
        nbr.setdefault(b, []).append(a)

    def pos_map():
        pm = {}
        for i, ns in lanes.items():
            for k, n in enumerate(ns):
                pm[n["id"]] = (k + 0.5) / max(1, len(ns))
        return pm
    for _ in range(3):
        for seq in (range(1, 5), range(3, -1, -1)):
            for i in seq:
                pm = pos_map()

                def key(n):
                    ys = [pm[x] for x in nbr.get(n["id"], []) if x in pm and g.nodes[x]["lane"] != i]
                    bary = sum(ys) / len(ys) if ys else pm[n["id"]]
                    return (n["kind"] == "more", n["kind"] == "batch" and i == 2, bary)
                lanes[i].sort(key=key)
    return lanes


def draw(g, png, dpi, title):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.lines import Line2D
        from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
    except ImportError:
        print("Thieu matplotlib. Chay: pip install matplotlib", file=sys.stderr)
        return 2
    lanes = order_lanes(g)
    nmax = max([len(v) for v in lanes.values()] + [1])
    BW, BH, GAPX, STEP = 3.0, 0.78, 2.6, 1.05
    H = max(5.5, nmax * STEP + 1.6)
    W = 5 * BW + 4 * GAPX
    pos = {}
    for i, ns in lanes.items():
        x = i * (BW + GAPX) + BW / 2
        tot = len(ns) * STEP
        y0 = (H - 1.2) / 2 + tot / 2 - STEP / 2
        for k, n in enumerate(ns):
            pos[n["id"]] = (x, y0 - k * STEP)
    fig_w = min(30, max(14, W * 0.85))
    fig_h = min(60, max(6, H * 0.62))
    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    ax.set_axis_off()
    for i, name in enumerate(LANES):
        x = i * (BW + GAPX)
        ax.add_patch(FancyBboxPatch((x - 0.35, -0.4), BW + 0.7, H - 0.4,
                                    boxstyle="round,pad=0.0,rounding_size=0.2",
                                    fc=LANE_FILL[i], ec="none", zorder=0))
        ax.text(x + BW / 2, H - 0.45, "%s (%d)" % (name, len(lanes[i])), ha="center", va="center",
                fontsize=11, fontweight="bold", color=INK, zorder=5)
    outd, ind = {}, {}
    for (a, b) in g.edges:
        outd[a] = outd.get(a, 0) + 1
        ind[b] = ind.get(b, 0) + 1
    for (a, b), e in sorted(g.edges.items(), key=lambda kv: kv[1]["doc"] > 0):
        (x1, y1), (x2, y2) = pos[a], pos[b]
        doc = e["doc"] > 0
        sx, tx = x1 + BW / 2, x2 - BW / 2
        if g.nodes[a]["lane"] > g.nodes[b]["lane"]:
            sx, tx = x1 - BW / 2, x2 + BW / 2
        ax.add_patch(FancyArrowPatch((sx, y1), (tx, y2), arrowstyle="-|>", mutation_scale=10,
                                     linewidth=1.3 if doc else 1.1,
                                     linestyle="-" if doc else (0, (4, 3)),
                                     color=DOC_COLOR if doc else INF_COLOR, alpha=0.85,
                                     shrinkA=0, shrinkB=1, zorder=2))
        lab = str(e["doc"]) + (" (+%d?)" % e["inf"] if e["inf"] else "") if doc else "%d?" % e["inf"]
        # nhan dat o dau it canh hon de khong chong nhau
        t = 0.72 if outd.get(a, 0) > ind.get(b, 0) else 0.28
        if abs(g.nodes[b]["lane"] - g.nodes[a]["lane"]) > 1:
            t = 0.85 if t > 0.5 else 0.15
        ax.text(sx + (tx - sx) * t, y1 + (y2 - y1) * t, lab, fontsize=7, ha="center", va="center",
                color=DOC_COLOR if doc else INF_COLOR, zorder=4,
                bbox=dict(boxstyle="round,pad=0.12", fc="white", ec="none", alpha=0.9))
    for nid, n in g.nodes.items():
        x, y = pos[nid]
        fc = BATCH_FILL if n["kind"] == "batch" else "#FFFFFF"
        ax.add_patch(FancyBboxPatch((x - BW / 2, y - BH / 2), BW, BH,
                                    boxstyle="round,pad=0.02,rounding_size=0.12",
                                    linewidth=1.2, fc=fc, ec=INK if n["kind"] != "more" else "#888888",
                                    linestyle="-" if n["kind"] != "more" else "--", zorder=3))
        lines = [ln[:32] for ln in n["label"].split("\n")][:2]
        ax.text(x, y + (0.13 if len(lines) > 1 else 0), lines[0], ha="center", va="center",
                fontsize=8.5, fontweight="bold", color=INK, zorder=5)
        if len(lines) > 1:
            ax.text(x, y - 0.17, lines[1], ha="center", va="center", fontsize=7.2, color="#555555", zorder=5)
    handles = [Line2D([0], [0], color=DOC_COLOR, lw=1.5, label="Lien = ghi nhan trong inventory (documented)"),
               Line2D([0], [0], color=INF_COLOR, lw=1.3, ls=(0, (4, 3)),
                      label="Dut = suy ra tu code, can kiem chung (inferred)"),
               Line2D([0], [0], color="none", marker="s", markersize=10, markerfacecolor=BATCH_FILL,
                      markeredgecolor=INK, label="Nhom Batch / Queue"),
               Line2D([0], [0], color="none", label="So tren canh = so API/job · '?' = suy ra · '*' = bang chua co trong 03_DB_Tables")]
    ax.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.5, 0.0), ncol=2, fontsize=9,
              frameon=False)
    if title:
        ax.set_title(title, fontsize=15, fontweight="bold", pad=14)
    ax.set_xlim(-0.6, W + 0.6)
    ax.set_ylim(-0.6, H)
    fig.tight_layout()
    os.makedirs(os.path.dirname(os.path.abspath(png)), exist_ok=True)
    fig.savefig(png, dpi=dpi, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return 0


# ---------------------------------------------------------------- markdown
def mid(nid):
    return "n_" + re.sub(r"[^A-Za-z0-9_]", "_", nid)


def mlabel(s):
    return s.replace('"', "'").replace("\n", "<br/>")


def write_md(g, detail, batch, hidden, md, png, title):
    lanes = order_lanes(g)
    L = ["# %s" % title, "",
         "![Code map](%s)" % os.path.relpath(png, os.path.dirname(os.path.abspath(md))), "",
         "Lien (`-->`) = ghi nhan trong inventory · Dut (`-.->`) = suy ra tu code (fe-api-calls, "
         "table-refs), **can kiem chung** truoc khi ghi Confirmed. Node la nhom; so tren canh = so API/job.", "",
         "```mermaid", "flowchart LR"]
    for i, name in enumerate(LANES):
        L.append('  subgraph L%d["%s"]' % (i, name))
        for n in lanes[i]:
            L.append('    %s["%s"]' % (mid(n["id"]), mlabel(n["label"])))
        L.append("  end")
    for (a, b), e in g.edges.items():
        if e["doc"]:
            lab = str(e["doc"]) + (" +%d?" % e["inf"] if e["inf"] else "")
            L.append('  %s -->|"%s"| %s' % (mid(a), lab, mid(b)))
        else:
            L.append('  %s -.->|"%d?"| %s' % (mid(a), e["inf"], mid(b)))
    L += ["```", "", "## Man hinh -> API -> Bang", "",
          "| Screen ID | Site | Route | API (documented) | API (inferred) | Tables (documented) | Tables (inferred) |",
          "|---|---|---|---|---|---|---|"]

    def j(x):
        return ", ".join(x) if x else "—"
    for d in detail:
        L.append("| %s | %s | %s | %s | %s | %s | %s |" % (
            d["screen"], d["site"] or S.UNKNOWN, (d["route"] or S.UNKNOWN).replace("|", "/"),
            j(d["api_doc"]), j(d["api_inf"]), j(d["tables_doc"]), j(d["tables_inf"])))
    if not detail:
        L.append("| — | — | — | — | — | — | — |")
    if batch:
        L += ["", "## Batch / Queue -> Bang", "",
              "| API ID | Kind | Schedule | Tables (documented) | Tables (inferred) |", "|---|---|---|---|---|"]
        for b in batch:
            L.append("| %s | %s | %s | %s | %s |" % (b["api"], b["kind"], b["schedule"].replace("|", "/"),
                                                    j(b["tables_doc"]), j(b["tables_inf"])))
    if hidden:
        L += ["", "Bang an khoi hinh (it ket noi): %s" % ", ".join(h.replace("tbl:", "") for h in hidden)]
    with open(md, "w", encoding="utf8") as fh:
        fh.write("\n".join(L) + "\n")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--inventory", required=True)
    ap.add_argument("--code", default="")
    ap.add_argument("--png", required=True)
    ap.add_argument("--md", required=True)
    ap.add_argument("--json", default="")
    ap.add_argument("--max-tables", type=int, default=25)
    ap.add_argument("--max-groups", type=int, default=60)
    ap.add_argument("--dpi", type=int, default=150)
    ap.add_argument("--title", default="")
    a = ap.parse_args()
    if not os.path.isfile(a.inventory):
        print("Khong thay inventory: %s" % a.inventory, file=sys.stderr)
        return 1
    data = load_inventory(a.inventory)
    recon = load_recon(a.code)
    g, detail, batch, hidden = build(data, recon, max(3, a.max_tables), max(5, a.max_groups))
    if not g.nodes:
        print("Inventory rong (10_Site/11_Repo/07_API/03_DB_Tables) — khong co gi de ve", file=sys.stderr)
        return 1
    sysname = val(S.meta(data).get("system_name"))
    title = a.title or "Code map" + (" — %s" % sysname if sysname else "")
    rc = draw(g, a.png, a.dpi, title)
    if rc:
        return rc
    os.makedirs(os.path.dirname(os.path.abspath(a.md)), exist_ok=True)
    write_md(g, detail, batch, hidden, a.md, a.png, title)
    if a.json:
        with open(a.json, "w", encoding="utf8") as fh:
            json.dump({"nodes": list(g.nodes.values()),
                       "edges": [{"from": x, "to": y, "documented": e["doc"], "inferred": e["inf"]}
                                 for (x, y), e in g.edges.items()],
                       "screens": detail, "batch": batch, "hidden_tables": hidden},
                      fh, indent=2, ensure_ascii=False)
    nd = sum(1 for e in g.edges.values() if e["doc"])
    ni = sum(1 for e in g.edges.values() if not e["doc"])
    print(json.dumps({"png": a.png, "md": a.md, "nodes": len(g.nodes),
                      "nodes_per_lane": {LANES[i]: sum(1 for n in g.nodes.values() if n["lane"] == i) for i in range(5)},
                      "edges_documented": nd, "edges_inferred_only": ni,
                      "hidden_tables": len(hidden), "screens": len(detail)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
