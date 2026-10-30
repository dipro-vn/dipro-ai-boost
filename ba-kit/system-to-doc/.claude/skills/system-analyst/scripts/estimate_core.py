"""Loi dung chung cho 見積書 Luong 2 (estimate.json -> sheet 見積書 + 前提条件・留意事項, va gate V-EST).

estimate.json = file trung gian, 1 dong = 1 CHUC NANG (man x thao tac), KHONG phai 1 impact. Song ngu: moi chuoi hien
thi la {"ja": ..., "vi": ...}. Ban "ja" gui KH (khong ma noi bo, khong tieng Viet); ban "vi" noi bo (them cot Impact,
can cu cong so). Cong so 実装 = function-rates.json [code].md x qty (hoac impl_override + override_note).

{
 "meta": {"cr_id", "title": {ja, vi}, "customer": {ja, vi}, "date": "YYYY-MM-DD", "valid_days": 30},
 "patterns": [{"id": "1B", "label": {ja, vi}}],                  # phuong an (S4); dong khong thuoc phuong an -> "pattern": null
 "combos":   [{"label": {ja, vi}, "patterns": ["1B", "2"]}],      # tong theo to hop
 "rows": [{"section": "user|sa|admin|common", "big": {ja, vi}, "mid": {..}, "small": {..}, "content": {..},
           "confirm": {..} | null, "reuse": "新規開発|一部流用|要確認", "must": "◉|◯", "step": "Step 1|Step 2",
           "pattern": "1B" | null, "code": "SCR-UPD-L", "qty": 1, "ui": "L|M|-" (mac dinh theo code), "req": true,
           "impacts": ["IMP-001"], "impl_override": null, "override_note": null}],
 "premises": {"前提条件": [{ja, vi}], "既存システムへの影響": [..], "クライアント側で事前に決定・準備が必要な事項": [..], "想定されるリスク": [..]},
 "questions": [{"id", "group": {ja, vi}, "question": {..}, "why": {..}, "proposal": {..}, "alt": {..}, "due": "YYYY/MM/DD"}]  # tuy chon:
     # cau hoi NGHIEP VU gui khach -> sheet 確認事項 / Cau hoi nghiep vu; id khong in ra (khong ma noi bo)
}
Van phong moi chuoi hien thi: .claude/sys-agent/writing-style.md (gate E4 / E9 / E10).
"""
import copy
import datetime
import io
import json
import re

SECTIONS = ["user", "sa", "admin", "common"]
REUSE = ["新規開発", "一部流用", "要確認"]
MUST = ["◉", "◯"]
STEPS = ["Step 1", "Step 2"]
PREMISE_KEYS = ["前提条件", "既存システムへの影響", "クライアント側で事前に決定・準備が必要な事項", "想定されるリスク"]
HDR_KEYS = {"no": ["No."], "big": ["項目（大）"], "mid": ["項目（中）"], "small": ["項目（小）"], "content": ["改修内容"],
            "reuse": ["流用可否"], "must": ["必須"], "step": ["STEP"], "req": ["要件定義"], "ui": ["UI"], "impl": ["実装"],
            "test": ["テスト"], "mgmt": ["管理"], "sub": ["小計"]}
PHASE = ["req", "ui", "impl", "test", "mgmt", "sub"]
HDR_VI = {"no": "No.", "big": "Hạng mục (lớn)", "mid": "Hạng mục (vừa)", "small": "Hạng mục (nhỏ)", "content": "Nội dung chỉnh sửa",
          "reuse": "Khả năng tận dụng", "must": "Bắt buộc: ◉\nTuỳ chọn: ◯", "step": "STEP", "req": "Định nghĩa\nyêu cầu",
          "ui": "Thiết kế\nUI / UX", "impl": "Lập trình", "test": "Kiểm thử", "mgmt": "Quản lý", "sub": "Tổng"}
for _k, _v in HDR_VI.items():  # do cot ca ban vi
    HDR_KEYS[_k] = HDR_KEYS[_k] + [_v]

LBL = {
    "ja": {"sheet": "詳細見積書", "notes_sheet": "前提条件・留意事項",
           "sections": {"user": "ユーザー予約サイト", "sa": "運営管理サイト（SUPER ADMIN）", "admin": "施設管理サイト（ADMIN）",
                        "common": "全システム"},
           "reuse": {k: k for k in REUSE}, "total": "合計(税別）", "pd": "（人日）", "pm": "（人月）",
           "step1": "Step1 の機能のみ", "step2": "Step２ の機能のみ", "all": "全ての機能", "pattern": "パターン",
           "ptable": "■パターン別合計（代替案は1つを選択）", "pm_col": "人月", "remarks": "■備考", "confirm": "【見積もり前提・確認事項】",
           "notes": ["※ 本見積は機能別の標準工数に基づく概算です。仕様確定後に再見積もりさせていただく場合がございます。",
                     "※ 「パターン」列が同じグループの代替案（例：①-A / ①-B）は、いずれか1つを選択いただく前提です。合計は「パターン別合計」をご参照ください。",
                     "※ 要件定義 = 実装 × {req}・UI/UXデザイン = 実装 × {uil}（新規の大規模画面）／ × {ui}（その他）・テスト = 実装 × {test}・管理 = (要件定義〜テスト) × {mgmt}・人月 = 人日 / {pm}"],
           "premise_heads": {k: "◆" + k for k in PREMISE_KEYS}, "valid": "見積日より{d}日間",
           "q_sheet": "確認事項", "q_title": "ご確認いただきたい事項",
           "q_cols": ["No.", "分類", "確認事項", "確認が必要な理由", "弊社からのご提案（ご指定がない場合）", "ご提案と異なる場合の影響", "ご回答期限", "ご回答"]},
    "vi": {"sheet": "ETA_vi", "notes_sheet": "Tiền đề · Lưu ý",
           "sections": {"user": "Trang đặt chỗ phía người dùng (WEB-02)", "sa": "Trang quản trị — Super Admin (WEB-01)",
                        "admin": "Trang quản trị — Admin cửa hàng (WEB-01)", "common": "Toàn hệ thống (DB)"},
           "reuse": {"新規開発": "Phát triển mới", "一部流用": "Tận dụng một phần", "要確認": "Cần xác nhận"},
           "total": "Tổng (chưa thuế)", "pd": "(người-ngày)", "pm": "(người-tháng)", "step1": "Chỉ Step 1",
           "step2": "Chỉ Step 2", "all": "Toàn bộ chức năng", "pattern": "Phương án",
           "ptable": "■ Tổng theo phương án (phương án thay thế chọn 1)", "pm_col": "người-tháng", "remarks": "■ Ghi chú",
           "confirm": "【Tiền đề ước tính · cần xác nhận】", "impact": "Impact (nội bộ)", "basis": "Căn cứ",
           "notes": ["※ Ước tính theo công số chuẩn cho từng loại chức năng (bảng công số {status} v{version} — PM / Tech Lead duyệt trước khi gửi khách hàng).",
                     "※ Các phương án cùng nhóm (VD ① phương án A / B) thay thế cho nhau, chỉ chọn 1 — tổng thực tế xem bảng 「Tổng theo phương án」.",
                     "※ Định nghĩa yêu cầu = Lập trình × {req} · Thiết kế UI / UX = Lập trình × {uil} (màn mới lớn) / × {ui} (còn lại) · Kiểm thử = Lập trình × {test} · Quản lý = (Định nghĩa yêu cầu … Kiểm thử) × {mgmt} · người-tháng = người-ngày / {pm}"],
           "premise_heads": {"前提条件": "◆ Tiền đề", "既存システムへの影響": "◆ Ảnh hưởng tới hệ thống hiện có",
                             "クライアント側で事前に決定・準備が必要な事項": "◆ Việc KH cần quyết định / chuẩn bị trước",
                             "想定されるリスク": "◆ Rủi ro dự kiến"},
           "valid": "{d} ngày kể từ ngày báo giá",
           "q_sheet": "Câu hỏi nghiệp vụ", "q_title": "Câu hỏi nghiệp vụ cần khách hàng trả lời",
           "q_cols": ["No.", "Nhóm", "Câu hỏi", "Vì sao cần hỏi", "Đề xuất của Dipro (nếu khách hàng không có ý kiến)", "Nếu khách hàng chọn khác", "Hạn trả lời", "Câu trả lời của khách hàng"]},
}
VI_CHARS = re.compile(r"[àáạảãầấậẩẫằắặẳẵèéẹẻẽềếệểễìíịỉĩòóọỏõồốộổỗờớợởỡùúụủũừứựửữỳýỵỷỹđăơư]", re.I)
INTERNAL = re.compile(r"\b(IMP|API|CQ|SC|EXT|FL|DS)-\d{2,3}\b|column:|table:|DS-component|cr\.json|\bTBD\b")


def t(v, lang):
    if isinstance(v, dict):
        return v.get(lang) or ""
    return "" if v is None else str(v)


def load_rates(path):
    d = json.load(open(path, encoding="utf8"))
    return d, {r["code"]: r for r in d["rates"]}


def load_template(path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    if "_meta" not in wb.sheetnames:
        raise ValueError("template %s thieu _meta — boc bang extract-mitsumori-template.py" % path)
    meta = {k: v for k, v in wb["_meta"].iter_rows(values_only=True) if k}
    return wb, meta


def find_layout(ws):
    for r in range(1, min(ws.max_row, 80) + 1):
        if any(str(ws.cell(row=r, column=c).value or "").strip() == "No." for c in range(1, 6)):
            cols = {}
            for rr in (r, r + 1):
                for c in range(1, ws.max_column + 1):
                    v = str(ws.cell(row=rr, column=c).value or "")
                    for k, keys in HDR_KEYS.items():
                        if k not in cols and any(x in v for x in keys):
                            cols[k] = c
            return r, cols
    return None, {}


# ---------------------------------------------------------------- tinh toan
def row_numbers(row, rates, tm):
    """-> dict req/ui/impl/test/mgmt/sub (so) cua 1 dong estimate theo he so template."""
    rate = rates.get(row.get("code")) or {}
    impl = row["impl_override"] if row.get("impl_override") is not None else float(rate.get("md", 0)) * float(row.get("qty", 1))
    fixed = rate.get("fixed") or {}
    ui_kind = row.get("ui") or rate.get("ui") or "M"
    req = fixed.get("req", impl * tm["ratio_req"]) if row.get("req", True) else 0.0
    ui = 0.0 if ui_kind == "-" else impl * (tm["ratio_ui_large"] if ui_kind == "L" else tm["ratio_ui"])
    test = fixed.get("test", impl * tm["ratio_test"])
    mgmt = fixed.get("mgmt", (req + ui + impl + test) * tm["ratio_mgmt"])
    return {"req": req, "ui": ui, "impl": impl, "test": test, "mgmt": mgmt, "sub": req + ui + impl + test + mgmt,
            "ui_kind": ui_kind, "fixed": fixed}


def validate(est, rates, cr=None):
    err = []
    m = est.get("meta") or {}
    for k in ("cr_id", "title", "customer", "date"):
        if not m.get(k):
            err.append("meta.%s rong" % k)
    pats = {p["id"] for p in est.get("patterns") or []}
    for c in est.get("combos") or []:
        for p in c.get("patterns") or []:
            if p not in pats:
                err.append("combo %s dung pattern %s khong khai bao" % (t(c.get("label"), "vi"), p))
    rows = est.get("rows") or []
    if not rows:
        err.append("rows rong")
    for i, r in enumerate(rows, 1):
        tag = "rows[%d]" % i
        if r.get("section") not in SECTIONS:
            err.append("%s section=%r (chi %s)" % (tag, r.get("section"), "/".join(SECTIONS)))
        for k in ("big", "content"):
            for lang in ("ja", "vi"):
                if not t(r.get(k), lang).strip():
                    err.append("%s thieu %s.%s" % (tag, k, lang))
        if r.get("reuse") not in REUSE:
            err.append("%s reuse=%r" % (tag, r.get("reuse")))
        if r.get("must") not in MUST:
            err.append("%s must=%r" % (tag, r.get("must")))
        if r.get("step") not in STEPS:
            err.append("%s step=%r" % (tag, r.get("step")))
        if r.get("pattern") and r["pattern"] not in pats:
            err.append("%s pattern=%r khong khai bao" % (tag, r["pattern"]))
        if r.get("code") not in rates:
            err.append("%s code=%r khong co trong function-rates" % (tag, r.get("code")))
        if r.get("impl_override") is not None and not t(r.get("override_note"), "vi").strip():
            err.append("%s impl_override khong co override_note" % tag)
        if (r.get("ui") or "M") not in ("L", "M", "-"):
            err.append("%s ui=%r" % (tag, r.get("ui")))
    for k in PREMISE_KEYS:
        if not (est.get("premises") or {}).get(k):
            err.append("premises.%s rong" % k)
    for i, q in enumerate(est.get("questions") or [], 1):
        for k in ("group", "question", "why", "proposal", "alt"):
            for lang in ("ja", "vi"):
                v = t(q.get(k), lang)
                if not v.strip():
                    err.append("questions[%d] thieu %s.%s" % (i, k, lang))
                elif INTERNAL.search(v):
                    err.append("questions[%d] %s.%s co ma noi bo %r — viet cho khach (writing-style.md)" % (i, k, lang, INTERNAL.search(v).group(0)))
    if cr is not None:
        ids = {i["id"] for i in cr.get("impacts") or []}
        for i, r in enumerate(rows, 1):
            for x in r.get("impacts") or []:
                if x not in ids:
                    err.append("rows[%d] impact %s khong co trong cr.json" % (i, x))
    return err


def trace_missing(est, cr):
    """Impact NEW / UPD (tru Mockup) chua thuoc dong chuc nang nao."""
    used = {x for r in est.get("rows") or [] for x in r.get("impacts") or []}
    return [i["id"] for i in cr.get("impacts") or []
            if i.get("change_type") in ("NEW", "UPD") and i.get("axis") != "Mockup" and i["id"] not in used]


# ---------------------------------------------------------------- ve sheet
def _styles(ws, r, maxc):
    out = {}
    for c in range(1, maxc + 1):
        x = ws.cell(row=r, column=c)
        out[c] = (copy.copy(x.font), copy.copy(x.fill), copy.copy(x.border), copy.copy(x.alignment), x.number_format)
    return out, ws.row_dimensions[r].height


def _apply(ws, r, st, maxc):
    sty, h = st
    for c in range(1, maxc + 1):
        if c in sty:
            x = ws.cell(row=r, column=c)
            x.font, x.fill, x.border, x.alignment, x.number_format = sty[c]
    if h:
        ws.row_dimensions[r].height = h


def copy_sheet(src, dst):
    """Copy gia tri + style + merge + do rong + anh tu sheet src sang dst (khac workbook)."""
    from openpyxl.drawing.image import Image as XLImage
    for row in src.iter_rows():
        for c in row:
            d = dst.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                d.font, d.fill, d.border = copy.copy(c.font), copy.copy(c.fill), copy.copy(c.border)
                d.alignment, d.number_format = copy.copy(c.alignment), c.number_format
    for k, dim in src.column_dimensions.items():
        dst.column_dimensions[k].width = dim.width
    for k, dim in src.row_dimensions.items():
        if dim.height:
            dst.row_dimensions[k].height = dim.height
    for m in src.merged_cells.ranges:
        dst.merge_cells(str(m))
    for im in getattr(src, "_images", []):
        n = XLImage(io.BytesIO(im._data()))
        n.anchor = im.anchor
        dst.add_image(n)
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines


RISK_RANK = {"High": 3, "Medium": 2, "Low": 1}
RISK_VI = {"High": "Lớn", "Medium": "Trung bình", "Low": "Nhỏ"}


def basis_detail(x, imp_by_id, rate_code, qty, impl, rate_item):
    """-> bullet nhieu dong (feedback 30/10/2026): '· N API mới\\n· Mức ảnh hưởng: Nhỏ'. Bo phan code × qty = MD (da co o cot MD rieng)."""
    imps = [imp_by_id[i] for i in (x.get("impacts") or []) if i in imp_by_id]
    bits = []
    for axis, name in (("System", "API"), ("DB", "bảng/cột DB"), ("Screen", "màn"), ("ThirdParty", "liên kết ngoài"), ("Mockup", "mockup"), ("Business", "quy tắc")):
        nw = sum(1 for i in imps if i["axis"] == axis and i.get("change_type") == "NEW")
        up = sum(1 for i in imps if i["axis"] == axis and i.get("change_type") == "UPD")
        if nw: bits.append("%d %s mới" % (nw, name))
        if up: bits.append("%d %s cập nhật" % (up, name))
    risk = max((RISK_RANK.get(i.get("risk"), 0) for i in imps), default=0)
    if risk: bits.append("Mức ảnh hưởng: " + RISK_VI[{3: "High", 2: "Medium", 1: "Low"}[risk]])
    return "\n".join("· " + b for b in bits) if bits else ""


def write_sheet(ws, est, rates, rates_meta, tm, lang, internal=False, cr=None):
    """Ghi bang chi tiet vao ws (da chua noi dung template). -> vi tri {first, last, row_step1, row_all, row_month, col}."""
    from openpyxl.styles import Alignment
    from openpyxl.utils import column_index_from_string as CI, get_column_letter as CL
    Lb = LBL[lang]
    col = {k[4:]: tm[k] for k in tm if str(k).startswith("col_")}
    ci = {k: CI(v) for k, v in col.items()}
    maxc = max(ci.values()) + (3 if internal else 1)
    sec_r, item_r, tot_r, notes_r = (int(tm[k]) for k in ("row_section", "row_item", "row_total", "row_notes"))
    st_sec, st_item = _styles(ws, sec_r, maxc), _styles(ws, item_r, maxc)
    st_tot = [_styles(ws, tot_r + k, maxc) for k in range(4)]
    st_note = _styles(ws, notes_r, maxc)
    for m in [m for m in list(ws.merged_cells.ranges) if m.min_row >= sec_r]:
        ws.unmerge_cells(str(m))
    for row in ws.iter_rows(min_row=sec_r, max_row=max(ws.max_row, sec_r)):
        for c in row:
            c.value = None
    m = est["meta"]
    date = datetime.date.fromisoformat(m["date"])
    repl = {"{{TITLE}}": t(m["title"], lang), "{{CUSTOMER}}": t(m["customer"], lang), "{{DATE}}": date,
            "{{VALID_UNTIL}}": date + datetime.timedelta(days=int(m.get("valid_days", 30)))}
    for row in ws.iter_rows(min_row=1, max_row=sec_r - 1):
        for c in row:
            if isinstance(c.value, str) and c.value in repl:
                c.value = repl[c.value]
                if isinstance(c.value, datetime.date):
                    c.number_format = "yyyy/mm/dd"
    if lang == "vi":  # dich header
        hdr = int(tm["row_header"])
        for k, v in HDR_VI.items():
            if k in ci:
                ws.cell(row=hdr + (0 if k in ("no", "big", "mid", "small", "content") else 1), column=ci[k], value=v)
        ws.cell(row=hdr, column=ci["req"], value="Công số (người-ngày)")
        for r0, c0, v in ((8, 5, "Kính gửi"), (10, 2, "■ Tiêu đề"), (11, 2, "■ Ngày báo giá"), (12, 2, "■ Hiệu lực đến"),
                          (5, 2, "BẢNG BÁO GIÁ CHI TIẾT")):
            if ws.cell(row=r0, column=c0).value:
                ws.cell(row=r0, column=c0, value=v)
    hdr = int(tm["row_header"])
    imp_by_id = {i["id"]: i for i in (cr or {}).get("impacts") or []}
    extra = {"pattern": max(ci.values()) + 1}
    if internal:
        extra.update(impact=extra["pattern"] + 1, basis=extra["pattern"] + 2)
    for k, c in extra.items():
        h = ws.cell(row=hdr + 1, column=c, value=Lb["pattern"] if k == "pattern" else Lb[k])
        src = ws.cell(row=hdr + 1, column=ci["sub"])
        h.font, h.fill, h.border, h.alignment = copy.copy(src.font), copy.copy(src.fill), copy.copy(src.border), copy.copy(src.alignment)
        ws.column_dimensions[CL(c)].width = 12 if k == "pattern" else 34
    pats = {p["id"]: t(p["label"], lang) for p in est.get("patterns") or []}
    L = lambda k, r: "%s%d" % (col[k], r)  # noqa: E731
    r, n, first = sec_r, 0, None
    for s in SECTIONS:
        rows = [x for x in est["rows"] if x["section"] == s]
        if not rows:
            continue
        _apply(ws, r, st_sec, maxc)
        ws.cell(row=r, column=ci["no"], value=" ◆ " + (est["meta"].get("section_labels", {}).get(s, {}).get(lang)
                                                       or Lb["sections"][s]))
        r += 1
        prev = (None, None)
        for x in rows:
            _apply(ws, r, st_item, maxc)
            n += 1
            nums = row_numbers(x, rates, tm)
            big, mid = t(x.get("big"), lang), t(x.get("mid"), lang)
            content = t(x.get("content"), lang)
            if x.get("confirm") and t(x["confirm"], lang):
                content += "\n\n" + Lb["confirm"] + "\n" + t(x["confirm"], lang)
            vals = {"no": n, "big": big if big != prev[0] else None, "mid": mid if (big, mid) != prev else None,
                    "small": t(x.get("small"), lang) or None, "content": content,
                    "reuse": Lb["reuse"][x["reuse"]], "must": x["must"], "step": x["step"], "impl": nums["impl"],
                    "req": (nums["fixed"]["req"] if "req" in nums["fixed"] else "=%s*%s" % (L("impl", r), tm["ratio_req"])) if x.get("req", True) else "-",
                    "ui": "-" if nums["ui_kind"] == "-" else "=%s*%s" % (L("impl", r), tm["ratio_ui_large"] if nums["ui_kind"] == "L" else tm["ratio_ui"]),
                    "test": nums["fixed"].get("test", "=%s*%s" % (L("impl", r), tm["ratio_test"])),
                    "mgmt": nums["fixed"].get("mgmt", "=SUM(%s:%s)*%s" % (L("req", r), L("test", r), tm["ratio_mgmt"])),
                    "sub": "=SUM(%s:%s)" % (L("req", r), L("mgmt", r))}
            prev = (big, mid)
            for k, v in vals.items():
                ws.cell(row=r, column=ci[k], value=v)
            ws.cell(row=r, column=extra["pattern"], value=pats.get(x.get("pattern")) or None)
            if internal:
                ws.cell(row=r, column=extra["impact"], value=", ".join(x.get("impacts") or []) or None)
                rate = rates[x["code"]]
                b = basis_detail(x, imp_by_id, x["code"], x.get("qty", 1), nums["impl"], rate["item"])
                if x.get("impl_override") is not None:
                    b = (b + "\n" if b else "") + "· override: " + t(x.get("override_note"), "vi")
                ws.cell(row=r, column=extra["basis"], value=b)
                for k in ("impact", "basis"):
                    ws.cell(row=r, column=extra[k]).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r, column=extra["pattern"]).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            lines = max(len(content) // 46 + content.count("\n") + 1, 2)
            ws.row_dimensions[r].height = min(15 * lines + 6, 400)
            first = first or r
            r += 1
    last = r - 1
    # tong: Step1 / Step2 / toan bo / nguoi-thang (phu dung first..last — gate kiem)
    tr = r
    for k in range(4):
        _apply(ws, tr + k, st_tot[k], maxc)
    ws.cell(row=tr, column=ci["no"], value=Lb["total"])
    for k, (u, lab) in enumerate(((Lb["pd"], Lb["step1"]), (Lb["pd"], Lb["step2"]), (Lb["pd"], Lb["all"]), (Lb["pm"], Lb["all"]))):
        ws.cell(row=tr + k, column=ci["mid"], value=u)
        ws.cell(row=tr + k, column=ci["small"], value=lab)
    stp = "$%s$%d:$%s$%d" % (col["step"], first, col["step"], last)
    for k in PHASE:
        c = col[k]
        ws.cell(row=tr, column=ci[k], value='=SUMIF(%s,"Step 1",%s%d:%s%d)' % (stp, c, first, c, last))
        ws.cell(row=tr + 1, column=ci[k], value='=SUMIF(%s,"Step 2",%s%d:%s%d)' % (stp, c, first, c, last))
        ws.cell(row=tr + 2, column=ci[k], value="=SUM(%s%d:%s%d)" % (c, first, c, last))
        ws.cell(row=tr + 3, column=ci[k], value="=%s%d/%s" % (c, tr + 2, tm["md_per_month"]))
        for j in range(4):
            ws.cell(row=tr + j, column=ci[k]).number_format = "0.00"
    ws.merge_cells(start_row=tr, start_column=ci["no"], end_row=tr + 3, end_column=ci["big"])
    r = tr + 5
    # tong theo to hop phuong an (S4)
    if est.get("combos"):
        pc = "$%s$%d:$%s$%d" % (CL(extra["pattern"]), first, CL(extra["pattern"]), last)
        _apply(ws, r, st_sec, maxc)
        ws.cell(row=r, column=ci["no"], value=Lb["ptable"])
        ws.cell(row=r, column=extra["pattern"], value=Lb["pm_col"])
        r += 1
        for cb in est["combos"]:
            _apply(ws, r, st_tot[2], maxc)
            ws.cell(row=r, column=ci["mid"], value=Lb["pd"])
            ws.cell(row=r, column=ci["small"], value=t(cb["label"], lang))
            for k in PHASE:
                c = col[k]
                ws.cell(row=r, column=ci[k], value="=" + "+".join('SUMIF(%s,"%s",%s%d:%s%d)' % (pc, pats[p], c, first, c, last)
                                                                   for p in cb["patterns"])).number_format = "0.00"
            x = ws.cell(row=r, column=extra["pattern"], value="=%s%d/%s" % (col["sub"], r, tm["md_per_month"]))
            x.number_format = "0.00"
            r += 1
        r += 1
    _apply(ws, r, st_note, maxc)
    ws.cell(row=r, column=ci["no"], value=Lb["remarks"])
    fmt = dict(req=tm["ratio_req"], uil=tm["ratio_ui_large"], ui=tm["ratio_ui"], test=tm["ratio_test"], mgmt=tm["ratio_mgmt"],
               pm=tm["md_per_month"], status=rates_meta.get("status"), version=rates_meta.get("version"))
    for k, line in enumerate(Lb["notes"]):
        ws.cell(row=r + 1 + k, column=ci["no"], value=line.format(**fmt))
    ws.print_area = "A1:%s%d" % (CL(maxc), r + 1 + len(Lb["notes"]))
    return {"first": first, "last": last, "row_step1": tr, "row_all": tr + 2, "row_month": tr + 3, "col": col,
            "col_pattern": CL(extra["pattern"]), "rows": n}


def write_premises(ws, est, lang):
    """Sheet 前提条件・留意事项: dung lai style tieu de ◆ + dong noi dung cua template, ghi lai tu dau."""
    from openpyxl.styles import Alignment
    head = next((c for row in ws.iter_rows() for c in row if str(c.value or "").startswith("◆")), None)
    hs = (copy.copy(head.font), copy.copy(head.fill)) if head else None
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))
    for row in ws.iter_rows():
        for c in row:
            c.value = None
    ws.column_dimensions["C"].width = 140
    r = 3
    for k in PREMISE_KEYS:
        h = ws.cell(row=r, column=2, value=LBL[lang]["premise_heads"][k])
        if hs:
            h.font, h.fill = copy.copy(hs[0]), copy.copy(hs[1])
        r += 1
        for it in est["premises"][k]:
            c = ws.cell(row=r, column=3, value="• " + t(it, lang))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 18 * (len(t(it, lang)) // 110 + 1)
            r += 1
        r += 2
    ws.title = LBL[lang]["notes_sheet"]


def write_questions(wb, est, lang, font_name):
    """Sheet cau hoi nghiep vu gui khach (ja: 確認事項 / vi: Cau hoi nghiep vu). Khong in ma noi bo (id chi o estimate.json)."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as CL
    Lb = LBL[lang]
    ws = wb.create_sheet(Lb["q_sheet"])
    ws.cell(row=1, column=1, value=Lb["q_title"]).font = Font(name=font_name, size=14, bold=True)
    thin = Side(style="thin", color="FF999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    for k, (h, w) in enumerate(zip(Lb["q_cols"], (6, 16, 60, 34, 40, 36, 13, 40)), 1):
        c = ws.cell(row=3, column=k, value=h)
        c.font, c.border = Font(name=font_name, bold=True), box
        c.fill = PatternFill("solid", fgColor="FFDAEEF3")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[CL(k)].width = w
    ws.row_dimensions[3].height = 36
    for n, q in enumerate(est["questions"], 1):
        vals = [n, t(q["group"], lang), t(q["question"], lang), t(q["why"], lang), t(q["proposal"], lang), t(q["alt"], lang),
                q.get("due") or "", ""]
        for k, v in enumerate(vals, 1):
            c = ws.cell(row=3 + n, column=k, value=v)
            c.font, c.border = Font(name=font_name, size=10), box
            c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if k in (1, 7) else "left")
        ws.cell(row=3 + n, column=8).fill = PatternFill("solid", fgColor="FFFFF9EB")
        ws.row_dimensions[3 + n].height = 70


def build_workbook(est, rates, rates_meta, template, lang, internal=False, cr=None):
    """-> (workbook moi, vi tri). Dung cho build-estimate.py va gen ETA noi bo. cr=cr.json (co 'impacts') de ghi basis detail."""
    wb, tm = load_template(template)
    del wb["_meta"]
    ws = wb[[s for s in wb.sheetnames if s != tm.get("notes_sheet")][0]]
    pos = write_sheet(ws, est, rates, rates_meta, tm, lang, internal, cr=cr)
    ws.title = LBL[lang]["sheet"]
    if tm.get("notes_sheet") in wb.sheetnames:
        write_premises(wb[tm["notes_sheet"]], est, lang)
    if est.get("questions"):
        write_questions(wb, est, lang, ws.cell(row=int(tm["row_header"]), column=2).font.name)
    for s in wb.worksheets:
        s.freeze_panes = None  # template co the freeze san — kit khong freeze (R13)
        for row in s.iter_rows():
            for c in row:
                if isinstance(c.value, float):
                    c.value = round(c.value, 6)
    return wb, pos, tm


def plan_totals(est, rates, tm):
    """-> {pattern id | None: {req, ui, impl, test, mgmt, sub}} — de ve lich / bao cao."""
    out = {}
    for x in est["rows"]:
        d = out.setdefault(x.get("pattern"), {k: 0.0 for k in PHASE})
        for k, v in row_numbers(x, rates, tm).items():
            if k in d:
                d[k] += v
    return out


# ---------------------------------------------------------------- gate V-EST
RANGE = re.compile(r"\$?([A-Z]{1,2})\$?(\d+):\$?([A-Z]{1,2})\$?(\d+)")


def scan_sheet(ws):
    """Doc 1 sheet 見積書 bat ky (khong can estimate.json) -> (layout, item rows, total rows info, loi cau truc)."""
    hdr, cols = find_layout(ws)
    if not hdr or "impl" not in cols:
        return None, [], [], ["khong thay header 見積書 (No. / 実装)"]
    from openpyxl.utils import get_column_letter as CL
    items, totals, bad = [], [], []
    for r in range(hdr + 2, ws.max_row + 1):
        lead = " ".join(str(ws.cell(row=r, column=c).value or "") for c in range(cols["no"], cols["small"] + 1))
        v = ws.cell(row=r, column=cols["impl"]).value
        if isinstance(v, (int, float)) and "合計" not in lead and "Tổng" not in lead:
            items.append({"row": r, **{k: ws.cell(row=r, column=c).value for k, c in cols.items()}})
        elif isinstance(v, str) and v.startswith("=") and RANGE.search(v) and items:
            totals.append({"row": r, "label": lead.strip(), **{k: ws.cell(row=r, column=c).value for k, c in cols.items()}})
    return {"row": hdr, "cols": {k: CL(c) for k, c in cols.items()}}, items, totals, bad


def check_ranges(layout, items, totals):
    """Moi cong thuc tong phai phu DUNG dai dong hang muc (bat loi kieu SUM(J19:J106))."""
    bad = []
    if not items:
        return ["khong co dong hang muc"]
    f, l = items[0]["row"], items[-1]["row"]
    # dong hang muc lien tuc theo section -> chi so sanh dau / cuoi
    for tr in totals:
        for k in PHASE:
            v = str(tr.get(k) or "")
            for m in RANGE.finditer(v):
                a, r1, b, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
                if a == b and (r1, r2) != (f, l) and r2 - r1 > 2:
                    bad.append("dong %d (%s) cot %s: %s:%s chi phu %d..%d, hang muc o %d..%d" % (
                        tr["row"], tr["label"][:20], k, "%s%d" % (a, r1), "%s%d" % (b, r2), r1, r2, f, l))
    return bad


JP_CHARS = re.compile(r"[\u3040-\u30ff\u4e00-\u9fff]")


def check_language(wb, lang):
    bad = []
    for s in wb.worksheets:
        if s.sheet_state != "visible":
            continue
        for row in s.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str) or v.startswith("="):
                    continue
                if lang == "ja" and VI_CHARS.search(v):
                    bad.append("%s!%s co chu tieng Viet: %r" % (s.title, c.coordinate, v[:40]))
                if INTERNAL.search(v):
                    bad.append("%s!%s co ma noi bo: %r" % (s.title, c.coordinate, INTERNAL.search(v).group(0)))
    return bad[:30]


def check_japanese_left(wb, header_row):
    """Ban vi: chu Nhat con sot ngoai letterhead (dong tren header cua sheet chinh) -> WARN."""
    out = []
    for i, s in enumerate(wb.worksheets):
        if s.sheet_state != "visible":
            continue
        for row in s.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and not v.startswith("=") and JP_CHARS.search(v) and not (i == 0 and c.row < header_row):
                    out.append("%s!%s con chu Nhat: %r" % (s.title, c.coordinate, v[:40]))
    return out[:20]


def check_style(items, est, rates, tm, lang):
    """Dong co cong lap trinh >= 1 ngay ma noi dung < 40 ky tu hoac khong co cau truc (nhan / gach dau dong) -> WARN."""
    out = []
    for it, x in zip(items, [x for s in SECTIONS for x in est["rows"] if x["section"] == s]):
        txt = str(it.get("content") or "")
        if row_numbers(x, rates, tm)["impl"] >= 1 and (len(txt) < 40 or not re.search(r"【|・|•|^\s*[-\d]+[.)]|Nội dung", txt, re.M)):
            out.append("dong %d: noi dung qua ngan / khong co cau truc (writing-style.md §2): %r" % (it["row"], txt[:40]))
    return out


def verify(wb, sheet, est, rates, rates_meta, tm, lang, cr=None):
    """-> list (id, title, bad, warn_only). Dung chung cho verify-estimate.py va verify-cr-impact.py."""
    ws = wb[sheet]
    layout, items, totals, bad0 = scan_sheet(ws)
    res = [("E1", "Cau truc: header 見積書 + cot du (流用可否 · 必須 · STEP · cong doan)", bad0 or
            (["thieu cot: %s" % k for k in HDR_KEYS if k not in (layout or {}).get("cols", {})] if layout else []), False)]
    if not layout:
        return res
    reuse_ok = set(LBL[lang]["reuse"].values())
    b = []
    for it in items:
        if it.get("reuse") not in reuse_ok:
            b.append("dong %d 流用可否=%r" % (it["row"], it.get("reuse")))
        if it.get("must") not in MUST:
            b.append("dong %d 必須/オプション=%r" % (it["row"], it.get("must")))
        if it.get("step") not in STEPS:
            b.append("dong %d STEP=%r" % (it["row"], it.get("step")))
    res.append(("E2", "Moi dong co 流用可否 · ◉/◯ · STEP hop le", b, False))
    res.append(("E3", "Cong thuc tong phu dung dai dong hang muc (Step1 · Step2 · toan bo · 人月)", check_ranges(layout, items, totals), False))
    internal = any(str(c.value or "") == LBL["vi"]["impact"] for c in ws[layout["row"] + 1]) if lang == "vi" else False
    res.append(("E4", "Ngon ngu / ma noi bo: ban %s%s" % (lang, " noi bo — bo qua" if internal else ""),
                [] if internal else check_language(wb, lang), False))
    if est is None:
        return res
    b = []
    if len(items) != len(est["rows"]):
        b.append("so dong %d != estimate.json %d" % (len(items), len(est["rows"])))
    for it, x in zip(items, [x for s in SECTIONS for x in est["rows"] if x["section"] == s]):
        nums = row_numbers(x, rates, tm)
        if round(float(it["impl"]), 2) != round(nums["impl"], 2):
            b.append("dong %d 実装=%s != %s x %s = %s" % (it["row"], it["impl"], x["code"], x.get("qty", 1), round(nums["impl"], 2)))
    s1 = sum(row_numbers(x, rates, tm)["sub"] for x in est["rows"] if x["step"] == "Step 1")
    s2 = sum(row_numbers(x, rates, tm)["sub"] for x in est["rows"] if x["step"] == "Step 2")
    al = sum(row_numbers(x, rates, tm)["sub"] for x in est["rows"])
    if round(s1 + s2, 4) != round(al, 4):
        b.append("Step1 + Step2 = %s != toan bo %s" % (round(s1 + s2, 2), round(al, 2)))
    res.append(("E5", "実装 = function-rates x qty; Step1 + Step2 = toan bo", b, False))
    w = ["dong %s: impl_override (%s)" % (i + 1, t(x.get("override_note"), "vi")) for i, x in enumerate(est["rows"]) if x.get("impl_override") is not None]
    if str(rates_meta.get("status", "")).upper() != "APPROVED":
        w.append("function-rates status=%s — PM / Tech Lead phai duyet truoc khi gui KH" % rates_meta.get("status"))
    res.append(("E6", "Bang cong so chuc nang da duyet, khong override", w, True))
    notes = tm.get("notes_sheet")
    b = []
    if lang == "ja":
        nm = LBL["ja"]["notes_sheet"]
        if nm not in wb.sheetnames:
            b.append("thieu sheet %s" % nm)
        else:
            txt = "\n".join(str(c.value or "") for row in wb[nm].iter_rows() for c in row)
            b += ["thieu khoi %s" % h for h in LBL["ja"]["premise_heads"].values() if h not in txt]
    res.append(("E7", "Sheet 前提条件・留意事項 du 4 khoi (ban KH)", b, False))
    if lang == "vi" and not internal:
        res.append(("E9", "Ban vi: khong con chu Nhat ngoai letterhead (writing-style.md quy tac 9)", check_japanese_left(wb, layout["row"]), True))
    res.append(("E10", "Van phong: dong chuc nang co noi dung cu the, co cau truc (writing-style.md)", check_style(items, est, rates, tm, lang), True))
    if cr is not None:
        res.append(("E8", "Moi impact NEW / UPD (tru Mockup) thuoc >= 1 dong chuc nang",
                    ["%s chua thuoc dong nao" % i for i in trace_missing(est, cr)], False))
    return res
