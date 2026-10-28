#!/usr/bin/env python3
"""Sinh workbook rong dung schema (header + khung Key/Value).

  python3 build-inventory.py --out <ver>/_internal/inventory.xlsx
  python3 build-inventory.py --out ... --bug-list        # workbook Bug List (O7) — 1 sheet "Bug List"
  python3 build-inventory.py --out <ver moi>/_internal/inventory.xlsx --carry-from <ver cu>/_internal/inventory.xlsx
      # carry: chep moi dong cua inventory cu (doc qua inv_schema.load -> chu cu tu doi theo D1:
      # cot trang thai -> "To verify", o noi dung -> "Cần xác nhận"), 00_Meta giu khung moi

CR Impact (Luong 2): dung build-cr-impact.py.
Artifact noi bo cua agent — KHONG phai deliverable cho user.
"""
import argparse
import os
import sys

import inv_schema as S


def build(path, sheets, meta_keys=None, carry_from=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("Thieu openpyxl. Chay: pip install openpyxl", file=sys.stderr)
        return 2

    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    head_font = Font(bold=True)

    for name, cols in sheets.items():
        ws = wb.create_sheet(name)
        for i, col in enumerate(cols, start=1):
            c = ws.cell(row=1, column=i, value=col)
            c.fill = head_fill
            c.font = head_font
            c.alignment = Alignment(vertical="center", wrap_text=True)
            width = 40 if col in ("Description", "Repro Steps", "Meaning",
                                  "Reason / Evidence", "Business Impact",
                                  "Locator", "Why suspicious", "Summary",
                                  "Question") else 18
            ws.column_dimensions[get_column_letter(i)].width = width

    if meta_keys:
        ws = wb[next(iter(sheets))]
        for i, k in enumerate(meta_keys, start=2):
            ws.cell(row=i, column=1, value=k)
            ws.cell(row=i, column=2, value=S.NEED_CONFIRM)

    if carry_from:   # S.load doi chu cu cua version truoc D1 (inv_schema.legacy)
        old = S.load(carry_from, {k: v for k, v in sheets.items() if k != "00_Meta"})
        for name, rows in old.items():
            for r in rows:
                wb[name].append([r.get(h, "") for h in sheets[name]])

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)  # vd 07_BugList/ chua co
    wb.save(path)
    print("Da tao %s (%d sheet)" % (path, len(sheets)))
    return 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--bug-list", action="store_true",
                    help="Sinh workbook Bug List thay vi Inventory")
    ap.add_argument("--carry-from", default=None,
                    help="inventory (hoac bug list) cua version truoc — chep dong, doi chu cu theo D1")
    a = ap.parse_args()
    if a.bug_list:
        return build(a.out, S.BUG_SHEETS, None, a.carry_from)   # 1 sheet, khong co 00_Meta
    return build(a.out, S.SHEETS, S.META_KEYS, a.carry_from)


if __name__ == "__main__":
    sys.exit(main())
