#!/usr/bin/env python3
"""Sinh muc luc cua 1 thu muc version: <ver>/README.md + <ver>/_internal/index.json.

  python3 build-version-index.py <ver folder> [--skip "O7=user khong yeu cau"] [--skip "O5=..."]

Tu nhan dien output theo duong dan chuan (O1..O7, hoac output CR neu co CR-*_Impact.xlsx: so hang muc + tong
MD + trang thai don gia, Figma CR (gate v-cr-figma), input/). O4 = 04_DesignSystem/project/ (1 bo) hoac
04_DesignSystem/WEB-xx/project/ (moi site 1 dong); link artifact + trang thai + so TBD lay tu STATUS.md
canh project/ (dong "- Artifact:" / "- Trạng thái:" / bang "## Thiếu (TBD)"),
dem so luong, doc ket qua gate tu _internal/gates/*.md (dong "N checks · X PASS · Y FAIL · Z WARN").
index.json la dau vao cua render-overview-docx.py (chuong 2). Chay lai sau moi gate.
README KHONG nhac toi _internal/ (artifact noi bo).
"""
import argparse
import glob
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import inv_schema as S  # noqa: E402
import cr_common as C  # noqa: E402

DASH = "—"
SUMMARY_RX = re.compile(r"(\d+) checks · (\d+) PASS · (\d+) FAIL · (\d+) WARN")
FIGMA_RX = re.compile(r"https?://(?:www\.)?figma\.com/[^\s)>\]]+")
ARTIFACT_RX = re.compile(r"https://claude\.ai/(?:code/)?artifact/[A-Za-z0-9_-]+")
VER_RX = re.compile(r"^ver(\d+)_(\d{2})(\d{2})(\d{2})_(.+)$")
BD_SYSTEM_SHEETS = {"Common mesage", "Screen Error message", "Screen Index",
                    "Change History", "Sample"}
ICON = {"ok": "✅", "warn": "⚠️", "fail": "❌", "skip": "⬜"}
# tien to ten file gate -> output (thu tu quan trong: dai truoc)
GATE_MAP = [("v-overview", "O6"), ("v-bd", "O1"), ("v-api", "O2"), ("v-db", "O3"),
            ("v-ds", "O4"), ("v-figma", "O5"), ("v-cr-figma", "O5"), ("v-cr", "CR"), ("v1", "INV"), ("v2", "INV"),
            ("v3", "O6"), ("v4", "O6"), ("v5", "O6"), ("v6", "O6"), ("v7", "O1"), ("v8", "O7")]


def rel(ver, p):
    return os.path.relpath(p, ver).replace(os.sep, "/")


def load_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("Thieu openpyxl. Chay: pip install openpyxl")
    return load_workbook(path, data_only=True)


def sheet_rows(ws, key_col=1):
    return sum(1 for r in ws.iter_rows(min_row=2, values_only=True)
               if len(r) >= key_col and r[key_col - 1] not in (None, ""))


def parse_gates(ver):
    """{output_id: [(file, P, F, W)]} — output O1 co the kem site: 'O1:WEB-01'."""
    out = {}
    for p in sorted(glob.glob(os.path.join(ver, "_internal", "gates", "*.md"))):
        stem = os.path.basename(p)[:-3].lower()
        target = next((o for pre, o in GATE_MAP if stem.startswith(pre)), None)
        if not target:
            continue
        hits = SUMMARY_RX.findall(open(p, encoding="utf8", errors="ignore").read())
        if not hits:
            continue
        _, pa, fa, wa = (int(x) for x in hits[-1])
        m = re.search(r"web-\d{2,}", stem)
        if target in ("O1", "O4") and m:
            target = target + ":" + m.group(0).upper()
        out.setdefault(target, []).append((os.path.basename(p), pa, fa, wa))
    return out


def gate_of(gates, key):
    items = gates.get(key, [])
    if key.startswith(("O1:", "O4:")):
        items = items + gates.get(key[:2], [])
    if not items:
        return DASH, "chua chay gate"
    f = sum(x[2] for x in items)
    w = sum(x[3] for x in items)
    detail = " · ".join("%s %d/%d PASS" % (n, pa, pa + fa + wa) for n, pa, fa, wa in items)
    if f:
        return "FAIL", detail
    return ("WARN" if w else "PASS"), detail


def finish(e, gates, skips, auto_skip=None):
    """Gan gate + trang thai cho 1 entry."""
    oid = e["id"]
    key = oid + ":" + e["site"] if oid in ("O1", "O4") and e.get("site") else oid
    e["gate"], e["gate_detail"] = gate_of(gates, key)
    reason = skips.get(oid) or skips.get(key.replace(":", "=")) or None
    if not e["exists"] and (reason or auto_skip):
        e["status"], e["reason"] = "skip", reason or auto_skip
        e["files"], e["count"] = [], DASH
        e["status_text"] = "%s Không chạy — %s" % (ICON["skip"], e["reason"])
        e["gate"] = DASH
        return e
    if not e["exists"]:
        e["status"], e["reason"] = "fail", "chưa có file"
    elif e["gate"] == "FAIL":
        e["status"], e["reason"] = "fail", "gate FAIL"
    elif e["gate"] == "WARN":
        e["status"], e["reason"] = "warn", "gate có cảnh báo"
    elif e["gate"] == DASH:
        e["status"], e["reason"] = "warn", "chưa chạy gate"
    else:
        e["status"], e["reason"] = "ok", ""
    e["status_text"] = ICON[e["status"]] + (" " + e["reason"] if e["reason"] else "")
    return e


def entry(oid, name, content, files, count, exists, **kw):
    d = {"id": oid, "name": name, "content": content, "files": files,
         "count": count, "exists": exists}
    d.update(kw)
    return d


def ds_summary(ds_dir, label):
    """1 bo design system = <ds_dir>/project/ (design-system.json, tokens.json, components/) + STATUS.md."""
    proj = os.path.join(ds_dir, "project")
    exists = os.path.isfile(os.path.join(proj, "design-system.json"))
    n_col = n_sty = n_comp = n_icon = 0
    if exists:
        try:
            tk = json.load(open(os.path.join(proj, "tokens.json"), encoding="utf8"))
            n_col = len((tk.get("color") or {}).get("tokens") or [])
            n_sty = sum(len(g.get("styles") or []) for g in (tk.get("type") or {}).get("groups") or [])
        except (OSError, ValueError, AttributeError):
            pass
        cdir = os.path.join(proj, "components")
        if os.path.isdir(cdir):
            n_comp = sum(1 for d in os.listdir(cdir) if os.path.isdir(os.path.join(cdir, d))
                         and d not in ("Cover", "lib", "src"))
        try:
            idx = json.load(open(os.path.join(proj, "design-system.json"), encoding="utf8"))
            n_icon = len(((idx.get("assetGroups") or {}).get("Icons") or {}).get("files") or {})
        except (OSError, ValueError, AttributeError):
            pass
        if not n_icon and os.path.isdir(os.path.join(proj, "assets", "Icons")):
            n_icon = sum(1 for f in os.listdir(os.path.join(proj, "assets", "Icons"))
                         if not f.lower().endswith((".md", ".json")))
    link, status, n_tbd = "", DASH, 0
    sp = os.path.join(ds_dir, "STATUS.md")
    if os.path.isfile(sp):
        sec = ""
        for line in open(sp, encoding="utf8").read().splitlines():
            m = re.match(r"^\s*-\s*(Artifact|Trạng thái|Trang thai)\s*:\s*(.*)$", line)
            if m and m.group(1) == "Artifact":
                a = ARTIFACT_RX.search(m.group(2))
                link = a.group(0) if a else link
            elif m:
                status = m.group(2).strip() or DASH
            if line.startswith("## "):
                sec = line[3:].strip()
            elif sec.startswith(("Thiếu", "Thieu")) and line.strip().startswith("|"):
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                if len(cells) >= 2 and not set(cells[0]) <= set("-: ") and not cells[1].startswith("Token"):
                    n_tbd += 1
    return {"exists": exists, "link": link, "status": status,
            "files": [label + "project/", label + "STATUS.md"] + ([link] if link else []),
            "count": "%d màu · %d style chữ · %d component · %d icon · %d TBD · %s" % (
                n_col, n_sty, n_comp, n_icon, n_tbd, status),
            "count_value": n_col}


def ds_entries(ver, gates, skips):
    """O4: 04_DesignSystem/project/ (1 bo) va/hoac 04_DesignSystem/WEB-xx/project/ (moi site 1 bo)."""
    root = os.path.join(ver, "04_DesignSystem")
    sets = [(site, d) for site, d in C.ds_dirs(root)]
    if not sets:
        sets = [("", root)]
    outs = []
    for site, d in sets:
        ds = ds_summary(d, "04_DesignSystem/" + (site + "/" if site else ""))
        kw = {"site": site} if site else {}
        if ds["link"]:
            kw["link"] = ds["link"]
        outs.append(finish(entry(
            "O4", "Design System" + (" %s" % site if site else ""),
            "Token, kiểu chữ, component, icon (format artifact Design System)",
            ds["files"], ds["count"], ds["exists"], count_value=ds["count_value"],
            ds_status=ds["status"], **kw), gates, skips))
    return outs


def baseline_outputs(ver, data, meta, sysname, n, gates, skips):
    outs = []
    sites = [r.get("Site ID", "") for r in (data or {}).get("10_Site", [])]
    screens_by_site = {}
    for r in (data or {}).get("02_Screen", []):
        screens_by_site[r.get("Site", "")] = screens_by_site.get(r.get("Site", ""), 0) + 1
    bd = {}
    for p in sorted(glob.glob(os.path.join(ver, "01_Screens", "BasicDesign_*.xlsx"))):
        m = re.match(r"BasicDesign_(.+?)_ver\d+\.xlsx$", os.path.basename(p))
        bd[m.group(1) if m else os.path.basename(p)] = p
    for site in sites + [s for s in bd if s not in sites]:
        p = bd.get(site)
        cnt = None
        if p:
            wb = load_xlsx(p)
            cnt = sum(1 for ws in wb.worksheets if ws.title not in BD_SYSTEM_SHEETS
                      and str(ws["F6"].value or "").strip() == "WORKING_SCREEN")
        count = ("%d sheet màn hình" % cnt) if cnt is not None else \
            ("%d màn hình (inventory)" % screens_by_site.get(site, 0) if data else S.UNKNOWN)
        outs.append(finish(entry(
            "O1", "Basic Design %s" % site, "Thiết kế cơ bản từng màn hình của site %s" % site,
            [rel(ver, p)] if p else ["01_Screens/BasicDesign_%s_ver%s.xlsx" % (site, n)],
            count, bool(p), site=site, count_value=cnt), gates, skips))

    apis = (data or {}).get("07_API", [])
    n_api = sum(1 for r in apis if r.get("Kind") == "API")
    n_other = len(apis) - n_api
    api_files = sorted(glob.glob(os.path.join(ver, "02_API", "API_Doc_*.xlsx")))
    cm = sorted(glob.glob(os.path.join(ver, "02_API", "CodeMap_*.png")))
    outs.append(finish(entry(
        "O2", "API & Code map", "Danh sách API/batch + sơ đồ FE→API→BE→DB",
        [rel(ver, p) for p in api_files + cm] or ["02_API/API_Doc_%s_ver%s.xlsx" % (sysname, n)],
        ("%d API · %d batch/khác" % (n_api, n_other)) if data else S.UNKNOWN,
        bool(api_files), count_value=len(apis) if data else None), gates, skips))

    tables = (data or {}).get("03_DB_Tables", [])
    db_files = sorted(glob.glob(os.path.join(ver, "03_DB", "DB_Doc_*.xlsx")))
    erd = sorted(glob.glob(os.path.join(ver, "03_DB", "ERD_*.png")))
    outs.append(finish(entry(
        "O3", "Database", "Danh sách bảng/cột/quan hệ + ERD",
        [rel(ver, p) for p in db_files + erd] or ["03_DB/DB_Doc_%s_ver%s.xlsx" % (sysname, n)],
        ("%d bảng · %d cột" % (len(tables), len((data or {}).get("04_DB_Columns", []))))
        if data else S.UNKNOWN,
        bool(db_files), count_value=len(tables) if data else None), gates, skips,
        auto_skip="không có DB (db_mode=NONE)" if meta.get("db_mode") == "NONE" else None))

    outs += ds_entries(ver, gates, skips)

    fl = os.path.join(ver, "05_Figma", "figma-links.md")
    urls = FIGMA_RX.findall(open(fl, encoding="utf8").read()) if os.path.isfile(fl) else []
    fig_out = meta.get("figma_output_url", "")
    outs.append(finish(entry(
        "O5", "Figma", "Link node Figma của màn hình/thành phần",
        ["05_Figma/figma-links.md"], "%d link" % len(urls), os.path.isfile(fl),
        count_value=len(urls)), gates, skips,
        auto_skip="không có Figma output" if fig_out in ("", S.UNKNOWN, DASH, "NO") else None))

    ov = sorted(glob.glob(os.path.join(ver, "06_Overview", "Overview_*.docx")))
    outs.append(finish(entry(
        "O6", "Overview", "Tài liệu tóm tắt khảo sát (docx)",
        [rel(ver, p) for p in ov] or ["06_Overview/Overview_%s_ver%s.docx" % (sysname, n)],
        "1 tài liệu", bool(ov), count_value=1 if ov else 0), gates, skips))

    bl = sorted(glob.glob(os.path.join(ver, "07_BugList", "BugList_*.xlsx")))
    nb = None
    if bl:
        wb = load_xlsx(bl[0])
        nb = sheet_rows(wb["Bugs"]) if "Bugs" in wb.sheetnames else 0
    outs.append(finish(entry(
        "O7", "Bug List", "Lỗi Medium/High đang tồn tại",
        [rel(ver, p) for p in bl] or ["07_BugList/BugList_%s_ver%s.xlsx" % (sysname, n)],
        ("%d bug" % nb) if nb is not None else DASH, bool(bl), count_value=nb), gates, skips,
        auto_skip=("bug_list=%s" % meta.get("bug_list")) if meta.get("bug_list") != "YES" else None))
    return outs


def cr_outputs(ver, impact, gates, skips):
    """CR folder: CR-<id>_Impact.xlsx (Summary · Estimation · Screen · API · Database · Figma · Q&A) · Figma CR · input/."""
    outs = []
    wb = load_xlsx(impact)
    n_imp, md = 0, 0.0
    for name, _axes in S.CR_AXIS_SHEETS + [(S.CR_SHEET_ESTIMATION, None)]:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        if name == S.CR_SHEET_ESTIMATION:      # B = Impact ID, M = 実装 (MD)
            ids = [(str(r[1] or ""), r[12]) for r in ws.iter_rows(values_only=True) if len(r) > 12]
            ids = [(i, m) for i, m in ids if re.match(S.CR_IMPACT_ID, i)]
            if ids:                              # Estimation du moi hang muc -> dung lam nguon dem
                n_imp = len(ids)
                md = sum(float(m) for _i, m in ids if isinstance(m, (int, float)))
            continue
        hdr = [str(c.value or "").strip() for c in ws[1]]
        ci, cm = hdr.index("Impact ID") if "Impact ID" in hdr else 1, hdr.index("MD") if "MD" in hdr else None
        for r in ws.iter_rows(min_row=2, values_only=True):
            if re.match(S.CR_IMPACT_ID, str(r[ci] or "")):
                n_imp += 1
                try:
                    md += float(r[cm]) if cm is not None and r[cm] is not None else 0
                except (TypeError, ValueError):
                    pass
    summ, rate_st = {}, "?"
    if S.CR_SHEET_SUMMARY in wb.sheetnames:
        for r in wb[S.CR_SHEET_SUMMARY].iter_rows(values_only=True):
            if r and r[0] is not None and len(r) > 1:
                summ.setdefault(str(r[0]).strip(), "" if r[1] is None else str(r[1]))
        rate_st = (summ.get(C.L_RATES) or "?").split()[0]
    cj = os.path.join(ver, "_internal", "cr.json")
    meta = {}
    if os.path.isfile(cj):
        try:
            meta = C.load_json(cj).get("meta") or {}
        except (OSError, ValueError):
            pass
    meta.setdefault("baseline_version", summ.get(C.L_BASELINE, ""))
    meta.setdefault("cr_id", summ.get(C.L_CR, ""))
    outs.append(finish(entry("CR", "CR Impact", "Công số + ảnh hưởng 6 trục (%s)" % " · ".join(S.CR_SHEETS),
                             [rel(ver, impact)], "%d hạng mục · %s MD (đơn giá %s)" % (
                                 n_imp, C.fmt_md(round(md, 2)), rate_st),
                             True, count_value=n_imp, md=round(md, 2), rates_status=rate_st), gates, skips))
    fl = os.path.join(ver, "05_Figma", "figma-links.md")
    urls = FIGMA_RX.findall(open(fl, encoding="utf8").read()) if os.path.isfile(fl) else []
    # CR khong bat buoc ve: khong co figma-links + khong co gate v-cr-figma -> skip "không vẽ"
    outs.append(finish(entry("O5", "Figma (CR)", "Màn/luồng thay đổi theo CR (badge NEW/UPD/DEL/IMPACT/AS-IS)",
                             ["05_Figma/figma-links.md"], "%d link" % len(urls),
                             os.path.isfile(fl) or bool(gates.get("O5")), count_value=len(urls)),
                       gates, skips, auto_skip="không vẽ"))
    inp = [p for p in glob.glob(os.path.join(ver, "input", "**", "*"), recursive=True)
           if os.path.isfile(p)]
    outs.append(finish(entry("IN", "Input CR", "Tài liệu yêu cầu gốc", ["input/"],
                             "%d file" % len(inp), bool(inp), count_value=len(inp)),
                       gates, skips))
    # input khong co gate -> trang thai theo su ton tai
    for e in outs:
        if e["id"] == "IN" and e["gate"] == DASH and e["exists"]:
            e["status"], e["reason"], e["status_text"] = "ok", "", ICON["ok"]
    return outs, meta


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("version_folder")
    ap.add_argument("--skip", action="append", default=[],
                    help='"O7=ly do" (lap lai duoc; rieng 1 site: "O1=WEB-02=ly do" / "O4=WEB-02=ly do")')
    a = ap.parse_args()

    ver = os.path.abspath(a.version_folder)
    if not os.path.isdir(ver):
        print("Khong thay thu muc %s" % ver, file=sys.stderr)
        return 1
    skips = {}
    for s in a.skip:
        k, _, v = s.partition("=")
        if k in ("O1", "O4") and re.match(r"^WEB-\d+=", v):
            site, _, v = v.partition("=")
            k = k + "=" + site
        skips[k.strip()] = v.strip() or "không yêu cầu"

    folder = os.path.basename(ver)
    m = VER_RX.match(folder)
    n = m.group(1) if m else "?"
    date = ("%s/%s/20%s" % (m.group(2), m.group(3), m.group(4))) if m else DASH

    inv_path = os.path.join(ver, "_internal", "inventory.xlsx")
    data = S.load(inv_path) if os.path.isfile(inv_path) else None
    meta = S.meta(data) if data else {}
    sysname = meta.get("system_name") or S.UNKNOWN
    if sysname == S.UNKNOWN:
        for p in glob.glob(os.path.join(ver, "0*", "*_ver*.*")):
            mm = re.match(r"(?:API_Doc|DB_Doc|Overview|BugList)_(.+?)_ver\d+\.", os.path.basename(p))
            if mm:
                sysname = mm.group(1)
                break

    gates = parse_gates(ver)
    impacts = sorted(glob.glob(os.path.join(ver, "CR-*_Impact.xlsx")))
    baseline_ref = meta.get("previous_version") or DASH
    if impacts:
        vtype = "CR"
        outs, summ = cr_outputs(ver, impacts[0], gates, skips)
        baseline_ref = summ.get("baseline_version") or baseline_ref
        base_inv = os.path.join(os.path.dirname(ver), baseline_ref, "_internal", "inventory.xlsx")
        if sysname == S.UNKNOWN and os.path.isfile(base_inv):
            sysname = S.meta(S.load(base_inv)).get("system_name") or S.UNKNOWN
    else:
        vtype = meta.get("version_type") if meta.get("version_type") in S.VERSION_TYPE else "BASELINE"
        outs = baseline_outputs(ver, data, meta, sysname, n, gates, skips)

    ig, igd = gate_of(gates, "INV")
    index = {"system": sysname, "version": "ver%s" % n, "version_folder": folder,
             "version_type": vtype, "date": date,
             "baseline_ref": baseline_ref if vtype == "CR" else DASH,
             "outputs": outs, "inventory_gate": {"gate": ig, "detail": igd}}

    os.makedirs(os.path.join(ver, "_internal"), exist_ok=True)
    with open(os.path.join(ver, "_internal", "index.json"), "w", encoding="utf8") as fh:
        json.dump(index, fh, ensure_ascii=False, indent=2)

    lines = ["# %s — ver%s (%s)" % (sysname, n, vtype), "",
             "- Thư mục: `%s`" % folder, "- Ngày: %s" % date, "- Loại: %s" % vtype]
    if vtype == "CR":
        lines.append("- Baseline tham chiếu: %s" % baseline_ref)
    lines += ["", "| # | Output | File | Số lượng | Gate | Trạng thái |",
              "|---|---|---|---|---|---|"]
    for e in outs:
        lines.append("| %s | %s | %s | %s | %s | %s |" % (
            e["id"], e["name"], "<br>".join("`%s`" % f for f in e["files"]) or DASH,
            e["count"], e["gate"], e["status_text"]))
    lines += ["", "Kiểm chứng dữ liệu nguồn (V1/V2): %s" % ig, "",
              "## Cách dùng khi có yêu cầu mới", "",
              "Chạy `/change-request <file | link | mô tả>` — agent đọc baseline mới nhất "
              "và sinh thư mục version mới (không sửa thư mục này).", ""]
    with open(os.path.join(ver, "README.md"), "w", encoding="utf8") as fh:
        fh.write("\n".join(lines))

    print("Da ghi README.md + _internal/index.json — %d output · %s" % (
        len(outs), " · ".join("%s=%s" % (e["id"] + (":" + e["site"] if e.get("site") else ""),
                                         e["status"]) for e in outs)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
