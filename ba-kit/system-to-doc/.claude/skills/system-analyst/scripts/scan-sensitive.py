#!/usr/bin/env python3
"""S1 — quet secret / du lieu that TRUOC KHI agent doc nguon (repo, dump, file input, CR).

  python3 scan-sensitive.py --path <repo|dump|inputs> [--path ...] --out <_internal>/gates/sensitive.md
          [--json <file>] [--threshold 5] [--max-mb 2] [--patterns <pii-patterns.json>]

Exit: 0 = sach (chi LOW/INFO) · 3 = co HIGH -> agent DUNG, bao user, hoi cach xu ly (POLICIES §5 muc 1)
      · 1 = sai tham so / loi.
Pattern dung chung .claude/config/pii-patterns.json (denyPatterns + allowValues + allowPathPatterns).
KHONG BAO GIO in gia tri: chi file#Lnn + pattern id + mau da che (ab****yz).
File .env / key / keystore... chi bao TEN, khong mo.
"""
import argparse
import importlib.util
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_PATTERNS = os.path.normpath(os.path.join(HERE, "..", "..", "..", "config", "pii-patterns.json"))

SKIP_DIRS = {".git", "node_modules", "vendor", "dist", "build", ".next", ".nuxt", ".svelte-kit",
             "target", "coverage", "__pycache__", ".venv", "venv", ".idea", ".gradle", "bower_components",
             "Pods", ".terraform", ".cache", ".turbo"}
ENV_OK = {".env.example", ".env.sample", ".env.template"}
KEY_EXT = (".pem", ".key", ".p8", ".p12", ".pfx", ".cer", ".keystore", ".jks", ".mobileprovision",
           ".provisionprofile")
DB_BIN_EXT = (".db", ".sqlite", ".sqlite3", ".mdb", ".accdb")
SQL_EXT = (".sql", ".dump", ".psql", ".bak")
CRED_NAMES = {".npmrc", ".yarnrc", ".yarnrc.yml", ".netrc", ".gitconfig", ".pgpass", ".htpasswd",
              "google-services.json", "googleservice-info.plist", "key.properties", "test-users.json",
              "credentials", "credentials.json", "known_hosts"}
NAME_WORDS = ("secret", "credential", "token", "password", "passwd", "apikey", "api_key",
              "service-account", "service_account", "serviceaccount", "adminsdk", "private_key",
              "privatekey")  # = scan-repo.py SENSITIVE_WORDS
SOURCE_EXT = (".ts", ".tsx", ".js", ".jsx", ".mjs", ".cjs", ".vue", ".svelte", ".py", ".php", ".rb",
              ".java", ".kt", ".go", ".cs", ".swift", ".dart", ".scala", ".rs", ".css", ".scss", ".html")
DATA_EXT = (".csv", ".tsv", ".json", ".jsonl", ".ndjson", ".xml", ".txt", ".log", ".xlsx", ".xlsm",
            ".yaml", ".yml") + SQL_EXT
PARTIAL_OK = (".sql", ".dump", ".csv", ".tsv", ".json", ".jsonl", ".ndjson", ".txt", ".log", ".xml")

PLACEHOLDER = re.compile(r"(?i)^(?:x{3,}|\*{3,}|\.{3,}|<.*>|\$\{.*\}|\{\{.*\}\}|%\(?\w*\)?s|"
                         r".*(?:change.?me|example|placeholder|your[_-]|dummy|sample|fake|todo|"
                         r"replace|redacted|secret_here|password_here|test(?:ing)?|null|none|undefined|"
                         r"true|false).*)$")
ENV_LOOKUP = re.compile(r"(?i)process\.env|import\.meta\.env|os\.(?:environ|getenv)|getenv\(|ENV\[|"
                        r"env\(|config\(|System\.getenv|@Value|Deno\.env|settings\.|\$_ENV|\$_SERVER")
CRED_KV = re.compile(r"""(?i)\b([\w.-]*(?:password|passwd|pwd|secret|api[_-]?key|access[_-]?key|"""
                     r"""client[_-]?secret|auth[_-]?token|private[_-]?key|token))\b["']?\s*(?:=>|:=|[:=])\s*"""
                     r"""(["'])([^"'\n]{4,})\2""")
INSERT_RX = re.compile(r"(?im)^\s*INSERT\s+INTO\b")
COPY_RX = re.compile(r"(?im)^\s*COPY\s+\S+.*\bFROM\s+stdin\b")
SA_RX = re.compile(r'"type"\s*:\s*"service_account"')
PII_IDS = ("email", "phone_vn", "phone_jp", "credit_card")


def _load_stmt_splitter():
    """Dung lai tach statement cua read-schema.py (bo qua than TRIGGER/PROCEDURE/FUNCTION
    nho DELIMITER va $$...$$). Khong load duoc -> None, roi ve INSERT_RX theo dong."""
    try:
        sys.path.insert(0, HERE)
        spec = importlib.util.spec_from_file_location("read_schema", os.path.join(HERE, "read-schema.py"))
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod.scan_statements, mod.DATA_STMT
    except Exception:
        return None


STMT = _load_stmt_splitter()


def mask(s):
    t = str(s).strip()
    if len(t) <= 4:
        return "*" * len(t)
    return t[:2] + "*" * min(len(t) - 4, 12) + t[-2:]


def load_cfg(path):
    try:
        with open(path, encoding="utf8") as fh:
            cfg = json.load(fh)
    except (OSError, ValueError) as e:
        print("Khong doc duoc pattern %s (%s) -> dung bo toi thieu" % (path, e), file=sys.stderr)
        cfg = {"denyPatterns": [
            {"id": "email", "re": r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"},
            {"id": "private_key", "re": r"-----BEGIN (?:RSA |EC |OPENSSH |PGP )?PRIVATE KEY-----"},
            {"id": "aws_key", "re": r"\b(?:AKIA|ASIA)[0-9A-Z]{16}\b"},
            {"id": "jwt", "re": r"eyJ[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}"},
            {"id": "conn_string", "re": r"(?:mysql|postgres(?:ql)?|mongodb(?:\+srv)?|redis|amqp)://[^\s:@/]+:[^\s:@/]+@"}],
            "allowValues": [], "allowPathPatterns": []}
    deny = []
    for p in cfg.get("denyPatterns", []):
        try:
            deny.append((p["id"], re.compile(p["re"], re.M)))
        except re.error as e:
            print("Bo qua pattern %s: %s" % (p.get("id"), e), file=sys.stderr)
    allow_paths = []
    for r in cfg.get("allowPathPatterns", []):
        try:
            allow_paths.append(re.compile(r))
        except re.error:
            pass
    return deny, sorted(cfg.get("allowValues", []), key=len, reverse=True), allow_paths


def file_reason(rel):
    """Ly do file nhay cam theo TEN (khong mo). Tra ve (level, id) hoac None."""
    name = os.path.basename(rel)
    low = name.lower()
    if low in ENV_OK:
        return None
    if low == ".env" or low.startswith(".env.") or low.endswith(".env"):
        return "HIGH", "env-file"
    if low.startswith(("id_rsa", "id_ed25519", "id_dsa", "id_ecdsa")) and not low.endswith(".pub"):
        return "HIGH", "ssh-private-key-file"
    if low.endswith(KEY_EXT):
        return "HIGH", "key/cert-file"
    if low.endswith(DB_BIN_EXT):
        return "HIGH", "db-file"
    if low in CRED_NAMES:
        return "HIGH", "credential-config-file"
    p = "/" + rel.replace(os.sep, "/")
    if "/.auth/" in p:
        return "HIGH", "auth-state-file"
    # dong bo scan-repo.py: ten chua secret/token/... -> khong doc, ke ca source code
    if any(w in low for w in NAME_WORDS) and not low.endswith(SQL_EXT):
        return "HIGH", "name-keyword-file"
    return None


def luhn(num):
    d = [int(c) for c in re.sub(r"\D", "", num)][::-1]
    if len(d) < 13:
        return False
    s = sum(d[0::2]) + sum(sum(divmod(2 * x, 10)) for x in d[1::2])
    return s % 10 == 0


class Scan:
    def __init__(self, deny, allow_vals, allow_paths, threshold):
        self.deny, self.allow_vals, self.allow_paths = deny, allow_vals, allow_paths
        self.threshold = threshold
        self.findings = []
        self.skipped = []
        self.files = 0
        self.multi = False

    def add(self, level, pid, where, sample="", note=""):
        self.findings.append({"level": level, "pattern": pid, "where": where,
                              "sample": mask(sample) if sample else "", "note": note})

    def strip_allowed(self, text):
        for v in self.allow_vals:
            if v in text:
                text = text.replace(v, " ")
        return text

    def text_checks(self, rel, text, ext, partial):
        low = ext.lower()
        is_data = low.endswith(DATA_EXT)
        is_sql = low.endswith(SQL_EXT) or re.search(r"(?i)dump|backup|export", os.path.basename(rel)) and low.endswith((".sql", ".txt"))
        clean = self.strip_allowed(text)

        def line_of(pos):
            return clean.count("\n", 0, pos) + 1
        # SQL dump co dong du lieu
        if is_sql:
            if STMT:
                # chi dem statement bat dau bang INSERT/REPLACE/COPY (INSERT trong than trigger bi bo qua)
                data = [(ln, st) for ln, st in STMT[0](clean) if STMT[1].match(st)]
                n_copy = sum(1 for _, st in data if st[:4].upper() == "COPY")
                n_ins = len(data) - n_copy
                first = data[0][0] if data else 0
            else:
                n_ins = len(INSERT_RX.findall(clean))
                n_copy = len(COPY_RX.findall(clean))
                m = INSERT_RX.search(clean) or COPY_RX.search(clean)
                first = line_of(m.start()) if m else 0
            if n_ins or n_copy:
                self.add("HIGH", "sql-data-rows", "%s#L%d" % (rel, first), "",
                         "%d INSERT, %d COPY ... FROM stdin%s" % (n_ins, n_copy, " (doc mot phan)" if partial else ""))
            else:
                self.add("INFO", "sql-schema-only", rel, "", "khong thay INSERT/COPY")
        if SA_RX.search(clean) and '"private_key"' in clean:
            m = SA_RX.search(clean)
            self.add("HIGH", "gcp-service-account-json", "%s#L%d" % (rel, line_of(m.start())), "", "")
        pii = {k: set() for k in PII_IDS}
        pii_first = {}
        seen_lines = set()
        for pid, rx in self.deny:
            for m in rx.finditer(clean):
                v = m.group(0).strip()
                ln = line_of(m.start())
                if pid in PII_IDS:
                    if pid == "credit_card" and not luhn(v):
                        continue
                    pii[pid].add(v)
                    pii_first.setdefault(pid, (ln, v))
                    continue
                if pid == "password_kv" and (ENV_LOOKUP.search(v) or PLACEHOLDER.match(re.split(r"[:=]", v, 1)[-1].strip(" \"'"))):
                    continue
                key = (ln, pid)
                if key in seen_lines:
                    continue
                seen_lines.add(key)
                self.add("HIGH", pid, "%s#L%d" % (rel, ln), v, "")
        for m in CRED_KV.finditer(clean):
            value = m.group(3).strip()
            ln = line_of(m.start())
            if PLACEHOLDER.match(value) or ENV_LOOKUP.search(value) or value.lower() == m.group(1).lower() \
                    or re.match(r"^[A-Z0-9_]+$", value) and "_" in value or " " in value.strip():
                continue
            if any(k[0] == ln for k in seen_lines):
                continue
            seen_lines.add((ln, "hardcoded-credential"))
            self.add("HIGH", "hardcoded-credential", "%s#L%d" % (rel, ln), value, "key " + m.group(1))
        many = self.threshold if is_data else self.threshold * 3
        for pid, vals in pii.items():
            if not vals:
                continue
            ln, v = pii_first[pid]
            n = len(vals)
            if pid == "credit_card" and is_data:
                self.add("HIGH", "credit_card", "%s#L%d" % (rel, ln), v, "%d so the hop le Luhn" % n)
            elif n >= many:
                self.add("HIGH", "bulk-" + pid, "%s#L%d" % (rel, ln), v,
                         "%d gia tri khac nhau -> nghi du lieu that" % n)
            else:
                self.add("LOW", pid, "%s#L%d" % (rel, ln), v, "%d gia tri (thuong la email tac gia / vi du)" % n)

    def xlsx_text(self, fp, max_cells=200000):
        try:
            from openpyxl import load_workbook
        except ImportError:
            return None
        try:
            wb = load_workbook(fp, read_only=True, data_only=True)
        except Exception:
            return None
        out, n = [], 0
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                out.append("\t".join("" if c is None else str(c) for c in row))
                n += len(row)
                if n > max_cells:
                    return "\n".join(out)
        return "\n".join(out)

    def scan_file(self, fp, rel, max_bytes, prefix=""):
        if any(r.search(fp.replace(os.sep, "/")) for r in self.allow_paths):
            return
        why = file_reason(rel)
        low = rel.lower()
        rel = prefix + rel  # ten hien thi (kem thu muc goc khi quet nhieu --path)
        if why:
            level, pid = why
            self.add(level, pid, rel, "", "chi bao TEN, khong mo file")
            if level == "HIGH":
                return
        self.files += 1
        try:
            size = os.path.getsize(fp)
        except OSError:
            return
        if low.endswith((".xlsx", ".xlsm")):
            if size > max_bytes * 10:
                self.skipped.append((rel, "qua lon (%d MB)" % (size // 1048576)))
                return
            text = self.xlsx_text(fp)
            if text is None:
                self.skipped.append((rel, "khong doc duoc xlsx"))
                return
            self.text_checks(rel, text, low, False)
            return
        if low.endswith((".gz", ".tgz", ".zip", ".7z", ".rar", ".jar", ".docx", ".pptx", ".pdf")):
            self.skipped.append((rel, "file nen/nhi phan — khong quet"))
            return
        partial = False
        if size > max_bytes:
            if not low.endswith(PARTIAL_OK):
                self.skipped.append((rel, "qua lon (%.1f MB) — khong quet" % (size / 1048576.0)))
                return
            partial = True
        try:
            with open(fp, "rb") as fh:
                raw = fh.read(max_bytes)
        except OSError as e:
            self.skipped.append((rel, "khong doc duoc: %s" % e.__class__.__name__))
            return
        if b"\x00" in raw[:8192]:
            return  # nhi phan (anh, font...) — bo qua im lang
        text = raw.decode("utf8", errors="ignore")
        if partial:
            self.skipped.append((rel, "chi quet %d MB dau" % (max_bytes // 1048576)))
        self.text_checks(rel, text, low, partial)

    def scan_path(self, root, max_bytes):
        root = os.path.abspath(root)
        if os.path.isfile(root):
            self.scan_file(root, os.path.basename(root), max_bytes)
            return
        if os.path.isfile(os.path.join(root, ".git", "config")):
            self.add("INFO", "git-config-present", (os.path.basename(root) + "/" if self.multi else "") + ".git/config", "", "khong doc; co the chua token o remote URL")
        for dirpath, dirnames, filenames in os.walk(root):
            dirnames[:] = sorted(d for d in dirnames if d not in SKIP_DIRS)
            for fn in sorted(filenames):
                fp = os.path.join(dirpath, fn)
                if os.path.islink(fp):
                    continue
                rel = os.path.relpath(fp, root).replace(os.sep, "/")
                self.scan_file(fp, rel, max_bytes, os.path.basename(root) + "/" if self.multi else "")


def render(scan, paths):
    hi = [f for f in scan.findings if f["level"] == "HIGH"]
    lo = [f for f in scan.findings if f["level"] != "HIGH"]
    L = ["# S1 — Sensitive scan", "",
         "Nguon: %s" % ", ".join("`%s`" % p for p in paths), "",
         "**%d file quet · %d HIGH · %d LOW/INFO · %d bo qua**" % (scan.files, len(hi), len(lo), len(scan.skipped)), ""]
    if hi:
        L += ["## HIGH — DUNG doc nguon nay (POLICIES.md §5 muc 1)", "",
              "Bao user: file nao, loai gi (KHONG kem gia tri). Hoi cach xu ly: go file / dua ban "
              "schema-only hoac da lam sach / xac nhan la du lieu gia. Credential that -> huong dan rotate.", "",
              "| # | Vi tri | Pattern | Mau (da che) | Ghi chu |", "|---|---|---|---|---|"]
        for i, f in enumerate(hi, 1):
            L.append("| %d | `%s` | %s | %s | %s |" % (i, f["where"], f["pattern"], f["sample"] or "—", f["note"] or "—"))
        L.append("")
    if lo:
        L += ["## LOW / INFO", "", "| Vi tri | Muc | Pattern | Mau (da che) | Ghi chu |", "|---|---|---|---|---|"]
        for f in lo[:80]:
            L.append("| `%s` | %s | %s | %s | %s |" % (f["where"], f["level"], f["pattern"], f["sample"] or "—", f["note"] or "—"))
        if len(lo) > 80:
            L.append("| ... | | | | +%d dong |" % (len(lo) - 80))
        L.append("")
    if scan.skipped:
        L += ["## Khong quet / quet mot phan", ""] + ["- `%s` — %s" % s for s in scan.skipped[:100]] + [""]
    L.append("KET QUA: %s" % ("HIGH -> exit 3, DUNG" if hi else "sach (chi LOW/INFO) -> exit 0"))
    return "\n".join(L) + "\n", hi, lo


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--path", action="append", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--json", default="")
    ap.add_argument("--threshold", type=int, default=5, help="so gia tri PII khac nhau / file de coi la du lieu that")
    ap.add_argument("--max-mb", type=float, default=2.0)
    ap.add_argument("--patterns", default=DEFAULT_PATTERNS)
    a = ap.parse_args()
    missing = [p for p in a.path if not os.path.exists(p)]
    if missing:
        print("Khong thay: %s" % ", ".join(missing), file=sys.stderr)
        return 1
    deny, allow_vals, allow_paths = load_cfg(a.patterns)
    scan = Scan(deny, allow_vals, allow_paths, max(1, a.threshold))
    scan.multi = len(a.path) > 1
    max_bytes = int(a.max_mb * 1024 * 1024)
    for p in a.path:
        scan.scan_path(p, max_bytes)
    text, hi, lo = render(scan, a.path)
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    with open(a.out, "w", encoding="utf8") as fh:
        fh.write(text)
    if a.json:
        with open(a.json, "w", encoding="utf8") as fh:
            json.dump({"paths": a.path, "files_scanned": scan.files, "high": hi, "low": lo,
                       "skipped": [{"path": s[0], "reason": s[1]} for s in scan.skipped]}, fh, indent=2, ensure_ascii=False)
    print(text)
    print("SUMMARY: %d file · %d HIGH · %d LOW/INFO · %d bo qua -> exit %d"
          % (scan.files, len(hi), len(lo), len(scan.skipped), 3 if hi else 0), file=sys.stderr)
    return 3 if hi else 0


if __name__ == "__main__":
    sys.exit(main())
