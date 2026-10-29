#!/usr/bin/env python3
"""Quan ly version folder trong <project>/outputs/ (Luong 1 baseline + Luong 2 CR).

  python3 version-tool.py next --outputs <project>/outputs --slug baseline [--date DDMMYY] [--create]
  python3 version-tool.py next --outputs <project>/outputs --cr-id CR-001 --slug them-coupon [--date DDMMYY] [--create]
  python3 version-tool.py snapshot --ver <folder version> [--files 02_Screens/a.xlsx ...]
  python3 version-tool.py promote --outputs <project>/outputs --cr CR-001 [--cr CR-002] --approved --coded \
          [--slug baseline] [--date DDMMYY] [--create]
  python3 version-tool.py latest-baseline --outputs <project>/outputs
  python3 version-tool.py list --outputs <project>/outputs [--json]
  python3 version-tool.py context --outputs <project>/outputs [--flow 1|2] [--cr-id CR-001]

Quy tac (versioning.md, R5 · D3 · D4):
  - Khao sat lan dau luon la ver1. Chua co CR nao tro toi baseline hien hanh -> `next` (Luong 1) tra ve CHINH
    baseline do (reuse:true) — chay lai = ghi vao ver1; truoc khi ghi de goi `snapshot` (luu file bi thay vao
    <ver>/_internal/history/<YYYYMMDD-HHMM>/).
  - Co folder CR tro toi baseline -> baseline KHOA: `next` Luong 1 exit 3 (locked). Baseline moi ver<N+1> chi tao
    bang `promote --cr ... --approved --coded` (CR da duyet + da code xong).
  - Folder CR moi: CR-<id>_<DDMMYY>_<slug>, _internal/version-meta.json ghi baseline goc. Chay lai cung CR ->
    tra ve folder CR do (reuse:true). Doc duoc layout cu: ver<N>_<DDMMYY>_<slug> va ver<N>_<DDMMYY>_CR-<id>-<slug>.

next            -> JSON {n, folder, path, reuse, created} ; --create mkdir _internal/{gates,recon,evidence} +
                   version-meta.json (folder co san: chi bo sung thu muc con thieu). exit 3 = baseline da khoa
snapshot        -> JSON {dir, files} ; copy file (mac dinh: moi file tru _internal/history) vao history truoc khi ghi de
promote         -> JSON {n, folder, path, previous_version, promoted_crs} — baseline moi
latest-baseline -> JSON {folder, path, n, inventory, outputs:{O1..O8: bool}} ; exit 1 neu khong co baseline
list            -> bang version + CR mo tren baseline moi nhat (canh bao xung dot CR-vs-CR)
context         -> JSON agent doc o DAU MOI LAN CHAY (user khong phai tu tim version):
                   latest (version moi nhat bat ky loai) · latest_baseline (+ inventory, outputs, locked, locked_by) ·
                   open_crs (CR tro toi baseline moi nhat, chua promote) · latest_by_cr · pending_feedback (muc
                   "Feedback cua user" trong run-log chua co version nao ghi "Feedback da xu ly") ·
                   suggest (flow1_action / mode de xuat cho Luong 1 / version doc lai cho Luong 2) · read (file can doc)
Folder CR cu va baseline da khoa: KHONG sua noi dung.
"""
import argparse
import datetime
import glob
import json
import os
import re
import shutil
import sys
import unicodedata

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import inv_schema as S  # noqa: E402

VER_RE = re.compile(r"^ver(\d+)_(\d{6})_(.+)$")
CR_DIR_RE = re.compile(r"^(CR-\d+)_(\d{6})_(.+)$", re.I)       # layout moi (D3)
CR_SLUG_RE = re.compile(r"^(CR-\d+)(?:-(.*))?$", re.I)          # layout cu: ver<N>_<date>_CR-<id>-<slug>
META_FILE = os.path.join("_internal", "version-meta.json")
HISTORY_DIR = os.path.join("_internal", "history")
EXIT_LOCKED = 3

# Duong dan output theo CONTRACT (glob tuong doi folder version)
OUTPUT_GLOBS = {
    "O1": ["01_Overview/Overview_*.docx"],
    "O2": ["02_Screens/BasicDesign_*.xlsx"],
    "O3": ["03_API/API_Doc_*.xlsx"],
    "O4": ["04_DB/DB_Doc_*.xlsx"],
    "O5": ["05_Figma/figma-links.md"],
    "O6": ["06_DesignSystem/project/design-system.json",            # 1 DS (single / themes)
           "06_DesignSystem/WEB-*/project/design-system.json"],     # 1 DS / website (per-site)
    "O7": ["07_BugList/BugList_*.xlsx"],
    "O8": ["08_DevSetup/DevSetup_*.md"],
}


def slugify(text):
    t = (text or "").replace("đ", "d").replace("Đ", "D")
    t = unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode("ascii")
    t = re.sub(r"[^a-z0-9]+", "-", t.lower()).strip("-")
    return re.sub(r"-{2,}", "-", t)


def norm_cr_id(cr_id):
    c = (cr_id or "").strip().upper()
    m = re.match(r"^(?:CR-?)?(\d+)$", c)
    if not m:
        raise SystemExit("--cr-id khong hop le: %r (vd CR-001)" % cr_id)
    return "CR-%03d" % int(m.group(1))


def read_kv(xlsx, sheet):
    """Doc sheet Key/Value -> dict (khong yeu cau du sheet nhu S.load)."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("Thieu openpyxl. Chay: pip install openpyxl", file=sys.stderr)
        sys.exit(2)
    try:
        wb = load_workbook(xlsx, read_only=True, data_only=True)
    except Exception:
        return {}
    if sheet not in wb.sheetnames:
        return {}
    out = {}
    for r in wb[sheet].iter_rows(min_row=2, values_only=True):
        if r and r[0] is not None:
            out[str(r[0]).strip()] = "" if len(r) < 2 or r[1] is None else str(r[1]).strip()
    wb.close()
    return out


def read_meta(path):
    f = os.path.join(path, META_FILE)
    try:
        return json.load(open(f, encoding="utf8")) if os.path.isfile(f) else {}
    except (OSError, ValueError):
        return {}


def write_meta(path, meta):
    with open(os.path.join(path, META_FILE), "w", encoding="utf8") as fh:
        json.dump(meta, fh, ensure_ascii=False, indent=2)


def scan(outputs):
    """Liet ke version folder (ca layout cu lan moi) -> list dict, sap theo thu tu thoi gian logic.

    n = so baseline (folder CR layout moi: n cua baseline no tro toi) — CR sap ngay sau baseline cua no.
    """
    if not os.path.isdir(outputs):
        return []
    vers = []
    for name in sorted(os.listdir(outputs)):
        p = os.path.join(outputs, name)
        m, mn = VER_RE.match(name), CR_DIR_RE.match(name)
        if not ((m or mn) and os.path.isdir(p)):
            continue
        layout = "ver" if m else "cr"
        n, date, slug = (int(m.group(1)), m.group(2), m.group(3)) if m else (None, mn.group(2), mn.group(3))
        vm = read_meta(p)
        inv = os.path.join(p, "_internal", "inventory.xlsx")
        vtype, cr_id, base_ref = (vm.get("version_type") or "").upper(), vm.get("cr_id") or "", vm.get("baseline") or ""
        kv = read_kv(inv, "00_Meta") if os.path.isfile(inv) else {}
        if vtype not in S.VERSION_TYPE:
            vtype = kv.get("version_type", "").upper()
        base_ref = base_ref or kv.get("baseline") or kv.get("baseline_version") or ""
        rl = os.path.join(p, "run-log.md")
        txt = open(rl, encoding="utf8", errors="ignore").read() if os.path.isfile(rl) else ""
        if vtype not in S.VERSION_TYPE:
            mm = re.search(r"version_type\W{0,6}(BASELINE|CR)\b", txt)
            vtype = mm.group(1) if mm else ""
        impacts = sorted(glob.glob(os.path.join(p, "CR-*_Impact.xlsx")))
        mc = CR_SLUG_RE.match(slug) if m else None
        if mc:
            cr_id = cr_id or mc.group(1).upper()
        if mn:
            cr_id = cr_id or mn.group(1).upper()
        cr_json = os.path.join(p, "_internal", "cr.json")
        if os.path.isfile(cr_json):
            try:
                cm = json.load(open(cr_json, encoding="utf8")).get("meta", {})
            except (OSError, ValueError):
                cm = {}
            cr_id = cm.get("cr_id") or cr_id
            base_ref = cm.get("baseline_version") or base_ref
        if not base_ref:
            mb = re.search(r"^baseline:\s*`?((?:ver\d+|CR-\d+)_\d{6}_[^\s`|]+)", txt, re.M)
            base_ref = mb.group(1) if mb else ""
        if impacts and not cr_id:
            cr_id = os.path.basename(impacts[0]).split("_")[0]
        if mn or impacts or mc:
            vtype = "CR"
        if not vtype:
            vtype = "BASELINE"   # folder ver<N> khong co dau hieu CR = baseline (quy tac moi)
        vers.append({"n": n, "folder": name, "path": os.path.abspath(p), "layout": layout,
                     "date": date, "type": vtype, "cr_id": cr_id.upper(),
                     "baseline_ref": base_ref, "promoted_crs": [c.upper() for c in vm.get("promoted_crs") or []],
                     "impact": os.path.abspath(impacts[0]) if impacts else ""})
    # CR layout moi: n = n cua baseline goc (khong ro -> sau cung)
    by_folder = {v["folder"]: v for v in vers}
    top = max([v["n"] for v in vers if v["n"] is not None] or [0])
    for v in vers:
        if v["n"] is None:
            b = by_folder.get(v["baseline_ref"])
            v["n"] = b["n"] if b and b["n"] is not None else top
    ddmmyy = lambda d: d[4:6] + d[2:4] + d[0:2]
    vers.sort(key=lambda v: (v["n"], 0 if v["type"] == "BASELINE" else 1, ddmmyy(v["date"]), v["folder"]))
    return vers


def promoted(vers):
    """CR da duoc promote vao baseline nao do."""
    return {c for v in vers if v["type"] == "BASELINE" for c in v["promoted_crs"]}


def crs_on(vers, b):
    """CR tro toi baseline b (khoa b). Layout cu khong ghi baseline -> CR co n > b.n."""
    if not b:
        return []
    out = []
    for v in vers:
        if v["type"] != "CR":
            continue
        if v["baseline_ref"]:
            if v["baseline_ref"] == b["folder"]:
                out.append(v)
        elif v["layout"] == "ver" and v["n"] > b["n"]:
            out.append(v)
    return out


FB_HEAD = re.compile(r"^##\s*Feedback c\w* user\s*$", re.I | re.M)
FB_DONE_HEAD = re.compile(r"^##\s*Feedback \S+ x\S+ l\S+\s*$", re.I | re.M)  # "## Feedback đã xử lý"
PLACEHOLDER = re.compile(r"^\s*(\(.*\)|<!--.*-->|-|—)?\s*$")


def _section(txt, head_re):
    m = head_re.search(txt)
    if not m:
        return ""
    rest = txt[m.end():]
    nxt = re.search(r"^##\s", rest, re.M)
    return rest[:nxt.start()] if nxt else rest


def read_runlog(path):
    """-> (feedback text user ghi, set folder da xu ly feedback)."""
    rl = os.path.join(path, "run-log.md")
    if not os.path.isfile(rl):
        return "", set()
    txt = open(rl, encoding="utf8", errors="ignore").read()
    fb = "\n".join(l for l in _section(txt, FB_HEAD).splitlines() if not PLACEHOLDER.match(l)).strip()
    done = set(re.findall(r"(?:ver\d+|CR-\d+)_\d{6}_[^\s|`)\]]+", _section(txt, FB_DONE_HEAD)))
    return fb, done


def outputs_exist(path):
    res = {}
    for k, pats in OUTPUT_GLOBS.items():
        res[k] = any(glob.glob(os.path.join(path, g)) for g in pats)
    return res


def _date_slug(a):
    date = a.date or datetime.date.today().strftime("%d%m%y")
    if not re.match(r"^\d{6}$", date):
        raise SystemExit("--date phai dang DDMMYY")
    slug = slugify(a.slug)
    if not slug:
        raise SystemExit("--slug rong sau khi chuan hoa")
    return date, slug


def _make(path, meta):
    """mkdir _internal/{gates,recon,evidence} (folder co san: chi bo sung) + version-meta.json neu chua co."""
    for sub in ("gates", "recon", "evidence"):
        os.makedirs(os.path.join(path, "_internal", sub), exist_ok=True)
    if not os.path.isfile(os.path.join(path, META_FILE)):
        write_meta(path, meta)


def cmd_next(a):
    vers = scan(a.outputs)
    date, slug = _date_slug(a)
    b = latest_baseline(vers)
    now = datetime.datetime.now().isoformat(timespec="seconds")
    if a.cr_id:
        cid = norm_cr_id(a.cr_id)
        if not b:
            print("Chua co baseline trong %s — chay Luong 1 truoc." % a.outputs, file=sys.stderr)
            return 1
        same = [v for v in vers if v["type"] == "CR" and v["layout"] == "cr" and v["cr_id"] == cid]
        if same:   # chay lai cung CR -> ghi vao folder CR do (snapshot truoc khi ghi de)
            v = same[-1]
            folder, path, reuse, base = v["folder"], v["path"], True, v["baseline_ref"] or b["folder"]
        else:
            folder = "%s_%s_%s" % (cid, date, slug)
            path, reuse, base = os.path.abspath(os.path.join(a.outputs, folder)), False, b["folder"]
        if a.create:
            _make(path, {"version_type": "CR", "cr_id": cid, "baseline": base, "created_at": now})
        print(json.dumps({"n": b["n"], "folder": folder, "path": path, "type": "CR", "cr_id": cid,
                          "baseline": base, "reuse": reuse, "created": bool(a.create),
                          **({"snapshot": "version-tool.py snapshot --ver %s" % path} if reuse else {})},
                         ensure_ascii=False))
        return 0
    if b:
        lock = crs_on(vers, b)
        if lock:
            print(json.dumps({"locked": True, "folder": b["folder"], "path": b["path"],
                              "locked_by": [v["cr_id"] or v["folder"] for v in lock],
                              "hint": ("%s da khoa (co CR tro toi). Khong ghi de. CR da duyet + da code xong -> "
                                       "version-tool.py promote --outputs %s --cr <CR-id> --approved --coded --create; "
                                       "chua co CR nao code xong -> chi doc lai %s.") % (b["folder"], a.outputs, b["folder"])},
                             ensure_ascii=False))
            return EXIT_LOCKED
        if a.create:
            _make(b["path"], {"version_type": "BASELINE", "created_at": now})
        print(json.dumps({"n": b["n"], "folder": b["folder"], "path": b["path"], "type": "BASELINE",
                          "reuse": True, "created": bool(a.create),
                          "snapshot": "version-tool.py snapshot --ver %s" % b["path"]}, ensure_ascii=False))
        return 0
    folder = "ver1_%s_%s" % (date, slug)
    path = os.path.abspath(os.path.join(a.outputs, folder))
    if a.create:
        _make(path, {"version_type": "BASELINE", "created_at": now})
    print(json.dumps({"n": 1, "folder": folder, "path": path, "type": "BASELINE", "reuse": False,
                      "created": bool(a.create)}, ensure_ascii=False))
    return 0


def cmd_snapshot(a):
    """Luu file sap bi ghi de vao <ver>/_internal/history/<YYYYMMDD-HHMM>/ (giu cau truc thu muc)."""
    ver = os.path.abspath(a.ver)
    if not os.path.isdir(ver):
        raise SystemExit("Khong thay folder version: %s" % ver)
    hist_root = os.path.join(ver, HISTORY_DIR)
    if a.files:
        # chap nhan ca duong dan tuyet doi nam trong folder version; ngoai folder -> bo (loc '..' ben duoi)
        rels = [os.path.relpath(os.path.abspath(f), ver) if os.path.isabs(f) else os.path.normpath(f) for f in a.files]
    else:
        rels = []
        for root, dirs, files in os.walk(ver):
            if os.path.abspath(root).startswith(hist_root):
                dirs[:] = []
                continue
            rels += [os.path.relpath(os.path.join(root, f), ver) for f in files]
    rels = [r for r in rels if not r.startswith(("..", HISTORY_DIR)) and os.path.isfile(os.path.join(ver, r))]
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M")
    dst, k = os.path.join(hist_root, stamp), 2
    while os.path.exists(dst):
        dst, k = os.path.join(hist_root, "%s-%d" % (stamp, k)), k + 1
    for r in rels:
        os.makedirs(os.path.dirname(os.path.join(dst, r)), exist_ok=True)
        shutil.copy2(os.path.join(ver, r), os.path.join(dst, r))
    if rels:
        with open(os.path.join(dst, "SNAPSHOT.md"), "w", encoding="utf8") as fh:
            fh.write("# Snapshot %s\n\n%d file truoc khi ghi de:\n\n" % (os.path.basename(dst), len(rels)) +
                     "".join("- %s\n" % r.replace(os.sep, "/") for r in sorted(rels)))
    print(json.dumps({"dir": dst if rels else None, "files": len(rels)}, ensure_ascii=False))
    return 0


def cmd_promote(a):
    """Baseline moi ver<N+1> = baseline hien hanh + CR da duyet + da code xong (lenh ro rang)."""
    if not (a.approved and a.coded):
        print("Tu choi: chi tao baseline moi khi CR da duyet VA da code xong — them --approved --coded "
              "sau khi user xac nhan.", file=sys.stderr)
        return 1
    vers = scan(a.outputs)
    b = latest_baseline(vers)
    if not b:
        print("Chua co baseline trong %s." % a.outputs, file=sys.stderr)
        return 1
    on_b = {v["cr_id"]: v for v in crs_on(vers, b) if v["cr_id"]}
    cids = [norm_cr_id(c) for c in a.cr]
    miss = [c for c in cids if c not in on_b]
    if miss:
        print("CR %s khong tro toi baseline hien hanh %s (CR tren baseline nay: %s)."
              % (", ".join(miss), b["folder"], ", ".join(sorted(on_b)) or "khong co"), file=sys.stderr)
        return 1
    date, slug = _date_slug(a)
    n = b["n"] + 1
    folder = "ver%d_%s_%s" % (n, date, slug)
    path = os.path.abspath(os.path.join(a.outputs, folder))
    if a.create:
        if os.path.exists(path):
            raise SystemExit("Da ton tai, khong ghi de: %s" % path)
        _make(path, {"version_type": "BASELINE", "previous_version": b["folder"], "promoted_crs": cids,
                     "approved": True, "coded": True,
                     "created_at": datetime.datetime.now().isoformat(timespec="seconds")})
    print(json.dumps({"n": n, "folder": folder, "path": path, "type": "BASELINE", "previous_version": b["folder"],
                      "promoted_crs": cids, "created": bool(a.create),
                      "next_step": "Luong 1 mode DELTA ghi vao %s (carry tu %s + phan CR cham)" % (folder, b["folder"])},
                     ensure_ascii=False))
    return 0


def latest_baseline(vers):
    bases = [v for v in vers if v["type"] == "BASELINE"]
    return bases[-1] if bases else None


def cmd_latest(a):
    b = latest_baseline(scan(a.outputs))
    if not b:
        print("Khong tim thay version BASELINE nao trong %s "
              "(00_Meta.version_type / run-log 'version_type: BASELINE'). "
              "Chay Luong 1 truoc." % a.outputs, file=sys.stderr)
        return 1
    inv = os.path.join(b["path"], "_internal", "inventory.xlsx")
    print(json.dumps({"folder": b["folder"], "path": b["path"], "n": b["n"],
                      "inventory": inv, "inventory_exists": os.path.isfile(inv),
                      "outputs": outputs_exist(b["path"])}, ensure_ascii=False, indent=2))
    return 0


def cmd_list(a):
    vers = scan(a.outputs)
    b = latest_baseline(vers)
    done = promoted(vers)
    open_crs = [v for v in crs_on(vers, b) if v["cr_id"] not in done]
    if a.json:
        print(json.dumps({"versions": vers, "latest_baseline": b and b["folder"],
                          "open_crs": [v["folder"] for v in open_crs]},
                         ensure_ascii=False, indent=2))
        return 0
    print("| N | Folder | Type | Date | CR ID | Baseline ref |")
    print("|---|---|---|---|---|---|")
    for v in vers:
        print("| %d | %s | %s | %s | %s | %s |" % (
            v["n"], v["folder"], v["type"], v["date"], v["cr_id"] or "—",
            v["baseline_ref"] or "—"))
    print("")
    print("Latest baseline: %s" % (b["folder"] if b else "(khong co)"))
    if open_crs:
        print("Open CRs on latest baseline (%d) — kiem tra xung dot CR-vs-CR "
              "(verify-cr-impact.py --other-cr):" % len(open_crs))
        for v in open_crs:
            warn = ""
            if v["baseline_ref"] and b and v["baseline_ref"] != b["folder"]:
                warn = "  [CANH BAO: tham chieu baseline %s]" % v["baseline_ref"]
            print("  - %s (%s) %s%s" % (v["folder"], v["cr_id"] or "?", v["impact"] or "", warn))
    else:
        print("Open CRs on latest baseline: 0")
    return 0


def cmd_context(a):
    vers = scan(a.outputs)
    b = latest_baseline(vers)
    by_cr = {}
    for v in vers:
        if v["type"] == "CR" and v["cr_id"]:
            by_cr[v["cr_id"].upper()] = v["folder"]
    # CR tren baseline moi nhat, chua promote — moi CR chi lay version MOI NHAT (CR chay lai khong dem trung)
    done = promoted(vers)
    lock = crs_on(vers, b)
    open_crs = [v for v in lock if v["cr_id"] not in done
                and (not v["cr_id"] or by_cr.get(v["cr_id"].upper()) == v["folder"])]
    logs = {v["folder"]: read_runlog(v["path"]) for v in vers}
    pending = []
    for i, v in enumerate(vers):
        fb = logs[v["folder"]][0]
        if not fb:
            continue
        # chay lai tai cho (ver1) -> chinh run-log do co the ghi "Feedback da xu ly"
        handled = any(v["folder"] in logs[w["folder"]][1] for w in vers[i:])
        if not handled:
            pending.append({"folder": v["folder"], "type": v["type"], "cr_id": v["cr_id"], "feedback": fb})
    ctx = {
        "outputs": os.path.abspath(a.outputs),
        "versions": len(vers),
        "latest": vers[-1]["folder"] if vers else None,
        "latest_type": vers[-1]["type"] if vers else None,
        "latest_baseline": None,
        "open_crs": [{"folder": v["folder"], "cr_id": v["cr_id"], "impact": v["impact"],
                      "baseline_ref": v["baseline_ref"]} for v in open_crs],
        "latest_by_cr": by_cr,
        "pending_feedback": pending,
        "next_n": max([v["n"] for v in vers] or [0]) + 1,
    }
    if b:
        inv = os.path.join(b["path"], "_internal", "inventory.xlsx")
        ctx["latest_baseline"] = {"folder": b["folder"], "path": b["path"], "n": b["n"],
                                  "inventory": inv, "inventory_exists": os.path.isfile(inv),
                                  "outputs": outputs_exist(b["path"]),
                                  "locked": bool(lock), "locked_by": [v["cr_id"] or v["folder"] for v in lock],
                                  "promoted_crs": b["promoted_crs"],
                                  "history": sorted(os.listdir(os.path.join(b["path"], HISTORY_DIR)))
                                  if os.path.isdir(os.path.join(b["path"], HISTORY_DIR)) else []}
    flow = str(a.flow or "")
    read, suggest = [], {}
    if b:
        read += [os.path.join(b["path"], "_internal", "inventory.xlsx"), os.path.join(b["path"], "run-log.md")]
    if flow != "2":
        # flow1_action: create_v1 (chua co gi) · rerun_in_place (ghi vao baseline, snapshot truoc) ·
        #               locked (baseline da khoa — chi promote khi CR da duyet + code xong, hoac chi doc lai)
        if not b:
            suggest["flow1_action"] = "create_v1"
        elif lock:
            suggest["flow1_action"] = "locked"
            suggest["locked_hint"] = ("%s da khoa boi %s. CR nao da duyet + da code xong -> promote; "
                                      "chua co -> chi doc lai, khong ghi." % (
                                          b["folder"], ", ".join(v["cr_id"] or v["folder"] for v in lock)))
        else:
            suggest["flow1_action"] = "rerun_in_place"
            suggest["flow1_target"] = b["folder"]
        if not b:
            suggest["flow1_mode"] = "FULL"
            suggest["why"] = "chua co baseline"
        elif open_crs or pending:
            suggest["flow1_mode"] = "DELTA"
            why = []
            if open_crs:
                why.append("%d CR tren baseline (%s) — CR nao da duyet + code xong thi promote roi Delta cap nhat AS-IS"
                           % (len(open_crs), ", ".join(v["cr_id"] or v["folder"] for v in open_crs)))
            if pending:
                why.append("%d feedback chua xu ly" % len(pending))
            suggest["why"] = " · ".join(why)
        else:
            suggest["flow1_mode"] = "DELTA"
            suggest["why"] = "da co baseline, khong co CR / feedback moi"
    if flow != "1":
        cid = norm_cr_id(a.cr_id) if a.cr_id else None
        if cid and cid in by_cr:
            prev = next(v for v in vers if v["folder"] == by_cr[cid])
            suggest["flow2_rerun_of"] = prev["folder"]
            read += [os.path.join(prev["path"], "_internal", "cr.json"), os.path.join(prev["path"], "run-log.md")]
            if prev["baseline_ref"] and b and prev["baseline_ref"] != b["folder"]:
                suggest["flow2_warn"] = ("%s doi chieu %s nhung baseline moi nhat la %s — dung baseline moi nhat"
                                         % (prev["folder"], prev["baseline_ref"], b["folder"]))
        suggest["flow2_baseline"] = b["folder"] if b else None
    for p in pending:
        read.append(os.path.join(os.path.abspath(a.outputs), p["folder"], "run-log.md"))
    ctx["suggest"] = suggest
    ctx["read"] = [r for i, r in enumerate(read) if os.path.exists(r) and r not in read[:i]]
    print(json.dumps(ctx, ensure_ascii=False, indent=2))
    return 0


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("next")
    p.add_argument("--outputs", required=True)
    p.add_argument("--slug", required=True)
    p.add_argument("--cr-id", default=None)
    p.add_argument("--date", default=None)
    p.add_argument("--create", action="store_true")
    p = sub.add_parser("snapshot")
    p.add_argument("--ver", required=True)
    p.add_argument("--files", nargs="*", default=None, help="duong dan tuong doi folder version (mac dinh: moi file)")
    p = sub.add_parser("promote")
    p.add_argument("--outputs", required=True)
    p.add_argument("--cr", action="append", required=True)
    p.add_argument("--approved", action="store_true", help="user xac nhan CR da duyet")
    p.add_argument("--coded", action="store_true", help="user xac nhan CR da code xong / da len he thong")
    p.add_argument("--slug", default="baseline")
    p.add_argument("--date", default=None)
    p.add_argument("--create", action="store_true")
    p = sub.add_parser("latest-baseline")
    p.add_argument("--outputs", required=True)
    p = sub.add_parser("list")
    p.add_argument("--outputs", required=True)
    p.add_argument("--json", action="store_true")
    p = sub.add_parser("context")
    p.add_argument("--outputs", required=True)
    p.add_argument("--flow", choices=["1", "2"], default=None)
    p.add_argument("--cr-id", default=None)
    a = ap.parse_args()
    return {"next": cmd_next, "snapshot": cmd_snapshot, "promote": cmd_promote, "latest-baseline": cmd_latest,
            "list": cmd_list, "context": cmd_context}[a.cmd](a)


if __name__ == "__main__":
    sys.exit(main())
