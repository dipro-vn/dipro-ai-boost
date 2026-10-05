#!/usr/bin/env python3
"""Render O6 — Overview docx (tom tat khao sat) tu inventory + narrative + index.

  python3 render-overview-docx.py --inventory <_internal>/inventory.xlsx \
      --narrative <_internal>/narrative.json --index <_internal>/index.json \
      [--flow <_internal>/flow/flow.json] [--flow-png flow.png] [--codemap-png CodeMap.png] \
      [--erd-png ERD.png] [--template templates/high-level-template.docx] \
      --out <ver>/06_Overview/Overview_<sys>_ver<N>.docx

Chay build-version-index.py TRUOC (sinh index.json) va SAU (cap nhat gate O6).
BANG/SO LIEU -> may sinh tu inventory + index; VAN XUOI -> agent viet trong narrative.json:
{ "project_name": "...", "customer": "...", "product_overview": "3-5 cau",
  "actors": [{"actor": "User", "role": "...", "main_usage": "..."}],
  "system_composition": ["Frontend: ...", "Backend: ..."],
  "flow_explanation": ["1. ...", "2. ..."] }
Tai lieu ngan: danh sach day du nam o O1..O3 xlsx, o day chi dem.
"""
import argparse
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import inv_schema as S  # noqa: E402

TITLE = "TỔNG QUAN KHẢO SÁT HỆ THỐNG"
HEADER = "LABO / MAINTAIN PROJECT — SYSTEM OVERVIEW"
FOOTER = "AI-generated analysis must be reviewed before use"
DASH = "—"
NO_IMG = "Hình không có sẵn:"
DEFAULT_TPL = os.path.abspath(os.path.join(HERE, "..", "..", "..", "..",
                                           "templates", "high-level-template.docx"))
FK_EMPTY = ("", "—", "-", "No", "NO", "N/A", S.UNKNOWN)
SCREEN_TYPES = S.SCREEN_TYPE
API_STATUS = S.API_STATUS

CHAPTERS = [
    "0. Phạm vi & nguồn khảo sát", "1. Tổng quan sản phẩm", "2. Danh sách output đã sinh",
    "3. Website & màn hình", "4. API & Batch", "5. Database", "6. Liên kết bên thứ 3",
    "7. Luồng tổng quan", "Phụ lục A. Open Questions", "Phụ lục B. Điều kiện khảo sát",
]
APPENDIX_B_KEYS = ["scope", "websites", "accounts", "access_approved_by", "crawl_mode",
                   "forbidden_zones", "db_mode", "figma_input_url", "figma_output_url",
                   "bug_list", "run_mode", "env_observed", "crawl_budget_used",
                   "sensitive_scan", "lang", "audience", "previous_version"]


# ---------- so lieu dung chung (verify-overview.py import ham nay) ----------
def role_names(meta, sites):
    """Chi lay TEN vai tro — khong bao gio in user/password."""
    roles = []
    for tok in re.split(r"[;,]", meta.get("accounts", "") or ""):
        name = re.split(r"[:=/@\s(]", tok.strip())[0] if tok.strip() else ""
        if name and name != S.UNKNOWN and name not in roles:
            roles.append(name)
    for r in sites:
        for tok in re.split(r"[;,]", r.get("Roles Observed", "")):
            t = tok.strip()
            if t and t not in (S.UNKNOWN, DASH) and t not in roles:
                roles.append(t)
    return roles


def summary_counts(data):
    screens, funcs, apis = data["02_Screen"], data["01_Function"], data["07_API"]
    site_ids = [r.get("Site ID", "") for r in data["10_Site"]]
    site_ids += sorted({r.get("Site", "") or S.UNKNOWN for r in screens} - set(site_ids))
    sites = []
    for sid in site_ids:
        rows = [r for r in screens if (r.get("Site", "") or S.UNKNOWN) == sid]
        sites.append([sid, len(rows)] + [sum(1 for r in rows if r.get("Type") == t)
                                         for t in SCREEN_TYPES])
    modules = []
    for mod in sorted({r.get("Module", "") or S.UNKNOWN for r in funcs}):
        rows = [r for r in funcs if (r.get("Module", "") or S.UNKNOWN) == mod]
        modules.append([mod, len(rows)] + [sum(1 for r in rows if r.get("Status") == s)
                                           for s in S.FUNCTION_STATUS])
    groups = []
    for grp in sorted({r.get("Group", "") or S.UNKNOWN for r in apis}):
        rows = [r for r in apis if (r.get("Group", "") or S.UNKNOWN) == grp]
        k = [r.get("Kind") for r in rows]
        groups.append([grp, len(rows), k.count("API"), k.count("BATCH"),
                       k.count("WEBHOOK") + k.count("QUEUE")] +
                      [sum(1 for r in rows if r.get("Status") == s) for s in API_STATUS])
    cols = data["04_DB_Columns"]
    if cols:
        rels = sum(1 for c in cols if str(c.get("FK", "")).strip() not in FK_EMPTY)
    else:
        rels = sum(len(S.split_ids(t.get("FK", ""))) for t in data["03_DB_Tables"])
    db = [["Số bảng", len(data["03_DB_Tables"])], ["Số cột", len(cols)],
          ["Số quan hệ (FK)", rels]]
    return {"sites": sites, "modules": modules, "api_groups": groups, "db": db}


# ---------- docx helpers ----------
def cell_text(v):
    s = "" if v is None else str(v).strip()
    return s if s else DASH


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i in range(len(headers)):
            cells[i].text = cell_text(row[i] if i < len(row) else None)
    doc.add_paragraph()
    return t


def bullet(doc, text):
    try:
        doc.add_paragraph(text, style="List Bullet")
    except KeyError:
        doc.add_paragraph("• " + text)


def add_image(doc, path, caption, why_missing):
    from docx.shared import Inches
    if path and os.path.isfile(path):
        sec = doc.sections[-1]
        w = sec.page_width - sec.left_margin - sec.right_margin
        pic = doc.add_picture(path, width=w)
        max_h = Inches(8)
        if pic.height > max_h:
            ratio = max_h / pic.height
            pic.height = int(pic.height * ratio)
            pic.width = int(pic.width * ratio)
        doc.add_paragraph(caption)
    else:
        doc.add_paragraph("%s %s — %s" % (NO_IMG, caption, why_missing))


def new_document(template):
    from docx import Document
    if template and os.path.isfile(template):
        doc = Document(template)
        body = doc.element.body
        for el in list(body):
            if not el.tag.endswith("}sectPr"):
                body.remove(el)
    else:
        doc = Document()
    sec = doc.sections[0]
    for part, text in ((sec.header, HEADER), (sec.footer, FOOTER)):
        ps = part.paragraphs
        if ps:
            ps[0].text = text
            for p in ps[1:]:
                p.text = ""
        else:
            part.add_paragraph(text)
    return doc


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--inventory", required=True)
    ap.add_argument("--narrative", required=True)
    ap.add_argument("--index", required=True)
    ap.add_argument("--flow", default=None, help="flow.json (lay explanation)")
    ap.add_argument("--flow-png", default=None)
    ap.add_argument("--codemap-png", default=None)
    ap.add_argument("--erd-png", default=None)
    ap.add_argument("--template", default=DEFAULT_TPL)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    try:
        import docx  # noqa: F401
    except ImportError:
        print("Thieu python-docx. Chay: pip install python-docx", file=sys.stderr)
        return 2

    data = S.load(a.inventory)
    meta = S.meta(data)
    nar = json.load(open(a.narrative, encoding="utf8"))
    idx = json.load(open(a.index, encoding="utf8"))
    flow = json.load(open(a.flow, encoding="utf8")) if a.flow else {}
    ver_dir = os.path.dirname(os.path.dirname(os.path.abspath(a.index)))
    cnt = summary_counts(data)
    db_mode = meta.get("db_mode", S.UNKNOWN)

    doc = new_document(a.template)
    doc.add_paragraph("LABO / MAINTAIN PROJECT")
    doc.add_heading(TITLE, level=0)
    add_table(doc, ["Mục", "Giá trị"], [
        ["Dự án", nar.get("project_name") or meta.get("system_name")],
        ["Khách hàng", nar.get("customer") or meta.get("customer")],
        ["Phiên bản", "%s (%s)" % (idx.get("version", DASH), idx.get("version_folder", DASH))],
        ["Ngày", idx.get("date") or meta.get("generated_date")],
        ["Loại chạy", idx.get("version_type") or meta.get("version_type") or "BASELINE"],
    ])

    # 0
    doc.add_heading(CHAPTERS[0], level=1)
    add_table(doc, ["Site ID", "Tên site", "URL", "Môi trường", "Crawl Mode"],
              [[r.get("Site ID"), r.get("Site Name"), r.get("URL"), r.get("Env"),
                r.get("Crawl Mode")] for r in data["10_Site"]] or [[DASH, "Không có website"]])
    add_table(doc, ["Repo ID", "Loại", "Stack", "Cho site"],
              [[r.get("Repo ID"), r.get("Kind"), r.get("Stack"), r.get("For Site")]
               for r in data["11_Repo"]] or [[DASH, "Không có source code"]])
    figma = meta.get("figma_input_url", "")
    add_table(doc, ["Hạng mục", "Giá trị"], [
        ["Database", db_mode],
        ["Figma đầu vào", "Có" if figma and figma not in (S.UNKNOWN, DASH, "NO") else "Không"],
        ["Crawl mode", meta.get("crawl_mode")],
        ["Tài khoản (vai trò)", ", ".join(role_names(meta, data["10_Site"])) or S.UNKNOWN],
    ])

    # 1
    doc.add_heading(CHAPTERS[1], level=1)
    doc.add_paragraph(nar.get("product_overview") or S.UNKNOWN)
    add_table(doc, ["Actor", "Vai trò", "Sử dụng chính"],
              [[x.get("actor"), x.get("role"), x.get("main_usage")] for x in nar.get("actors", [])]
              or [[S.UNKNOWN, S.UNKNOWN, S.UNKNOWN]])
    for line in nar.get("system_composition") or [S.UNKNOWN]:
        bullet(doc, line)

    # 2
    doc.add_heading(CHAPTERS[2], level=1)
    out_rel = os.path.relpath(os.path.abspath(a.out), ver_dir).replace(os.sep, "/")
    rows = []
    for e in idx.get("outputs", []):
        files = list(e.get("files") or []) if e.get("exists") else []
        gate = "%s · %s" % (e.get("gate") or DASH, e.get("status_text") or DASH)
        if e.get("id") == "O6":
            files, gate = [out_rel], "%s · tài liệu này" % (e.get("gate") or DASH)
        rows.append(["%s — %s" % (e.get("id"), e.get("name")), e.get("content"),
                     "; ".join(files) or DASH, e.get("count"), gate])
    add_table(doc, ["Output", "Nội dung", "File / Link", "Số lượng", "Gate"], rows)

    # 3
    doc.add_heading(CHAPTERS[3], level=1)
    doc.add_paragraph("Số màn hình theo site và loại (danh sách chi tiết: O1 Basic Design).")
    add_table(doc, ["Site", "Số màn hình"] + SCREEN_TYPES, cnt["sites"])
    doc.add_paragraph("Số chức năng theo module và trạng thái bằng chứng.")
    add_table(doc, ["Module", "Số chức năng"] + S.FUNCTION_STATUS, cnt["modules"])

    # 4
    doc.add_heading(CHAPTERS[4], level=1)
    apis = data["07_API"]
    if apis:
        add_table(doc, ["Nhóm API", "Tổng", "API", "Batch", "Webhook/Queue"] + API_STATUS,
                  cnt["api_groups"])
        doc.add_paragraph("Tổng: %d (Confirmed %d · cần xác minh %d). Chi tiết: O2 API doc." % (
            len(apis), sum(1 for r in apis if r.get("Status") == "Confirmed"),
            sum(1 for r in apis if r.get("Status") != "Confirmed")))
    else:
        doc.add_paragraph("Không phát hiện API/batch trong source code đã khảo sát.")
    add_image(doc, a.codemap_png, "Hình 1. Code map FE → API → BE → DB",
              "chưa sinh code map (không có source code hoặc chưa chạy bước O2)")

    # 5
    doc.add_heading(CHAPTERS[5], level=1)
    if db_mode == "NONE":
        doc.add_paragraph("Database not available.")
        doc.add_paragraph("Không được cấp dump/kết nối DB và không có migration/ORM để suy ra "
                          "(db_mode=NONE).")
    else:
        doc.add_paragraph("Nguồn: %s. Danh sách bảng/cột đầy đủ: O3 DB doc." % db_mode)
        add_table(doc, ["Chỉ số DB", "Giá trị"], cnt["db"])
        add_image(doc, a.erd_png, "Hình 2. ERD", "chưa sinh ERD")

    # 6
    doc.add_heading(CHAPTERS[6], level=1)
    exts = data["09_Integration"]
    if exts:
        add_table(doc, ["EXT ID", "Tên", "Loại", "Mục đích", "Chiều", "Trạng thái"],
                  [[r.get("EXT ID"), r.get("Name"), r.get("Kind"), r.get("Purpose"),
                    r.get("Direction"), r.get("Status")] for r in exts])
    else:
        doc.add_paragraph("Không phát hiện liên kết bên thứ 3.")

    # 7
    doc.add_heading(CHAPTERS[7], level=1)
    add_image(doc, a.flow_png, "Hình 3. Luồng tổng quan", "chưa vẽ sơ đồ flow")
    for line in flow.get("explanation") or nar.get("flow_explanation") or [S.UNKNOWN]:
        doc.add_paragraph(line)

    # Phu luc A
    doc.add_heading(CHAPTERS[8], level=1)
    opens = [q for q in data["06_OpenQuestions"] if q.get("Status") == "Open"]
    if opens:
        add_table(doc, ["Q ID", "Loại", "Hạng mục", "Lý do / Bằng chứng", "Cần làm"],
                  [[q.get("Q ID"), q.get("Type"), q.get("Item"), q.get("Reason / Evidence"),
                    q.get("Required Action")] for q in opens])
    else:
        doc.add_paragraph("Không còn câu hỏi mở.")

    # Phu luc B
    doc.add_heading(CHAPTERS[9], level=1)
    add_table(doc, ["Khóa", "Giá trị"],
              [[k, ", ".join(role_names(meta, [])) if k == "accounts" else meta.get(k)]
               for k in APPENDIX_B_KEYS])

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    doc.save(a.out)
    print("Da render %s — %d output · %d site · %d function · %d API · %d EXT · %d open question"
          % (a.out, len(idx.get("outputs", [])), len(data["10_Site"]), len(data["01_Function"]),
             len(apis), len(exts), len(opens)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
