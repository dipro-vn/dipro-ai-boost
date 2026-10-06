#!/usr/bin/env python3
"""Ham dung chung cho Luong 2 (CR): doc cr.json + bang don gia, tinh MD, doc baseline, resolve ref.

Dung boi build-cr-impact.py · verify-cr-impact.py · verify-cr-figma.py · build-version-index.py.
Schema cr.json: xem docstring build-cr-impact.py.
"""
import json
import os
import re

import inv_schema as S

NONE_REFS = ("—", "-", "–", "")
ID_PREFIX = {"SC": "02_Screen", "F": "01_Function", "API": "07_API",
             "EXT": "09_Integration", "WEB": "10_Site"}
SITE_RX = re.compile(r"^WEB-\d{2,}$")

# --- nhan hien thi tren sheet Summary (verify + index doc lai theo cac nhan nay) ---
AXIS_VI = {"System": "Hệ thống", "DB": "Database", "Business": "Nghiệp vụ",
           "Screen": "Màn hình", "ThirdParty": "Bên thứ 3", "Mockup": "Thiết kế giao diện"}
L_CR = "Change Request"
L_TITLE = "Tiêu đề"
L_BASELINE = "Baseline"
L_RATES = "Đơn giá"
L_AXIS_HDR = "Trục"
L_TOTAL = "Tổng"
L_TOTAL_MD = "Tổng MD"

META_FIELDS = list(S.CR_META_KEYS)
JUST_FIELDS = ["item", "request", "baseline_ref", "baseline_quote", "criteria", "why", "not_feedback"]
NOT_CR_FIELDS = ["item", "label", "reason", "baseline_ref"]
Q_FIELDS = ["id", "question", "why", "axis", "owner", "status"]
IMPACT_FIELDS = ["id", "axis", "change_type", "baseline_ref", "item", "change", "why_change",
                 "impact_on_current", "conflict", "conflict_detail", "risk", "rate_code", "qty",
                 "md_note", "evidence", "question", "cr_item", "option"]
# cr_item = item cua justification ma hang muc phuc vu (CR-001.3) -> nhom section tren sheet Estimation.
# option  = true khi hang muc thuoc 拡張案 (phuong an mo rong) -> cot option tren Estimation.

# --- nhan tren Summary / Estimation (verify + index doc lai theo cac nhan nay) ---
L_IMPL = "実装 (MD — đơn giá × số lượng)"
L_TOTAL_PD = "総工数 (人日 — gồm 要件定義・UI/UX・テスト・管理)"
L_TOTAL_PM = "人月"
L_MIN_PLAN = "trong đó 最小改修案 (option = FALSE) — 実装"
L_EXT_PLAN = "trong đó 拡張案 (option = TRUE) — 実装"
L_OBJ_HDR = "Đối tượng"
L_LINK_EST = "→ Xem chi tiết công số: sheet Estimation"

# Doi tuong dem tren Summary + CR Change Table Figma. Dem = qty (so doi tuong theo don vi cua dong don gia).
OBJECT_ORDER = ["Màn hình", "API", "Batch / job", "Module / cấu hình", "Bảng DB", "Cột DB",
                "Đối tượng DB khác", "Rule nghiệp vụ", "Liên kết bên thứ 3", "Mockup màn",
                "Token / component DS"]
OBJECT_BY_RATE = [("SCR-", "Màn hình"), ("API-", "API"), ("BATCH-", "Batch / job"),
                  ("SYS-", "Module / cấu hình"), ("DB-TBL-", "Bảng DB"), ("DB-COL-", "Cột DB"),
                  ("BIZ-", "Rule nghiệp vụ"), ("EXT-", "Liên kết bên thứ 3"), ("UI-TOKEN-", "Token / component DS"),
                  ("UI-IMP", "Token / component DS"), ("UI-", "Mockup màn")]


def refs_of(value):
    return [r.strip() for r in str(value or "").split(";")
            if r.strip() and r.strip() not in NONE_REFS]


def txt(v):
    return "" if v is None else str(v).strip()


def load_json(path):
    with open(path, encoding="utf8") as fh:
        return json.load(fh)


def load_rates(path):
    """-> (meta dict, {code: rate})."""
    d = load_json(path)
    rates = {}
    for r in d.get("rates") or []:
        if isinstance(r, dict) and r.get("code"):
            rates[r["code"]] = r
    return d, rates


def qty_of(imp):
    try:
        return float(imp.get("qty"))
    except (TypeError, ValueError):
        return None


def md_of(imp, rates):
    """MD = rate.md x qty (round 2); None neu khong tinh duoc."""
    r = rates.get(txt(imp.get("rate_code")))
    q = qty_of(imp)
    if not r or q is None:
        return None
    try:
        return round(float(r.get("md")) * q, 2)
    except (TypeError, ValueError):
        return None


def md_matrix(impacts, rates):
    """-> {axis: {type: md}} + tong hang/cot/tong chung (round 2)."""
    m = {a: {t: 0.0 for t in S.CR_CHANGE_TYPE} for a in S.CR_AXES}
    for imp in impacts:
        v = md_of(imp, rates)
        if v is not None and imp.get("axis") in m and imp.get("change_type") in S.CR_CHANGE_TYPE:
            m[imp["axis"]][imp["change_type"]] += v
    row_tot = {a: round(sum(m[a].values()), 2) for a in S.CR_AXES}
    col_tot = {t: round(sum(m[a][t] for a in S.CR_AXES), 2) for t in S.CR_CHANGE_TYPE}
    for a in m:
        for t in m[a]:
            m[a][t] = round(m[a][t], 2)
    return m, row_tot, col_tot, round(sum(row_tot.values()), 2)


def object_of(imp, rates):
    """Doi tuong dem cua 1 hang muc. rate co key 'object' thi dung; DB-DEL / DB-IMP xet loai ref."""
    code = txt(imp.get("rate_code"))
    r = rates.get(code) or {}
    if txt(r.get("object")):
        return txt(r["object"])
    if code in ("DB-DEL", "DB-IMP"):
        refs = [x.lower() for x in refs_of(imp.get("baseline_ref"))]
        if refs and all(x.startswith("table:") for x in refs):
            return "Bảng DB"
        if refs and all(x.startswith("column:") for x in refs):
            return "Cột DB"
        return "Đối tượng DB khác"
    for pre, label in OBJECT_BY_RATE:
        if code.startswith(pre):
            return label
    return "Đối tượng DB khác" if imp.get("axis") == "DB" else txt(imp.get("axis"))


def object_stats(impacts, rates):
    """-> [(doi tuong, {NEW,UPD,DEL,IMPACT: so luong, 'total': so luong, 'md': MD 実装})] theo OBJECT_ORDER
    + dong cuoi ('Tổng', ...). So luong = tong qty."""
    st = {}
    for imp in impacts:
        o = object_of(imp, rates)
        d = st.setdefault(o, {t: 0.0 for t in S.CR_CHANGE_TYPE + ["total", "md"]})
        q = qty_of(imp) or 0.0
        if imp.get("change_type") in S.CR_CHANGE_TYPE:
            d[imp["change_type"]] += q
        d["total"] += q
        d["md"] += md_of(imp, rates) or 0.0
    order = [o for o in OBJECT_ORDER if o in st] + sorted(o for o in st if o not in OBJECT_ORDER)
    out = [(o, {k: round(v, 2) for k, v in st[o].items()}) for o in order]
    tot = {k: round(sum(d[k] for _, d in out), 2) for k in S.CR_CHANGE_TYPE + ["total", "md"]}
    return out + [(L_TOTAL, tot)]


def plan_md(impacts, rates):
    """-> (MD 最小改修案, MD 拡張案) theo co option."""
    ext = sum(md_of(i, rates) or 0.0 for i in impacts if i.get("option") is True)
    tot = sum(md_of(i, rates) or 0.0 for i in impacts)
    return round(tot - ext, 2), round(ext, 2)


def load_estimation_template(path):
    """-> (workbook openpyxl, meta {row_item, ratio_K...}) tu templates/template_estimation.xlsx."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    if "Estimation" not in wb.sheetnames or "_meta" not in wb.sheetnames:
        raise ValueError("template %s thieu sheet Estimation / _meta — boc lai bang "
                         "extract-estimation-template.py" % path)
    meta = {}
    for k, v in wb["_meta"].iter_rows(values_only=True):
        if k:
            key = str(k).split("_要件")[0].split("_UI")[0].split("_テスト")[0].split("_管理")[0]
            meta[key] = v
    for k in ("row_header", "row_section", "row_item", "row_total", "row_month", "row_notes",
              "ratio_K", "ratio_L", "ratio_N", "ratio_O", "md_per_month"):
        if meta.get(k) is None:
            raise ValueError("template _meta thieu %s" % k)
    return wb, meta


def find_template(arg, kit_root):
    """--template > ./templates/... > <kit>/templates/..."""
    for p in (arg, S.ESTIMATION_TEMPLATE_PATH, os.path.join(kit_root, S.ESTIMATION_TEMPLATE_PATH)):
        if p and os.path.isfile(p):
            return p
    return None


def fmt_md(v):
    return ("%.2f" % v).rstrip("0").rstrip(".") if v is not None else ""


def validate_cr(cr, rates):
    """Kiem hinh dang + enum + don gia cua cr.json. -> list loi (rong = OK)."""
    err = []
    if not isinstance(cr, dict):
        return ["cr.json khong phai object"]
    meta = cr.get("meta") or {}
    for k in META_FIELDS:
        if not txt(meta.get(k)) or txt(meta.get(k)) == S.UNKNOWN:
            err.append("meta.%s rong/UNKNOWN" % k)
    if txt(meta.get("source_type")) and meta["source_type"] not in S.CR_SOURCE_TYPE:
        err.append("meta.source_type=%s (chi %s)" % (meta["source_type"], "/".join(S.CR_SOURCE_TYPE)))
    if txt(meta.get("overall_risk")) and meta["overall_risk"] not in S.CR_RISK:
        err.append("meta.overall_risk=%s (chi %s)" % (meta["overall_risk"], "/".join(S.CR_RISK)))
    for key in ("justification", "not_cr", "no_impact_axes", "questions", "impacts"):
        if not isinstance(cr.get(key, []), list):
            err.append("%s phai la list" % key)
            cr[key] = []
    for i, j in enumerate(cr.get("justification") or []):
        crit = j.get("criteria") if isinstance(j, dict) else None
        if not isinstance(crit, list) or not crit:
            err.append("justification[%d] criteria phai la list C1..C6 khong rong" % i)
        else:
            err += ["justification[%d] criteria %s khong co trong C1..C6" % (i, c)
                    for c in crit if c not in S.CR_CRITERIA]
    for i, n in enumerate(cr.get("not_cr") or []):
        if txt(n.get("label")) not in S.CR_NOT_CR_LABEL:
            err.append("not_cr[%d] label=%s (chi %s)" % (i, n.get("label"), "/".join(S.CR_NOT_CR_LABEL)))
    for i, x in enumerate(cr.get("no_impact_axes") or []):
        if txt(x.get("axis")) not in S.CR_AXES:
            err.append("no_impact_axes[%d] axis=%s" % (i, x.get("axis")))
    for i, q in enumerate(cr.get("questions") or []):
        if not re.match(S.CR_QUESTION_ID, txt(q.get("id"))):
            err.append("questions[%d] id=%r sai dang CQ-NNN" % (i, q.get("id")))
    seen = set()
    for i, imp in enumerate(cr.get("impacts") or []):
        iid = txt(imp.get("id"))
        tag = iid or "impacts[%d]" % i
        if not re.match(S.CR_IMPACT_ID, iid):
            err.append("%s id sai dang IMP-NNN" % tag)
        if iid in seen:
            err.append("%s trung id" % tag)
        seen.add(iid)
        for f, enum in (("axis", S.CR_AXES), ("change_type", S.CR_CHANGE_TYPE),
                        ("conflict", S.CR_CONFLICT), ("risk", S.CR_RISK)):
            if txt(imp.get(f)) not in enum:
                err.append("%s %s=%s (chi %s)" % (tag, f, imp.get(f), "/".join(enum)))
        err += rate_errors(imp, rates, tag)
        if "option" in imp and not isinstance(imp.get("option"), bool):
            err.append("%s option=%r phai la true/false" % (tag, imp.get("option")))
    items = {txt(j.get("item")) for j in cr.get("justification") or [] if isinstance(j, dict)}
    for i, imp in enumerate(cr.get("impacts") or []):
        ci = txt(imp.get("cr_item"))
        if not ci:
            err.append("%s thieu cr_item (item justification ma hang muc phuc vu)" % (txt(imp.get("id")) or i))
        elif items and ci not in items:
            err.append("%s cr_item=%s khong co trong justification (%s)" % (
                txt(imp.get("id")), ci, ",".join(sorted(items))))
    return err


def rate_errors(imp, rates, tag):
    err = []
    code = txt(imp.get("rate_code"))
    r = rates.get(code)
    if not r:
        err.append("%s rate_code=%r khong co trong bang don gia" % (tag, code))
    elif (r.get("axis"), r.get("change_type")) != (imp.get("axis"), imp.get("change_type")):
        err.append("%s rate_code %s danh cho %s/%s, dong la %s/%s" % (
            tag, code, r.get("axis"), r.get("change_type"), imp.get("axis"), imp.get("change_type")))
    q = qty_of(imp)
    if q is None or q <= 0:
        err.append("%s qty=%r phai > 0" % (tag, imp.get("qty")))
    return err


# ---------------- baseline ----------------
def load_baseline(bdir):
    """ids/bang/cot tu inventory + design system (1 bo project/ hoac moi site WEB-xx/project/)."""
    inv = os.path.join(bdir, "_internal", "inventory.xlsx")
    if not os.path.isfile(inv):
        raise SystemExit("Khong thay baseline inventory: %s" % inv)
    data = S.load(inv)
    ids = {}
    for pre, sheet in ID_PREFIX.items():
        col = S.ID_PATTERNS[sheet][0]
        ids[pre] = {r.get(col, "") for r in data.get(sheet, []) if r.get(col)}
    tables = {r.get("Table", "").lower() for r in data.get("03_DB_Tables", []) if r.get("Table")}
    cols = set()
    for r in data.get("04_DB_Columns", []):
        t, c = r.get("Table", "").lower(), r.get("Column", "").lower()
        if t and c:
            tables.add(t)
            cols.add("%s.%s" % (t, c))
    ds = load_ds_sites(os.path.join(bdir, "04_DesignSystem"))
    return {"ids": ids, "tables": tables, "cols": cols, "ds": ds, "has_ds": bool(ds),
            "links": build_links(data), "data": data}


def ds_dirs(ds_root):
    """-> [(site, thu muc chua STATUS.md + project/)] — site '' = layout 1 bo."""
    out = []
    if os.path.isdir(os.path.join(ds_root, "project")):
        out.append(("", ds_root))
    if os.path.isdir(ds_root):
        for d in sorted(os.listdir(ds_root)):
            if SITE_RX.match(d) and os.path.isdir(os.path.join(ds_root, d, "project")):
                out.append((d, os.path.join(ds_root, d)))
    return out


def load_ds_sites(ds_root):
    """{site: {"tokens": {ten: gia tri}, "comps": set}}."""
    out = {}
    for site, d in ds_dirs(ds_root):
        proj = os.path.join(d, "project")
        tokens = {}
        tp = os.path.join(proj, "tokens.json")
        if os.path.isfile(tp):
            try:
                tokens = artifact_tokens(load_json(tp))
            except ValueError:
                tokens = {}
        out[site] = {"tokens": tokens, "comps": artifact_components(os.path.join(proj, "components"))}
    return out


def artifact_tokens(tk):
    out = {}
    if not isinstance(tk, dict):
        return out
    for v in tk.values():
        if isinstance(v, dict) and isinstance(v.get("tokens"), list):
            for t in v["tokens"]:
                if isinstance(t, dict) and t.get("name"):
                    out[str(t["name"]).lower()] = json.dumps(t.get("value"), sort_keys=True)
    ty = tk.get("type") if isinstance(tk.get("type"), dict) else {}
    for gr in ty.get("groups") or []:
        for st in (gr.get("styles") or []) if isinstance(gr, dict) else []:
            if isinstance(st, dict) and st.get("name"):
                out[str(st["name"]).lower()] = json.dumps(
                    {k: st.get(k) for k in ("fontSize", "fontWeight", "lineHeight", "family")}, sort_keys=True)
    return out


def artifact_components(cdir):
    out = set()
    if os.path.isdir(cdir):
        out |= {d.lower() for d in os.listdir(cdir)
                if os.path.isdir(os.path.join(cdir, d)) and d not in ("Cover", "lib", "src")}
        dts = os.path.join(cdir, "index.d.ts")
        if os.path.isfile(dts):
            out |= {m.lower() for m in re.findall(
                r"export\s+(?:declare\s+)?(?:function|const|class)\s+([A-Za-z_$][\w$]*)",
                open(dts, encoding="utf8").read())}
    return out


def build_links(data):
    """Do thi 1-hop: Entry From (2 chieu), Function.Screen IDs <-> SC, API.Called By Screens <-> SC."""
    g = {}

    def link(a, b):
        if a and b and a != b:
            g.setdefault(a, set()).add(b)
            g.setdefault(b, set()).add(a)
    for r in data.get("02_Screen", []):
        for x in S.split_ids(r.get("Entry From")):
            link(r.get("Screen ID"), x)
    for r in data.get("01_Function", []):
        for x in S.split_ids(r.get("Screen IDs")):
            if x.startswith("SC-"):
                link(r.get("Function ID"), x)
    for r in data.get("07_API", []):
        for x in S.split_ids(r.get("Called By Screens")):
            if x.startswith("SC-"):
                link(r.get("API ID"), x)
    return g


def _ds_parse(body):
    """'WEB-01:primary' -> ('WEB-01','primary'); 'primary' -> (None,'primary')."""
    head, sep, rest = body.partition(":")
    if sep and SITE_RX.match(head.strip()):
        return head.strip(), rest.strip()
    return None, body.strip()


def resolve(ref, B):
    """-> (ok, ly do loi, canh bao). Canh bao = ref DS khong ghi site ma nhieu site khac gia tri."""
    low = ref.lower()
    if low.startswith("table:"):
        return (low[6:].strip() in B["tables"], "bang khong co trong baseline", "")
    if low.startswith("column:"):
        return (low[7:].strip() in B["cols"], "cot khong co trong baseline", "")
    if low.startswith("ds-component:") or low.startswith("ds:"):
        comp = low.startswith("ds-component:")
        site, name = _ds_parse(ref[13:] if comp else ref[3:])
        name = re.sub(r"[`*_]", "", name).strip().lower() if comp else name.lower()
        if not B["ds"]:
            return (False, "baseline khong co design system", "")
        if site is not None and site not in B["ds"]:
            return (False, "baseline khong co design system cho %s" % site, "")
        sites = [site] if site is not None else list(B["ds"])
        if comp:
            hit = [s for s in sites if name in B["ds"][s]["comps"]]
            return (bool(hit), "component khong co trong design system baseline", "")
        hit = [s for s in sites if name in B["ds"][s]["tokens"]]
        warn = ""
        if site is None and len({B["ds"][s]["tokens"][name] for s in hit}) > 1:
            warn = "%s co o %s voi gia tri khac nhau — ghi ro DS:<WEB-xx>:%s" % (
                ref, ",".join(hit), name)
        return (bool(hit), "token khong co trong tokens.json baseline", warn)
    m = re.match(r"^(SC|F|API|EXT|WEB)-\d+$", ref)
    if m:
        return (ref in B["ids"][m.group(1)], "id khong co trong baseline", "")
    return (False, "cu phap ref khong hop le", "")
