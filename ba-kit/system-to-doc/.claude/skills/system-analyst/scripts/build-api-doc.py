#!/usr/bin/env python3
"""O2 — API Documentation (xlsx) tu inventory: API + webhook + BATCH/QUEUE.

  python3 build-api-doc.py --inventory <_internal>/inventory.xlsx --system <name> --version ver1 \\
      --out <ver>/02_API/API_Doc_<sys>_ver1.xlsx

Nguon: 07_API, 08_API_Fields (+02_Screen ten man hinh, 11_Repo, 05_Evidence locator).
Sheet: 00_Index (tong hop + link) · G_<group> (moi API 1 block: thong tin + Request + Response)
       · Batch (moi BATCH/QUEUE). O trong -> UNKNOWN (hoac — neu khong ap dung).
"""
import argparse
import datetime
import json
import os
import re
import sys

import inv_schema as S

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Thieu openpyxl. Chay: pip install openpyxl", file=sys.stderr)
    sys.exit(2)
import bd_styles as ST

NA = "—"
NO_FIELD = "UNKNOWN — chua xac dinh tu code"
INDEX_COLS = ["No", "API ID", "Group", "Kind", "Method", "Path / Schedule", "Summary", "Auth",
              "Called By Screens", "Related Tables", "Status"]
BATCH_COLS = ["No", "API ID", "Schedule", "Handler", "Summary", "Related Tables", "Status", "Evidence"]
REQ_COLS = ["In", "Field", "Type", "Required", "Format / Constraint", "Description", "Evidence"]
RES_COLS = ["In", "Field or status code", "Type", "Format / Constraint", "Description", "Evidence"]
BLOCK_LABELS = ["Summary", "Kind", "Auth", "Handler", "Repo", "Called By Screens",
                "Related Tables", "Status", "Evidence"]
BATCH_KINDS = ("BATCH", "QUEUE")
LINK_FONT = Font(name=ST.FONT, size=10, color="FF0563C1", underline="single")
SECTION_FILL = "FFFFF2CC"


def v(x, empty=S.UNKNOWN):
    x = "" if x is None else str(x).strip()
    return x if x else empty


def sheet_title(prefix, name, used):
    base = re.sub(r"[\[\]:*?/\\']", "_", "%s%s" % (prefix, name)).strip()[:31]
    t, i = base, 2
    while t.lower() in used:
        suf = "~%d" % i
        t = base[:31 - len(suf)] + suf
        i += 1
    used.add(t.lower())
    return t


def link(cell, sheet, ref, text=None):
    if text is not None:
        cell.value = text
    cell.hyperlink = "#'%s'!%s" % (sheet, ref)
    cell.font = LINK_FONT


def api_key(r):
    m = re.match(r"^API-(\d+)$", r.get("API ID", ""))
    return (int(m.group(1)) if m else 10 ** 9, r.get("API ID", ""))


def title_of(r):
    return "%s · %s %s" % (r["API ID"], v(r.get("Method")), v(r.get("Path / Schedule")))


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_grid(ws, r0, header, rows, center_cols=()):
    for i, h in enumerate(header, start=1):
        ST.header(ws.cell(r0, i, h))
    for ri, row in enumerate(rows, start=r0 + 1):
        for ci, val in enumerate(row, start=1):
            ST.data(ws.cell(ri, ci, val), size=10, center=ci in center_cols)
    ws.auto_filter.ref = "A%d:%s%d" % (r0, get_column_letter(len(header)), r0 + max(len(rows), 1))
    ws.freeze_panes = ws.cell(r0 + 1, 3)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--inventory", required=True)
    ap.add_argument("--system", required=True)
    ap.add_argument("--version", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    data = S.load(a.inventory)
    apis = sorted([r for r in data["07_API"] if r.get("API ID")], key=api_key)
    if not apis:
        print("07_API rong — khong co API/batch nao de sinh O2", file=sys.stderr)
        return 1
    fields = {}
    for f in data["08_API_Fields"]:
        fields.setdefault(f.get("API ID", ""), []).append(f)
    screens = {s.get("Screen ID"): s.get("Screen Name", "") for s in data["02_Screen"]}
    repos = {r.get("Repo ID"): r for r in data["11_Repo"]}
    evloc = {e.get("EV ID"): e for e in data["05_Evidence"]}

    def screens_txt(val, kind=""):
        ids = S.split_ids(val)
        if not ids:  # batch/queue/webhook khong do man hinh goi -> khong ap dung
            return NA if (v(val, "") in ("—", "-", "N/A") or kind in ("BATCH", "QUEUE", "WEBHOOK")) else S.UNKNOWN
        return "; ".join("%s %s" % (i, screens[i]) if screens.get(i) else i for i in ids)

    def repo_txt(val):
        ids = S.split_ids(val)
        out = []
        for i in ids:
            r = repos.get(i)
            if r:
                extra = ", ".join(x for x in (r.get("Kind"), r.get("Stack")) if x)
                out.append("%s (%s)" % (i, extra) if extra else i)
            else:
                out.append(i)
        return "; ".join(out) or S.UNKNOWN

    def ev_txt(val):
        ids = S.split_ids(val)
        out = []
        for i in ids:
            e = evloc.get(i)
            out.append("%s (%s: %s)" % (i, e.get("Type", ""), e.get("Locator", "")) if e else i)
        return "\n".join(out) or S.UNKNOWN

    wb = Workbook()
    wb.remove(wb.active)
    used = {"00_index", "batch"}
    wsi = wb.create_sheet("00_Index")
    groups = []
    for r in apis:
        g = v(r.get("Group"))
        if g not in groups:
            groups.append(g)
    gsheet = {g: sheet_title("G_", g, used) for g in groups}
    gws = {g: wb.create_sheet(gsheet[g]) for g in groups}
    wsb = wb.create_sheet("Batch")

    # --- group sheets: block / API
    anchor = {}
    NC = 7
    for g in groups:
        ws = gws[g]
        ws["A1"] = "API GROUP — %s" % g
        ws["A1"].font = Font(name=ST.FONT, size=13, bold=True, color="FF1F4E78")
        link(ws["G1"], "00_Index", "A1", "↑ Index")
        r = 3
        for api in [x for x in apis if v(x.get("Group")) == g]:
            aid = api["API ID"]
            anchor[aid] = (gsheet[g], "A%d" % r)
            for c in range(1, NC + 1):
                ST.title(ws.cell(r, c))
                ws.cell(r, c).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(r, 1, title_of(api))
            link(ws.cell(r, NC), "00_Index", "A1", "↑ Index")
            ws.cell(r, NC).font = Font(name=ST.FONT, size=10, color="FFFFFFFF", underline="single", bold=True)
            r += 1
            vals = {"Summary": v(api.get("Summary")),
                    "Kind": "%s · %s" % (v(api.get("Kind")), v(api.get("Method"))),
                    "Auth": v(api.get("Auth")), "Handler": v(api.get("Handler")),
                    "Repo": repo_txt(api.get("Repo")),
                    "Called By Screens": screens_txt(api.get("Called By Screens"), api.get("Kind")),
                    "Related Tables": v(api.get("Related Tables")),
                    "Status": v(api.get("Status")) + (("  ·  Open Q: " + api["Open Q"]) if api.get("Open Q") else ""),
                    "Evidence": ev_txt(api.get("Evidence"))}
            labels = list(BLOCK_LABELS) + (["Note"] if api.get("Note") else [])
            vals["Note"] = v(api.get("Note"), NA)
            for lab in labels:
                ST.header(ws.cell(r, 1, lab))
                ws.cell(r, 1).alignment = Alignment(horizontal="left", vertical="center")
                ST.data(ws.cell(r, 2, vals[lab]), size=10)
                for c in range(3, NC + 1):
                    ST.data(ws.cell(r, c), size=10)
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=NC)
                lines = max(1, str(vals[lab]).count("\n") + 1, len(str(vals[lab])) // 110 + 1)
                if lines > 1:
                    ws.row_dimensions[r].height = 14 * lines
                r += 1
            for direction, cols in (("REQUEST", REQ_COLS), ("RESPONSE", RES_COLS)):
                for c in range(1, NC + 1):
                    ST.banner(ws.cell(r, c))
                    ws.cell(r, c).fill = PatternFill("solid", fgColor=SECTION_FILL)
                ws.cell(r, 1, direction)
                r += 1
                for i, h in enumerate(cols, start=1):
                    ST.header(ws.cell(r, i, h))
                r += 1
                rows = [f for f in fields.get(aid, []) if f.get("Direction") == direction]
                if not rows:
                    vals_row = [NO_FIELD] + [NA] * (len(cols) - 1)
                    for i, val in enumerate(vals_row, start=1):
                        ST.data(ws.cell(r, i, val), size=10)
                    r += 1
                for f in rows:
                    if direction == "REQUEST":
                        vals_row = [v(f.get("In")), v(f.get("Field")), v(f.get("Type")),
                                    v(f.get("Required")), v(f.get("Format / Constraint"), NA),
                                    v(f.get("Description")), v(f.get("Evidence"), v(api.get("Evidence")))]
                    else:
                        vals_row = [v(f.get("In")), v(f.get("Field")), v(f.get("Type"), NA),
                                    v(f.get("Format / Constraint"), NA), v(f.get("Description")),
                                    v(f.get("Evidence"), v(api.get("Evidence")))]
                    for i, val in enumerate(vals_row, start=1):
                        ST.data(ws.cell(r, i, val), size=10, center=(i == 1))
                    r += 1
            r += 1  # dong trong ngan cach block
        set_widths(ws, [20, 28, 16, 12, 30, 44, 24])
        ws.freeze_panes = "A2"

    # --- 00_Index
    n_batch = sum(1 for x in apis if x.get("Kind") in BATCH_KINDS)
    st = {}
    for x in apis:
        st[v(x.get("Status"))] = st.get(v(x.get("Status")), 0) + 1
    kinds = {}
    for x in apis:
        kinds[v(x.get("Kind"))] = kinds.get(v(x.get("Kind")), 0) + 1
    wsi["A1"] = "API DOCUMENTATION — %s" % a.system
    wsi["A1"].font = Font(name=ST.FONT, size=14, bold=True, color="FF1F4E78")
    meta_rows = [("System", a.system), ("Version", a.version),
                 ("Generated", datetime.date.today().isoformat()),
                 ("#Entries", len(apis)), ("#APIs", len(apis) - n_batch), ("#Batch", n_batch),
                 ("#Groups", len(groups)),
                 ("Kind", " · ".join("%s %d" % kv for kv in sorted(kinds.items()))),
                 ("Status", " · ".join("%s %d" % (k, st[k]) for k in S.API_STATUS if k in st) +
                  "".join(" · %s %d" % (k, n) for k, n in st.items() if k not in S.API_STATUS)),
                 ("Ghi chu", "Bam API ID de mo block chi tiet. UNKNOWN = chua co bang chung tu code; "
                             "— = khong ap dung.")]
    r = 3
    for k, val in meta_rows:
        ST.header(wsi.cell(r, 1, k))
        wsi.cell(r, 1).alignment = Alignment(horizontal="left", vertical="center")
        ST.data(wsi.cell(r, 2, val), size=10)
        r += 1
    hdr = r + 1
    rows = []
    for i, x in enumerate(apis, start=1):
        rows.append([i, x["API ID"], v(x.get("Group")), v(x.get("Kind")), v(x.get("Method")),
                     v(x.get("Path / Schedule")), v(x.get("Summary")), v(x.get("Auth")),
                     screens_txt(x.get("Called By Screens"), x.get("Kind")), v(x.get("Related Tables")),
                     v(x.get("Status"))])
    write_grid(wsi, hdr, INDEX_COLS, rows, center_cols=(1, 4, 5, 11))
    for i, x in enumerate(apis):
        sh, ref = anchor[x["API ID"]]
        link(wsi.cell(hdr + 1 + i, 2), sh, ref)
        link(wsi.cell(hdr + 1 + i, 3), sh, "A1")
    set_widths(wsi, [12, 12, 16, 10, 9, 34, 44, 18, 28, 22, 11])
    wsi.column_dimensions["A"].width = 14

    # --- Batch
    wsb["A1"] = "BATCH / QUEUE JOBS — %s" % a.system
    wsb["A1"].font = Font(name=ST.FONT, size=13, bold=True, color="FF1F4E78")
    link(wsb["H1"], "00_Index", "A1", "↑ Index")
    batches = [x for x in apis if x.get("Kind") in BATCH_KINDS]
    brows = [[i, x["API ID"], v(x.get("Path / Schedule")), v(x.get("Handler")), v(x.get("Summary")),
              v(x.get("Related Tables")), v(x.get("Status")), v(x.get("Evidence"))]
             for i, x in enumerate(batches, start=1)]
    if not brows:
        brows = [[1, NA, NA, NA, "Khong tim thay batch/queue job trong source", NA, NA, NA]]
    write_grid(wsb, 3, BATCH_COLS, brows, center_cols=(1, 7))
    for i, x in enumerate(batches):
        sh, ref = anchor[x["API ID"]]
        link(wsb.cell(4 + i, 2), sh, ref)
    set_widths(wsb, [6, 12, 22, 28, 44, 22, 11, 16])

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    wb.save(a.out)
    print(json.dumps({"out": a.out, "entries": len(apis), "batch": n_batch, "groups": len(groups),
                      "fields": sum(len(fields.get(x["API ID"], [])) for x in apis)}, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
