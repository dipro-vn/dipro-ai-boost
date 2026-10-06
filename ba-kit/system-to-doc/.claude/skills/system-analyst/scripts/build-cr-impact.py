#!/usr/bin/env python3
"""Render CR Impact workbook (Luong 2) tu cr.json: 1 file, DUNG 7 sheet theo inv_schema.CR_SHEETS.

  python3 build-cr-impact.py --cr-json <ver>/_internal/cr.json \
      --rates .claude/config/md-unit-rates.json --out <ver>/CR-<id>_Impact.xlsx \
      [--template templates/template_estimation.xlsx]

Sheet: Summary (cong so thay doi + so doi tuong doi theo NEW/UPD/DEL/IMPACT + link Estimation) ·
Estimation (dung template boc tu 見積書 bang extract-estimation-template.py: 実装 = MD, 要件定義/UI/テスト/管理
= cong thuc he so cua template) · Screen / API / Database / Figma (hang muc truc Screen / System / DB / Mockup) ·
Q&A (de xuat, vi sao la CR, cau hoi KH, truc khong anh huong, khong thuoc CR).
MD do script tinh = rate.md x qty (AI chi chon rate_code + qty, KHONG go so MD).
cr.json sai (enum, rate_code la, rate khac truc/loai, qty <= 0, thieu cr_item) -> exit 1, khong ghi gi.

Schema <ver>/_internal/cr.json (agent viet):
{"meta": {cr_id, cr_title, baseline_version, cr_version, source_type FILE|CHAT|LINK, source_ref,
          received_date, requested_by, overall_risk High|Medium|Low, recommendation},
 "justification": [{item:"CR-001.1", request, baseline_ref:"SC-014", baseline_quote, criteria:["C1","C3"],
                    why:"<=2 cau", not_feedback:"1 cau"}],              # >=1 dong, ref phai co trong baseline
 "not_cr": [{item, label: BUG|QUESTION, reason, baseline_ref}],
 "no_impact_axes": [{axis, reason >=15 ky tu}],                      # truc khong co dong impact nao
 "questions": [{id:"CQ-001", question, why, axis, owner, status}],
 "impacts": [{id:"IMP-001", axis System|DB|Business|Screen|ThirdParty|Mockup,
              change_type NEW|UPD|DEL|IMPACT, baseline_ref "SC-014;DS:primary" | "—", item, change,
              why_change, impact_on_current, conflict Yes|No|UNKNOWN, conflict_detail,
              risk High|Medium|Low, rate_code (ma trong md-unit-rates.json), qty > 0, md_note,
              evidence, question "CQ-001;CQ-002" | "", cr_item "CR-001.1", option true|false}]}
"""
import argparse
import math
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
KIT = os.path.abspath(os.path.join(HERE, "..", "..", "..", ".."))
sys.path.insert(0, HERE)
import inv_schema as S  # noqa: E402
import cr_common as C  # noqa: E402

# nhan ngan co dau cho khach (ma -> nhan); thieu ma -> dung CR_CRITERIA
CRIT_SHORT = {
    "C1": "Chức năng mới, chưa có trong hệ thống",
    "C2": "Đổi hành vi đang chạy đúng",
    "C3": "Thêm/bỏ màn, actor, chức năng, liên kết ngoài",
    "C4": "Đổi business rule / validation",
    "C5": "Đổi nền tảng / yêu cầu phi chức năng",
    "C6": "Kéo theo đổi DB / API / tài liệu đã bàn giao",
}
SRC_VI = {"FILE": "File", "CHAT": "Chat", "LINK": "Link"}
RISK_VI = {"High": "Cao (High)", "Medium": "Trung bình (Medium)", "Low": "Thấp (Low)"}
WIDTH_SUM = [26, 30, 36, 13, 30, 14, 14]      # cot A..G sheet Summary / Q&A
WIDTH_IMPACT = {"No": 5, "Impact ID": 10, "Trục": 11, "Loại": 8, "Baseline Ref": 20,
                "Hạng mục": 24, "Nội dung thay đổi": 36, "Vì sao phải sửa": 36,
                "Ảnh hưởng tới hiện tại": 34, "Xung đột": 9, "Chi tiết xung đột": 26,
                "Rủi ro": 9, "Mã đơn giá": 13, "Số lượng": 8, "MD": 7, "Ghi chú MD": 36,
                "Evidence": 30, "Câu hỏi": 12}
WARN_RED = "FFC00000"


def est_lines(text, width):
    n = 0
    for part in str(text or "").split("\n"):
        n += max(1, math.ceil(len(part) * 1.1 / max(width, 1)))
    return n


class Sheet:
    """Ghi sheet Summary tung dong, tu uoc luong chieu cao dong."""

    def __init__(self, ws):
        self.ws, self.r = ws, 0

    def put(self, cells, kind="data", merge_to=None, center=(), bold=(), color=None):
        from openpyxl.styles import Font
        import bd_styles as BD
        self.r += 1
        r = self.r
        lines = 1
        for col, val in cells:
            c = self.ws.cell(row=r, column=col, value=val)
            fn = {"header": BD.header, "title": BD.title, "banner": BD.banner}.get(kind)
            if fn:
                fn(c)
            else:
                BD.data(c, size=10, center=col in center)
                if col in bold:
                    c.font = Font(name=BD.FONT, size=10, bold=True)
            if color:
                c.font = Font(name=BD.FONT, size=10, bold=True, color=color)
            end = merge_to.get(col, col) if merge_to else col
            w = sum(WIDTH_SUM[i - 1] for i in range(col, end + 1))
            lines = max(lines, est_lines(val, w))
        for col, end in (merge_to or {}).items():
            for cc in range(col + 1, end + 1):
                fn2 = {"header": BD.header, "title": BD.title, "banner": BD.banner}.get(kind, BD.data)
                fn2(self.ws.cell(row=r, column=cc))
            if end > col:
                self.ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=end)
        self.ws.row_dimensions[r].height = max(16, 14 * lines + 2)
        return r

    def gap(self):
        self.r += 1
        self.ws.row_dimensions[self.r].height = 8


def sheet_setup(ws, ncols, last_row):
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(WIDTH_SUM[:ncols], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:%s%d" % (get_column_letter(ncols), last_row)


def link(cell, sheet, ref="A1"):
    from openpyxl.styles import Font
    import bd_styles as BD
    cell.hyperlink = "#'%s'!%s" % (sheet, ref)
    cell.font = Font(name=BD.FONT, size=10, bold=True, color="FF0563C1", underline="single")


def build_summary(ws, cr, rates_meta, rates, est):
    """Summary: chi cong so thay doi + so doi tuong doi + link sang Estimation / sheet chi tiet."""
    meta = cr["meta"]
    sh = Sheet(ws)
    full = {1: 7}
    status = str(rates_meta.get("status", "")).upper()
    sh.put([(1, "%s — %s" % (meta["cr_id"], meta["cr_title"]))], "title", merge_to=full)
    for k, v in ((C.L_CR, meta["cr_id"]), (C.L_BASELINE, meta["baseline_version"]),
                 ("Nguồn", "%s — %s" % (SRC_VI.get(meta["source_type"], meta["source_type"]), meta["source_ref"])),
                 ("Ngày nhận", meta["received_date"]), ("Người yêu cầu (vai trò)", meta["requested_by"]),
                 ("Mức ảnh hưởng chung", RISK_VI.get(meta["overall_risk"], meta["overall_risk"]))):
        sh.put([(1, k), (2, v)], merge_to={2: 7}, bold=(1,))

    sh.gap()
    sh.put([(1, "1. Công số thay đổi")], "banner", merge_to=full)
    grand = round(sum(C.md_of(i, rates) or 0.0 for i in cr["impacts"]), 2)
    mn, ex = C.plan_md(cr["impacts"], rates)
    sh.put([(1, C.L_IMPL), (2, grand), (3, "MD · %d hạng mục" % len(cr["impacts"]))],
           merge_to={3: 7}, bold=(1, 2), center=(2,))
    r = sh.put([(1, C.L_TOTAL_PD), (2, "='%s'!P%d" % (S.CR_SHEET_ESTIMATION, est["row_total"])),
                (3, "人日 — theo hệ số của template Estimation")], merge_to={3: 7}, bold=(1, 2), center=(2,))
    ws.cell(row=r, column=2).number_format = "0.00"
    r = sh.put([(1, C.L_TOTAL_PM), (2, "='%s'!P%d" % (S.CR_SHEET_ESTIMATION, est["row_month"])),
                (3, "人月 = 人日 / %s" % C.fmt_md(est["md_per_month"]))], merge_to={3: 7}, bold=(1, 2), center=(2,))
    ws.cell(row=r, column=2).number_format = "0.00"
    sh.put([(1, C.L_MIN_PLAN), (2, mn), (3, "MD")], merge_to={3: 7}, center=(2,))
    sh.put([(1, C.L_EXT_PLAN), (2, ex), (3, "MD")], merge_to={3: 7}, center=(2,))
    r = sh.put([(1, C.L_LINK_EST)], merge_to=full)
    link(ws.cell(row=r, column=1), S.CR_SHEET_ESTIMATION, "B%d" % est["row_header"])
    sh.put([(1, C.L_RATES), (2, "%s v%s — %s" % (status or "?", rates_meta.get("version", "?"),
                                                  C.txt(rates_meta.get("scope_note"))))],
           merge_to={2: 7}, bold=(1,))
    if status != "APPROVED":
        sh.put([(1, "Lưu ý"), (2, "Ước lượng sơ bộ — chưa được PM/Tech Lead duyệt")],
               merge_to={2: 7}, color=WARN_RED)

    sh.gap()
    sh.put([(1, "2. Số đối tượng thay đổi")], "banner", merge_to=full)
    sh.put([(1, C.L_OBJ_HDR)] + [(i + 2, t) for i, t in enumerate(S.CR_CHANGE_TYPE)] +
           [(6, C.L_TOTAL), (7, "実装 MD")], "header")
    for o, d in C.object_stats(cr["impacts"], rates):
        vals = [d[t] for t in S.CR_CHANGE_TYPE] + [d["total"], d["md"]]
        sh.put([(1, o)] + [(i + 2, int(v) if v == int(v) else v) for i, v in enumerate(vals)],
               center=tuple(range(2, 8)), bold=(1, 2, 3, 4, 5, 6, 7) if o == C.L_TOTAL else (1,))

    sh.gap()
    sh.put([(1, "3. Sheet chi tiết")], "banner", merge_to=full)
    for name, axes in S.CR_AXIS_SHEETS:
        n = sum(1 for i in cr["impacts"] if i["axis"] in axes)
        r = sh.put([(1, "→ " + name), (2, "%d hạng mục — trục %s" % (
            n, ", ".join("%s (%s)" % (C.AXIS_VI[a], a) for a in axes)))], merge_to={2: 7})
        link(ws.cell(row=r, column=1), name)
    r = sh.put([(1, "→ " + S.CR_SHEET_QA), (2, "Vì sao là CR · %d câu hỏi cần KH trả lời · không thuộc CR"
                                            % len(cr["questions"]))], merge_to={2: 7})
    link(ws.cell(row=r, column=1), S.CR_SHEET_QA)
    sh.put([(1, "Ghi chú"), (2, "Trục %s chỉ có trong Estimation (luồng nghiệp vụ đã thể hiện trên Figma CR-1)."
                            % ", ".join(C.AXIS_VI[a] for a in S.CR_ESTIMATION_ONLY_AXES))], merge_to={2: 7})
    sheet_setup(ws, 7, sh.r)


def build_qa(ws, cr, rates_meta, rates):
    """Q&A: de xuat + vi sao la CR + cau hoi KH + truc khong anh huong + khong thuoc CR."""
    meta = cr["meta"]
    sh = Sheet(ws)
    full = {1: 6}
    sh.put([(1, "%s — Hỏi đáp với khách hàng" % meta["cr_id"])], "title", merge_to=full)
    sh.put([(1, "Mức ảnh hưởng chung"), (2, RISK_VI.get(meta["overall_risk"], meta["overall_risk"]))],
           merge_to={2: 6}, bold=(1,))
    sh.put([(1, "Đề xuất"), (2, meta["recommendation"])], merge_to={2: 6}, bold=(1,))

    sh.gap()
    sh.put([(1, "1. Vì sao đây là Change Request")], "banner", merge_to=full)
    sh.put([(1, "Hạng mục"), (2, "Yêu cầu"), (3, "Hệ thống hiện tại (baseline)"),
            (4, "Tiêu chí"), (5, "Giải trình")], "header", merge_to={5: 6})
    for j in cr["justification"]:
        crit = "\n".join("%s · %s" % (c, CRIT_SHORT.get(c, S.CR_CRITERIA.get(c, ""))) for c in j["criteria"])
        why = C.txt(j.get("why"))
        if C.txt(j.get("not_feedback")):
            why += "\nKhông phải lỗi: " + C.txt(j["not_feedback"])
        sh.put([(1, j.get("item")), (2, j.get("request")),
                (3, "%s: “%s”" % (j.get("baseline_ref"), C.txt(j.get("baseline_quote")))),
                (4, crit), (5, why)], merge_to={5: 6}, bold=(1,))

    sh.gap()
    sh.put([(1, "2. Câu hỏi cần khách hàng trả lời")], "banner", merge_to=full)
    sh.put([(1, "Mã"), (2, "Câu hỏi"), (3, "Vì sao cần"), (4, "Trục"), (5, "Người trả lời"),
            (6, "Trạng thái")], "header")
    for q in cr["questions"] or [{"id": "—", "question": "Không có câu hỏi"}]:
        sh.put([(1, q.get("id")), (2, q.get("question")), (3, q.get("why")), (4, q.get("axis")),
                (5, q.get("owner")), (6, q.get("status"))], bold=(1,), center=(4, 6))

    sh.gap()
    sh.put([(1, "3. Trục không ảnh hưởng")], "banner", merge_to=full)
    sh.put([(1, "Trục"), (2, "Lý do")], "header", merge_to={2: 6})
    for x in cr["no_impact_axes"] or [{"axis": "", "reason": "Không có — cả 6 trục đều có hạng mục"}]:
        ax = x.get("axis")
        sh.put([(1, "%s (%s)" % (C.AXIS_VI.get(ax, ax), ax) if ax else "—"), (2, x.get("reason"))],
               merge_to={2: 6})

    if cr["not_cr"]:
        sh.gap()
        sh.put([(1, "4. Không thuộc CR")], "banner", merge_to=full)
        sh.put([(1, "Hạng mục"), (2, "Phân loại"), (3, "Lý do"), (6, "Baseline Ref")], "header",
               merge_to={3: 5})
        for n in cr["not_cr"]:
            sh.put([(1, n.get("item")), (2, "Lỗi hệ thống (BUG)" if n.get("label") == "BUG"
                                          else "Câu hỏi (QUESTION)"),
                    (3, n.get("reason")), (6, n.get("baseline_ref"))], merge_to={3: 5}, bold=(1,))
    sheet_setup(ws, 6, sh.r)


def build_impact(ws, impacts, rates, empty_note=""):
    """1 sheet hang muc (Screen / API / Database / Figma): header CR_IMPACT_COLS + dong Tong MD."""
    import bd_styles as BD
    from openpyxl.utils import get_column_letter
    cols = S.CR_IMPACT_COLS
    for i, h in enumerate(cols, start=1):
        BD.header(ws.cell(row=1, column=i, value=h))
        ws.column_dimensions[get_column_letter(i)].width = WIDTH_IMPACT.get(h, 14)
    ws.row_dimensions[1].height = 30
    center = {"No", "Impact ID", "Trục", "Loại", "Xung đột", "Rủi ro", "Mã đơn giá", "Số lượng", "MD"}
    total = 0.0
    for n, imp in enumerate(impacts, start=1):
        r = rates[imp["rate_code"]]
        md = C.md_of(imp, rates)
        total += md
        q = C.qty_of(imp)
        note = "%s (%s MD/%s × %s)" % (r.get("item"), C.fmt_md(float(r["md"])), r.get("unit"), C.fmt_md(q))
        if imp.get("option") is True:
            note = "拡張案 · " + note
        if C.txt(imp.get("md_note")):
            note += "; " + C.txt(imp["md_note"])
        vals = {"No": n, "Impact ID": imp["id"], "Trục": imp["axis"], "Loại": imp["change_type"],
                "Baseline Ref": C.txt(imp.get("baseline_ref")) or "—", "Hạng mục": imp.get("item"),
                "Nội dung thay đổi": imp.get("change"), "Vì sao phải sửa": imp.get("why_change"),
                "Ảnh hưởng tới hiện tại": imp.get("impact_on_current"), "Xung đột": imp["conflict"],
                "Chi tiết xung đột": C.txt(imp.get("conflict_detail")) or "—", "Rủi ro": imp["risk"],
                "Mã đơn giá": imp["rate_code"], "Số lượng": q if q != int(q) else int(q), "MD": md,
                "Ghi chú MD": note, "Evidence": imp.get("evidence"),
                "Câu hỏi": C.txt(imp.get("question")) or "—"}
        lines = 1
        for i, h in enumerate(cols, start=1):
            c = ws.cell(row=n + 1, column=i, value=vals[h])
            BD.data(c, size=10, code=h == "Impact ID", center=h in center)
            lines = max(lines, est_lines(vals[h], WIDTH_IMPACT.get(h, 14)))
        ws.row_dimensions[n + 1].height = min(14 * lines + 4, 160)
    last = len(impacts) + 1
    if not impacts:
        c = ws.cell(row=2, column=2, value=empty_note or "Không có hạng mục")
        BD.data(c, size=10)
        last = 2
    tr = last + 1
    for i, h in enumerate(cols, start=1):
        BD.banner(ws.cell(row=tr, column=i))
    ws.cell(row=tr, column=cols.index("Impact ID") + 1, value=C.L_TOTAL_MD)
    ws.cell(row=tr, column=cols.index("MD") + 1, value=round(total, 2))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), last)
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    return round(total, 2)


def _style_of(ws, row, ncol=17):
    import copy
    out = {}
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        out[c] = (copy.copy(cell.font), copy.copy(cell.fill), copy.copy(cell.border),
                  copy.copy(cell.alignment), cell.number_format, ws.cell(row=row, column=c).value)
    return out, ws.row_dimensions[row].height


def _apply_style(ws, row, st, height=None):
    styles, h = st
    for c, (font, fill, border, align, nf, _v) in styles.items():
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border, cell.alignment, cell.number_format = font, fill, border, align, nf
    ws.row_dimensions[row].height = height if height is not None else h


def build_estimation(ws, cr, rates, rates_meta, est):
    """Dien sheet Estimation (da copy tu template): section = cr_item, dong = hang muc, M = MD, K/L/N/O/P = cong thuc."""
    import datetime
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import FormulaRule
    meta = cr["meta"]
    R = {k: int(est[k]) for k in ("row_header", "row_section", "row_item", "row_total", "row_month", "row_notes")}
    st_sec, st_item = _style_of(ws, R["row_section"]), _style_of(ws, R["row_item"])
    st_tot, st_mon = _style_of(ws, R["row_total"]), _style_of(ws, R["row_month"])
    notes = [_style_of(ws, r) for r in range(R["row_notes"], R["row_notes"] + 12)]
    note_merges = [m for m in list(ws.merged_cells.ranges) if m.min_row >= R["row_total"]]
    import copy
    m_fill = copy.copy(ws.cell(row=R["row_item"], column=13).fill)
    dv_src = [(dv.type, dv.formula1, str(dv.sqref).split(":")[0].rstrip("0123456789"))
              for dv in ws.data_validations.dataValidation]
    for m in note_merges:
        ws.unmerge_cells(str(m))
    ws.data_validations.dataValidation = []
    ws.conditional_formatting = type(ws.conditional_formatting)()
    for r in range(R["row_section"], ws.max_row + 1):
        for c in range(1, 18):
            ws.cell(row=r, column=c).value = None
    # placeholder dau trang
    repl = {"{{CUSTOMER}}": C.txt(meta.get("customer")) or C.txt(meta.get("requested_by")),
            "{{TITLE}}": "%s — %s" % (meta["cr_id"], meta["cr_title"]),
            "{{DATE}}": datetime.date.today().strftime("%Y/%m/%d"),
            "{{VALID_UNTIL}}": C.txt(meta.get("valid_until")) or "—"}
    for row in ws.iter_rows(min_row=1, max_row=R["row_header"] - 1):
        for c in row:
            if isinstance(c.value, str) and c.value in repl:
                c.value = repl[c.value]
    rk, rl, rn, ro = (est["ratio_K"], est["ratio_L"], est["ratio_N"], est["ratio_O"])
    req = {j["item"]: C.txt(j.get("request")) for j in cr["justification"]}
    order = [j["item"] for j in cr["justification"]]
    r = R["row_section"]
    first = None
    for item in order:
        rows = [i for i in cr["impacts"] if C.txt(i.get("cr_item")) == item]
        if not rows:
            continue
        _apply_style(ws, r, st_sec)
        ws.cell(row=r, column=2, value=" ◆ %s — %s" % (item, req.get(item, "")[:90]))
        r += 1
        for imp in rows:
            _apply_style(ws, r, st_item)
            md = C.md_of(imp, rates)
            rate = rates[imp["rate_code"]]
            note = []
            if C.txt(imp.get("question")):
                note.append("Chờ " + C.txt(imp["question"]))
            if imp.get("conflict") != "No" and C.txt(imp.get("conflict_detail")):
                note.append("Xung đột %s: %s" % (imp["conflict"], C.txt(imp["conflict_detail"])))
            if C.txt(imp.get("md_note")):
                note.append(C.txt(imp["md_note"]))
            vals = {2: imp["id"], 3: "%s (%s)" % (C.AXIS_VI.get(imp["axis"], imp["axis"]), imp["change_type"]),
                    4: C.txt(imp.get("baseline_ref")) or "—", 5: imp.get("item"), 6: imp.get("change"),
                    7: "\n".join(note) or None,
                    8: "Phát triển mới" if imp["change_type"] == "NEW" else "Tái sử dụng ",
                    9: None, 10: bool(imp.get("option")), 13: md,
                    11: "=M%d*%s" % (r, rk), 12: "=M%d*%s" % (r, rl), 14: "=M%d*%s" % (r, rn),
                    15: "=SUM(K%d:N%d)*%s" % (r, r, ro), 16: "=SUM(K%d:O%d)" % (r, r),
                    17: "%s × %s = %s MD (%s MD/%s)" % (imp["rate_code"], C.fmt_md(C.qty_of(imp)),
                                                          C.fmt_md(md), C.fmt_md(float(rate["md"])), rate.get("unit"))}
            for c, v in vals.items():
                ws.cell(row=r, column=c, value=v)
            lines = max(est_lines(vals[6], 70), est_lines(vals[5], 34), est_lines(vals[7], 34))
            ws.row_dimensions[r].height = min(max(30, 15 * lines + 6), 240)
            first = first or r
            r += 1
    last = r - 1
    tr, mr = r, r + 1
    _apply_style(ws, tr, st_tot)
    _apply_style(ws, mr, st_mon)
    for c, v in st_tot[0].items():
        if isinstance(v[5], str) and not v[5].startswith("="):
            ws.cell(row=tr, column=c, value=v[5])
    for c, v in st_mon[0].items():
        if isinstance(v[5], str) and not v[5].startswith("="):
            ws.cell(row=mr, column=c, value=v[5])
    for col in "KLMNOP":
        ws["%s%d" % (col, tr)] = "=SUM(%s%d:%s%d)" % (col, first, col, last)
        ws["%s%d" % (col, mr)] = "=%s%d/%s" % (col, tr, C.fmt_md(est["md_per_month"]))
    ws.merge_cells(start_row=tr, start_column=2, end_row=mr, end_column=3)
    nr = mr + 2
    for k, st in enumerate(notes):
        _apply_style(ws, nr + k, st)
        for c, v in st[0].items():
            if isinstance(v[5], str) and v[5] != "{{ASSUMPTIONS}}":
                ws.cell(row=nr + k, column=c, value=v[5])
    for m in note_merges:
        if m.min_row >= R["row_notes"]:
            off = nr - R["row_notes"]
            ws.merge_cells(start_row=m.min_row + off, start_column=m.min_col,
                           end_row=m.max_row + off, end_column=m.max_col)
    assume = ["実装 = MD theo bảng đơn giá %s v%s (mã đơn giá × số lượng, cột Q). %s" % (
                  rates_meta.get("status"), rates_meta.get("version"), C.txt(rates_meta.get("scope_note"))),
              "要件定義 = 実装 × %s · UI・UXデザイン = 実装 × %s · テスト = 実装 × %s · 管理 = (要件定義…テスト) × %s · "
              "人月 = 人日 / %s (hệ số của template Estimation)." % (rk, rl, rn, ro, C.fmt_md(est["md_per_month"])),
              "option = TRUE: hạng mục thuộc 拡張案 — chỉ tính khi KH chọn phương án mở rộng.",
              "Các dòng có 'Chờ CQ-xxx' ở cột QA/Note phụ thuộc câu trả lời của KH (sheet Q&A)."]
    ws.cell(row=nr + 3, column=2, value=assume[0])
    ws.cell(row=nr + 5, column=2, value="\n".join(assume[1:]))
    ws.row_dimensions[nr + 3].height = 45
    # validation + conditional format cho toan bo dong hang muc
    rng = lambda col: "%s%d:%s%d" % (col, first, col, last)  # noqa: E731
    for typ, f1, col in dv_src:
        dv = DataValidation(type=typ, formula1=f1, allow_blank=True)
        dv.add(rng(col))
        ws.add_data_validation(dv)
    ws.conditional_formatting.add(rng("M"), FormulaRule(formula=["LEN(TRIM(M%d))=0" % first], fill=m_fill))
    ws.print_area = "A1:Q%d" % (nr + 10)
    return {"row_header": R["row_header"], "row_total": tr, "row_month": mr, "first": first, "last": last,
            "md_per_month": est["md_per_month"]}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cr-json", required=True)
    ap.add_argument("--rates", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--template", default=None, help="mac dinh %s" % S.ESTIMATION_TEMPLATE_PATH)
    a = ap.parse_args()
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("Thieu openpyxl. Chay: pip install openpyxl", file=sys.stderr)
        return 2
    try:
        cr = C.load_json(a.cr_json)
        rates_meta, rates = C.load_rates(a.rates)
    except (OSError, ValueError) as e:
        print("Khong doc duoc input: %s" % e, file=sys.stderr)
        return 1
    err = C.validate_cr(cr, rates)
    if not err and not cr.get("justification"):
        err.append("justification rong — phai giai trinh vi sao la CR")
    tpl = C.find_template(a.template, KIT)
    if not tpl:
        err.append("khong thay template Estimation (%s) — boc bang extract-estimation-template.py"
                   % S.ESTIMATION_TEMPLATE_PATH)
    if err:
        print("cr.json KHONG hop le — khong ghi file:", file=sys.stderr)
        for e in err:
            print("  - " + e, file=sys.stderr)
        return 1
    for k in ("not_cr", "no_impact_axes", "questions", "impacts"):
        cr.setdefault(k, [])
    try:
        wb, est = C.load_estimation_template(tpl)
    except (OSError, ValueError) as e:
        print("Template Estimation loi: %s" % e, file=sys.stderr)
        return 1
    del wb["_meta"]
    ws_est = wb["Estimation"]
    ws_est.title = S.CR_SHEET_ESTIMATION
    pos = build_estimation(ws_est, cr, rates, rates_meta, est)
    build_summary(wb.create_sheet(S.CR_SHEET_SUMMARY, 0), cr, rates_meta, rates, pos)
    total = 0.0
    noi = {x.get("axis"): x.get("reason") for x in cr["no_impact_axes"]}
    for name, axes in S.CR_AXIS_SHEETS:
        rows = [i for i in cr["impacts"] if i["axis"] in axes]
        total += build_impact(wb.create_sheet(name), rows, rates,
                              "; ".join("%s: %s" % (x, noi[x]) for x in axes if x in noi))
    build_qa(wb.create_sheet(S.CR_SHEET_QA), cr, rates_meta, rates)
    wb.active = 0
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    wb.save(a.out)
    grand = sum(C.md_of(i, rates) or 0.0 for i in cr["impacts"])
    print("Da ghi %s — %d hang muc · 実装 %s MD (don gia %s v%s) · sheet %s" % (
        a.out, len(cr["impacts"]), C.fmt_md(grand), rates_meta.get("status"), rates_meta.get("version"),
        " · ".join(wb.sheetnames)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
