#!/usr/bin/env python3
"""Boc 1 file 見積書 (estimate) cua cong ty thanh templates/template_estimation.xlsx cho sheet Estimation (Luong 2).

  python3 extract-estimation-template.py --src <見積書.xlsx> [--sheet "見積書(VN)"] \
      [--out templates/template_estimation.xlsx]

Giu nguyen: letterhead (ten / dia chi cong ty, logo), tieu de 御見積書, khoi ■件名 / ■見積日 / ■有効期限,
2 dong header bang (No. · 項目(大/中/小) · 改修内容 · QA/Note · Tái sử dụng/phát triển mới · step · option ·
工数(人日) 要件定義 / UI・UXデザイン / 実装 / テスト / 管理 / 小計), style, do rong cot, data validation,
conditional formatting, khoi ■備考 / ◆前提条件.

Bo di (du lieu cua 1 du an cu the): ten khach hang, ten du an, ngay, tung dong hang muc, so tien. Thay bang
placeholder {{CUSTOMER}} {{TITLE}} {{DATE}} {{VALID_UNTIL}} {{SECTION}} {{ASSUMPTIONS}}.

He so cong doan doc tu CONG THUC cua dong hang muc dau tien trong file nguon (VD K=M*0.2, N=M*0.5,
O=SUM(K:N)*0.2) va 人月 = 人日/20 -> ghi vao sheet an `_meta`. build-cr-impact.py chi doc template —
PM doi he so thi sua template (hoac file nguon roi boc lai), KHONG sua trong code.

Layout template (dong):  1..21 = header nguon · 22 = dong section mau · 23 = dong hang muc mau
                         24 = 合計(人日) · 25 = 人月 · 27.. = ■備考 / ◆前提条件
Exit 0 = OK · 1 = file nguon khong dung dang (khong thay header / cong thuc) · 2 = thieu openpyxl.
"""
import argparse
import copy
import os
import re
import sys

HDR_ROW_TEXT = "No."
PHASE_COLS = {"K": "要件定義", "L": "UI・UXデザイン", "M": "実装", "N": "テスト", "O": "管理", "P": "小計"}
OUT_ROW = {"section": 22, "item": 23, "total": 24, "month": 25, "notes": 27}


def find_rows(ws):
    """-> (row header 1 'No.', row section dau, row hang muc dau, row 合計, row 人月, row ■備考)."""
    hdr = sec = item = tot = month = notes = None
    for r in range(1, ws.max_row + 1):
        b = str(ws.cell(row=r, column=2).value or "").strip()
        k = ws.cell(row=r, column=11).value
        if hdr is None and b == HDR_ROW_TEXT:
            hdr = r
        elif hdr and sec is None and b.startswith("◆"):
            sec = r
        elif sec and item is None and isinstance(k, str) and k.startswith("=M"):
            item = r
        elif item and tot is None and b.startswith("合計"):
            tot = r
            month = r + 1
        elif tot and notes is None and b.startswith("■"):
            notes = r
    return hdr, sec, item, tot, month, notes


def ratios(ws, item, month):
    """Doc he so tu cong thuc dong hang muc dau + dong 人月."""
    out = {}
    pat = {"K": r"^=M%d\*([\d.]+)$", "L": r"^=M%d\*([\d.]+)$", "N": r"^=M%d\*([\d.]+)$",
           "O": r"^=SUM\(K%d:N%d\)\*([\d.]+)$"}
    for col, p in pat.items():
        v = str(ws["%s%d" % (col, item)].value or "").replace(" ", "")
        m = re.match(p % ((item,) * p.count("%d")), v)
        if not m:
            raise ValueError("cong thuc %s%d = %r khong dung dang %s" % (col, item, v, p))
        out["ratio_" + col] = float(m.group(1))
    v = str(ws["P%d" % month].value or "").replace(" ", "")
    m = re.match(r"^=P\d+/([\d.]+)$", v)
    if not m:
        raise ValueError("cong thuc P%d (人月) = %r khong dung dang =P<n>/<so ngay>" % (month, v))
    out["md_per_month"] = float(m.group(1))
    return out


def copy_row(src, dst, rs, rd, max_col):
    dst.row_dimensions[rd].height = src.row_dimensions[rs].height
    for c in range(1, max_col + 1):
        a, b = src.cell(row=rs, column=c), dst.cell(row=rd, column=c)
        b.value = a.value
        if a.has_style:
            copy_style(a, b)


def copy_style(a, b):
    """Copy style giua 2 workbook (khong copy _style — chi so style la cua tung workbook)."""
    b.font = copy.copy(a.font)
    b.fill = copy.copy(a.fill)
    b.border = copy.copy(a.border)
    b.alignment = copy.copy(a.alignment)
    b.number_format = a.number_format
    b.protection = copy.copy(a.protection)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--sheet", default="見積書(VN)")
    ap.add_argument("--out", default="templates/template_estimation.xlsx")
    a = ap.parse_args()
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        print("Thieu openpyxl", file=sys.stderr)
        return 2
    wb_src = load_workbook(a.src)
    if a.sheet not in wb_src.sheetnames:
        print("Khong co sheet %r trong %s (co: %s)" % (a.sheet, a.src, wb_src.sheetnames), file=sys.stderr)
        return 1
    src = wb_src[a.sheet]
    hdr, sec, item, tot, month, notes = find_rows(src)
    if not all((hdr, sec, item, tot, notes)):
        print("Khong nhan ra cau truc 見積書: header=%s section=%s item=%s 合計=%s ■備考=%s"
              % (hdr, sec, item, tot, notes), file=sys.stderr)
        return 1
    try:
        meta = ratios(src, item, month)
    except ValueError as e:
        print("He so: %s" % e, file=sys.stderr)
        return 1
    max_col = max(src.max_column, 17)

    wb = Workbook()
    ws = wb.active
    ws.title = "Estimation"
    # 1. header 1..hdr+1 giu nguyen vi tri
    for r in range(1, hdr + 2):
        copy_row(src, ws, r, r, max_col)
    shift = OUT_ROW["section"] - sec
    # 2. dong section / item / tong / nhan cong thang
    copy_row(src, ws, sec, OUT_ROW["section"], max_col)
    copy_row(src, ws, item, OUT_ROW["item"], max_col)
    copy_row(src, ws, tot, OUT_ROW["total"], max_col)
    copy_row(src, ws, month, OUT_ROW["month"], max_col)
    n_notes = 0
    for r in range(notes, src.max_row + 1):
        if not any(src.cell(row=r, column=c).has_style or src.cell(row=r, column=c).value
                   for c in range(2, 8)):
            if r > notes + 12:
                break
        copy_row(src, ws, r, OUT_ROW["notes"] + (r - notes), max_col)
        n_notes = r - notes + 1
    # 3. xoa du lieu du an cu the -> placeholder
    ws["B8"].value = "{{CUSTOMER}}"
    ws["C10"].value = "{{TITLE}}"
    ws["C11"].value = "{{DATE}}"
    ws["C12"].value = "{{VALID_UNTIL}}"
    for ref in ("D15", "K15"):         # so tien — kit khong tinh tien
        ws[ref].value = None
    ws["B%d" % OUT_ROW["section"]].value = "{{SECTION}}"
    ir = OUT_ROW["item"]
    for col in "BCDEFGHIJ":
        ws["%s%d" % (col, ir)].value = None
    ws["Q%d" % ir].value = None
    ws["M%d" % ir].value = 0
    ws["K%d" % ir].value = "=M%d*%s" % (ir, meta["ratio_K"])
    ws["L%d" % ir].value = "=M%d*%s" % (ir, meta["ratio_L"])
    ws["N%d" % ir].value = "=M%d*%s" % (ir, meta["ratio_N"])
    ws["O%d" % ir].value = "=SUM(K%d:N%d)*%s" % (ir, ir, meta["ratio_O"])
    ws["P%d" % ir].value = "=SUM(K%d:O%d)" % (ir, ir)
    ws.row_dimensions[ir].height = 45
    tr, mr = OUT_ROW["total"], OUT_ROW["month"]
    for col in "KLMNOP":
        ws["%s%d" % (col, tr)].value = "=SUM(%s%d:%s%d)" % (col, ir, col, ir)
        ws["%s%d" % (col, mr)].value = "=%s%d/%s" % (col, tr, int(meta["md_per_month"]))
    # cot Q: can cu MD (kit them — file nguon de trong)
    ws["Q%d" % hdr].value = "Căn cứ 実装 (mã đơn giá × số lượng)"
    copy_style(ws["K%d" % hdr], ws["Q%d" % hdr])
    for r in range(OUT_ROW["notes"] + 1, OUT_ROW["notes"] + n_notes):
        for c in range(2, 8):
            if ws.cell(row=r, column=c).value and r > OUT_ROW["notes"] + 1:
                ws.cell(row=r, column=c).value = None
    ws["B%d" % (OUT_ROW["notes"] + 3)].value = "{{ASSUMPTIONS}}"
    # 4. merge: header giu, phan duoi doi dong
    for rng in src.merged_cells.ranges:
        if rng.max_row <= hdr + 1:
            ws.merge_cells(str(rng))
        elif rng.min_row >= notes:
            ws.merge_cells(start_row=rng.min_row - notes + OUT_ROW["notes"], start_column=rng.min_col,
                           end_row=rng.max_row - notes + OUT_ROW["notes"], end_column=rng.max_col)
        elif rng.min_row == tot:
            ws.merge_cells(start_row=tr, start_column=rng.min_col, end_row=mr, end_column=rng.max_col)
    # 5. do rong cot, freeze, view
    for k, d in src.column_dimensions.items():
        if d.width:
            ws.column_dimensions[k].width = d.width
    ws.freeze_panes = "A%d" % (hdr + 2)
    ws.sheet_view.showGridLines = bool(src.sheet_view.showGridLines)
    # 6. validation + conditional format (lay danh sach gia tri tu nguon)
    for dv in src.data_validations.dataValidation:
        cols = {str(r).split(":")[0].rstrip("0123456789") for r in str(dv.sqref).split()}
        for col in cols:
            nv = DataValidation(type=dv.type, formula1=dv.formula1, allow_blank=True)
            nv.add("%s%d:%s%d" % (col, ir, col, ir))
            ws.add_data_validation(nv)
    ws.conditional_formatting.add("M%d:M%d" % (ir, ir),
                                  FormulaRule(formula=["LEN(TRIM(M%d))=0" % ir],
                                              fill=copy.copy(src["M%d" % item].fill)))
    # 7. logo / anh letterhead
    for im in getattr(src, "_images", []):
        if im.anchor._from.row < hdr:
            ws.add_image(copy.copy(im))
    # 8. _meta (an)
    mt = wb.create_sheet("_meta")
    rows = [("source_file", "見積書 (ten file nguon khong luu — co ten khach hang)"), ("source_sheet", a.sheet),
            ("row_header", hdr), ("row_section", OUT_ROW["section"]), ("row_item", ir),
            ("row_total", tr), ("row_month", mr), ("row_notes", OUT_ROW["notes"]),
            ("ratio_K_要件定義", meta["ratio_K"]), ("ratio_L_UI", meta["ratio_L"]),
            ("ratio_N_テスト", meta["ratio_N"]), ("ratio_O_管理", meta["ratio_O"]),
            ("md_per_month", meta["md_per_month"]),
            ("note", "M = 実装 (MD tu bang don gia). K/L/N = M x he so; O = SUM(K:N) x he so; P = SUM(K:O).")]
    for k, v in rows:
        mt.append([k, v])
    mt.sheet_state = "hidden"
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    wb.save(a.out)
    print("Da ghi %s — header dong %d, he so 要件定義 %s · UI %s · テスト %s · 管理 %s · 人月 = 人日/%s"
          % (a.out, hdr, meta["ratio_K"], meta["ratio_L"], meta["ratio_N"], meta["ratio_O"],
             int(meta["md_per_month"])))
    return 0


if __name__ == "__main__":
    sys.exit(main())
