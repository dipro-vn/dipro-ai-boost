#!/usr/bin/env python3
"""GATE V-CR — CR Impact (Luong 2): cr.json + CR-<id>_Impact.xlsx (7 sheet inv_schema.CR_SHEETS).

  python3 verify-cr-impact.py <ver>/CR-001_Impact.xlsx --cr-json <ver>/_internal/cr.json \
      --baseline <outputs>/verK_... [--rates .claude/config/md-unit-rates.json] \
      [--other-cr <ver khac>/_internal/cr.json ...] [--out <ver>/_internal/gates/v-cr.md]

Baseline Ref: SC-/F-/API-/EXT-/WEB- · table:orders · column:orders.status · DS:<token> · DS-component:<Comp>
· DS:WEB-01:<token> / DS-component:WEB-01:<Comp> (chon site) · — ; nhieu ref cach nhau ';'.
Design system baseline: 04_DesignSystem/project/ (1 bo) hoac 04_DesignSystem/WEB-xx/project/ (moi site 1 bo).
--rates mac dinh: <kit>/.claude/config/md-unit-rates.json. --template mac dinh templates/template_estimation.xlsx
(he so cong doan cua Estimation doc tu _meta cua template). Cong thuc Excel duoc kiem bang TEXT (khong can
Excel / LibreOffice tinh lai). Tom tat impact in ra stderr.
"""
import argparse
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import inv_schema as S  # noqa: E402
import cr_common as C  # noqa: E402
from gate_report import Gate  # noqa: E402

KIT = os.path.abspath(os.path.join(HERE, "..", "..", "..", ".."))


def read_axis_sheet(ws):
    hdr = [C.txt(c.value) for c in ws[1]]
    rows, total = [], None
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        d = {h: r[k] if k < len(r) else None for k, h in enumerate(hdr) if h}
        iid = C.txt(d.get("Impact ID"))
        if iid == C.L_TOTAL_MD:
            total = d.get("MD")
        elif re.match(S.CR_IMPACT_ID, iid):
            d["__row__"] = i
            rows.append(d)
    return hdr, rows, total


def read_xlsx(path):
    """-> dict: names, axis {sheet: (hdr, rows, total)}, summary rows (gia tri + cell), estimation rows, qa text."""
    from openpyxl import load_workbook
    wb = load_workbook(path)                      # cong thuc giu dang text
    out = {"names": wb.sheetnames, "axis": {}, "summ": [], "summ_links": [], "est": None, "qa": set()}
    for name, _axes in S.CR_AXIS_SHEETS:
        if name in wb.sheetnames:
            out["axis"][name] = read_axis_sheet(wb[name])
    if S.CR_SHEET_SUMMARY in wb.sheetnames:
        ws = wb[S.CR_SHEET_SUMMARY]
        out["summ"] = [list(r) for r in ws.iter_rows(values_only=True)]
        out["summ_links"] = [(C.txt(c.value), c.hyperlink.location if c.hyperlink and c.hyperlink.location
                              else (c.hyperlink.target if c.hyperlink else ""))
                             for r in ws.iter_rows() for c in r if c.hyperlink]
    if S.CR_SHEET_ESTIMATION in wb.sheetnames:
        ws = wb[S.CR_SHEET_ESTIMATION]
        items, tot = [], None
        for r in range(1, ws.max_row + 1):
            b = C.txt(ws.cell(row=r, column=2).value)
            if re.match(S.CR_IMPACT_ID, b):
                items.append({"row": r, "id": b, "J": ws.cell(row=r, column=10).value,
                              **{col: ws["%s%d" % (col, r)].value for col in "KLMNOP"}})
            elif b.startswith("合計") and tot is None and items:
                tot = r
        out["est"] = {"items": items, "total_row": tot,
                      "total": {col: ws["%s%d" % (col, tot)].value for col in "KLMNOP"} if tot else {},
                      "month": {col: ws["%s%d" % (col, tot + 1)].value for col in "KLMNOP"} if tot else {}}
    if S.CR_SHEET_QA in wb.sheetnames:
        out["qa"] = {C.txt(r[0]) for r in wb[S.CR_SHEET_QA].iter_rows(values_only=True) if r and r[0]}
    return out


def num(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def summary_value(summ, label):
    for r in summ:
        if r and C.txt(r[0]) == label and len(r) > 1:
            return r[1]
    return None


def summary_objects(summ):
    """Bang 'Doi tuong x NEW/UPD/DEL/IMPACT/Tong/MD' tren Summary -> [(doi tuong, [6 so])]."""
    out, start = [], None
    for i, r in enumerate(summ):
        if r and C.txt(r[0]) == C.L_OBJ_HDR and [C.txt(x) for x in r[1:5]] == S.CR_CHANGE_TYPE:
            start = i
            break
    if start is None:
        return None
    for r in summ[start + 1:]:
        a = C.txt(r[0] if r else "")
        if not a:
            break
        out.append((a, [num(x) for x in (list(r[1:7]) + [None] * 6)[:6]]))
        if a == C.L_TOTAL:
            break
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("impact")
    ap.add_argument("--cr-json", required=True)
    ap.add_argument("--baseline", required=True)
    ap.add_argument("--rates", default=os.path.join(KIT, S.MD_RATES_PATH))
    ap.add_argument("--other-cr", action="append", default=[])
    ap.add_argument("--template", default=None)
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("Thieu openpyxl. Chay: pip install openpyxl", file=sys.stderr)
        return 2

    cr = C.load_json(a.cr_json)
    rates_meta, rates = C.load_rates(a.rates)
    B = C.load_baseline(a.baseline)
    X = read_xlsx(a.impact)
    names, summ = X["names"], X["summ"]
    tpl = C.find_template(a.template, KIT)
    est_meta = C.load_estimation_template(tpl)[1] if tpl else None
    g = Gate("GATE V-CR — %s" % os.path.basename(a.impact))
    meta = cr.get("meta") or {}
    J = cr.get("justification") or []
    IMP = cr.get("impacts") or []
    NOI = cr.get("no_impact_axes") or []
    Q = cr.get("questions") or []
    base_name = os.path.basename(os.path.normpath(a.baseline))

    # 1. Workbook
    bad = []
    if names != S.CR_SHEETS:
        bad.append("sheet = %s (phai dung %s)" % (names, S.CR_SHEETS))
    for name, (hdr, _r, _t) in X["axis"].items():
        if [h for h in hdr if h] != S.CR_IMPACT_COLS:
            bad.append("header %s khac CR_IMPACT_COLS: %s" % (name, [h for h in hdr if h]))
    g.check(1, "Workbook dung %d sheet %s, header sheet hang muc dung schema" % (
        len(S.CR_SHEETS), " · ".join(S.CR_SHEETS)), bad)

    # 2. Meta
    bad = ["meta.%s rong/UNKNOWN" % k for k in S.CR_META_KEYS
           if not C.txt(meta.get(k)) or C.txt(meta.get(k)) == S.UNKNOWN]
    if C.txt(meta.get("source_type")) and meta["source_type"] not in S.CR_SOURCE_TYPE:
        bad.append("source_type=%s" % meta["source_type"])
    if C.txt(meta.get("overall_risk")) and meta["overall_risk"] not in S.CR_RISK:
        bad.append("overall_risk=%s" % meta["overall_risk"])
    if C.txt(meta.get("baseline_version")) and meta["baseline_version"] != base_name:
        bad.append("baseline_version=%s != %s" % (meta["baseline_version"], base_name))
    g.check(2, "meta du key, source_type/overall_risk hop le, baseline_version khop", bad)

    ds_warn = []

    def res(ref):
        ok, why, w = C.resolve(ref, B)
        if w:
            ds_warn.append(w)
        return ok, why

    # 3. Justification
    bad = [] if J else ["justification rong — phai giai trinh vi sao la CR"]
    for i, j in enumerate(J):
        tag = C.txt(j.get("item")) or "justification[%d]" % i
        miss = [f for f in ("item", "request", "baseline_ref", "baseline_quote", "why", "not_feedback")
                if not C.txt(j.get(f))]
        if miss:
            bad.append("%s thieu %s" % (tag, ",".join(miss)))
        crit = j.get("criteria") if isinstance(j.get("criteria"), list) else []
        if not crit:
            bad.append("%s khong co tieu chi C1..C6" % tag)
        bad += ["%s tieu chi %s khong co trong C1..C6" % (tag, c) for c in crit if c not in S.CR_CRITERIA]
        for ref in C.refs_of(j.get("baseline_ref")):
            ok, why = res(ref)
            if not ok:
                bad.append("%s %s: %s" % (tag, ref, why))
    g.check(3, "Giai trinh CR: >=1 dong, tieu chi C1..C6, baseline_ref co that + trich dan", bad)

    # 4. Phu 6 truc
    bad = []
    noi = {}
    for x in NOI:
        noi.setdefault(C.txt(x.get("axis")), []).append(x)
    for ax in noi:
        if ax not in S.CR_AXES:
            bad.append("no_impact_axes axis=%s khong hop le" % ax)
    for ax in S.CR_AXES:
        has = any(i.get("axis") == ax for i in IMP)
        if has and ax in noi:
            bad.append("%s vua co impact vua khai khong anh huong" % ax)
        elif not has and ax not in noi:
            bad.append("%s bo trong — them impact hoac no_impact_axes kem ly do" % ax)
        elif not has and len(C.txt(noi[ax][0].get("reason"))) < 15:
            bad.append("%s khong anh huong nhung ly do < 15 ky tu" % ax)
    g.check(4, "Ca 6 truc: co impact HOAC no_impact_axes (ly do >= 15 ky tu), khong ca hai", bad)

    # 5. Impact ID
    bad, seen = [], set()
    for i, imp in enumerate(IMP):
        iid = C.txt(imp.get("id"))
        if not re.match(S.CR_IMPACT_ID, iid):
            bad.append("impacts[%d] id=%r sai dang" % (i, iid))
        if iid in seen:
            bad.append("%s trung" % iid)
        seen.add(iid)
    g.check(5, "Impact ID dung dang IMP-NNN, duy nhat", bad)

    def tag(imp):
        return C.txt(imp.get("id")) or "?"

    # 6. Enum
    bad = []
    for imp in IMP:
        for f, enum in (("axis", S.CR_AXES), ("change_type", S.CR_CHANGE_TYPE),
                        ("conflict", S.CR_CONFLICT), ("risk", S.CR_RISK)):
            if C.txt(imp.get(f)) not in enum:
                bad.append("%s %s=%s" % (tag(imp), f, imp.get(f)))
    g.check(6, "Enum axis / change_type / conflict / risk", bad)

    # 7. UPD/DEL/IMPACT ref ton tai
    bad = []
    for imp in IMP:
        if imp.get("change_type") not in ("UPD", "DEL", "IMPACT"):
            continue
        refs = C.refs_of(imp.get("baseline_ref"))
        if not refs:
            bad.append("%s %s khong co baseline_ref" % (tag(imp), imp.get("change_type")))
        for ref in refs:
            ok, why = res(ref)
            if not ok:
                bad.append("%s %s: %s" % (tag(imp), ref, why))
    g.check(7, "UPD/DEL/IMPACT: moi baseline_ref ton tai trong baseline", bad)

    # 8. NEW
    bad = []
    for imp in IMP:
        if imp.get("change_type") != "NEW":
            continue
        for ref in C.refs_of(imp.get("baseline_ref")):
            low = ref.lower()
            if low.startswith("column:"):
                tc = low[7:].strip()
                if "." not in tc or tc.split(".")[0] not in B["tables"]:
                    bad.append("%s %s: bang cha khong co trong baseline" % (tag(imp), ref))
                elif tc in B["cols"] and imp.get("conflict") != "Yes":
                    bad.append("%s %s: cot da ton tai — trung ten, phai conflict=Yes" % (tag(imp), ref))
                continue
            ok, why = res(ref)
            if not ok:
                bad.append("%s %s: %s (NEW chi gan vao cha da ton tai)" % (tag(imp), ref, why))
    g.check(8, "NEW: ref la — hoac cha ton tai; cot moi trung cot cu phai conflict=Yes", bad)

    # 9. Noi dung bat buoc
    bad = []
    for imp in IMP:
        miss = [f for f in ("item", "change", "impact_on_current") if not C.txt(imp.get(f))]
        if len(C.txt(imp.get("why_change"))) < 10:
            miss.append("why_change(>=10 ky tu)")
        if miss:
            bad.append("%s thieu %s" % (tag(imp), ",".join(miss)))
    g.check(9, "Moi impact co hang muc, noi dung, vi sao phai sua, anh huong", bad)

    # 10. Don gia
    bad = []
    for imp in IMP:
        bad += C.rate_errors(imp, rates, tag(imp))
    g.check(10, "rate_code co trong bang don gia, dung truc + loai; qty > 0", bad)

    # 11. Sheet Screen / API / Database / Figma khop cr.json, MD = rate x qty
    bad = []
    for name, axes in S.CR_AXIS_SHEETS:
        hdr, xrows, xtotal = X["axis"].get(name, ([], [], None))
        want_imp = [i for i in IMP if i.get("axis") in axes]
        by_id = {C.txt(r.get("Impact ID")): r for r in xrows}
        if [C.txt(r.get("Impact ID")) for r in xrows] != [tag(i) for i in want_imp]:
            bad.append("%s: Impact ID %s != cr.json %s — chay lai build-cr-impact.py" % (
                name, list(by_id), [tag(i) for i in want_imp]))
        for imp in want_imp:
            r = by_id.get(tag(imp))
            if not r:
                continue
            loc = "%s!%s" % (name, r["__row__"])
            for col, f in (("Trục", "axis"), ("Loại", "change_type"), ("Mã đơn giá", "rate_code")):
                if C.txt(r.get(col)) != C.txt(imp.get(f)):
                    bad.append("%s %s=%s != cr.json %s" % (loc, col, r.get(col), imp.get(f)))
            if num(r.get("Số lượng")) != num(imp.get("qty")):
                bad.append("%s Số lượng=%s != qty %s" % (loc, r.get("Số lượng"), imp.get("qty")))
            want = C.md_of(imp, rates)
            if want is not None and num(r.get("MD")) != want:
                bad.append("%s %s MD=%s != %s x %s = %s" % (loc, tag(imp), r.get("MD"),
                           C.fmt_md(rates[imp["rate_code"]].get("md")), imp.get("qty"), C.fmt_md(want)))
        sub = round(sum(C.md_of(i, rates) or 0.0 for i in want_imp), 2)
        if num(xtotal) != sub:
            bad.append("%s: dong Tong MD=%s != %s" % (name, xtotal, C.fmt_md(sub)))
    m, row_tot, col_tot, grand = C.md_matrix(IMP, rates)
    g.check(11, "Sheet Screen/API/Database/Figma khop cr.json; MD = don gia x so luong; dong Tong MD dung", bad)

    # 12. Summary: cong so + so doi tuong + link Estimation
    bad = []
    if num(summary_value(summ, C.L_IMPL)) != grand:
        bad.append("%s=%s != tong MD %s" % (C.L_IMPL, summary_value(summ, C.L_IMPL), C.fmt_md(grand)))
    mn, ex = C.plan_md(IMP, rates)
    if num(summary_value(summ, C.L_MIN_PLAN)) != mn or num(summary_value(summ, C.L_EXT_PLAN)) != ex:
        bad.append("最小改修案 / 拡張案 = %s / %s != %s / %s (theo co option)" % (
            summary_value(summ, C.L_MIN_PLAN), summary_value(summ, C.L_EXT_PLAN), mn, ex))
    est = X["est"] or {}
    for label, key in ((C.L_TOTAL_PD, 0), (C.L_TOTAL_PM, 1)):
        f = C.txt(summary_value(summ, label)).replace(" ", "")
        want_row = (est.get("total_row") or 0) + key
        if f != "='%s'!P%d" % (S.CR_SHEET_ESTIMATION, want_row):
            bad.append("%s phai la cong thuc ='%s'!P%d (dang: %r)" % (label, S.CR_SHEET_ESTIMATION, want_row, f))
    objs = summary_objects(summ)
    want_objs = [(o, [d[t] for t in S.CR_CHANGE_TYPE] + [d["total"], d["md"]])
                 for o, d in C.object_stats(IMP, rates)]
    if objs is None:
        bad.append("Summary khong co bang Doi tuong x NEW/UPD/DEL/IMPACT")
    elif objs != want_objs:
        bad.append("bang Doi tuong %s != %s" % (objs, want_objs))
    targets = [loc for _t, loc in X["summ_links"]]
    for sheet in [S.CR_SHEET_ESTIMATION] + [n for n, _ in S.CR_AXIS_SHEETS] + [S.CR_SHEET_QA]:
        if not any(("'%s'!" % sheet) in (loc or "") for loc in targets):
            bad.append("Summary thieu link toi sheet %s" % sheet)
    g.check(12, "Summary: 実装 = tong MD, 総工数/人月 link Estimation, bang so doi tuong dung, du link sheet", bad)

    # 22. Estimation theo template
    bad = []
    if not est:
        bad.append("khong co sheet Estimation")
    elif not est_meta:
        bad.append("khong thay template Estimation de doi chieu he so (%s)" % S.ESTIMATION_TEMPLATE_PATH)
    else:
        got_ids = [x["id"] for x in est["items"]]
        if sorted(got_ids) != sorted(tag(i) for i in IMP) or len(set(got_ids)) != len(got_ids):
            bad.append("Estimation co dong hang muc %d (%d khac nhau) != %d impact cua cr.json" % (
                len(got_ids), len(set(got_ids)), len(IMP)))
        byid = {i["id"]: i for i in IMP}
        for x in est["items"]:
            imp, r = byid.get(x["id"]), x["row"]
            if not imp:
                bad.append("Estimation!B%d %s khong co trong cr.json" % (r, x["id"]))
                continue
            if num(x["M"]) != C.md_of(imp, rates):
                bad.append("Estimation!M%d (実装) = %s != MD %s" % (r, x["M"], C.fmt_md(C.md_of(imp, rates))))
            want = {"K": "=M%d*%s" % (r, est_meta["ratio_K"]), "L": "=M%d*%s" % (r, est_meta["ratio_L"]),
                    "N": "=M%d*%s" % (r, est_meta["ratio_N"]),
                    "O": "=SUM(K%d:N%d)*%s" % (r, r, est_meta["ratio_O"]), "P": "=SUM(K%d:O%d)" % (r, r)}
            for col, f in want.items():
                if C.txt(x[col]).replace(" ", "") != f:
                    bad.append("Estimation!%s%d = %r != %s (he so template)" % (col, r, x[col], f))
            if bool(x["J"]) != bool(imp.get("option")):
                bad.append("Estimation!J%d option=%s != cr.json %s" % (r, x["J"], bool(imp.get("option"))))
        if est["items"]:
            f0, f1 = est["items"][0]["row"], est["items"][-1]["row"]
            for col in "KLMNOP":
                if C.txt(est["total"].get(col)).replace(" ", "") != "=SUM(%s%d:%s%d)" % (col, f0, col, f1):
                    bad.append("Estimation dong 合計 %s = %r khong phu du %d..%d" % (col, est["total"].get(col), f0, f1))
                want = "=%s%d/%s" % (col, est["total_row"], C.fmt_md(est_meta["md_per_month"]))
                if C.txt(est["month"].get(col)).replace(" ", "") != want:
                    bad.append("Estimation dong 人月 %s = %r != %s" % (col, est["month"].get(col), want))
    g.check(22, "Estimation: 1 dong / hang muc, 実装 = MD, cong thuc he so dung template, 合計/人月 phu du", bad)

    # 23. cr_item / option
    bad = []
    items = {C.txt(j.get("item")) for j in J}
    for imp in IMP:
        if C.txt(imp.get("cr_item")) not in items:
            bad.append("%s cr_item=%r khong co trong justification" % (tag(imp), imp.get("cr_item")))
        if "option" in imp and not isinstance(imp.get("option"), bool):
            bad.append("%s option=%r phai true/false" % (tag(imp), imp.get("option")))
    used = {C.txt(i.get("cr_item")) for i in IMP}
    bad += ["justification %s khong co hang muc impact nao" % x for x in sorted(items - used)]
    g.check(23, "Moi impact gan cr_item co that; moi item CR co >= 1 impact; option la bool", bad)

    # 24. Q&A
    bad = []
    qa = X["qa"]
    bad += ["Q&A thieu giai trinh %s" % C.txt(j.get("item")) for j in J if C.txt(j.get("item")) not in qa]
    bad += ["Q&A thieu cau hoi %s" % C.txt(q.get("id")) for q in Q if C.txt(q.get("id")) not in qa]
    bad += ["Q&A thieu muc khong thuoc CR %s" % C.txt(n.get("item")) for n in cr.get("not_cr") or []
            if C.txt(n.get("item")) not in qa]
    g.check(24, "Q&A: du giai trinh CR, cau hoi KH, muc khong thuoc CR", bad)

    q_ids = [C.txt(q.get("id")) for q in Q]
    q_set = set(q_ids)

    def has_q(imp):
        return any(x in q_set for x in S.split_ids(imp.get("question")))

    # 13. Conflict
    bad = []
    for imp in IMP:
        if imp.get("conflict") != "Yes":
            continue
        if not C.txt(imp.get("conflict_detail")) or C.txt(imp.get("conflict_detail")) in C.NONE_REFS:
            bad.append("%s conflict=Yes thieu conflict_detail" % tag(imp))
        if not (has_q(imp) or C.txt(imp.get("md_note"))):
            bad.append("%s conflict=Yes khong co CQ hoac md_note" % tag(imp))
    g.check(13, "Conflict=Yes: co chi tiet + (CQ hoac md_note)", bad)

    # 14. High risk
    bad = ["%s risk=High khong co CQ / md_note" % tag(imp) for imp in IMP
           if imp.get("risk") == "High" and not (has_q(imp) or C.txt(imp.get("md_note")))]
    g.check(14, "Risk=High: co cau hoi hoac giai phap (question / md_note)", bad)

    # 15. CQ
    bad = ["questions %r sai dang CQ-NNN" % x for x in q_ids if not re.match(S.CR_QUESTION_ID, x)]
    bad += ["%s trung" % d for d in sorted({x for x in q_ids if x and q_ids.count(x) > 1})]
    for q in Q:
        if not C.txt(q.get("question")):
            bad.append("%s thieu noi dung cau hoi" % q.get("id"))
    for imp in IMP:
        bad += ["%s question %s khong co trong questions" % (tag(imp), x)
                for x in S.split_ids(imp.get("question")) if x not in q_set]
    g.check(15, "CQ dung dang, duy nhat, moi question tro toi CQ co that", bad)

    # 16. Loai ref theo truc
    def axis_imps(ax, types):
        return [i for i in IMP if i.get("axis") == ax and i.get("change_type") in types]
    bad = []
    for imp in axis_imps("Screen", ("UPD", "DEL", "IMPACT")):
        refs = C.refs_of(imp.get("baseline_ref"))
        if not refs or any(not x.startswith("SC-") for x in refs):
            bad.append("%s ref=%s (Screen UPD/DEL/IMPACT phai la SC-)" % (tag(imp), imp.get("baseline_ref")))
    g.check("16a", "Screen UPD/DEL/IMPACT tham chieu SC-", bad)
    bad = ["%s ref=%s (chi table:/column:/—)" % (tag(i), x) for i in axis_imps("DB", S.CR_CHANGE_TYPE)
           for x in C.refs_of(i.get("baseline_ref")) if not x.lower().startswith(("table:", "column:"))]
    g.check("16b", "DB tham chieu table:/column:/—", bad)
    bad = ["%s ref=%s (chi EXT-/—)" % (tag(i), x) for i in axis_imps("ThirdParty", S.CR_CHANGE_TYPE)
           for x in C.refs_of(i.get("baseline_ref")) if not x.startswith("EXT-")]
    g.check("16c", "ThirdParty tham chieu EXT-/—", bad)
    bad = ["%s mockup khong bam DS:/DS-component: nao" % tag(i) for i in axis_imps("Mockup", ("NEW", "UPD"))
           if not any(x.lower().startswith(("ds:", "ds-component:")) for x in C.refs_of(i.get("baseline_ref")))]
    title = "Mockup NEW/UPD bam design system baseline (DS:/DS-component:)"
    if B["has_ds"]:
        g.check("16d", title, bad)
    else:
        g.check("16d", title + " — baseline khong co 04_DesignSystem", bad, level="WARN")

    # 17. Evidence
    bad = [tag(i) for i in IMP if not C.txt(i.get("evidence"))]
    g.check(17, "Evidence (trich CR / file baseline) cho moi impact", bad)

    # 18. DS mo ho
    g.check(18, "Ref DS khong mo ho giua cac site (DS:WEB-xx:<token>)", sorted(set(ds_warn)), level="WARN")

    # 19. Don gia da duyet chua
    st = C.txt(rates_meta.get("status")).upper()
    if st == "APPROVED":
        g.ok(19, "Bang don gia APPROVED", "v%s" % rates_meta.get("version"))
    else:
        g.warn(19, "Bang don gia APPROVED", "status=%s v%s — MD la uoc luong so bo, PM/Tech Lead phai duyet"
               % (st or "?", rates_meta.get("version")))

    # 20. CR-vs-CR
    mine = {}
    for imp in IMP:
        if imp.get("change_type") in ("UPD", "DEL"):
            for x in C.refs_of(imp.get("baseline_ref")):
                mine.setdefault(x.lower(), []).append(tag(imp))
    overlap = []
    for other in a.other_cr:
        try:
            o = C.load_json(other)
        except (OSError, ValueError) as e:
            overlap.append("%s: khong doc duoc (%s)" % (other, e))
            continue
        ocr = C.txt((o.get("meta") or {}).get("cr_id")) or other
        for imp in o.get("impacts") or []:
            if imp.get("change_type") not in ("UPD", "DEL"):
                continue
            for x in C.refs_of(imp.get("baseline_ref")):
                if x.lower() in mine:
                    overlap.append("%s: %s (%s) <-> %s %s" % (x, ",".join(mine[x.lower()]),
                                   meta.get("cr_id", "?"), ocr, imp.get("id")))
    if a.other_cr:
        g.check(20, "Xung dot CR-vs-CR (cung baseline_ref bi UPD/DEL)", overlap, level="WARN")
    else:
        g.ok(20, "Xung dot CR-vs-CR", "khong truyen --other-cr")

    # 21. Khong thuoc CR
    bad = []
    for n in cr.get("not_cr") or []:
        if C.txt(n.get("label")) not in S.CR_NOT_CR_LABEL:
            bad.append("%s label=%s" % (n.get("item"), n.get("label")))
        if not C.txt(n.get("reason")):
            bad.append("%s thieu reason" % n.get("item"))
    g.check(21, "not_cr: label BUG/QUESTION + ly do", bad)

    # Tom tat impact
    print("IMPACT %s (baseline %s):" % (meta.get("cr_id", "?"), base_name), file=sys.stderr)
    for ax in S.CR_AXES:
        cnt = {t: sum(1 for i in IMP if i.get("axis") == ax and i.get("change_type") == t)
               for t in S.CR_CHANGE_TYPE}
        print("  %-11s %s · %s MD" % (ax, " · ".join("%s %d" % kv for kv in cnt.items()),
                                       C.fmt_md(row_tot[ax])), file=sys.stderr)
    print("  Tong MD: %s (don gia %s) · Conflict=Yes: %d · Risk=High: %d · CQ: %d" % (
        C.fmt_md(grand), st or "?", sum(i.get("conflict") == "Yes" for i in IMP),
        sum(i.get("risk") == "High" for i in IMP), len(Q)), file=sys.stderr)

    if a.out:
        os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    return g.emit(a.out)


if __name__ == "__main__":
    sys.exit(main())
