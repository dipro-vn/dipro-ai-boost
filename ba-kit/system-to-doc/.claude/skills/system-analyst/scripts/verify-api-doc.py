#!/usr/bin/env python3
"""GATE — O2 API Documentation (doi chieu voi inventory).

  python3 verify-api-doc.py <API_Doc.xlsx> --inventory <_internal>/inventory.xlsx \\
      [--out <_internal>/gates/v-api.md]

Chan: API thieu/trung trong Index, thieu block, link noi bo gay, Confirmed khong co code-ref,
field mo coi, gia tri ngoai enum, batch thieu trong sheet Batch, o trong trong Index.
"""
import argparse
import re
import sys

import inv_schema as S
from gate_report import Gate

try:
    from openpyxl import load_workbook
except ImportError:
    print("Thieu openpyxl. Chay: pip install openpyxl", file=sys.stderr)
    sys.exit(2)

API_ID = re.compile(r"^(API-\d{3,})\b")
LINK = re.compile(r"^#?'?(.+?)'?!\$?([A-Z]{1,3})\$?(\d+)$")
NO_FIELD = "UNKNOWN — chua xac dinh tu code"
REQ_IN = {"path", "query", "header", "cookie", "body"}
RES_IN = {"status", "response-body", "header", "cookie"}


def link_target(cell):
    h = cell.hyperlink
    if not h:
        return None
    raw = h.location or h.target or ""
    if h.location is None and not raw.startswith("#"):
        return ("EXTERNAL", raw)
    m = LINK.match(raw.lstrip("#") if raw.startswith("#") else raw)
    return (m.group(1), "%s%s" % (m.group(2), m.group(3))) if m else ("BAD", raw)


def find_header(ws, first, second, max_row=60):
    for r in range(1, min(ws.max_row, max_row) + 1):
        if ws.cell(r, 1).value == first and ws.cell(r, 2).value == second:
            return r
    return None


def parse_blocks(ws):
    """-> {api_id: {row, labels{}, REQUEST:n, RESPONSE:n}}, list trung."""
    blocks, dups, cur, mode = {}, [], None, None
    r = 1
    while r <= ws.max_row:
        a = ws.cell(r, 1).value
        a = "" if a is None else str(a)
        m = API_ID.match(a)
        if m and " · " in a:
            aid = m.group(1)
            if aid in blocks:
                dups.append(aid)
            cur = {"row": r, "title": a, "labels": {}, "REQUEST": 0, "RESPONSE": 0,
                   "sheet": ws.title}
            blocks[aid] = cur
            mode = None
        elif cur is not None:
            if a in ("REQUEST", "RESPONSE"):
                mode = a
                r += 1  # bo qua dong header bang
            elif a == "":
                mode = None
            elif mode:
                if not a.startswith("UNKNOWN —"):
                    cur[mode] += 1
            else:
                cur["labels"][a] = "" if ws.cell(r, 2).value is None else str(ws.cell(r, 2).value)
        r += 1
    return blocks, dups


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("doc")
    ap.add_argument("--inventory", required=True)
    ap.add_argument("--out")
    a = ap.parse_args()

    data = S.load(a.inventory)
    apis = [r for r in data["07_API"] if r.get("API ID")]
    ids = [r["API ID"] for r in apis]
    idset = set(ids)
    fields = data["08_API_Fields"]
    evtype = {e.get("EV ID"): e.get("Type", "") for e in data["05_Evidence"]}
    wb = load_workbook(a.doc)
    g = Gate("GATE O2 — API Documentation")

    # 1 sheet bat buoc
    g.check(1, "Co sheet 00_Index va Batch", [s for s in ("00_Index", "Batch") if s not in wb.sheetnames])
    if "00_Index" not in wb.sheetnames:
        return g.emit(a.out)
    wsi = wb["00_Index"]
    hdr = find_header(wsi, "No", "API ID")
    idx_rows = []
    if hdr:
        ncol = 0
        while wsi.cell(hdr, ncol + 1).value not in (None, ""):
            ncol += 1
        for r in range(hdr + 1, wsi.max_row + 1):
            vals = [wsi.cell(r, c).value for c in range(1, ncol + 1)]
            if all(x in (None, "") for x in vals):
                continue
            idx_rows.append((r, vals))
    g.check(2, "00_Index co bang (header No · API ID)", [] if hdr else ["khong thay header"])

    # 3 moi API dung 1 lan trong Index
    seen = [str(v[1]) for _, v in idx_rows]
    bad = ["%s thieu" % i for i in ids if seen.count(i) == 0]
    bad += ["%s trung %d lan" % (i, seen.count(i)) for i in sorted(set(seen)) if seen.count(i) > 1]
    bad += ["%s khong co trong 07_API" % i for i in sorted(set(seen) - idset)]
    g.check(3, "Moi API ID trong 07_API xuat hien dung 1 lan o 00_Index", bad)

    # 4 so luong khop title block
    title = {}
    for r in range(1, (hdr or 1)):
        k = wsi.cell(r, 1).value
        if k:
            title[str(k)] = wsi.cell(r, 2).value
    n_batch = sum(1 for x in apis if x.get("Kind") in ("BATCH", "QUEUE"))
    bad = []
    if str(title.get("#Entries")) != str(len(apis)):
        bad.append("#Entries=%s, 07_API=%d" % (title.get("#Entries"), len(apis)))
    if str(title.get("#Batch")) != str(n_batch):
        bad.append("#Batch=%s, thuc te=%d" % (title.get("#Batch"), n_batch))
    if len(idx_rows) != len(apis):
        bad.append("so dong Index=%d, 07_API=%d" % (len(idx_rows), len(apis)))
    g.check(4, "So luong o title block / Index khop inventory", bad)

    # 5 block: moi API dung 1 block
    blocks, dups = {}, []
    for ws in wb.worksheets:
        if ws.title.startswith("G_"):
            b, d = parse_blocks(ws)
            dups += d
            for k, val in b.items():
                if k in blocks:
                    dups.append(k)
                blocks[k] = val
    bad = ["%s thieu block" % i for i in ids if i not in blocks]
    bad += ["%s block trung" % i for i in sorted(set(dups))]
    bad += ["%s block khong co trong 07_API" % i for i in sorted(set(blocks) - idset)]
    g.check(5, "Moi API co dung 1 block trong sheet G_<group>", bad)

    # 6 so dong Request/Response trong block = 08_API_Fields
    bad = []
    for i in ids:
        if i not in blocks:
            continue
        for d in ("REQUEST", "RESPONSE"):
            n = sum(1 for f in fields if f.get("API ID") == i and f.get("Direction") == d)
            if blocks[i][d] != n:
                bad.append("%s %s: doc=%d, 08=%d" % (i, d, blocks[i][d], n))
    g.check(6, "So dong Request/Response moi block khop 08_API_Fields", bad)

    # 7 hyperlink noi bo phan giai duoc
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                t = link_target(c)
                if not t:
                    continue
                where = "%s!%s" % (ws.title, c.coordinate)
                if t[0] in ("BAD", "EXTERNAL"):
                    bad.append("%s -> link khong hop le %s" % (where, t[1]))
                    continue
                if t[0] not in wb.sheetnames:
                    bad.append("%s -> sheet '%s' khong ton tai" % (where, t[0]))
                    continue
                tv = wb[t[0]][t[1]].value
                sv = "" if c.value is None else str(c.value)
                m = API_ID.match(sv)
                if m and not str(tv or "").startswith(m.group(1)):
                    bad.append("%s (%s) -> %s!%s = '%s'" % (where, sv, t[0], t[1], str(tv)[:30]))
                elif tv in (None, ""):
                    bad.append("%s -> %s!%s rong" % (where, t[0], t[1]))
    g.check(7, "Moi hyperlink noi bo tro toi sheet+o ton tai (API ID khop)", bad)

    # 8 Confirmed: Summary != UNKNOWN + co code-ref
    bad = []
    for x in apis:
        if x.get("Status") != "Confirmed":
            continue
        i = x["API ID"]
        summ = (x.get("Summary") or "").strip()
        if not summ or summ == S.UNKNOWN:
            bad.append("%s Summary UNKNOWN" % i)
        elif i in blocks and blocks[i]["labels"].get("Summary", "") in ("", S.UNKNOWN):
            bad.append("%s block Summary rong" % i)
        if not any(evtype.get(e) == "code-ref" for e in S.split_ids(x.get("Evidence"))):
            bad.append("%s khong co EV code-ref" % i)
    g.check(8, "API Confirmed co Summary va >=1 Evidence code-ref", bad)

    # 9 field mo coi
    g.check(9, "Moi dong 08_API_Fields tro toi API ID ton tai",
            ["row%s %s" % (f["__row__"], f.get("API ID") or "(trong)") for f in fields
             if f.get("API ID") not in idset])

    # 10 enum
    bad = []
    for x in apis:
        for col, enum in (("Kind", S.API_KIND), ("Method", S.API_METHOD), ("Status", S.API_STATUS)):
            if x.get(col) not in enum:
                bad.append("%s %s=%s" % (x["API ID"], col, x.get(col) or "(trong)"))
    for f in fields:
        d, i = f.get("Direction"), f.get("In")
        if d not in S.API_FIELD_DIRECTION:
            bad.append("08 row%s Direction=%s" % (f["__row__"], d or "(trong)"))
        if i not in S.API_FIELD_IN:
            bad.append("08 row%s In=%s" % (f["__row__"], i or "(trong)"))
        elif (d == "REQUEST" and i not in REQ_IN) or (d == "RESPONSE" and i not in RES_IN):
            bad.append("08 row%s In=%s khong hop voi %s" % (f["__row__"], i, d))
    g.check(10, "Kind/Method/Status/Direction/In thuoc enum S.API_*", bad)

    # 11 batch co trong sheet Batch
    bad = []
    if "Batch" in wb.sheetnames:
        wsb = wb["Batch"]
        bids = [str(wsb.cell(r, 2).value) for r in range(1, wsb.max_row + 1)]
        bad = [x["API ID"] for x in apis if x.get("Kind") in ("BATCH", "QUEUE") and x["API ID"] not in bids]
    else:
        bad = ["khong co sheet Batch"]
    g.check(11, "Moi BATCH/QUEUE co dong trong sheet Batch", bad)

    # 12 khong o trong trong Index
    g.check(12, "00_Index khong co o trong",
            ["row%d col%d" % (r, c + 1) for r, vals in idx_rows for c, val in enumerate(vals)
             if val is None or str(val).strip() == ""])

    # 13 WARN: Confirmed API khong co RESPONSE
    g.check(13, "API Confirmed co mo ta Response",
            [x["API ID"] for x in apis if x.get("Status") == "Confirmed" and x.get("Kind") == "API"
             and not any(f.get("API ID") == x["API ID"] and f.get("Direction") == "RESPONSE" for f in fields)],
            level="WARN")
    return g.emit(a.out)


if __name__ == "__main__":
    sys.exit(main())
