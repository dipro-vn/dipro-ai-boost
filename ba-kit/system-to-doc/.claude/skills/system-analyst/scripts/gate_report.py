#!/usr/bin/env python3
"""Khung bao cao chung cho moi gate script cua kit system-to-doc.

Exit code thong nhat:  0 = PASS  ·  1 = co FAIL  ·  2 = thieu dependency
"""
import sys


class Gate:
    def __init__(self, name):
        self.name = name
        self.results = []   # (check_no, title, level, detail)

    def ok(self, no, title, detail=""):
        self.results.append((no, title, "PASS", detail))

    def fail(self, no, title, detail=""):
        self.results.append((no, title, "FAIL", detail))

    def warn(self, no, title, detail=""):
        self.results.append((no, title, "WARN", detail))

    def check(self, no, title, bad_items, fmt=lambda x: str(x), level="FAIL"):
        """Tien ich: bad_items rong -> PASS, nguoc lai -> FAIL/WARN kem chi tiet."""
        if not bad_items:
            self.ok(no, title)
        else:
            detail = "; ".join(fmt(b) for b in bad_items[:15])
            if len(bad_items) > 15:
                detail += " ... (+%d)" % (len(bad_items) - 15)
            (self.fail if level == "FAIL" else self.warn)(no, title, detail)

    @property
    def counts(self):
        p = sum(1 for r in self.results if r[2] == "PASS")
        f = sum(1 for r in self.results if r[2] == "FAIL")
        w = sum(1 for r in self.results if r[2] == "WARN")
        return p, f, w

    def render(self):
        p, f, w = self.counts
        lines = ["# %s" % self.name, "",
                 "| # | Check | Ket qua | Chi tiet |",
                 "|---|---|---|---|"]
        for no, title, level, detail in self.results:
            icon = {"PASS": "PASS", "FAIL": "**FAIL**", "WARN": "WARN"}[level]
            lines.append("| %s | %s | %s | %s |" % (no, title, icon, detail.replace("|", "/")))
        lines += ["", "**%d checks · %d PASS · %d FAIL · %d WARN**" %
                  (len(self.results), p, f, w)]
        return "\n".join(lines)

    def emit(self, out_path=None):
        text = self.render()
        print(text)
        if out_path:
            with open(out_path, "w", encoding="utf8") as fh:
                fh.write(text + "\n")
        p, f, w = self.counts
        print("\nSUMMARY: %d checks · %d PASS · %d FAIL · %d WARN" %
              (len(self.results), p, f, w), file=sys.stderr)
        return 1 if f else 0


def frozen_sheets(path):
    """-> [ten sheet con freeze panel] cua 1 file xlsx (R13: kit khong freeze panel)."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    return [ws.title for ws in wb.worksheets if ws.freeze_panes is not None]


def legacy_word_hits(path):
    """-> [vi tri] con chu cu inv_schema.LEGACY_WORD trong xlsx / docx gui nguoi doc (D1: 'Cần xác nhận — Q-xxx')."""
    import inv_schema as S
    if path.lower().endswith(".docx"):
        from docx_read import Docx
        return ["docx"] if S.has_legacy(Docx(path).all_text) else []
    from openpyxl import load_workbook
    wb = load_workbook(path)
    return ["%s!%s" % (ws.title, c.coordinate) for ws in wb.worksheets for row in ws.iter_rows()
            for c in row if S.has_legacy(c.value)]
