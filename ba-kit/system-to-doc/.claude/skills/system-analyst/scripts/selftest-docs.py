#!/usr/bin/env python3
"""Self-test cho O2 (API doc) + O3 (DB doc) + read-schema — chung minh gate KHONG pass rong.

  python3 selftest-docs.py [--keep]

Dung inventory gia lap trong thu muc tam: read-schema -> inventory -> build 2 doc -> 2 gate PASS,
roi tiem tung loi da biet vao ban copy va xac nhan gate FAIL dung check. In `N ca · N PASS`.
"""
import argparse
import csv
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import inv_schema as S  # noqa: E402

try:
    import openpyxl
except ImportError:
    print("Thieu openpyxl. Chay: pip install openpyxl", file=sys.stderr)
    sys.exit(2)

SECRET = "selftest_secret_value@example.com"
DUMP = """-- MySQL dump (schema-only, gia lap)
CREATE TABLE `users` (
  `id` int unsigned NOT NULL AUTO_INCREMENT,
  `email` varchar(191) NOT NULL COMMENT 'Email dang nhap',
  `role` enum('admin','staff') NOT NULL DEFAULT 'staff',
  `created_at` datetime NOT NULL,
  PRIMARY KEY (`id`),
  UNIQUE KEY `uq_email` (`email`)
) ENGINE=InnoDB COMMENT='Tai khoan';

CREATE TABLE `orders` (
  `id` bigint unsigned NOT NULL AUTO_INCREMENT,
  `user_id` int unsigned NOT NULL,
  `total` decimal(12,2) NOT NULL,
  `status` enum('new','paid') NOT NULL,
  PRIMARY KEY (`id`),
  CONSTRAINT `fk_u` FOREIGN KEY (`user_id`) REFERENCES `users` (`id`)
) ENGINE=InnoDB;

CREATE TABLE `payments` (
  `id` bigint unsigned NOT NULL AUTO_INCREMENT,
  `order_id` bigint unsigned NOT NULL,
  `amount` decimal(12,2) NOT NULL,
  PRIMARY KEY (`id`)
) ENGINE=InnoDB;
"""


def run(args, **kw):
    p = subprocess.run([sys.executable] + args, cwd=HERE, capture_output=True, text=True, **kw)
    return p.returncode, p.stdout + p.stderr


def failed_checks(md_path):
    if not os.path.exists(md_path):
        return set()
    return {int(m.group(1)) for m in re.finditer(r"^\| (\d+) \|[^\n]*\| \*\*FAIL\*\* \|", open(md_path).read(), re.M)}


def put(wb, sheet, rows):
    ws = wb[sheet]
    for r in rows:
        ws.append([r.get(h, "") for h in S.SHEETS[sheet]])


def make_inventory(tmp, recon, db_mode="DUMP", n_extra_tables=0):
    inv = os.path.join(tmp, "inventory_%s_%d.xlsx" % (db_mode, n_extra_tables))
    rc, out = run(["build-inventory.py", "--out", inv])
    assert rc == 0, out
    wb = openpyxl.load_workbook(inv)
    for row in wb["00_Meta"].iter_rows(min_row=2):
        if row[0].value == "db_mode":
            row[1].value = db_mode
    rd = lambda f: list(csv.DictReader(open(os.path.join(recon, f), encoding="utf8")))
    tables = rd("tables.csv")
    for t in tables:
        if t["Table"] == "payments":
            t["Note"] += "\nlogical FK: payments.order_id → orders.id — inferred from EV-0001"
        if t["Table"] == "orders":
            t["Confidence"], t["Evidence"], t["Purpose"] = "High", t["Evidence"] + ";EV-0001", "Don hang"
    cols = rd("columns.csv")
    for i in range(n_extra_tables):   # bang phu cho test ERD > 60 bang
        name = "t%03d" % i
        tables.append({"Table": name, "Purpose": "UNKNOWN", "PK": "id", "FK": "ref_id->users.id" if i % 3 == 0 else "—",
                       "Est Rows": "UNKNOWN", "Related Function IDs": "UNKNOWN", "Evidence": "EV-0001",
                       "Confidence": "Low", "Note": "—"})
        cols.append({"Table": name, "Column": "id", "Type": "int", "PK": "Yes", "FK": "—", "Nullable": "NO",
                     "Evidence": "EV-0001", "Confidence": "Low"})
        if i % 3 == 0:
            cols.append({"Table": name, "Column": "ref_id", "Type": "int", "PK": "No", "FK": "users.id",
                         "Nullable": "YES", "Evidence": "EV-0001", "Confidence": "Low"})
    put(wb, "03_DB_Tables", tables)
    put(wb, "04_DB_Columns", cols)
    put(wb, "05_Evidence", rd("evidence.csv"))
    put(wb, "05_Evidence", [
        {"EV ID": "EV-0001", "Type": "code-ref", "Locator": "REPO-01:src/order.controller.ts#L10-L30"},
        {"EV ID": "EV-0002", "Type": "code-ref", "Locator": "REPO-01:src/jobs/cleanup.ts#L5"},
        {"EV ID": "EV-0003", "Type": "screenshot", "Locator": "https://app.example/orders",
         "Artifact": "evidence/EV-0003.png", "Captured At": "2026-10-01", "Actor/Role": "admin"}])
    put(wb, "02_Screen", [{"Screen ID": "SC-001", "Site": "WEB-01", "Screen Name": "Danh sach don", "Status": "Confirmed"}])
    put(wb, "11_Repo", [{"Repo ID": "REPO-01", "Path": "/x/api", "Kind": "BE", "Stack": "NestJS"}])
    put(wb, "07_API", [
        {"API ID": "API-001", "Group": "Order", "Kind": "API", "Method": "GET", "Path / Schedule": "/api/orders",
         "Summary": "Danh sach don", "Auth": "JWT", "Handler": "OrderController.list", "Repo": "REPO-01",
         "Called By Screens": "SC-001", "Related Tables": "orders", "Evidence": "EV-0001", "Status": "Confirmed"},
        {"API ID": "API-002", "Group": "Order", "Kind": "API", "Method": "POST", "Path / Schedule": "/api/orders",
         "Summary": "Tao don", "Auth": "JWT", "Handler": "OrderController.create", "Repo": "REPO-01",
         "Called By Screens": "SC-001", "Related Tables": "orders", "Evidence": "EV-0001", "Status": "Confirmed"},
        {"API ID": "API-003", "Group": "Batch", "Kind": "BATCH", "Method": "CRON", "Path / Schedule": "0 2 * * *",
         "Summary": "Don dep don cu", "Auth": "—", "Handler": "CleanupJob.run", "Repo": "REPO-01",
         "Called By Screens": "—", "Related Tables": "orders", "Evidence": "EV-0002", "Status": "Confirmed"}])
    F = lambda *a: dict(zip(S.SHEETS["08_API_Fields"], a))
    put(wb, "08_API_Fields", [
        F("API-001", "REQUEST", "query", "page", "integer", "No", "min 1", "Trang", "EV-0001"),
        F("API-001", "RESPONSE", "status", "200", "", "", "", "OK", "EV-0001"),
        F("API-002", "REQUEST", "body", "total", "number", "Yes", "decimal(12,2)", "Tong tien", "EV-0001"),
        F("API-002", "RESPONSE", "response-body", "id", "integer", "", "", "Id don moi", "EV-0001")])
    wb.save(inv)
    return inv


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--keep", action="store_true", help="giu thu muc tam de xem")
    a = ap.parse_args()
    tmp = tempfile.mkdtemp(prefix="selftest-docs-")
    results = []

    def case(name, ok, detail=""):
        results.append((name, bool(ok), detail))
        print("%s  %s%s" % ("PASS" if ok else "FAIL", name, ("  — " + detail) if detail and not ok else ""))

    try:
        # --- read-schema
        dump = os.path.join(tmp, "schema.sql")
        open(dump, "w").write(DUMP)
        bad = os.path.join(tmp, "with_data.sql")
        open(bad, "w").write(DUMP + "INSERT INTO `users` VALUES (1,'%s','admin',NOW());\n" % SECRET)
        rc, out = run(["read-schema.py", "--dump", bad, "--out", os.path.join(tmp, "x")])
        case("read-schema tu choi dump co INSERT (exit 3, khong in du lieu)",
             rc == 3 and SECRET not in out and not os.path.exists(os.path.join(tmp, "x", "recon", "db", "columns.csv")),
             "rc=%d" % rc)
        internal = os.path.join(tmp, "_internal")
        rc, out = run(["read-schema.py", "--dump", dump, "--out", internal, "--ev-start", "10"])
        recon = os.path.join(internal, "recon", "db")
        cols = list(csv.DictReader(open(os.path.join(recon, "columns.csv")))) if rc == 0 else []
        c = {(r["Table"], r["Column"]): r for r in cols}
        ok = (rc == 0 and json.loads(out[out.index("{"):out.rindex("}") + 1])["next_ev"] == 13
              and c[("users", "email")]["Max Length"] == "191"
              and c[("users", "email")]["Meaning"] == "Email dang nhap"
              and c[("users", "email")]["Confidence"] == "Medium"
              and "UNIQUE" in c[("users", "email")]["Constraint"]
              and c[("users", "role")]["Constraint"].startswith("ENUM: admin | staff")
              and c[("orders", "total")]["Format"] == "12,2"
              and c[("orders", "user_id")]["FK"] == "users.id"
              and c[("users", "created_at")]["Format"] == "YYYY-MM-DD HH:MM:SS"
              and c[("orders", "status")]["Meaning"] == S.UNKNOWN)
        case("read-schema parse dump schema-only dung (max length/format/enum/comment/FK/next_ev)", ok, out[-300:])

        inv = make_inventory(tmp, recon)
        api = os.path.join(tmp, "API.xlsx")
        db = os.path.join(tmp, "DB.xlsx")
        erd = os.path.join(tmp, "ERD.png")
        rc1, o1 = run(["build-api-doc.py", "--inventory", inv, "--system", "T", "--version", "ver1", "--out", api])
        rc2, o2 = run(["build-db-doc.py", "--inventory", inv, "--system", "T", "--version", "ver1",
                       "--out", db, "--erd-png", erd])
        case("build-api-doc + build-db-doc chay OK", rc1 == 0 and rc2 == 0 and os.path.exists(erd), o1[-200:] + o2[-200:])

        inv_none = make_inventory(tmp, recon, db_mode="NONE")
        rc, out = run(["build-db-doc.py", "--inventory", inv_none, "--system", "T", "--version", "ver1",
                       "--out", os.path.join(tmp, "none.xlsx")])
        case("build-db-doc tu choi db_mode=NONE (exit 1)", rc == 1 and not os.path.exists(os.path.join(tmp, "none.xlsx")))

        inv_mig = make_inventory(tmp, recon, db_mode="MIGRATION")
        rc, out = run(["build-db-doc.py", "--inventory", inv_mig, "--system", "T", "--version", "ver1",
                       "--out", os.path.join(tmp, "mig.xlsx")])
        case("build-db-doc chap nhan db_mode=MIGRATION", rc == 0)

        def gate_api(doc, inventory=inv):
            md = os.path.join(tmp, "g.md")
            if os.path.exists(md):
                os.remove(md)
            rc, _ = run(["verify-api-doc.py", doc, "--inventory", inventory, "--out", md])
            return rc, failed_checks(md)

        def gate_db(doc, inventory=inv, png=erd):
            md = os.path.join(tmp, "g.md")
            if os.path.exists(md):
                os.remove(md)
            args = ["verify-db-doc.py", doc, "--inventory", inventory, "--out", md]
            if png:
                args += ["--erd-png", png]
            rc, _ = run(args)
            return rc, failed_checks(md)

        rc, f = gate_api(api)
        case("verify-api-doc PASS tren doc sach", rc == 0 and not f, str(f))
        rc, f = gate_db(db)
        case("verify-db-doc PASS tren doc sach", rc == 0 and not f, str(f))

        def mutate(src, fn):
            dst = os.path.join(tmp, "mut_%d.xlsx" % len(results))
            shutil.copy(src, dst)
            wb = openpyxl.load_workbook(dst)
            fn(wb)
            wb.save(dst)
            return dst

        def find(wb, pred):
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if pred(c):
                            return ws, c
            return None, None

        # --- loi API
        def drop_block(wb):
            ws, c = find(wb, lambda c: str(c.value or "").startswith("API-002 ·"))
            c.value = "(da xoa)"
        rc, f = gate_api(mutate(api, drop_block))
        case("API: mat block API-002 -> check 5 FAIL", rc == 1 and 5 in f, str(f))

        def break_link(wb):
            ws, c = find(wb, lambda c: c.hyperlink is not None and str(c.value or "") == "API-001")
            c.hyperlink = "#'G_KhongTonTai'!A3"
        rc, f = gate_api(mutate(api, break_link))
        case("API: hyperlink gay (sheet khong ton tai) -> check 7 FAIL", rc == 1 and 7 in f, str(f))

        def wrong_link(wb):
            ws, c = find(wb, lambda c: c.hyperlink is not None and str(c.value or "") == "API-001")
            _, t = find(wb, lambda c: str(c.value or "").startswith("API-002 ·"))
            c.hyperlink = "#'%s'!%s" % (t.parent.title, t.coordinate)
        rc, f = gate_api(mutate(api, wrong_link))
        case("API: link API-001 tro nham block API-002 -> check 7 FAIL", rc == 1 and 7 in f, str(f))

        def blank_index(wb):
            ws = wb["00_Index"]
            _, c = find(wb, lambda c: c.parent.title == "00_Index" and str(c.value or "") == "Tao don")
            c.value = None
        rc, f = gate_api(mutate(api, blank_index))
        case("API: o trong trong 00_Index -> check 12 FAIL", rc == 1 and 12 in f, str(f))

        def drop_field(wb):
            _, c = find(wb, lambda c: str(c.value or "") == "Id don moi")
            c.parent.delete_rows(c.row, 1)
        rc, f = gate_api(mutate(api, drop_field))
        case("API: block thieu 1 dong Response -> check 6 FAIL", rc == 1 and 6 in f, str(f))

        wbi = openpyxl.load_workbook(inv)
        ws = wbi["07_API"]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == "API-002":
                row[S.SHEETS["07_API"].index("Evidence")].value = "EV-0003"
        ws8 = wbi["08_API_Fields"]
        ws8.append(["API-999", "REQUEST", "query", "x", "string", "No", "", "", "EV-0001"])
        inv_bad = os.path.join(tmp, "inv_bad.xlsx")
        wbi.save(inv_bad)
        rc, f = gate_api(api, inv_bad)
        case("API: Confirmed chi co screenshot (khong code-ref) -> check 8 FAIL", rc == 1 and 8 in f, str(f))
        case("API: 08_API_Fields tro toi API ID khong ton tai -> check 9 FAIL", rc == 1 and 9 in f, str(f))

        # --- loi DB
        def fake_rel(wb):
            ws = wb["01_Relationships"]
            _, c = find(wb, lambda c: c.parent.title == "01_Relationships" and c.column == 4 and c.value == "users")
            c.value = "customers"
        rc, f = gate_db(mutate(db, fake_rel))
        case("DB: quan he tro toi bang bia 'customers' -> check 5 FAIL", rc == 1 and 5 in f, str(f))

        def blank_grid(wb):
            ws = [w for w in wb.worksheets if w.title.startswith("T_users")][0]
            _, c = find(wb, lambda c: c.parent is ws and c.value == "Email dang nhap")
            c.value = None
        rc, f = gate_db(mutate(db, blank_grid))
        case("DB: o trong trong luoi cot -> check 7 FAIL", rc == 1 and 7 in f, str(f))

        def break_db_link(wb):
            _, c = find(wb, lambda c: c.parent.title == "00_Overview" and c.value == "orders" and c.hyperlink)
            c.hyperlink = "#'T_users'!B2"
        rc, f = gate_db(mutate(db, break_db_link))
        case("DB: link 'orders' tro sang sheet users -> check 6 FAIL", rc == 1 and 6 in f, str(f))

        def drop_col(wb):
            ws = [w for w in wb.worksheets if w.title.startswith("T_orders")][0]
            _, c = find(wb, lambda c: c.parent is ws and c.value == "total")
            ws.delete_rows(c.row, 1)
        rc, f = gate_db(mutate(db, drop_col))
        case("DB: sheet bang thieu 1 cot -> check 4 FAIL", rc == 1 and 4 in f, str(f))

        def drop_sheet(wb):
            wb.remove([w for w in wb.worksheets if w.title.startswith("T_payments")][0])
        rc, f = gate_db(mutate(db, drop_sheet))
        case("DB: mat sheet bang payments -> check 2 FAIL", rc == 1 and 2 in f, str(f))

        def high_no_code(wb):
            ws = [w for w in wb.worksheets if w.title.startswith("T_users")][0]
            _, c = find(wb, lambda c: c.parent is ws and c.value == "Low")
            c.value = "High"
        rc, f = gate_db(mutate(db, high_no_code))
        case("DB: Confidence High chi co db-query -> check 8 FAIL", rc == 1 and 8 in f, str(f))

        def drop_img(wb):
            wb["02_ERD"]._images = []
        rc, f = gate_db(mutate(db, drop_img))
        case("DB: truyen --erd-png nhung anh khong nhung -> check 9 FAIL", rc == 1 and 9 in f, str(f))

        # --- ERD > 60 bang
        inv_big = make_inventory(tmp, recon, n_extra_tables=70)
        big_png = os.path.join(tmp, "ERD_big.png")
        rc, out = run(["build-db-doc.py", "--inventory", inv_big, "--system", "T", "--version", "ver1",
                       "--out", os.path.join(tmp, "big.xlsx"), "--erd-png", big_png])
        j = json.loads(out[out.index("{"):out.rindex("}") + 1]) if rc == 0 else {}
        case("ERD >60 bang: ve 60, bao so bang an", rc == 0 and j.get("erd_omitted") == 73 - 60, str(j))
        rc, f = gate_db(os.path.join(tmp, "big.xlsx"), inv_big, big_png)
        case("verify-db-doc PASS voi 73 bang", rc == 0 and not f, str(f))
    finally:
        if a.keep:
            print("Giu thu muc tam:", tmp)
        else:
            shutil.rmtree(tmp, ignore_errors=True)

    n = len(results)
    p = sum(1 for r in results if r[1])
    print("\n%d ca · %d PASS" % (n, p))
    return 0 if p == n else 1


if __name__ == "__main__":
    sys.exit(main())
