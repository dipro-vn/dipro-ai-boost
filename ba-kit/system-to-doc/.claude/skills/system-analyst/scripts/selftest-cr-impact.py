#!/usr/bin/env python3
"""Self-test Luong 2: build-cr-impact.py + GATE V-CR + GATE V-CR-FIGMA that su bat loi.

  python3 selftest-cr-impact.py [--keep]

Baseline gia (inventory + design system theo site WEB-01/WEB-02, va layout 1 bo project/) + cr.json DUNG
-> build -> V-CR PASS. Tiem tung loi vao ban copy -> dung check FAIL/WARN. Chay lai sau MOI lan sua
build-cr-impact.py / verify-cr-impact.py / verify-cr-figma.py / cr_common.py.
"""
import argparse
import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import inv_schema as S  # noqa: E402
import cr_common as C  # noqa: E402

BASE = "ver1_011026_baseline"
KIT_RATES = os.path.abspath(os.path.join(HERE, "..", "..", "..", "..", S.MD_RATES_PATH))


def sh(script, *args):
    p = subprocess.run([sys.executable, os.path.join(HERE, script)] + list(args),
                       capture_output=True, text=True)
    return p.returncode, p.stdout + p.stderr


def results(output):
    out = {}
    for line in output.splitlines():
        parts = [p.strip() for p in line.split("|")]
        if len(parts) > 4 and parts[3] in ("PASS", "**FAIL**", "WARN"):
            out[parts[1]] = parts[3].strip("*")
    return out


def fill(path, rows):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    for sheet, items in rows.items():
        ws = wb[sheet]
        hdr = [c.value for c in ws[1]]
        for d in items:
            ws.append([d.get(h, "") for h in hdr])
    wb.save(path)


def make_baseline(root):
    b = os.path.join(root, "outputs", BASE)
    os.makedirs(os.path.join(b, "_internal"))
    inv = os.path.join(b, "_internal", "inventory.xlsx")
    sh("build-inventory.py", "--out", inv)
    fill(inv, {
        "00_Meta": [{"Key": "version_type", "Value": "BASELINE"}],
        "01_Function": [{"Function ID": "F-001", "Function": "Dat hang", "Screen IDs": "SC-001"},
                        {"Function ID": "F-002", "Function": "Xem don", "Screen IDs": "SC-002"}],
        "02_Screen": [{"Screen ID": "SC-001", "Screen Name": "Gio hang", "Entry From": "—"},
                      {"Screen ID": "SC-002", "Screen Name": "Danh sach don", "Entry From": "SC-001"},
                      {"Screen ID": "SC-003", "Screen Name": "Thanh toan", "Entry From": "SC-001"},
                      {"Screen ID": "SC-009", "Screen Name": "Cai dat", "Entry From": "—"}],
        "03_DB_Tables": [{"Table": "orders"}, {"Table": "users"}],
        "04_DB_Columns": [{"Table": "orders", "Column": "id"}, {"Table": "orders", "Column": "status"},
                          {"Table": "users", "Column": "email"}],
        "07_API": [{"API ID": "API-001", "Method": "POST", "Path / Schedule": "/orders",
                    "Called By Screens": "SC-001"}],
        "09_Integration": [{"EXT ID": "EXT-001", "Name": "VNPay"}],
        "10_Site": [{"Site ID": "WEB-01"}, {"Site ID": "WEB-02"}],
    })
    write_ds_sites(b)
    return b


def ds_project(p, primary):
    os.makedirs(os.path.join(p, "project", "components", "Button"), exist_ok=True)
    os.makedirs(os.path.join(p, "project", "components", "Cover"), exist_ok=True)
    json.dump({"name": "Shop", "color": {"tokens": [{"name": "primary", "value": primary},
                                                    {"name": "status-error", "value": "#cf222e"}]},
               "type": {"groups": [{"name": "Text", "styles": [{"name": "body", "fontSize": "14px"}]}]},
               "spacing": {"tokens": [{"name": "space-4", "value": "16px"}]}},
              open(os.path.join(p, "project", "tokens.json"), "w"))
    open(os.path.join(p, "project", "components", "index.d.ts"), "w").write(
        "export declare function Button(p: {}): JSX.Element;\n"
        "export declare function TextField(p: {}): JSX.Element;\n")


def write_ds_sites(b):
    """Layout moi: 04_DesignSystem/WEB-xx/project/ — primary khac gia tri giua 2 site."""
    ds = os.path.join(b, "04_DesignSystem")
    shutil.rmtree(ds, ignore_errors=True)
    ds_project(os.path.join(ds, "WEB-01"), "#1a7f37")
    ds_project(os.path.join(ds, "WEB-02"), "#0969da")


def write_ds_single(b):
    ds = os.path.join(b, "04_DesignSystem")
    shutil.rmtree(ds, ignore_errors=True)
    ds_project(ds, "#1a7f37")


def imp(iid, axis, ct, ref, item, code, qty=1, **kw):
    d = {"id": iid, "axis": axis, "change_type": ct, "baseline_ref": ref, "item": item,
         "change": "Thay doi %s" % item, "why_change": "Yeu cau CR muc 1 bat buoc sua %s" % item,
         "impact_on_current": "Luong hien tai cua %s" % item, "conflict": "No", "conflict_detail": "",
         "risk": "Low", "rate_code": code, "qty": qty, "md_note": "", "evidence": "input/CR-001.md §1",
         "question": "", "cr_item": "CR-001.1", "option": False}
    d.update(kw)
    return d


def valid_cr():
    return {
        "meta": {"cr_id": "CR-001", "cr_title": "Them xac thuc OTP", "baseline_version": BASE,
                 "cr_version": "ver2_021026_CR-001-otp", "source_type": "FILE",
                 "source_ref": "input/CR-001.md", "received_date": "02/10/2026",
                 "requested_by": "Anh A (PM khach hang)", "overall_risk": "Medium",
                 "recommendation": "Tra loi CQ-001 truoc khi lam"},
        "justification": [{"item": "CR-001.1", "request": "Them buoc OTP truoc khi dat hang",
                           "baseline_ref": "SC-001", "baseline_quote": "Gio hang -> dat hang ngay, khong co OTP",
                           "criteria": ["C1", "C3"], "why": "Buoc OTP chua co trong he thong.",
                           "not_feedback": "He thong dang chay dung nhu baseline."}],
        "not_cr": [{"item": "CR-001.9", "label": "BUG", "reason": "Nut luu khong hoat dong — loi",
                    "baseline_ref": "SC-002"}],
        "no_impact_axes": [{"axis": "ThirdParty", "reason": "CR khong them / sua lien ket ngoai, EXT-001 giu nguyen"}],
        "questions": [{"id": "CQ-001", "question": "Gia tri status moi?", "why": "Bao cao loc theo status",
                       "axis": "DB", "owner": "KH", "status": "Open"}],
        "impacts": [
            imp("IMP-001", "System", "UPD", "API-001", "POST /orders", "API-UPD"),
            imp("IMP-002", "DB", "NEW", "column:orders.otp_code", "orders.otp_code", "DB-COL-NEW"),
            imp("IMP-003", "DB", "UPD", "column:orders.status", "orders.status", "DB-COL-UPD",
                conflict="Yes", conflict_detail="Them PENDING_OTP", risk="High", question="CQ-001"),
            imp("IMP-004", "Business", "UPD", "F-001", "Dat hang", "BIZ-UPD"),
            imp("IMP-005", "Screen", "NEW", "—", "Man OTP", "SCR-NEW-S"),
            imp("IMP-006", "Screen", "UPD", "SC-001", "Gio hang", "SCR-UPD-S"),
            imp("IMP-007", "Screen", "IMPACT", "SC-002", "Danh sach don", "SCR-IMP", qty=2),
            imp("IMP-008", "Mockup", "NEW", "DS:WEB-01:primary;DS-component:Button;DS:space-4",
                "Mockup man OTP", "UI-NEW", option=True),
        ]}


def find(cr, iid):
    return next(i for i in cr["impacts"] if i["id"] == iid)


def m_imp(iid, **kw):
    return lambda cr: find(cr, iid).update(kw)


def m_meta(**kw):
    return lambda cr: cr["meta"].update(kw)


def m_set(key, val):
    return lambda cr: cr.__setitem__(key, val)


def m_just(**kw):
    return lambda cr: cr["justification"][0].update(kw)


def x_extra_sheet(path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    wb.create_sheet("Notes")
    wb.save(path)


def x_cell(sheet, finder, col, value=99):
    def f(path):
        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb[sheet]
        for r in range(1, ws.max_row + 1):
            if finder(ws, r):
                ws.cell(row=r, column=col).value = value
                break
        wb.save(path)
    return f


def est_row(iid):
    return lambda ws, r: ws.cell(row=r, column=2).value == iid


MD_COL = S.CR_IMPACT_COLS.index("MD") + 1
# (ten, mutation cr.json, mutation xlsx, check, muc, build phai tu choi)
CASES = [
    ("workbook thua 1 sheet", None, x_extra_sheet, "1", "FAIL", False),
    ("truc ThirdParty bo trong", m_set("no_impact_axes", []), None, "4", "FAIL", False),
    ("truc vua impact vua khong anh huong",
     lambda cr: cr["no_impact_axes"].append({"axis": "DB", "reason": "khong anh huong gi toi DB ca"}),
     None, "4", "FAIL", False),
    ("tieu chi C9 khong ton tai", m_just(criteria=["C9"]), None, "3", "FAIL", True),
    ("justification baseline_ref SC-777", m_just(baseline_ref="SC-777"), None, "3", "FAIL", False),
    ("justification rong", m_set("justification", []), None, "3", "FAIL", True),
    ("UPD man SC-999 khong ton tai", m_imp("IMP-006", baseline_ref="SC-999"), None, "7", "FAIL", False),
    ("NEW cot trung ten khong Conflict=Yes", m_imp("IMP-002", baseline_ref="column:orders.status"),
     None, "8", "FAIL", False),
    ("rate_code khong ton tai", m_imp("IMP-001", rate_code="API-XXL"), None, "10", "FAIL", True),
    ("rate_code sai truc", m_imp("IMP-001", rate_code="SCR-UPD-S"), None, "10", "FAIL", True),
    ("qty = 0", m_imp("IMP-001", qty=0), None, "10", "FAIL", True),
    ("MD bi sua tay tren sheet API", None,
     x_cell("API", lambda ws, r: ws.cell(row=r, column=2).value == "IMP-001", MD_COL), "11", "FAIL", False),
    ("dong Tong MD sheet Screen bi sua", None,
     x_cell("Screen", lambda ws, r: ws.cell(row=r, column=2).value == "Tổng MD", MD_COL), "11", "FAIL", False),
    ("実装 tren Summary bi sua", None,
     x_cell("Summary", lambda ws, r: ws.cell(row=r, column=1).value == C.L_IMPL, 2), "12", "FAIL", False),
    ("bang so doi tuong tren Summary bi sua", None,
     x_cell("Summary", lambda ws, r: ws.cell(row=r, column=1).value == "Màn hình", 2), "12", "FAIL", False),
    ("総工数 tren Summary go so thay vi link Estimation", None,
     x_cell("Summary", lambda ws, r: ws.cell(row=r, column=1).value == C.L_TOTAL_PD, 2), "12", "FAIL", False),
    ("Estimation 実装 (M) bi sua tay", None, x_cell("Estimation", est_row("IMP-004"), 13), "22", "FAIL", False),
    ("Estimation cong thuc テスト doi he so", None,
     x_cell("Estimation", est_row("IMP-004"), 14, "=M30*0.9"), "22", "FAIL", False),
    ("Estimation option khac cr.json", None, x_cell("Estimation", est_row("IMP-008"), 10, False), "22", "FAIL", False),
    ("Q&A mat cau hoi CQ-001", None,
     x_cell("Q&A", lambda ws, r: ws.cell(row=r, column=1).value == "CQ-001", 1, "—"), "24", "FAIL", False),
    ("impact thieu cr_item", m_imp("IMP-001", cr_item=""), None, "23", "FAIL", True),
    ("cr_item khong co trong justification", m_imp("IMP-001", cr_item="CR-001.7"), None, "23", "FAIL", True),
    ("option khong phai bool", m_imp("IMP-001", option="yes"), None, "23", "FAIL", True),
    ("High risk khong CQ / md_note", m_imp("IMP-001", risk="High"), None, "14", "FAIL", False),
    ("CQ treo (CQ-009)", m_imp("IMP-004", question="CQ-009"), None, "15", "FAIL", False),
    ("Mockup khong bam DS", m_imp("IMP-008", baseline_ref="—"), None, "16d", "FAIL", False),
    ("DS ref khong ghi site, 2 site khac gia tri -> WARN", m_imp("IMP-008", baseline_ref="DS:primary"),
     None, "18", "WARN", False),
    ("DS token khong ton tai", m_imp("IMP-008", baseline_ref="DS:nope"), None, "8", "FAIL", False),
    ("Screen IMPACT ref F-", m_imp("IMP-007", baseline_ref="F-001"), None, "16a", "FAIL", False),
    ("DB ref SC-", m_imp("IMP-002", baseline_ref="SC-001"), None, "16b", "FAIL", False),
    ("Evidence trong", m_imp("IMP-005", evidence=""), None, "17", "FAIL", False),
    ("baseline_version lech", m_meta(baseline_version="ver0_010126_baseline"), None, "2", "FAIL", False),
    ("Conflict=Yes thieu detail", m_imp("IMP-004", conflict="Yes"), None, "13", "FAIL", False),
]


def valid_nodes(cr, rates_path):
    """Ket qua readback hop le cho fixture: CR-1 (F-001), CR-2 (SC-001/SC-002/IMP-005 + AS-IS SC-003), thong ke."""
    def n(i, badge, ref, kind, label="x", **kw):
        d = {"id": "1:%d" % i, "name": "%s · %s · %s" % (badge, ref, label), "badge": badge, "ref": ref, "kind": kind}
        d.update(kw)
        return d
    nodes = [n(1, "UPD", "SC-001", "screen", grp="G-1"), n(2, "NEW", "IMP-005", "screen", grp="G-1"),
             n(3, "IMPACT", "SC-002", "screen", grp="G-1"), n(4, "UPD", "F-001", "flow", grp="FL-01"),
             n(5, "AS-IS", "SC-003", "screen", grp="G-1"),
             n(6, "NEW", "IMP-005", "mockup", ds="DS:WEB-01:primary;DS-component:WEB-01:Button", site="WEB-01",
               colors=["#1a7f37", "#ffffff"]),
             # non-happy: NG tren CR-1 (lane FL-01) + popup loi tren CR-2 (group G-1)
             n(7, "UPD", "F-001", "flow", grp="FL-01", shape="ng"),
             n(8, "UPD", "SC-001", "screen", grp="G-1", shape="popup"),
             # man UPD tren CR-3 co to mau phan doi
             n(9, "UPD", "SC-001", "mockup", ds="DS-component:WEB-01:Button", site="WEB-01",
               colors=["#ffffff"], chg=["AS-IS", "UPD"])]
    arrows = [{"id": "9:1", "from": "1:5", "to": "1:1"}, {"id": "9:2", "from": "1:1", "to": "1:2"},
              {"id": "9:3", "from": "1:1", "to": "1:3"}, {"id": "9:4", "from": "1:4", "to": "1:1"},
              {"id": "9:5", "from": "1:4", "to": "1:7"}, {"id": "9:6", "from": "1:1", "to": "1:8"}]
    _m, rates = C.load_rates(rates_path)
    stats = []
    for o, d in C.object_stats(cr["impacts"], rates):
        for c, v in zip(S.CR_CHANGE_TYPE + ["Tổng", "実装 MD"], [d[t] for t in S.CR_CHANGE_TYPE] + [d["total"], d["md"]]):
            stats.append({"name": "STAT · %s · %s" % (o, c), "value": str(v)})
    stats.append({"name": "STAT · TOTAL · 実装 — MD", "value": str(C.object_stats(cr["impacts"], rates)[-1][1]["md"])})
    return {"nodes": nodes, "arrows": arrows, "stats": stats, "legend": 1, "unbadged": [], "overlap": [],
            "idx": [{"ref": "SC-001", "type": "List"}, {"ref": "IMP-005", "type": "Form"}, {"ref": "SC-001", "type": "Modal"}],
            "idxTotal": ["TỔNG = 2 màn chính + 1 popup / màn lỗi = 3"], "erridx": 1,
            "desc": [{"ref": "IMP-005", "items": 2, "errors": 1}, {"ref": "SC-001", "items": 3, "errors": 2}],
            "layout": {"cr2": [{"name": "CR-2 Screen Flow · Site (WEB-01)", "site": "WEB-01", "groups": ["G-1"],
                                "flows": [{"name": "Flow G-1 — x", "flow": "G-1", "groups": ["G-1"]}]}],
                       "cr3": [{"name": "CR-3 Screens · Site (WEB-01)", "site": "WEB-01", "fill": "#9da5ad",
                                "flows": [{"name": "Flow G-1 — x", "flow": "G-1", "mockups": 2}], "loose": 0}],
                       "mockups": 2}}


def f_nodes(fn):
    def m(d):
        fn(d["nodes"])
    return m


FIG_CASES = [
    ("AS-IS SC-009 ngoai pham vi", f_nodes(lambda ns: ns.append(
        {"id": "1:20", "name": "AS-IS · SC-009 · Cai dat", "badge": "AS-IS", "ref": "SC-009", "kind": "screen"})), "4"),
    ("badge NEW tren SC-001 (impact la UPD)", f_nodes(lambda ns: ns[0].update(name="NEW · SC-001 · x", badge="NEW")), "2"),
    ("man SC-002 bi anh huong chua ve", f_nodes(lambda ns: ns.pop(2)), "3"),
    ("man NEW IMP-005 chua ve tren CR-2", f_nodes(lambda ns: ns.pop(1)), "3"),
    ("thong ke Man hinh / NEW sai", lambda d: d["stats"][0].update(value="9"), "5"),
    ("con dong chi tiet table-row", f_nodes(lambda ns: ns.append(
        {"id": "1:30", "name": "UPD · SC-001 · IMP-006", "badge": "UPD", "ref": "SC-001", "kind": "table-row"})), "5"),
    ("node CR-2 khong co mui ten", lambda d: d.__setitem__("arrows", d["arrows"][:1]), "7"),
    ("thieu legend", lambda d: d.__setitem__("legend", 0), "8"),
    ("chong node cu", lambda d: d.__setitem__("overlap", ["Output 1"]), "9"),
    ("mockup dung mau ngoai DS", f_nodes(lambda ns: ns[5].update(colors=["#ff00ff"])), "10"),
    ("mockup khong ghi DS", f_nodes(lambda ns: ns[5].update(ds="")), "10"),
    ("mockup clone UI that (src) co mau ngoai DS -> chi WARN", f_nodes(lambda ns: ns[5].update(colors=["#ff00ff"], src="1:999")), "10"),
    ("lane CR-1 khong co NG", f_nodes(lambda ns: ns[6].update(shape="box")), "11"),
    ("group CR-2 khong co popup / man loi", f_nodes(lambda ns: ns[7].update(shape="box")), "12"),
    ("Screen Index Loai ngoai enum", lambda d: d["idx"][0].update(type="Inline"), "12"),
    ("Screen Index thieu TONG", lambda d: d.__setitem__("idxTotal", []), "12"),
    ("⑦ Error Index thieu dong", lambda d: d.__setitem__("erridx", 0), "12"),
    ("CR-3 man thieu panel mo ta", lambda d: d.__setitem__("desc", d["desc"][1:]), "13"),
    ("CR-3 mo ta thieu ERROR SCENARIOS", lambda d: d["desc"][0].update(errors=0), "13"),
    ("CR-3 man UPD khong to mau phan doi", f_nodes(lambda ns: ns[8].update(chg=["AS-IS"])), "13"),
    ("CR-3 thieu man UPD SC-001", f_nodes(lambda ns: ns.pop(8)), "13"),
    ("readback cu khong co layout", lambda d: d.pop("layout"), "14"),
    ("CR-2 khong chia section theo site", lambda d: d["layout"].update(cr2=[]), "14"),
    ("CR-2 group khong nam trong section flow rieng", lambda d: d["layout"]["cr2"][0].update(flows=[]), "14"),
    ("CR-2 group nam o 2 site", lambda d: d["layout"]["cr2"].append({"name": "y", "site": "WEB-02", "groups": ["G-1"]}), "14"),
    ("CR-3 nen section khong phai #9DA5AD", lambda d: d["layout"]["cr3"][0].update(fill="#ffffff"), "14"),
    ("CR-3 khong chia theo flow", lambda d: d["layout"]["cr3"][0].update(flows=[], loose=2), "14"),
    ("CR-3 2 section cung site", lambda d: d["layout"]["cr3"].append(dict(d["layout"]["cr3"][0])), "14"),
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--keep", action="store_true")
    a = ap.parse_args()
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("Thieu openpyxl", file=sys.stderr)
        return 2
    root = tempfile.mkdtemp(prefix="selftest-cr-")
    total = passed = 0
    lines = []

    def record(name, ok, detail=""):
        nonlocal total, passed
        total += 1
        passed += bool(ok)
        lines.append("%s %s%s" % ("PASS" if ok else "FAIL", name, (" — " + detail) if detail and not ok else ""))

    try:
        base = make_baseline(root)
        crdir = os.path.join(root, "outputs", "ver2_021026_CR-001-otp")
        os.makedirs(os.path.join(crdir, "_internal"))
        xl = os.path.join(crdir, "CR-001_Impact.xlsx")
        cj = os.path.join(crdir, "_internal", "cr.json")
        rates = os.path.join(root, "rates.json")
        shutil.copy(KIT_RATES, rates)
        rd = json.load(open(rates))
        rd["status"] = "DRAFT"
        json.dump(rd, open(rates, "w"))

        def build(cr):
            json.dump(cr, open(cj, "w"), ensure_ascii=False)
            if os.path.exists(xl):
                os.remove(xl)
            return sh("build-cr-impact.py", "--cr-json", cj, "--rates", rates, "--out", xl)

        def verify(*extra):
            return sh("verify-cr-impact.py", xl, "--cr-json", cj, "--baseline", base, "--rates", rates, *extra)

        c0 = valid_cr()
        code, out = build(c0)
        record("build cr.json dung -> exit 0", code == 0 and os.path.isfile(xl), out)
        code, out = verify()
        res = results(out)
        bad = {k: v for k, v in res.items() if v != "PASS"}
        record("fixture dung -> V-CR PASS (chi WARN 19 don gia DRAFT)", code == 0 and bad == {"19": "WARN"},
               "khong PASS: %s" % bad)
        from openpyxl import load_workbook
        wb = load_workbook(xl)
        flat = [str(c) for r in wb["Summary"].iter_rows(values_only=True) for c in r if c is not None]
        grand = sum(C.md_of(i, C.load_rates(rates)[1]) for i in c0["impacts"])
        record("Summary co canh bao 'Ước lượng sơ bộ' khi DRAFT + 実装 = tong MD + du %d sheet" % len(S.CR_SHEETS),
               any("Ước lượng sơ bộ" in x for x in flat) and wb.sheetnames == S.CR_SHEETS
               and any(r and r[0] == C.L_IMPL and r[1] == grand for r in wb["Summary"].iter_rows(values_only=True)),
               "sheets=%s" % wb.sheetnames)
        est = [r for r in wb["Estimation"].iter_rows(values_only=True) if r and len(r) > 2 and
               str(r[1] or "").startswith("IMP-")]
        record("Estimation 1 dong / hang muc (%d)" % len(c0["impacts"]), len(est) == len(c0["impacts"]),
               "nhan %d" % len(est))

        for name, mj, mx, chk, level, rejects in CASES:
            c = copy.deepcopy(c0)
            if mj:
                mj(c)
            code, out = build(c0)          # xlsx dung tu cr.json goc
            if rejects:
                json.dump(c, open(cj + ".bad", "w"))
                bcode, bout = sh("build-cr-impact.py", "--cr-json", cj + ".bad", "--rates", rates,
                                 "--out", xl + ".bad.xlsx")
                record("%s -> build tu choi (exit 1, khong ghi file)" % name,
                       bcode == 1 and not os.path.exists(xl + ".bad.xlsx"), "exit %d" % bcode)
            elif mj:
                code, out = build(c)
                if code:
                    record("%s -> build" % name, False, out)
                    continue
            json.dump(c, open(cj, "w"), ensure_ascii=False)
            if mx:
                mx(xl)
            code, out = verify()
            got = results(out).get(chk)
            want_code = 1 if level == "FAIL" else 0
            record("%s -> check %s %s" % (name, chk, level), got == level and code == want_code,
                   "nhan %s (exit %d)" % (got, code))

        # don gia APPROVED -> 19 PASS, khong con canh bao
        rd["status"] = "APPROVED"
        json.dump(rd, open(rates, "w"))
        build(c0)
        code, out = verify()
        flat = [str(c) for r in load_workbook(xl)["Summary"].iter_rows(values_only=True) for c in r if c]
        record("don gia APPROVED -> check 19 PASS, Summary khong canh bao",
               results(out).get("19") == "PASS" and code == 0 and not any("Ước lượng sơ bộ" in x for x in flat),
               "nhan %s" % results(out).get("19"))
        rd["status"] = "DRAFT"
        json.dump(rd, open(rates, "w"))

        # CR-vs-CR
        build(c0)
        other = os.path.join(root, "cr-002.json")
        c2 = copy.deepcopy(c0)
        c2["meta"]["cr_id"] = "CR-002"
        json.dump(c2, open(other, "w"))
        code, out = verify("--other-cr", other)
        record("CR-vs-CR cung SC-001 -> check 20 WARN", results(out).get("20") == "WARN" and code == 0,
               "nhan %s" % results(out).get("20"))

        # Figma
        nj = os.path.join(crdir, "_internal", "cr-figma-nodes.json")

        def fig(nodes):
            json.dump(nodes, open(nj, "w"), ensure_ascii=False)
            return sh("verify-cr-figma.py", "--nodes", nj, "--cr-json", cj, "--baseline", base, "--rates", rates)
        code, out = fig(valid_nodes(c0, rates))
        bad = {k: v for k, v in results(out).items() if v != "PASS"}
        record("figma hop le -> V-CR-FIGMA PASS", code == 0 and not bad, "khong PASS: %s" % bad)
        for name, mut, chk in FIG_CASES:
            d = valid_nodes(c0, rates)
            mut(d)
            code, out = fig(d)
            got = results(out).get(chk)
            want = "WARN" if "chi WARN" in name else "FAIL"
            record("figma %s -> check %s %s" % (name, chk, want), got == want and (code == 1) == (want == "FAIL"),
                   "nhan %s (exit %d)" % (got, code))

        # CR-1 lane dang use case (render-cr-figma.usecase_lane_errors)
        import importlib.util
        spec = importlib.util.spec_from_file_location("rcf", os.path.join(HERE, "render-cr-figma.py"))
        rcf = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(rcf)
        def lane():
            return {"lane": "FL-01", "nodes": [{"key": "a", "type": "actor", "ref": "FL-01"}, {"key": "s", "type": "start", "pos": 1, "ref": "SC-001"},
                    {"key": "d", "type": "decision", "pos": 2, "ref": "SC-001"}, {"key": "e", "type": "end", "pos": 3, "ref": "SC-001", "result": "ok"},
                    {"key": "n", "type": "ng", "pos": 2, "row": 1, "ref": "SC-001"}, {"key": "f", "type": "end", "pos": 3, "row": 1, "ref": "SC-001", "result": "fail"}],
                    "edges": [["a", "s"], ["s", "d"], ["d", "e"], ["d", "n", "", "ng"], ["n", "f", "", "ng"]]}
        record("CR-1 use case hop le -> khong loi", rcf.usecase_lane_errors(lane()) == [], str(rcf.usecase_lane_errors(lane())))
        for name, fn in [("cot technology cu", lambda L: L["nodes"][2].update(type="technology")),
                         ("ref API trong lane", lambda L: L["nodes"][2].update(ref="API-001")),
                         ("thieu ket thuc happy", lambda L: L["nodes"][3].update(row=1)),
                         ("NG khong di tiep", lambda L: L["edges"].pop()),
                         ("pos ngoai 1..6", lambda L: L["nodes"][1].update(pos=9))]:
            L = lane(); fn(L)
            record("CR-1 use case %s -> renderer tu choi" % name, bool(rcf.usecase_lane_errors(L)))
        record("Figma CR-1/CR-2 co chu UNKNOWN -> renderer tu choi",
               bool(rcf.unknown_text_errors({"flows": [{"nodes": [{"label": "Message: UNKNOWN"}]}]})) and not rcf.unknown_text_errors({"flows": [lane()]}))

        # Layout DS 1 bo (04_DesignSystem/project/)
        write_ds_single(base)
        c = copy.deepcopy(c0)
        find(c, "IMP-008")["baseline_ref"] = "DS:primary;DS-component:Button;DS:body"
        build(c)
        code, out = verify()
        bad = {k: v for k, v in results(out).items() if v not in ("PASS",) and k != "19"}
        record("DS layout 1 bo project/ -> PASS", code == 0 and not bad, "khong PASS: %s" % bad)
        c = copy.deepcopy(c0)
        build(c)
        code, out = verify()
        record("DS layout 1 bo + ref DS:WEB-01:... -> check 8 FAIL", results(out).get("8") == "FAIL",
               "nhan %s" % results(out).get("8"))

        # Baseline khong co design system -> 16d chi WARN
        shutil.rmtree(os.path.join(base, "04_DesignSystem"))
        c = copy.deepcopy(c0)
        find(c, "IMP-008")["baseline_ref"] = "—"
        build(c)
        code, out = verify()
        record("baseline khong DS -> check 16d WARN", results(out).get("16d") == "WARN" and code == 0,
               "nhan %s (exit %d)" % (results(out).get("16d"), code))
    finally:
        if a.keep:
            print("Giu fixture tai %s" % root)
        else:
            shutil.rmtree(root, ignore_errors=True)

    print("\n".join(lines))
    print("\n%d ca · %d PASS" % (total, passed))
    print("SUMMARY: %d ca · %d PASS · %d FAIL" % (total, passed, total - passed), file=sys.stderr)
    return 0 if total == passed else 1


if __name__ == "__main__":
    sys.exit(main())
