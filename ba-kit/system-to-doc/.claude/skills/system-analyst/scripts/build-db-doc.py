#!/usr/bin/env python3
"""O3 — Database Documentation (xlsx) + ERD (png) tu inventory.

  python3 build-db-doc.py --inventory <_internal>/inventory.xlsx --system <name> --version ver1 \\
      --out <ver>/03_DB/DB_Doc_<sys>_ver1.xlsx [--erd-png <ver>/03_DB/ERD_<sys>_ver1.png]

Nguon su that: 03_DB_Tables, 04_DB_Columns (+00_Meta db_mode, 01_Function, 05_Evidence).
Sheet: 00_Overview · 01_Relationships · 02_ERD · T_<table> (moi bang 1 sheet).
Quan he logic: dong trong 03_DB_Tables.Note dang `logical FK: a.col → b.id — inferred from EV-0102`.
⚠ db_mode=NONE -> exit 1 (khong co DB, O3 khong sinh). O trong -> UNKNOWN / — (khong ap dung).
"""
import argparse
import datetime
import math
import os
import re
import sys

import inv_schema as S

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Thieu openpyxl. Chay: pip install openpyxl", file=sys.stderr)
    sys.exit(2)
import bd_styles as ST

NA = "—"
OVERVIEW_COLS = ["No", "Table", "Purpose", "#Columns", "PK", "FK →", "Referenced by",
                 "Related Function IDs", "Confidence"]
REL_COLS = ["No", "From Table", "From Column", "To Table", "To Column", "Kind",
            "Cardinality", "Evidence"]
GRID_COLS = ["No", "Column", "Type", "Format", "Max Length", "Nullable", "Default", "PK", "FK",
             "Constraint", "Purpose/Meaning", "Evidence", "Confidence"]
LOGICAL = re.compile(r"logical\s+FK\s*:\s*([\w.$]+)\.([\w$]+)\s*(?:→|->|=>)\s*([\w.$]+)\.([\w$]+)"
                     r"(?:\s*[—–-]+\s*([^\n;]*))?", re.I)
EV_RE = re.compile(r"EV-\d{4,}")
LINK_FONT = Font(name=ST.FONT, size=10, color="FF0563C1", underline="single")


def v(x, empty=S.UNKNOWN):
    x = "" if x is None else str(x).strip()
    return x if x else empty


def sheet_title(prefix, name, used):
    base = re.sub(r"[\[\]:*?/\\']", "_", "%s%s" % (prefix, name))[:31]
    t, i = base, 2
    while t.lower() in used:
        suf = "~%d" % i
        t = base[:31 - len(suf)] + suf
        i += 1
    used.add(t.lower())
    return t


def link(cell, sheet, ref, text=None):
    if text is not None:
        cell.value = text
    cell.hyperlink = "#'%s'!%s" % (sheet, ref)
    cell.font = LINK_FONT


def parse_fk_text(s):
    """'col->t.c; col2→t2.c2' -> [(col, t, c)]."""
    out = []
    for part in re.split(r"[;\n]", s or ""):
        m = re.match(r"\s*([\w$]+)\s*(?:->|→|=>)\s*([\w.$]+)\.([\w$]+)\s*$", part)
        if m:
            out.append(m.groups())
    return out


def collect(data):
    tables = [r for r in data["03_DB_Tables"] if r.get("Table")]
    cols = {}
    for r in data["04_DB_Columns"]:
        if r.get("Table"):
            cols.setdefault(r["Table"], []).append(r)
    tnames = {t["Table"] for t in tables}
    rels, seen = [], set()

    def add(ft, fc, tt, tc, kind, ev):
        key = (ft, fc, tt, tc)
        if key in seen:
            return
        seen.add(key)
        rels.append({"from_t": ft, "from_c": fc, "to_t": tt, "to_c": tc, "kind": kind, "ev": ev})

    for t in tables:
        name, ev = t["Table"], v(t.get("Evidence"))
        for c in cols.get(name, []):
            fk = (c.get("FK") or "").strip()
            m = re.match(r"^([\w.$]+)\.([\w$]+)$", fk)
            if m:
                add(name, c["Column"], m.group(1), m.group(2), "FK", v(c.get("Evidence"), ev))
        for fc, tt, tc in parse_fk_text(t.get("FK")):
            add(name, fc, tt, tc, "FK", ev)
    for t in tables:
        for m in LOGICAL.finditer(t.get("Note") or ""):
            evs = EV_RE.findall(m.group(5) or "")
            add(m.group(1), m.group(2), m.group(3), m.group(4), "Logical (inferred)",
                ";".join(evs) or v(t.get("Evidence")))
    # cardinality
    for r in rels:
        src = [c for c in cols.get(r["from_t"], []) if c.get("Column") == r["from_c"]]
        pk_cols = [c["Column"] for c in cols.get(r["from_t"], []) if c.get("PK") == "Yes"]
        uniq = bool(src) and bool(re.search(r"(^|;)\s*UNIQUE\s*($|;)", src[0].get("Constraint") or ""))
        one = (pk_cols == [r["from_c"]]) or uniq
        r["card"] = ("1:1" if one else "N:1") + ("" if r["kind"] == "FK" else " (inferred)")
    return tables, cols, rels, tnames


def func_map(data):
    out = {}
    for f in data.get("01_Function", []):
        for t in S.split_ids(f.get("Related Tables")):
            out.setdefault(t, []).append(f.get("Function ID", ""))
    return out


# ---------------------------------------------------------------- ERD
def render_erd(tables, cols, rels, path, system, version, limit=60):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.patches import FancyBboxPatch
    except ImportError:
        print("Thieu matplotlib — bo qua ERD", file=sys.stderr)
        return None
    names = [t["Table"] for t in tables]
    deg = {n: 0 for n in names}
    adj = {n: set() for n in names}
    for r in rels:
        for a, b in ((r["from_t"], r["to_t"]), (r["to_t"], r["from_t"])):
            if a in deg and a != b:
                deg[a] += 1
                if b in adj:
                    adj[a].add(b)
    keep = sorted(names, key=lambda n: (-deg[n], n))[:limit]
    omitted = len(names) - len(keep)
    keepset = set(keep)
    order, done = [], set()          # BFS: bang lien quan dung canh nhau
    for seed in keep:
        if seed in done:
            continue
        queue = [seed]
        done.add(seed)
        while queue:
            n = queue.pop(0)
            order.append(n)
            for m in sorted(adj[n] & keepset, key=lambda x: (-deg[x], x)):
                if m not in done:
                    done.add(m)
                    queue.append(m)
    n = len(order)
    if n == 0:
        return None
    ncol = max(1, min(8, int(math.ceil(math.sqrt(n * 1.4)))))
    nrow = int(math.ceil(n / float(ncol)))
    MAXL = 9
    logical = {(r["from_t"], r["from_c"]) for r in rels if r["kind"] != "FK"}

    def lines_of(t):
        out = []
        for c in cols.get(t, []):
            tag = []
            if c.get("PK") == "Yes":
                tag.append("PK")
            if (c.get("FK") or NA) not in (NA, "", S.UNKNOWN):
                tag.append("FK")
            elif (t, c.get("Column")) in logical:
                tag.append("FK~")
            if tag:
                out.append(("%s %s" % ("/".join(tag), c["Column"]), "PK" in tag))
        if len(out) > MAXL:
            extra = len(out) - (MAXL - 1)
            out = out[:MAXL - 1] + [("... +%d cot khoa" % extra, False)]
        if not out:
            out = [("(%d cot, khong PK/FK)" % len(cols.get(t, [])), False)]
        return out

    W, LH, HH, GX, GY = 3.0, 0.26, 0.42, 1.2, 1.0
    boxes, cell, row_top, row_bot = {}, {}, [], []
    y = 0.0
    for ri in range(nrow):
        row = order[ri * ncol:(ri + 1) * ncol]
        rh = max(HH + LH * len(lines_of(t)) + 0.15 for t in row)
        row_top.append(-y)
        row_bot.append(-y - rh)
        for ci, t in enumerate(row):
            h = HH + LH * len(lines_of(t)) + 0.15
            boxes[t] = (ci * (W + GX), -y - h, W, h)
            cell[t] = (ri, ci)
        y += rh + GY
    total_w = ncol * (W + GX) - GX
    total_h = y - GY
    fig_w, fig_h = max(8, (total_w + 1.2) * 0.85), max(5, (total_h + 2.4) * 0.85)
    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    ax.set_xlim(-GX / 2 - 0.2, total_w + GX / 2 + 0.2)
    ax.set_ylim(-total_h - GY / 2 - 0.3, 1.25)
    ax.axis("off")
    ax.text(0, 0.95, "ERD — %s (%s)" % (system, version), fontsize=14, fontweight="bold", va="center")
    sub = ("%d bang, %d quan he · net lien = FK vat ly · net dut cam = quan he logic suy luan (FK~) · "
           "mui ten tro ve bang duoc tham chieu" % (len(names), len(rels)))
    if omitted:
        sub += "\nChi ve %d bang nhieu lien ket nhat — AN %d bang (xem sheet 01_Relationships)" % (len(keep), omitted)
    ax.text(0, 0.55, sub, fontsize=8.5, color="#555555", va="top")

    k = 0
    for r in rels:
        a, b = boxes.get(r["from_t"]), boxes.get(r["to_t"])
        if not a or not b:
            continue
        ls = "-" if r["kind"] == "FK" else "--"
        col = "#2F5597" if r["kind"] == "FK" else "#C55A11"
        off = ((k % 5) - 2) * 0.09
        k += 1
        ax_, ay, aw, ah = a
        bx, by, bw, bh = b
        (sr, sc), (tr, tc) = cell[r["from_t"]], cell[r["to_t"]]
        arrow = dict(arrowstyle="-|>", color=col, lw=1.3, ls=ls, shrinkA=0, shrinkB=0)
        if r["from_t"] == r["to_t"]:            # tu tham chieu: vong ben phai
            x1, ya, yb = ax_ + aw, ay + ah * 0.7, ay + ah * 0.3
            ax.plot([x1, x1 + 0.35, x1 + 0.35], [ya, ya, yb], color=col, ls=ls, lw=1.3, zorder=1)
            ax.annotate("", xy=(x1, yb), xytext=(x1 + 0.35, yb), arrowprops=arrow, zorder=1)
            continue
        if sr == tr and abs(sc - tc) == 1:      # canh nhau cung hang: duong thang
            yy = max(ay, by) + min(ah, bh) / 2 + off
            x1, x2 = (ax_ + aw, bx) if tc > sc else (ax_, bx + bw)
            ax.annotate("", xy=(x2, yy), xytext=(x1, yy), arrowprops=arrow, zorder=1)
            continue
        if sr == tr + 1 and sc == tc:            # ngay duoi cung cot
            xx = ax_ + aw / 2 + off
            ax.annotate("", xy=(xx, by), xytext=(xx, ay + ah), arrowprops=arrow, zorder=1)
            continue
        # tuyen vuong goc: canh box -> hanh lang cot -> hanh lang hang -> vao box dich
        right = tc >= sc
        x0 = ax_ + aw if right else ax_
        cx = x0 + (GX / 2 if right else -GX / 2) + off
        y0 = ay + ah / 2 + off
        if tr > sr:                              # dich o hang duoi: vao tu canh tren
            gy, ey = row_top[tr] + GY / 2, by + bh
        else:                                    # dich cung hang / hang tren: vao tu canh duoi
            gy, ey = row_bot[tr] - GY / 2, by
        gy += off
        ex = bx + bw / 2 + off * 2
        ax.plot([x0, cx, cx, ex], [y0, y0, gy, gy], color=col, ls=ls, lw=1.3, zorder=1)
        ax.annotate("", xy=(ex, ey), xytext=(ex, gy), arrowprops=arrow, zorder=1)
    for t, (x, y0, w, h) in boxes.items():
        ax.add_patch(FancyBboxPatch((x, y0), w, h, boxstyle="round,pad=0,rounding_size=0.06",
                                    fc="white", ec="#1F4E78", lw=1.3, zorder=2))
        ax.add_patch(FancyBboxPatch((x, y0 + h - HH), w, HH, boxstyle="round,pad=0,rounding_size=0.06",
                                    fc="#1F4E78", ec="#1F4E78", lw=1.3, zorder=3))
        label = t if len(t) <= 30 else t[:28] + ".."
        ax.text(x + 0.12, y0 + h - HH / 2, label, color="white", fontsize=9.5, fontweight="bold",
                va="center", zorder=4)
        for i, (txt, pk) in enumerate(lines_of(t)):
            ax.text(x + 0.12, y0 + h - HH - LH * (i + 0.65), txt[:34], fontsize=8.3,
                    color="#000000" if pk else "#333333", fontweight="bold" if pk else "normal",
                    va="center", zorder=4)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    dpi = 110 if fig_w * fig_h < 400 else 80
    fig.savefig(path, dpi=dpi, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return {"drawn": len(keep), "omitted": omitted}


# ---------------------------------------------------------------- workbook
def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_grid(ws, r0, header, rows, center_cols=()):
    for i, h in enumerate(header, start=1):
        ws.cell(r0, i, h)
        ST.header(ws.cell(r0, i))
    for ri, row in enumerate(rows, start=r0 + 1):
        for ci, val in enumerate(row, start=1):
            c = ws.cell(ri, ci, val)
            ST.data(c, size=10, center=ci in center_cols)
    last = r0 + max(len(rows), 1)
    ws.auto_filter.ref = "A%d:%s%d" % (r0, get_column_letter(len(header)), last)
    ws.freeze_panes = ws.cell(r0 + 1, 3)


def kv_block(ws, r, pairs):
    for k, val in pairs:
        a, b = ws.cell(r, 1, k), ws.cell(r, 2, val)
        ST.header(a)
        a.alignment = Alignment(horizontal="left", vertical="center")
        ST.data(b, size=10)
        r += 1
    return r


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--inventory", required=True)
    ap.add_argument("--system", required=True)
    ap.add_argument("--version", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--erd-png")
    a = ap.parse_args()

    data = S.load(a.inventory)
    meta = S.meta(data)
    mode = (meta.get("db_mode") or S.UNKNOWN).strip().upper()
    if mode == "NONE":
        print("db_mode=NONE: khong co DB — O3 khong sinh", file=sys.stderr)
        return 1
    tables, cols, rels, tnames = collect(data)
    if not tables:
        print("03_DB_Tables rong — khong co gi de sinh O3", file=sys.stderr)
        return 1
    fmap = func_map(data)

    erd = None
    if a.erd_png:
        erd = render_erd(tables, cols, rels, a.erd_png, a.system, a.version)

    wb = Workbook()
    wb.remove(wb.active)
    used = {"00_overview", "01_relationships", "02_erd"}
    ws0 = wb.create_sheet("00_Overview")
    ws1 = wb.create_sheet("01_Relationships")
    ws2 = wb.create_sheet("02_ERD")
    tsheet = {t["Table"]: sheet_title("T_", t["Table"], used) for t in tables}

    # --- 00_Overview
    ws0["A1"] = "DATABASE DOCUMENTATION — %s" % a.system
    ws0["A1"].font = Font(name=ST.FONT, size=14, bold=True, color="FF1F4E78")
    ncol = sum(len(cols.get(t["Table"], [])) for t in tables)
    n_fk = sum(1 for r in rels if r["kind"] == "FK")
    r = kv_block(ws0, 3, [
        ("System", a.system), ("Version", a.version),
        ("Generated", datetime.date.today().isoformat()), ("DB Mode", mode),
        ("#Tables", len(tables)), ("#Columns", ncol),
        ("#Relationships", "%d (FK %d · Logical %d)" % (len(rels), n_fk, len(rels) - n_fk)),
        ("Ghi chu", "UNKNOWN = chua co bang chung tu code/DB; — = khong ap dung. "
                    "Bam ten bang de mo sheet chi tiet."),
    ])
    hdr = r + 1
    out_by, in_by = {}, {}
    for rel in rels:
        tag = "" if rel["kind"] == "FK" else " (logical)"
        out_by.setdefault(rel["from_t"], []).append(rel["to_t"] + tag)
        in_by.setdefault(rel["to_t"], []).append(rel["from_t"] + tag)
    rows = []
    for i, t in enumerate(tables, start=1):
        name = t["Table"]
        funcs = list(dict.fromkeys(S.split_ids(t.get("Related Function IDs")) +
                                   [f for f in fmap.get(name, []) if f]))
        funcs = [f for f in funcs if f != S.UNKNOWN]
        rows.append([i, name, v(t.get("Purpose")), len(cols.get(name, [])), v(t.get("PK"), NA),
                     ", ".join(dict.fromkeys(out_by.get(name, []))) or NA,
                     ", ".join(dict.fromkeys(in_by.get(name, []))) or NA,
                     ", ".join(funcs) or S.UNKNOWN, v(t.get("Confidence"))])
    write_grid(ws0, hdr, OVERVIEW_COLS, rows, center_cols=(1, 4, 9))
    for i, t in enumerate(tables):
        link(ws0.cell(hdr + 1 + i, 2), tsheet[t["Table"]], "B2")
    set_widths(ws0, [16, 28, 46, 10, 16, 28, 28, 22, 12])
    ws0.column_dimensions["A"].width = 18

    # --- 01_Relationships
    ws1["A1"] = "RELATIONSHIPS — %s" % a.system
    ws1["A1"].font = Font(name=ST.FONT, size=13, bold=True, color="FF1F4E78")
    link(ws1["E1"], "00_Overview", "A1", "← Overview")
    rrows = [[i, x["from_t"], x["from_c"], x["to_t"], x["to_c"], x["kind"], x["card"], x["ev"]]
             for i, x in enumerate(rels, start=1)]
    if not rrows:
        rrows = [[1, NA, NA, NA, NA, NA, NA, "Khong tim thay FK vat ly hay quan he logic"]]
    write_grid(ws1, 3, REL_COLS, rrows, center_cols=(1, 6, 7))
    for i, x in enumerate(rels):
        if x["from_t"] in tsheet:
            link(ws1.cell(4 + i, 2), tsheet[x["from_t"]], "B2")
        if x["to_t"] in tsheet:
            link(ws1.cell(4 + i, 4), tsheet[x["to_t"]], "B2")
    set_widths(ws1, [6, 26, 22, 26, 18, 18, 14, 26])

    # --- 02_ERD
    ws2["A1"] = "ERD — %s" % a.system
    ws2["A1"].font = Font(name=ST.FONT, size=13, bold=True, color="FF1F4E78")
    link(ws2["H1"], "00_Overview", "A1", "← Overview")
    if erd and a.erd_png and os.path.exists(a.erd_png):
        note = "File: %s · ve %d bang%s" % (os.path.basename(a.erd_png), erd["drawn"],
                                            (" · an %d bang it lien ket" % erd["omitted"]) if erd["omitted"] else "")
        ws2["A2"] = note
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(a.erd_png)
            scale = min(1.0, 1500.0 / max(img.width, 1))
            img.width, img.height = int(img.width * scale), int(img.height * scale)
            ws2.add_image(img, "A4")
        except ImportError:
            ws2["A3"] = "Thieu Pillow — khong nhung duoc anh, xem file PNG di kem"
    else:
        ws2["A2"] = "ERD chua sinh (chay lai voi --erd-png)."

    # --- T_<table>
    for t in tables:
        name = t["Table"]
        ws = wb.create_sheet(tsheet[name])
        link(ws["A1"], "00_Overview", "A%d" % (hdr + 1 + tables.index(t)), "← Overview")
        ev = v(t.get("Evidence"))
        r = kv_block(ws, 2, [("TABLE", name), ("Purpose", v(t.get("Purpose"))),
                             ("PK", v(t.get("PK"), NA)), ("FK", v(t.get("FK"), NA)),
                             ("Evidence / Confidence", "%s · %s" % (ev, v(t.get("Confidence")))),
                             ("Note", v(t.get("Note"), NA))])
        ws["B2"].font = Font(name=ST.FONT, size=12, bold=True, color="FF1F4E78")
        for rr in range(2, r):
            ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=len(GRID_COLS))
        grow = []
        for i, c in enumerate(cols.get(name, []), start=1):
            fk = v(c.get("FK"), NA)
            grow.append([i, v(c.get("Column")), v(c.get("Type")), v(c.get("Format")),
                         v(c.get("Max Length"), NA), v(c.get("Nullable")), v(c.get("Default"), NA),
                         v(c.get("PK"), "No"), fk, v(c.get("Constraint"), NA), v(c.get("Meaning")),
                         v(c.get("Evidence"), ev), v(c.get("Confidence"))])
        if not grow:
            grow = [[1, S.UNKNOWN] + [S.UNKNOWN] * (len(GRID_COLS) - 2)]
        write_grid(ws, r + 1, GRID_COLS, grow, center_cols=(1, 5, 6, 8, 13))
        for i, c in enumerate(cols.get(name, [])):
            m = re.match(r"^([\w.$]+)\.([\w$]+)$", (c.get("FK") or "").strip())
            if m and m.group(1) in tsheet:
                link(ws.cell(r + 2 + i, 9), tsheet[m.group(1)], "B2")
        ws.freeze_panes = ws.cell(r + 2, 3)
        set_widths(ws, [22, 24, 22, 22, 12, 10, 18, 7, 20, 30, 44, 14, 12])

    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 100
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    wb.save(a.out)
    import json
    print(json.dumps({"out": a.out, "tables": len(tables), "columns": ncol,
                      "relationships": len(rels), "logical": len(rels) - n_fk,
                      "erd": a.erd_png if erd else None,
                      "erd_omitted": erd["omitted"] if erd else 0}, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
