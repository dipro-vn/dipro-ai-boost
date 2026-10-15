#!/usr/bin/env python3
"""Output US — export Danh sach User Story (USER-STORIES.md) ra workbook .xlsx.

Doc rule day du: .claude/ba-agent/user-stories.md  (+ story-splitting.md, ac-standard.md)

Script lam 2 viec:
  1. VALIDATE register — gate CAU TRUC (khong phai gate chat luong nghiep vu):
     du cot, US ID hop le/khong trung, truong bat buoc khong rong, AC ref dung
     dang, INVEST du 6 ky hieu, Epic nhat quan, FR coverage (neu co --expect-fr).
  2. EXPORT xlsx 3 sheet: Guideline · User Stories · Summary.

Khong co --spec: script KHONG doc SPEC.md -> khong kiem AC-NN co ton tai that.
Co --spec <SPEC.md>: doi chieu THAM CHIEU (AC-NN ton tai, AC mo coi, Screen Code,
FR No., OQ). Van KHONG cham chat luong AC — phan do BA tu ra checklist
ac-standard.md muc 9.

Usage:
  python3 export-user-stories.py <USER-STORIES.md> --out <user_stories.xlsx> \
      [--feature NAME] [--version v2] [--date DDMMYYYY] [--expect-fr "1-58"] \
      [--spec <SPEC.md>] [--report <file.md>]

Exit: 0 = PASS | 1 = co FAIL (khong ghi file) | 2 = thieu openpyxl / khong doc duoc input
"""
import argparse
import datetime
import re
import sys

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.stderr.write("FAIL: thieu openpyxl. Chay: pip install openpyxl\n")
    sys.exit(2)

# --- cot bat buoc trong register markdown ---------------------------------
REQUIRED = ["US ID", "Tieu de", "As a", "I want to", "So that", "FR No.", "AC ref", "INVEST", "Spike"]
OPTIONAL = ["Epic", "Screen Code", "Flow", "Ghi chu"]

HEADER_ALIASES = {
    "epic": "Epic",
    "us id": "US ID", "id": "US ID", "story id": "US ID",
    "tiêu đề": "Tieu de", "tieu de": "Tieu de", "title": "Tieu de",
    "as a": "As a", "persona": "As a",
    "i want to": "I want to", "i want": "I want to",
    "so that": "So that",
    "fr no.": "FR No.", "fr no": "FR No.", "fr": "FR No.",
    "screen code": "Screen Code", "screen": "Screen Code",
    "flow": "Flow",
    "ac ref": "AC ref", "ac": "AC ref",
    "invest": "INVEST",
    "spike": "Spike",
    "ghi chú": "Ghi chu", "ghi chu": "Ghi chu", "note": "Ghi chu",
}

US_ID_RE = re.compile(r"US-[A-Z0-9]{2,6}-\d{3}")
EPIC_RE = re.compile(r"^(EP-\d{2})\s*[—–:\-]*\s*(.*)$")
AC_TOKEN_RE = re.compile(r"(AC-\d{2,3}|PENDING-OQ-\d{2,3})")
INVEST_RE = re.compile(r"([INVEST])\s*(✅|⚠️?|❌)")
EMPTY_MARKS = {"", "-", "—", "–", "n/a", "na"}
GENERIC_PERSONA = {"user", "users", "người dùng", "nguoi dung", "người dùng cuối",
                   "mọi người", "moi nguoi", "everyone", "end user", "end-user"}
SPIKE_YES = {"có", "co", "yes", "y", "spike"}

THIN = Side(style="thin", color="FFD0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = "FF0969DA"
INPUT_FILL = "FFFFF8C5"      # vang — o Dev / PO dien
WARN_FILL = "FFFFF9EB"
EPIC_FILL = "FFE8F4FD"

GUIDELINE = [
    ("H", "Danh sách User Story"),
    ("P", "File này sinh tự động từ USER-STORIES.md (nguồn sự thật do BA viết). Xin đừng sửa trực tiếp "
          "trong file này — gửi feedback cho BA để cập nhật nguồn rồi xuất lại."),
    ("H2", "1. Cấu trúc"),
    ("P", "Tối đa 2 tầng: Epic → User Story. Epic là tuỳ chọn (ô Epic ghi — nếu không dùng)."),
    ("H2", "2. Ý nghĩa các cột"),
    ("P", "US ID — mã Story, không đánh lại số giữa các version."),
    ("P", "As a / I want to / So that — persona, hành động, giá trị nghiệp vụ."),
    ("P", "FR No. / Screen Code / Flow — trace về chức năng nguồn, màn hình, business flow (Output 1). "
          "Chỉ để tra cứu, không vẽ lên Figma."),
    ("P", "AC ref — mã Acceptance Criteria trong SPEC.md ## Acceptance Criteria. Nội dung AC KHÔNG chép "
          "vào đây (tránh lệch). PENDING-OQ-NN = AC chưa viết được vì còn Open Question."),
    ("P", "INVEST — BA tự đánh giá: ✅ đạt · ⚠️ có vấn đề · (❌ không xuất hiện: phải sửa trước khi xuất)."),
    ("P", "Spike — Có = Story chưa ước lượng được vì còn thiếu thông tin (xem cột Ghi chú)."),
    ("H2", "3. Hai cột vàng"),
    ("P", "Estimate (Dev điền) và Ưu tiên (PO điền) để trống có chủ ý — BA không ước lượng effort "
          "và không quyết định ưu tiên."),
    ("H2", "4. Dữ liệu"),
    ("P", "Mọi ví dụ trong file dùng dữ liệu mẫu — không có dữ liệu thật của khách hàng."),
]


class Report:
    def __init__(self):
        self.rows = []

    def add(self, name, status, detail=""):
        self.rows.append((name, status, detail))

    def counts(self):
        return tuple(sum(1 for r in self.rows if r[1] == s) for s in ("PASS", "FAIL", "WARN"))


# --- parse ------------------------------------------------------------------
def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse_register(path):
    try:
        with open(path, encoding="utf-8") as fh:
            lines = fh.read().splitlines()
    except OSError as exc:
        sys.stderr.write("FAIL: khong doc duoc %s (%s)\n" % (path, exc))
        sys.exit(2)

    header_idx, keys = None, None
    for i, line in enumerate(lines):
        if line.count("|") < 4:
            continue
        mapped = [HEADER_ALIASES.get(c.lower().strip()) for c in split_row(line)]
        if mapped.count("US ID") == 1 and mapped.count("As a") == 1:
            header_idx, keys = i, mapped
            break
    if header_idx is None:
        sys.stderr.write(
            "FAIL: khong tim thay bang User Story trong %s. "
            "Header phai co cot 'US ID' va 'As a' — xem user-stories.md muc 3\n" % path)
        sys.exit(2)

    rows = []
    for line in lines[header_idx + 2:]:
        if line.count("|") < 4:
            if rows:
                break
            continue
        cells = split_row(line)
        if set("".join(cells)) <= set("-: "):
            continue
        row = {k: v for k, v in zip(keys, cells) if k}
        if row.get("US ID"):
            rows.append(row)
    return keys, rows


def parse_fr_set(text):
    """'#12,13' / '#19-21' / '1-58' / '#12, #14 SYSTEM' -> set(int)."""
    out = set()
    for a, b in re.findall(r"(\d+)\s*[-–]\s*(\d+)", text):
        lo, hi = int(a), int(b)
        if lo <= hi:
            out.update(range(lo, hi + 1))
    stripped = re.sub(r"\d+\s*[-–]\s*\d+", " ", text)
    out.update(int(n) for n in re.findall(r"\d+", stripped))
    return out



# --- doc SPEC.md (tuy chon, --spec) ----------------------------------------
def _sections(lines):
    """Tach SPEC theo heading '## ' (bo qua heading nam trong code fence)."""
    secs, cur, buf, fence = {}, None, [], False
    for ln in lines:
        if ln.strip().startswith("```"):
            fence = not fence
        m = None if fence else re.match(r"^##\s+(.*\S)\s*$", ln)
        if m:
            if cur is not None:
                secs[cur] = buf
            cur, buf = m.group(1).strip().lower(), []
        else:
            buf.append(ln)
    if cur is not None:
        secs[cur] = buf
    return secs


def parse_spec(path):
    try:
        with open(path, encoding="utf-8") as fh:
            lines = fh.read().splitlines()
    except OSError as exc:
        sys.stderr.write("FAIL: khong doc duoc --spec %s (%s)\n" % (path, exc))
        sys.exit(2)
    secs = _sections(lines)

    def sec(prefix):
        for k, v in secs.items():
            if k.startswith(prefix):
                return v
        return None

    info = {"has_ac": False, "ac": set(), "ac_removed": set(), "pending_ac": set(),
            "oq": set(), "screens": set(), "frs": set(), "has_screens": False}
    ac = sec("acceptance criteria")
    if ac is not None:
        info["has_ac"] = True
        for i, ln in enumerate(ac):
            m = re.match(r"^#{3,4}\s*(AC-\d{2,3})\b(.*)$", ln.strip())
            if not m:
                continue
            ident, rest = m.group(1), m.group(2)
            body = "\n".join(ac[i:i + 6])
            if re.search(r"đã bỏ|da bo|removed", rest, re.I):
                info["ac_removed"].add(ident)
            else:
                info["ac"].add(ident)
                if "PENDING-OQ" in body:
                    info["pending_ac"].add(ident)
    oq = sec("open questions")
    if oq is not None:
        info["oq"] = set(re.findall(r"OQ-\d{2,3}", "\n".join(oq)))
    scr = sec("screens")
    if scr is not None:
        info["has_screens"] = True
        header = None
        for ln in scr:
            if ln.count("|") >= 3:
                cells = [c.strip() for c in ln.strip().strip("|").split("|")]
                low = [c.lower() for c in cells]
                if "screen code" in low:
                    header = low
                    continue
                if header and not set("".join(cells)) <= set("-: "):
                    row = dict(zip(header, cells))
                    code = row.get("screen code", "").strip("` ")
                    if code:
                        info["screens"].add(code)
                    info["frs"] |= parse_fr_set(row.get("fr no.", ""))
            elif ln.lstrip().startswith(">") and "#" in ln:
                info["frs"] |= {int(n) for n in re.findall(r"#(\d+)", ln)}
    return info


def is_empty(v):
    return (v or "").strip().lower() in EMPTY_MARKS


def parse_epic(v):
    if is_empty(v):
        return None
    m = EPIC_RE.match(v.strip())
    return (m.group(1), m.group(2).strip()) if m else False


# --- validate ---------------------------------------------------------------
def validate(keys, rows, expect_fr, rep, spec=None):
    missing = [c for c in REQUIRED if c not in keys]
    rep.add("Cot bat buoc", "FAIL" if missing else "PASS",
            ("thieu: %s" % ", ".join(missing)) if missing else "du %d cot" % len(REQUIRED))
    if missing:
        return {}
    if not rows:
        rep.add("So Story", "FAIL", "register khong co dong nao")
        return {}
    rep.add("So Story", "PASS", "%d Story" % len(rows))

    ids = [r["US ID"] for r in rows]
    dup = sorted({i for i in ids if ids.count(i) > 1})
    rep.add("US ID khong trung", "FAIL" if dup else "PASS",
            ("trung: %s" % ", ".join(dup)) if dup else "%d ID duy nhat" % len(set(ids)))

    bad = [i for i in ids if not US_ID_RE.fullmatch(i)]
    rep.add("Format US ID", "FAIL" if bad else "PASS",
            ("sai format (can US-FEAT-001): %s" % ", ".join(bad)) if bad else "US-<FEAT>-NNN")

    for col, label in (("As a", "As a"), ("I want to", "I want to"), ("So that", "So that")):
        empty = [r["US ID"] for r in rows if is_empty(r.get(col))]
        rep.add("%s khong rong" % label, "FAIL" if empty else "PASS",
                ("thieu: %s" % ", ".join(empty)) if empty else "day du")

    # FR No.
    frs, no_fr = [], []
    for r in rows:
        s = parse_fr_set(r.get("FR No.", ""))
        frs.append(s)
        if not s:
            no_fr.append(r["US ID"])
    rep.add("FR No. (trace)", "FAIL" if no_fr else "PASS",
            ("thieu so FR: %s" % ", ".join(no_fr)) if no_fr else "moi Story co FR")

    # AC ref
    bad_ac = []
    for r in rows:
        cell = r.get("AC ref", "")
        tokens = [t for t in re.split(r"[,\s;]+", cell) if t]
        if not tokens or any(not AC_TOKEN_RE.fullmatch(t) for t in tokens):
            bad_ac.append(r["US ID"])
    rep.add("AC ref hop le", "FAIL" if bad_ac else "PASS",
            ("sai/thieu (can AC-NN hoac PENDING-OQ-NN): %s" % ", ".join(bad_ac)) if bad_ac
            else "AC-NN / PENDING-OQ-NN")

    # INVEST
    invest, bad_inv, fail_inv = {}, [], []
    for r in rows:
        found = {}
        for letter, mark in INVEST_RE.findall(r.get("INVEST", "")):
            found.setdefault(letter, "⚠️" if mark.startswith("⚠") else mark)
        invest[r["US ID"]] = found
        if set(found) != set("INVEST"):
            bad_inv.append(r["US ID"])
        elif "❌" in found.values():
            fail_inv.append(r["US ID"])
    rep.add("INVEST du 6 ky hieu", "FAIL" if bad_inv else "PASS",
            ("thieu/sai (can I✅ N✅ V✅ E✅ S✅ T✅): %s" % ", ".join(bad_inv)) if bad_inv else "du 6")
    rep.add("INVEST khong con ❌", "FAIL" if fail_inv else "PASS",
            ("sua Story truoc khi export: %s" % ", ".join(fail_inv)) if fail_inv else "khong co ❌")

    # Epic
    epic_names, bad_epic = {}, []
    for r in rows:
        e = parse_epic(r.get("Epic", ""))
        if e is False:
            bad_epic.append(r["US ID"])
        elif e:
            epic_names.setdefault(e[0], set()).add(e[1])
    incon = sorted(k for k, v in epic_names.items() if len(v) > 1)
    rep.add("Epic dung format (EP-NN)", "FAIL" if bad_epic else "PASS",
            ("sai format: %s" % ", ".join(bad_epic)) if bad_epic
            else ("%d Epic" % len(epic_names)) if epic_names else "khong dung Epic")
    rep.add("Epic cung ma cung ten", "FAIL" if incon else "PASS",
            ("ten khac nhau: %s" % ", ".join(incon)) if incon else "nhat quan")

    # Coverage
    covered = set().union(*frs) if frs else set()
    cov = {"covered": covered}
    if expect_fr is not None:
        miss = sorted(expect_fr - covered)
        extra = sorted(covered - expect_fr)
        cov.update(expected=expect_fr, missing=miss)
        rep.add("US COVERAGE (FR)", "FAIL" if miss else "PASS",
                ("THIEU FR: %s" % ", ".join("#%d" % n for n in miss)) if miss
                else "phu %d/%d FR" % (len(expect_fr & covered), len(expect_fr)))
        if extra:
            rep.add("FR ngoai pham vi --expect-fr", "WARN",
                    ", ".join("#%d" % n for n in extra))

    # ---- WARN -------------------------------------------------------------
    generic = [r["US ID"] for r in rows if (r.get("As a", "").strip().lower() in GENERIC_PERSONA)]
    if generic:
        rep.add("Persona cu the", "WARN", "persona chung chung: %s" % ", ".join(generic))
    echo = [r["US ID"] for r in rows
            if r.get("I want to", "").strip() and
            r.get("I want to", "").strip().lower() in r.get("So that", "").strip().lower()]
    if echo:
        rep.add("So that khac I want to", "WARN", "lap lai goal: %s" % ", ".join(echo))
    compound = [r["US ID"] for r in rows if re.search(r"\svà\s", r.get("Tieu de", ""), re.I)]
    if compound:
        rep.add("Tieu de 1 muc tieu", "WARN", "co 'va' — xem xet tach: %s" % ", ".join(compound))
    spike_bad = [r["US ID"] for r in rows
                 if invest.get(r["US ID"], {}).get("E") == "⚠️"
                 and r.get("Spike", "").strip().lower() not in SPIKE_YES]
    if spike_bad:
        rep.add("E⚠️ co Spike", "WARN", "Estimable ⚠️ nhung Spike != Co: %s" % ", ".join(spike_bad))
    spike_odd = [r["US ID"] for r in rows
                 if r.get("Spike", "").strip().lower() in SPIKE_YES
                 and invest.get(r["US ID"], {}).get("E") == "✅"]
    if spike_odd:
        rep.add("Spike nhat quan voi E", "WARN", "Spike = Co nhung E✅: %s" % ", ".join(spike_odd))
    small = [r["US ID"] for r in rows if invest.get(r["US ID"], {}).get("S") == "⚠️"]
    if small:
        rep.add("Small", "WARN", "S⚠️ — xem story-splitting.md: %s" % ", ".join(small))

    # ---- doi chieu voi SPEC (chi khi co --spec) ---------------------------
    if spec is not None:
        if not spec["has_ac"]:
            rep.add("SPEC co ## Acceptance Criteria", "FAIL",
                    "khong tim thay section — viet AC theo ac-standard.md truoc")
        else:
            used, dangling, removed = set(), [], []
            for r in rows:
                for t in re.findall(r"AC-\d{2,3}", r.get("AC ref", "")):
                    used.add(t)
                    if t in spec["ac_removed"]:
                        removed.append("%s→%s" % (r["US ID"], t))
                    elif t not in spec["ac"]:
                        dangling.append("%s→%s" % (r["US ID"], t))
            rep.add("AC ref ton tai trong SPEC", "FAIL" if (dangling or removed) else "PASS",
                    ("khong co/da bo: %s" % ", ".join(dangling + removed))
                    if (dangling or removed) else "%d AC duoc tham chieu deu co that" % len(used))
            orphan = sorted(spec["ac"] - used - spec["pending_ac"])
            if orphan:
                rep.add("AC mo coi (khong Story nao tro toi)", "WARN", ", ".join(orphan))
        # PENDING-OQ phai co OQ that
        pend = set()
        for r in rows:
            pend |= {t.replace("PENDING-", "") for t in re.findall(r"PENDING-OQ-\d{2,3}", r.get("AC ref", ""))}
        miss_oq = sorted(pend - spec["oq"])
        if pend:
            rep.add("PENDING-OQ ton tai trong SPEC ## Open Questions", "WARN" if miss_oq else "PASS",
                    ("khong thay: %s" % ", ".join(miss_oq)) if miss_oq else "%d OQ" % len(pend))
        if spec["has_screens"]:
            bad_sc = []
            for r in rows:
                for t in re.split(r"[,\s;]+", r.get("Screen Code", "")):
                    if t and not is_empty(t) and t not in spec["screens"]:
                        bad_sc.append("%s→%s" % (r["US ID"], t))
            rep.add("Screen Code ton tai trong SPEC ## Screens", "WARN" if bad_sc else "PASS",
                    ("khong thay: %s" % ", ".join(bad_sc[:8])) if bad_sc else "khop")
            if spec["frs"]:
                unk = sorted(covered - spec["frs"])
                rep.add("FR No. co trong SPEC ## Screens", "WARN" if unk else "PASS",
                        ("khong thay: %s" % ", ".join("#%d" % n for n in unk)) if unk else "khop")
    return {"invest": invest, "epics": epic_names, "cov": cov}


# --- build workbook ---------------------------------------------------------
def _hdr(cell, text, fill=HEADER_FILL):
    cell.value = text
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Arial", size=10, bold=True,
                     color="FFFFFFFF" if fill == HEADER_FILL else "FF1F2328")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def build_workbook(rows, info, feature, version, date_str, out):
    wb = openpyxl.Workbook()

    # Sheet 1 — Guideline
    ws = wb.active
    ws.title = "Guideline"
    ws.column_dimensions["A"].width = 118
    r = 1
    for kind, text in GUIDELINE:
        c = ws.cell(row=r, column=1, value=text)
        if kind == "H":
            c.font = Font(name="Arial", size=14, bold=True, color="FF0969DA")
            r += 1
            ws.cell(row=r, column=1,
                    value="Feature: %s · Version: %s · Ngày: %s" % (feature, version, date_str)
                    ).font = Font(name="Arial", size=9, italic=True, color="FF57606A")
        elif kind == "H2":
            c.font = Font(name="Arial", size=11, bold=True)
        else:
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 30
        r += 2 if kind == "H2" else 1

    # Sheet 2 — User Stories
    us = wb.create_sheet("User Stories")
    headers = [
        ("Epic ID", 9), ("Epic", 26), ("US ID", 14), ("Tiêu đề", 30),
        ("As a", 22), ("I want to", 40), ("So that", 40),
        ("FR No.", 11), ("Screen Code", 15), ("Flow", 9), ("AC ref", 16),
        ("INVEST", 24), ("Spike", 8), ("Ghi chú", 32),
        ("Estimate (Dev điền)", 16), ("Ưu tiên (PO điền)", 16),
    ]
    for i, (name, width) in enumerate(headers, start=1):
        fill = INPUT_FILL if name.endswith("điền)") else HEADER_FILL
        _hdr(us.cell(row=1, column=i), name, fill)
        us.column_dimensions[get_column_letter(i)].width = width
    us.row_dimensions[1].height = 34

    def sort_key(r):
        e = parse_epic(r.get("Epic", ""))
        return (0, e[0]) if e else (1, ""), r.get("US ID", "")

    ordered = sorted(rows, key=lambda r: (sort_key(r)[0], sort_key(r)[1]))
    for i, row in enumerate(ordered, start=2):
        e = parse_epic(row.get("Epic", ""))
        eid, ename = (e if e else ("—", "—"))
        vals = [
            eid, ename, row.get("US ID", ""), row.get("Tieu de", ""),
            row.get("As a", ""), row.get("I want to", ""), row.get("So that", ""),
            row.get("FR No.", ""),
            "—" if is_empty(row.get("Screen Code")) else row.get("Screen Code"),
            "—" if is_empty(row.get("Flow")) else row.get("Flow"),
            row.get("AC ref", ""), row.get("INVEST", ""),
            "—" if is_empty(row.get("Spike")) else row.get("Spike"),
            row.get("Ghi chu", ""), "", "",
        ]
        warn = "⚠" in row.get("INVEST", "")
        for j, v in enumerate(vals, start=1):
            c = us.cell(row=i, column=j, value=v)
            c.border = BORDER
            c.font = Font(name="Arial", size=10, bold=(j == 3))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if j >= 15:
                c.fill = PatternFill("solid", fgColor=INPUT_FILL)
            elif j <= 2 and e:
                c.fill = PatternFill("solid", fgColor=EPIC_FILL)
            elif j in (12, 13) and (warn or v == "Có"):
                c.fill = PatternFill("solid", fgColor=WARN_FILL)
        us.row_dimensions[i].height = 52
    last = len(ordered) + 1
    us.freeze_panes = "E2"
    us.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), last)

    # Sheet 3 — Summary
    sm = wb.create_sheet("Summary")
    sm.column_dimensions["A"].width = 34
    sm.column_dimensions["B"].width = 14
    sm.cell(row=1, column=1, value="User Stories — %s (%s)" % (feature, version)
            ).font = Font(name="Arial", size=13, bold=True, color="FF0969DA")

    def line(r, label, value, bold=False):
        sm.cell(row=r, column=1, value=label).font = Font(name="Arial", size=10, bold=bold)
        c = sm.cell(row=r, column=2, value=value)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER

    r = 3
    line(r, "Tổng số Story", len(rows), True)
    r += 1
    line(r, "Số Epic", len(info.get("epics", {})) or "không dùng Epic")
    r += 1
    line(r, "Story cần Spike", sum(1 for x in rows if x.get("Spike", "").strip().lower() in SPIKE_YES))
    r += 1
    line(r, "Story chờ OQ (PENDING-OQ)", sum(1 for x in rows if "PENDING-OQ" in x.get("AC ref", "")))

    r += 2
    _hdr(sm.cell(row=r, column=1), "INVEST — số Story ⚠️")
    _hdr(sm.cell(row=r, column=2), "Số Story")
    names = {"I": "I — Independent", "N": "N — Negotiable", "V": "V — Valuable",
             "E": "E — Estimable", "S": "S — Small", "T": "T — Testable"}
    for letter in "INVEST":
        r += 1
        n = sum(1 for v in info.get("invest", {}).values() if v.get(letter) == "⚠️")
        sm.cell(row=r, column=1, value=names[letter]).border = BORDER
        c = sm.cell(row=r, column=2, value=n)
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")

    if info.get("epics"):
        r += 2
        _hdr(sm.cell(row=r, column=1), "Epic")
        _hdr(sm.cell(row=r, column=2), "Số Story")
        counts = {}
        for x in rows:
            e = parse_epic(x.get("Epic", ""))
            if e:
                counts[e[0]] = counts.get(e[0], 0) + 1
        for eid, n in sorted(counts.items()):
            r += 1
            sm.cell(row=r, column=1, value="%s %s" % (eid, " / ".join(sorted(info["epics"][eid])))
                    ).border = BORDER
            c = sm.cell(row=r, column=2, value=n)
            c.border = BORDER
            c.alignment = Alignment(horizontal="center")

    cov = info.get("cov", {})
    r += 2
    if "expected" in cov:
        line(r, "FR cần phủ", len(cov["expected"]))
        r += 1
        line(r, "FR đã phủ", len(cov["expected"] & cov["covered"]))
        r += 1
        line(r, "FR THIẾU", ", ".join("#%d" % n for n in cov["missing"]) or "[]")
    else:
        line(r, "FR coverage", "N/A — không có --expect-fr")

    wb.save(out)
    return ordered


def render(rep, src, out, n, cov):
    lines = ["# Output US — Export Report", "",
             "Source: `%s`" % src, "Output: `%s`" % out, "",
             "| Check | Status | Chi tiết |", "|---|---|---|"]
    for name, status, detail in rep.rows:
        icon = {"PASS": "✅ PASS", "FAIL": "❌ FAIL", "WARN": "⚠️ WARN"}[status]
        lines.append("| %s | %s | %s |" % (name, icon, detail.replace("|", "/")))
    ok, fail, warn = rep.counts()
    lines += ["", "**%d checks · %d PASS · %d FAIL · %d WARN**" % (len(rep.rows), ok, fail, warn), "",
              "US COUNT: %d" % n]
    if "expected" in cov:
        miss = ", ".join("#%d" % x for x in cov["missing"])
        lines.append("US COVERAGE — phủ %d/%d · THIẾU: [%s]"
                     % (len(cov["expected"] & cov["covered"]), len(cov["expected"]), miss))
    else:
        lines.append("US COVERAGE: N/A — không có --expect-fr")
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("register", help="duong dan USER-STORIES.md")
    ap.add_argument("--out", required=True, help="duong dan xlsx xuat ra")
    ap.add_argument("--feature", default="<feature>", help="ten feature")
    ap.add_argument("--version", default="v1", help="version BA dang chay (VD v2)")
    ap.add_argument("--date", default="", help="DDMMYYYY — mac dinh hom nay")
    ap.add_argument("--expect-fr", default=None,
                    help="tap FR can phu, VD '1-58' hoac '1-20,25,30-35'. Bo neu nguon khong danh so")
    ap.add_argument("--spec", default=None,
                    help="duong dan SPEC.md — bat doi chieu AC ref / Screen Code / FR / OQ voi SPEC that")
    ap.add_argument("--report", help="ghi report markdown ra file")
    args = ap.parse_args()

    date_str = args.date or datetime.date.today().strftime("%d%m%Y")
    expect = parse_fr_set(args.expect_fr) if args.expect_fr else None
    if args.expect_fr and not expect:
        sys.stderr.write("FAIL: --expect-fr khong doc duoc: %r\n" % args.expect_fr)
        sys.exit(2)

    keys, rows = parse_register(args.register)
    rep = Report()
    spec = parse_spec(args.spec) if args.spec else None
    info = validate(keys, rows, expect, rep, spec)
    ok, fail, warn = rep.counts()
    cov = info.get("cov", {}) if info else {}

    if fail:
        out = render(rep, args.register, "(khong ghi — con FAIL)", len(rows), cov)
        print(out)
        if args.report:
            with open(args.report, "w", encoding="utf-8") as fh:
                fh.write(out)
        print("\nFAIL > 0 -> sua USER-STORIES.md roi chay lai. Khong ghi xlsx.")
        sys.exit(1)

    rows = build_workbook(rows, info, args.feature, args.version, date_str, args.out)
    out = render(rep, args.register, args.out, len(rows), cov)
    print(out)
    if args.report:
        with open(args.report, "w", encoding="utf-8") as fh:
            fh.write(out)
        print("Report: %s" % args.report)
    sys.exit(0)


if __name__ == "__main__":
    main()
