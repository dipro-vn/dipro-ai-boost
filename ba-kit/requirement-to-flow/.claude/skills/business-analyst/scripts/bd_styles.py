#!/usr/bin/env python3
"""Style chuan cho Basic Design workbook — dinh nghia TUONG MINH.

Vi sao co file nay:
  Truoc day populate logic copy style tu "hang xom" (row truoc do trong sheet).
  Khi master workbook sach (chua co row du lieu nao), no roi ve copy tu HEADER
  -> data row bi chu trang nen xanh, hoac mat het duong ke.
  Style phai duoc dinh nghia tuong minh, khong phu thuoc noi dung san co.

Spec: .claude/ba-agent/basic-design/workbook-structure.md muc 2.8
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "Arial"

# --- mau ---
HEADER_FILL = "FFDAEEF3"   # xanh nhat — header bang du lieu
TITLE_FILL = "FF1F4E78"    # xanh dam — hang title cua Screen Index
BANNER_FILL = "FFF2F2F2"   # xam nhat — banner phan nhom trong Screen Error message
DATA_FILL = "FFFFFFFF"     # trang — moi data row
ERR_TITLE_FILL = "FFFCE4E4"  # hong nhat — title ERROR SCENARIOS

BLACK = "FF000000"
WHITE = "FFFFFFFF"
CODE_GREEN = "FF1F5C1F"    # ma code o cot dau (theo Common mesage)
ERR_RED = "FFC00000"

_thin = Side(style="thin", color="FF000000")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _apply(cell, fill, font, align, border=BORDER):
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = font
    cell.alignment = align
    cell.border = border


def header(cell, size=10):
    """Header cua bang du lieu — xanh nhat, chu den dam, can giua."""
    _apply(cell, HEADER_FILL,
           Font(name=FONT, size=size, bold=True, color=BLACK),
           Alignment(horizontal="center", vertical="center", wrap_text=True))


def title(cell, size=10):
    """Hang title dam (Screen Index row 5) — xanh dam, chu trang."""
    _apply(cell, TITLE_FILL,
           Font(name=FONT, size=size, bold=True, color=WHITE),
           Alignment(horizontal="center", vertical="center", wrap_text=True))


def data(cell, size=11, code=False, center=False):
    """Data row — nen trang, chu den, ke vien day du.

    code=True   -> cot ma (Message Code / Screen ID): xanh la dam, in dam
    center=True -> can giua (cot so thu tu, type, status...)
    """
    _apply(cell, DATA_FILL,
           Font(name=FONT, size=size, bold=bool(code),
                color=CODE_GREEN if code else BLACK),
           Alignment(horizontal="center" if center else "left",
                     vertical="center", wrap_text=True))


def banner(cell, size=11):
    """Banner phan nhom trong Screen Error message — xam nhat, chu dam."""
    _apply(cell, BANNER_FILL,
           Font(name=FONT, size=size, bold=True, color=BLACK),
           Alignment(horizontal="left", vertical="center", wrap_text=True))


def err_title(cell, size=11):
    """Title bang ERROR SCENARIOS trong screen sheet — hong nhat, chu do dam."""
    _apply(cell, ERR_TITLE_FILL,
           Font(name=FONT, size=size, bold=True, color=ERR_RED),
           Alignment(horizontal="left", vertical="center", wrap_text=True))


def row(ws, r, cols, kind="data", **kw):
    """Ap style cho ca 1 hang. cols = iterable chi so cot (1-based)."""
    fn = {"header": header, "title": title, "data": data,
          "banner": banner, "err_title": err_title}[kind]
    for c in cols:
        fn(ws.cell(row=r, column=c), **kw)
