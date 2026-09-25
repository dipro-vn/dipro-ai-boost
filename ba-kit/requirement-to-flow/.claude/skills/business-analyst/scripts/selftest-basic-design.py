#!/usr/bin/env python3
"""Self-test cho Quality Gate O5 — chung minh gate KHONG pass rong.

Cach dung:
    python3 selftest-basic-design.py <workbook-da-PASS.xlsx> [--before <backup.xlsx>]

Lay 1 workbook da PASS gate, tiem tung loi da biet vao ban copy, roi xac nhan
gate bat dung check ky vong. Gate PASS ma khong co self-test nay = chua chung
minh duoc gi: mot gate hong van "PASS" moi thu.

Bo test nay da tung tim ra 2 loi that cua gate:
  - check 14 CRASH khi border side = None (AttributeError, khong phai FAIL)
  - check 13 do im.width thay vi anchor.ext -> bo sot anh bi phong to trong Excel
"""
import argparse
import importlib.util
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.styles import Border, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
GATE = os.path.join(HERE, "verify-basic-design.py")

# Dung lai chinh logic cua gate de xac dinh working sheet.
# Tu viet lai se sai: sheet `Screen Index` co cot F = "Sheet Type", nen
# ws["F6"] cua no cung bang "WORKING_SCREEN" — phai loai sheet he thong.
_spec = importlib.util.spec_from_file_location("bd_gate", GATE)
_gate = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_gate)


def unmerge_at(ws, r):
    for rng in [str(x) for x in ws.merged_cells.ranges]:
        if ws[rng.split(":")[0]].row == r:
            ws.unmerge_cells(rng)


def screen_sheets(wb):
    return _gate.working_sheets(wb)


def first_screen_sheet(wb):
    ws = screen_sheets(wb)
    return ws[0] if ws else None


# moi case: (ten, check ky vong, ham tiem loi)
def _no_border(wb):
    for c in range(1, 11):
        wb["Screen Error message"].cell(5, c).border = Border()


def _header_fill(wb):
    for c in range(1, 13):
        wb["Screen Index"].cell(6, c).fill = PatternFill("solid", fgColor="FF1F4E78")


def _drop_err_table(wb):
    ws = first_screen_sheet(wb)
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 8).value
        if v and "ERROR SCENARIOS" in str(v):
            for rr in range(r, r + 12):
                unmerge_at(ws, rr)
                for c in range(8, 19):
                    ws.cell(rr, c).value = None
            return


def _touch_sample(wb):
    wb["Sample"]["B1"] = "BI GHI BAY"


def _blow_up_image(wb):
    ws = first_screen_sheet(wb)
    if ws and ws._images:
        im = ws._images[0]
        im.anchor.ext.cx = int(1600 * 9525)
        im.anchor.ext.cy = int(1200 * 9525)


def _skip_chg_id(wb):
    h = wb["Change History"]
    r = 6
    while h.cell(r, 1).value is not None:
        r += 1
    h.cell(r - 1, 1).value = "CHG-0099"


def _drop_index_row(wb):
    wb["Screen Index"].delete_rows(6)


def _ghost_code(wb):
    first_screen_sheet(wb)["Q9"] = "loi -> E_ZZ_999"


CASES = [
    ("xoa duong ke data row", 14, _no_border),
    ("data row to fill header", 14, _header_fill),
    ("xoa bang ERROR SCENARIOS", 15, _drop_err_table),
    ("ghi bay vao Sample", 1, _touch_sample),
    ("phong to anh vuot khung", 13, _blow_up_image),
    ("Change ID nhay coc", 12, _skip_chg_id),
    ("xoa row Screen Index", 5, _drop_index_row),
    ("tro Message Code ma", 6, _ghost_code),
]


def run_gate(path, before, screens):
    cmd = [sys.executable, GATE, path, "--expect-screens", screens]
    if before:
        cmd += ["--before", before]
    p = subprocess.run(cmd, capture_output=True, text=True)
    out = p.stdout + p.stderr
    if "Traceback" in out:
        return "CRASH", []
    fails = []
    for line in out.splitlines():
        if line.startswith("| ") and "**FAIL**" in line:
            try:
                fails.append(int(line.split("|")[1].strip()))
            except ValueError:
                pass
    return p.returncode, fails


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--before")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.workbook)
    screens = ",".join(str(ws["F1"].value) for ws in screen_sheets(wb))
    if not screens:
        sys.exit("FAIL: workbook khong co working sheet nao")

    code, fails = run_gate(args.workbook, args.before, screens)
    print(f"{'BASELINE (khong tiem loi)':34} exit={code} fails={fails or '[]'}")
    if code != 0:
        sys.exit("FAIL: workbook dau vao phai PASS gate truoc khi self-test")

    tmp = tempfile.mkdtemp()
    bad = 0
    print()
    print(f"{'LOI DA TIEM':34} {'KY VONG':8} {'BAT DUOC':12} KQ")
    print("-" * 74)
    for name, want, inject in CASES:
        p = os.path.join(tmp, f"t_{want}_{abs(hash(name)) % 9999}.xlsx")
        shutil.copy(args.workbook, p)
        w = openpyxl.load_workbook(p)
        inject(w)
        w.save(p)
        code, fails = run_gate(p, args.before, screens)
        ok = code == 1 and want in fails
        if not ok:
            bad += 1
        print(f"{name:34} #{want:<7} {str(fails) if fails != 'CRASH' else 'CRASH':12} "
              f"{'OK' if ok else 'THAT BAI'}")
    shutil.rmtree(tmp, ignore_errors=True)
    print()
    n = len(CASES)
    print(f"{n - bad}/{n} case bat dung check ky vong")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
